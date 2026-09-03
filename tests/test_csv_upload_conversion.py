"""CSV uploads become single-sheet workbooks so guided setup can scan them."""

import io
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

import web_server as ws
from sia.agent.job_manager import job_manager
from sia.modules.csv_workbook import (
    CsvConversionError,
    convert_csv_to_xlsx,
    excel_safe_sheet_name,
    read_csv_grid,
)


MESSY_CSV = (
    "Danone France Amazon DSP\n"
    "Exported 2026-08-12\n"
    "\n"
    "Date,Channel,Spend,Product Code\n"
    "2026-01-05,Display,1234.50,00734\n"
    "2026-01-06,Video,987,00812\n"
)


def _write(tmp_path: Path, name: str, text: str) -> Path:
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return path


def test_grid_preserves_title_and_blank_rows(tmp_path):
    rows, delimiter, _ = read_csv_grid(_write(tmp_path, "messy.csv", MESSY_CSV))

    assert delimiter == ","
    # Layout demarcation keys off narrow title rows and blank separators.
    assert rows[0] == ["Danone France Amazon DSP"]
    assert rows[2] == []
    assert rows[3] == ["Date", "Channel", "Spend", "Product Code"]


def test_convert_writes_single_sheet_named_after_file(tmp_path):
    result = convert_csv_to_xlsx(_write(tmp_path, "Danone France Amazon DSP.csv", MESSY_CSV))

    assert result["sheet_name"] == "Danone France Amazon DSP"
    assert result["path"].suffix == ".xlsx"

    workbook = load_workbook(result["path"])
    try:
        assert workbook.sheetnames == ["Danone France Amazon DSP"]
    finally:
        workbook.close()


def test_numbers_stay_numeric_and_codes_stay_text(tmp_path):
    result = convert_csv_to_xlsx(_write(tmp_path, "types.csv", MESSY_CSV))
    frame = pd.read_excel(result["path"], header=None)

    assert frame.iloc[4, 2] == 1234.50
    assert frame.iloc[5, 2] == 987
    # Leading zeros identify a product, not a quantity.
    assert frame.iloc[4, 3] == "00734"


def test_semicolon_delimiter_is_detected(tmp_path):
    csv_path = _write(
        tmp_path,
        "euro.csv",
        "Date;Channel;Spend\n2026-01-05;Display;12\n2026-01-06;Video;15\n",
    )
    rows, delimiter, _ = read_csv_grid(csv_path)

    assert delimiter == ";"
    assert rows[0] == ["Date", "Channel", "Spend"]


def test_non_utf8_bytes_still_read(tmp_path):
    csv_path = tmp_path / "latin.csv"
    csv_path.write_bytes("Marché,Spend\nFrance,10\n".encode("cp1252"))

    rows, _, encoding = read_csv_grid(csv_path)

    assert rows[0][0] == "Marché"
    assert encoding in {"utf-8", "cp1252", "latin-1"}


def test_empty_csv_is_rejected(tmp_path):
    with pytest.raises(CsvConversionError):
        read_csv_grid(_write(tmp_path, "blank.csv", "\n\n"))


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("Q1:Q2 [draft]", "Q1 Q2 draft"),
        ("a" * 40, "a" * 31),
        ("   ", "Sheet1"),
    ],
)
def test_sheet_name_sanitized_for_excel(raw, expected):
    assert excel_safe_sheet_name(raw) == expected


def test_upload_converts_csv_and_registers_a_real_sheet(tmp_path, monkeypatch):
    monkeypatch.setattr(ws, "UPLOAD_FOLDER", tmp_path)
    job_id = "csvjob01"

    client = ws.app.test_client()
    res = client.post(
        "/api/upload",
        data={
            "file": (io.BytesIO(MESSY_CSV.encode("utf-8")), "Danone France Amazon DSP.csv"),
            "type": "data",
            "job_id": job_id,
        },
        content_type="multipart/form-data",
    )

    assert res.status_code == 200
    body = res.get_json()
    assert body["converted_from_csv"] is True
    assert body["sheets"] == ["Danone France Amazon DSP"]

    job = job_manager.get_job(job_id)
    source = job["source_registry"][0]
    # Previously this was a sheet-less ":default" source with nothing to scan.
    assert source["sheet_name"] == "Danone France Amazon DSP"
    assert not source["source_id"].endswith(":default")
    assert str(source["file_path"]).endswith(".xlsx")
    assert Path(source["file_path"]).exists()


def test_upload_rejects_unreadable_csv(tmp_path, monkeypatch):
    monkeypatch.setattr(ws, "UPLOAD_FOLDER", tmp_path)

    client = ws.app.test_client()
    res = client.post(
        "/api/upload",
        data={
            "file": (io.BytesIO(b"   \n"), "empty.csv"),
            "type": "data",
            "job_id": "csvjob02",
        },
        content_type="multipart/form-data",
    )

    assert res.status_code == 400
    assert "Could not read" in res.get_json()["error"]
    assert list(tmp_path.glob("*.csv")) == []
