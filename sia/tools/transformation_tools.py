"""
Transformation Tools for Data Flattening.
Contains all tools the LLM can call to transform hierarchical spreadsheet data.
"""
import json
import logging
import re
from typing import Dict, List, Any, Optional, Union, Tuple
import pandas as pd
import numpy as np
from dataclasses import dataclass

from sia.agent.date_column_inference import infer_date_cadence_fast

logger = logging.getLogger(__name__)


@dataclass
class ToolResult:
    """Result of a tool execution."""
    success: bool
    data: Any  # The transformed data
    message: str = ""
    changes_made: Dict[str, Any] = None


@dataclass
class DeletionPreview:
    """Preview of what a destructive operation would delete.
    
    Used for HITL approval before executing destructive tools.
    Users can review and selectively approve/reject deletions.
    """
    tool_name: str
    tool_description: str = ""
    rows_to_delete: List[int] = None  # Row indices to be deleted
    columns_to_delete: List[str] = None  # Column names to be deleted
    sample_deleted_data: List[Dict[str, Any]] = None  # First 10 rows of data to be deleted
    full_deleted_data: pd.DataFrame = None  # Full data for Excel export
    reason: str = ""  # Why this data is being deleted
    impact_summary: str = ""  # e.g., "Will remove 15 rows (12% of data)"
    
    def __post_init__(self):
        if self.rows_to_delete is None:
            self.rows_to_delete = []
        if self.columns_to_delete is None:
            self.columns_to_delete = []
        if self.sample_deleted_data is None:
            self.sample_deleted_data = []
    
    
    def to_dict(self) -> Dict[str, Any]:
        """Convert to dictionary for JSON serialization (without full DataFrame)."""
        # Sanitize sample data (handle NaNs, timestamps)
        sanitized_sample = []
        for row in self.sample_deleted_data[:10]:
            clean_row = {}
            for k, v in row.items():
                if pd.isna(v):
                    clean_row[k] = None
                elif hasattr(v, 'isoformat'):
                    clean_row[k] = v.isoformat()
                elif hasattr(v, 'item'): # numpy types
                    clean_row[k] = v.item()
                else:
                    clean_row[k] = v
            sanitized_sample.append(clean_row)

        return {
            "tool_name": self.tool_name,
            "tool_description": self.tool_description,
            "rows_to_delete": self.rows_to_delete[:100],  # Limit for JSON
            "columns_to_delete": self.columns_to_delete,
            "sample_deleted_data": sanitized_sample,  # Cleaned sample
            "total_rows_to_delete": len(self.rows_to_delete),
            "deleted_count": len(self.rows_to_delete),  # NEW: Match JS expectation
            "total_columns_to_delete": len(self.columns_to_delete),
            "reason": self.reason,
            "impact_summary": self.impact_summary
        }
    
    def to_excel(self, filepath: str) -> str:
        """Export full deletion preview to Excel for user review."""
        if self.full_deleted_data is not None and not self.full_deleted_data.empty:
            self.full_deleted_data.to_excel(filepath, index=True, index_label="Original Row")
            return filepath
        return ""



# ===== Tool Registry =====
# Built from TOOL_SCHEMAS (single source of truth); see sia/tools/tool_validator.py.

def build_available_tools_from_tool_schemas() -> List[Dict[str, Any]]:
    """Planner-style tool listing: names, descriptions, and param hints from schemas."""
    from sia.tools.tool_validator import TOOL_SCHEMAS

    out = []
    for name in sorted(TOOL_SCHEMAS.keys()):
        schema = TOOL_SCHEMAS[name]
        params_desc = {}
        for param in schema.params:
            type_hint = getattr(param.param_type, "__name__", str(param.param_type))
            bits = []
            if not param.required:
                bits.append("optional")
            if not param.required and param.default is not None:
                bits.append("default={!r}".format(param.default))
            head = ("[" + ", ".join(bits) + "] ") if bits else ""
            tail = (param.description or "").strip() or "type: {}".format(type_hint)
            params_desc[param.name] = (head + tail).strip()
        out.append(
            {
                "name": name,
                "description": (schema.description or "").strip(),
                "params": params_desc,
            }
        )
    return out


AVAILABLE_TOOLS = build_available_tools_from_tool_schemas()


class TransformationTools:
    """Collection of data transformation tools."""
    
    # _resolve_columns consolidated at line 1603
    

    @staticmethod
    def extract_headers(grid, header_row: int) -> ToolResult:
        """
        Extract column headers from specified row.
        
        Args:
            grid: VisualGrid object or pandas DataFrame
            header_row: Row index (0-based) containing headers
            
        Returns:
            ToolResult with list of header names
        """
        try:
            if hasattr(grid, 'cells'):
                # VisualGrid object
                headers = []
                for col in range(grid.total_cols):
                    cell = grid.get_cell(header_row, col)
                    value = cell.value if cell else ""
                    headers.append(str(value) if value else f"Column_{col}")
            else:
                # DataFrame
                headers = list(grid.iloc[header_row])
                headers = [str(h) if h and not pd.isna(h) else f"Column_{i}" 
                          for i, h in enumerate(headers)]
            
            return ToolResult(
                success=True,
                data=headers,
                message=f"Extracted {len(headers)} headers from row {header_row}"
            )
        except Exception as e:
            return ToolResult(success=False, data=None, message=f"Error: {e}")

    @staticmethod
    def _cell_nonempty(val: Any) -> bool:
        if val is None or (isinstance(val, float) and np.isnan(val)):
            return False
        if isinstance(val, pd.Timestamp):
            return True
        s = str(val).strip()
        if not s or s.lower() in ("nan", "none"):
            return False
        return True

    @staticmethod
    def _row_starts_block(row: pd.Series, block_cols: List[str]) -> bool:
        for c in block_cols:
            if TransformationTools._cell_nonempty(row[c]):
                return True
        return False

    @staticmethod
    def _block_segments(
        df: pd.DataFrame,
        block_start_columns: List[str],
    ) -> Tuple[List[Tuple[int, int]], List[str]]:
        """Contiguous row ranges using sparse block-start columns (pre-forward-fill)."""
        bs_resolved: List[str] = []
        for k in block_start_columns or []:
            c, _ = TransformationTools._fuzzy_find_column(df, k)
            if c and c in df.columns:
                bs_resolved.append(str(c))
        if not bs_resolved:
            return [], []

        n = len(df)
        block_starts: List[int] = []
        for i in range(n):
            if TransformationTools._row_starts_block(df.iloc[i], bs_resolved):
                block_starts.append(i)
        if not block_starts:
            return [], bs_resolved

        segments: List[Tuple[int, int]] = []
        if block_starts[0] > 0:
            segments.append((0, block_starts[0]))
        for b, next_b in zip(block_starts, block_starts[1:] + [n]):
            segments.append((b, next_b))
        return segments, bs_resolved

    @staticmethod
    def _child_numeric_fill_rate_in_segments(
        series: pd.Series,
        segments: List[Tuple[int, int]],
    ) -> float:
        """Mean share of continuation rows with a numeric value per multi-row segment."""
        rates: List[float] = []
        numeric = pd.to_numeric(series, errors="coerce")
        for s, e in segments:
            if e - s < 2:
                continue
            child = numeric.iloc[s + 1 : e]
            nrow_c = int(e - s - 1)
            if nrow_c <= 0:
                continue
            rates.append(float(child.notna().sum()) / float(nrow_c))
        return float(sum(rates) / len(rates)) if rates else 0.0

    @staticmethod
    def _metric_has_row_level_children(
        series: pd.Series,
        segments: List[Tuple[int, int]],
        *,
        child_sparse_threshold: float = 0.25,
    ) -> bool:
        """True when child rows already carry the metric (do not block-allocate)."""
        return (
            TransformationTools._child_numeric_fill_rate_in_segments(series, segments)
            >= float(child_sparse_threshold)
        )

    @staticmethod
    def _metric_segment_looks_like_repeated_merge_block(
        series: pd.Series,
        start: int,
        end: int,
        *,
        rtol: float = 1e-6,
    ) -> bool:
        """True when every numeric value in the segment is the same (Excel merge repeat)."""
        if end - start < 2:
            return False
        num = pd.to_numeric(series.iloc[start:end], errors="coerce")
        finite = num[np.isfinite(num)]
        if len(finite) < 2:
            return False
        ref = float(finite.iloc[0])
        if ref == 0.0:
            return bool(float((finite - ref).abs().max()) <= rtol)
        return bool(float(((finite - ref).abs() / max(abs(ref), 1.0)).max()) <= rtol)

    @staticmethod
    def _resolve_merged_metric_range_specs(
        df: pd.DataFrame,
        merged_metric_ranges: Optional[List[Dict[str, Any]]],
    ) -> List[Dict[str, Any]]:
        """Fuzzy-resolve column names and clamp row spans to the dataframe."""
        resolved: List[Dict[str, Any]] = []
        n = len(df)
        for spec in merged_metric_ranges or []:
            if not isinstance(spec, dict):
                continue
            col_raw = spec.get("column_name") or spec.get("metric_col")
            if not col_raw:
                continue
            cname, _ = TransformationTools._fuzzy_find_column(df, col_raw)
            if not cname or cname not in df.columns:
                continue
            try:
                s = int(spec.get("start_row", 0))
                e = int(spec.get("end_row", 0))
            except (TypeError, ValueError):
                continue
            if e <= s or s >= n:
                continue
            s = max(0, s)
            e = min(n, e)
            if e - s < 2:
                continue
            resolved.append(
                {
                    "excel_range": str(spec.get("excel_range") or ""),
                    "column_name": str(cname),
                    "start_row": s,
                    "end_row": e,
                    "rows_spanned": int(e - s),
                    "top_left_value": spec.get("top_left_value"),
                }
            )
        return resolved

    @staticmethod
    def _merged_segments_for_metric(
        merged_specs: List[Dict[str, Any]],
        metric_name: str,
    ) -> List[Tuple[int, int]]:
        segs: List[Tuple[int, int]] = []
        mnorm = str(metric_name or "").strip().lower()
        for spec in merged_specs:
            cname = str(spec.get("column_name") or "").strip()
            if cname.lower() != mnorm:
                continue
            segs.append((int(spec["start_row"]), int(spec["end_row"])))
        return segs

    @staticmethod
    def _block_total_from_segment(series: pd.Series, start: int, end: int) -> Optional[float]:
        num = pd.to_numeric(series.iloc[start:end], errors="coerce")
        if num.empty:
            return None
        parent = pd.to_numeric(num.iloc[0], errors="coerce")
        if np.isfinite(parent):
            return float(parent)
        finite = num[num.notna()]
        if finite.empty:
            return None
        return float(finite.iloc[0])

    @staticmethod
    def _allocate_metric_on_merged_segments(
        df: pd.DataFrame,
        metric_name: str,
        segments: List[Tuple[int, int]],
        method: str = "equal",
        weight_name: Optional[str] = None,
    ) -> Tuple[pd.DataFrame, Dict[str, Any]]:
        """Allocate a block total across explicit merged-cell row spans."""
        out = df.copy()
        mnorm = str(method or "equal").strip().lower()
        if mnorm not in ("equal", "by_weight"):
            mnorm = "equal"
        stats: Dict[str, Any] = {
            "blocks_processed": 0,
            "blocks_zero_total": 0,
            "blocks_skipped_single_row": 0,
            "blocks_skipped_row_level_children": 0,
            "merged_segments_used": len(segments),
        }
        numeric_series = pd.to_numeric(out[metric_name], errors="coerce")

        for s, e in segments:
            if s >= e:
                continue
            nrow = int(e - s)
            stats["blocks_processed"] += 1
            if nrow < 2:
                stats["blocks_skipped_single_row"] += 1
                continue

            repeated_merge = TransformationTools._metric_segment_looks_like_repeated_merge_block(
                numeric_series, s, e
            )
            child_slice = numeric_series.iloc[s + 1 : e]
            child_rate = float(child_slice.notna().sum()) / float(max(nrow - 1, 1))
            if child_rate >= 0.25 and not repeated_merge:
                stats["blocks_skipped_row_level_children"] += 1
                continue

            total_block = TransformationTools._block_total_from_segment(numeric_series, s, e)
            if total_block is None or not np.isfinite(total_block):
                stats["blocks_zero_total"] += 1
                continue

            if mnorm == "equal":
                alloc = pd.Series(total_block / float(nrow), index=out.index[s:e])
            else:
                wcol = weight_name or ""
                wser = (
                    pd.to_numeric(out.iloc[s:e][wcol], errors="coerce").clip(lower=0)
                    if wcol in out.columns
                    else pd.Series(0.0, index=out.index[s:e])
                )
                denom = float(wser.sum())
                if denom <= 0.0:
                    alloc = pd.Series(total_block / float(nrow), index=out.index[s:e])
                else:
                    alloc = (wser / denom) * total_block
            out.loc[out.index[s:e], metric_name] = alloc.values
            numeric_series = pd.to_numeric(out[metric_name], errors="coerce")

        return out, stats

    @staticmethod
    def _ffill_columns_within_segments(
        df: pd.DataFrame,
        columns: List[str],
        segments: List[Tuple[int, int]],
    ) -> pd.DataFrame:
        out = df.copy()
        for col in columns:
            if col not in out.columns:
                continue
            ser = out[col].copy()
            blank = ser.isna()
            if ser.dtype == object or str(ser.dtype) == "string":
                blank = blank | ser.astype(str).str.strip().eq("")
            ser = ser.mask(blank, np.nan)
            for s, e in segments:
                if s >= e:
                    continue
                seg = ser.iloc[s:e].ffill()
                ser.iloc[s:e] = seg.values
            out[col] = ser
        return out

    @staticmethod
    def _ffill_columns_within_merged_spans(
        df: pd.DataFrame,
        columns: List[str],
        merged_specs: List[Dict[str, Any]],
    ) -> Tuple[pd.DataFrame, List[str]]:
        """
        Forward-fill dimension values inside Excel vertical merge spans.

        ``merged_metric_ranges`` includes merges in **any** column (date, channel labels,
        spend). Spend columns use spans for allocation; dimension columns use spans here
        when block_start_columns are dense (e.g. channel on every row) so segment ffill
        cannot propagate parent-row dates.
        """
        out = df.copy()
        col_set = {str(c) for c in columns}
        filled: List[str] = []
        for spec in merged_specs or []:
            if not isinstance(spec, dict):
                continue
            cname = str(spec.get("column_name") or "").strip()
            if not cname or cname not in out.columns or cname not in col_set:
                continue
            if TransformationTools._spend_like_column_name(cname):
                continue
            try:
                s = int(spec.get("start_row", 0))
                e = int(spec.get("end_row", 0))
            except (TypeError, ValueError):
                continue
            if e <= s or s >= len(out):
                continue
            s = max(0, s)
            e = min(len(out), e)
            ser = out[cname].copy()
            blank = ser.isna()
            if ser.dtype == object or str(ser.dtype) == "string":
                blank = blank | ser.astype(str).str.strip().eq("")
            work = ser.mask(blank, np.nan)
            seg = work.iloc[s:e].ffill()
            if seg.isna().any() and seg.notna().any():
                seg = seg.bfill()
            ser.iloc[s:e] = seg.values
            out[cname] = ser
            if cname not in filled:
                filled.append(cname)
        return out, filled

    @staticmethod
    def _allocate_metric_on_segments(
        df: pd.DataFrame,
        metric_name: str,
        segments: List[Tuple[int, int]],
        method: str = "equal",
        weight_name: Optional[str] = None,
        *,
        child_sparse_threshold: float = 0.25,
    ) -> Tuple[pd.DataFrame, Dict[str, Any]]:
        out = df.copy()
        mnorm = str(method or "equal").strip().lower()
        stats: Dict[str, Any] = {
            "blocks_processed": 0,
            "blocks_zero_total": 0,
            "blocks_skipped_single_row": 0,
            "blocks_skipped_row_level_children": 0,
        }
        numeric_series = pd.to_numeric(out[metric_name], errors="coerce")

        for s, e in segments:
            if s >= e:
                continue
            nrow = int(e - s)
            stats["blocks_processed"] += 1
            if nrow < 2:
                stats["blocks_skipped_single_row"] += 1
                continue

            child_slice = numeric_series.iloc[s + 1 : e]
            child_rate = float(child_slice.notna().sum()) / float(max(nrow - 1, 1))
            repeated_merge = TransformationTools._metric_segment_looks_like_repeated_merge_block(
                numeric_series, s, e
            )
            if child_rate >= float(child_sparse_threshold) and not repeated_merge:
                stats["blocks_skipped_row_level_children"] += 1
                continue

            parent_val = pd.to_numeric(numeric_series.iloc[s], errors="coerce")
            if not np.isfinite(parent_val):
                stats["blocks_zero_total"] += 1
                continue
            total_block = float(parent_val)

            if mnorm == "equal":
                alloc = pd.Series(total_block / float(nrow), index=out.index[s:e])
            else:
                wcol = weight_name or ""
                wser = (
                    pd.to_numeric(out.iloc[s:e][wcol], errors="coerce").clip(lower=0)
                    if wcol in out.columns
                    else pd.Series(0.0, index=out.index[s:e])
                )
                denom = float(wser.sum())
                if denom <= 0.0:
                    alloc = pd.Series(total_block / float(nrow), index=out.index[s:e])
                else:
                    alloc = (wser / denom) * total_block
            out.loc[out.index[s:e], metric_name] = alloc.values
            numeric_series = pd.to_numeric(out[metric_name], errors="coerce")

        return out, stats

    @staticmethod
    def _spend_like_column_name(name: str) -> bool:
        n = str(name or "").strip().lower().replace(" ", "_")
        keys = ("spend", "spent", "budget", "cost", "investment", "fee", "billing")
        return any(k in n for k in keys)

    @staticmethod
    def _weight_like_column_name(name: str) -> bool:
        n = str(name or "").strip().lower().replace(" ", "_")
        keys = ("impression", "impr", "view", "grp", "reach", "click", "engagement")
        return any(k in n for k in keys)

    @staticmethod
    def fill_merged_cells(df: pd.DataFrame, direction: str = "down", 
                          columns: List[Union[int, str]] = None) -> ToolResult:
        """
        Fill merged cell values down or right.
        
        Args:
            df: DataFrame to modify
            direction: 'down' or 'right'
            columns: Optional list of column indices or names to fill
            
        Returns:
            ToolResult with modified DataFrame
        """
        try:
            df = df.copy()
            cols, _ = TransformationTools._resolve_columns(df, columns)
            
            if direction == "down":
                for col in cols:
                    df[col] = df[col].ffill()
            elif direction == "right":
                # Fill right across specified columns
                # We apply ffill horizontally across the dataframe snippet
                df.update(df[cols].ffill(axis=1))
            
            return ToolResult(
                success=True,
                data=df,
                message=f"Filled merged cells {direction} for {len(cols)} columns"
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in fill_merged_cells: {e}")

    @staticmethod
    def unmerge_and_fill(df: pd.DataFrame, direction: str = "down",
                        columns: List[Union[int, str]] = None) -> ToolResult:
        """
        Alias for fill_merged_cells with clearer naming.
        Unmerges cells and fills values down or right to propagate parent values.
        
        Args:
            df: DataFrame to modify
            direction: 'down' or 'right'
            columns: Optional list of column indices or names to fill
        """
        return TransformationTools.fill_merged_cells(df, direction=direction, columns=columns)

    @staticmethod
    def allocate_block_metric(
        df: pd.DataFrame,
        metric_col: Union[str, int],
        block_start_columns: List[Any],
        method: str = "equal",
        weight_col: Optional[Union[str, int]] = None,
        row_filters: Optional[List[Dict[str, Any]]] = None,
        merged_metric_ranges: Optional[List[Dict[str, Any]]] = None,
    ) -> ToolResult:
        """
        Split a **block-level** metric (one non-null cell per contiguous block—e.g. contract
        total on the parent row) across all rows in that block so weekly/daily rollup can SUM safely.

        A **block** is a maximal row range `[start, end)` where only the **first** row(s) marked
        by ``block_start_columns`` begin a block: each row with any non-empty value in one of those
        columns starts a new block; following rows belong to it until the next such row.

        **Do not** use ``transform.fill_merged`` on spend/budget columns for this pattern—that
        duplicates totals. Use allocate, then ``transform.aggregate_weekly`` with explicit
        ``metric_rules``.

        **Explicit split rules (same as “Expand grouped rows with metric semantics”):**

        - **Option A — equal split** (``method="equal"``): for each block,
          ``row_metric = block_total / number_of_rows_in_block`` (every post row gets the same share).
        - **Option B — weighted split** (``method="by_weight"``, ``weight_col`` e.g. impressions):
          ``row_metric = block_total * (row_weight / sum_of_weights_in_block)`` with non‑negative weights;
          if weights sum to 0 within a block, falls back to equal split for that block.

        Typical ``block_start_columns``: ``Ref``, ``Name``, or whatever is non-null only on the
        header row of each influencer/section block.
        """
        try:
            if df is None or df.empty:
                return ToolResult(success=True, data=df, message="Empty input; nothing to allocate.")

            work, filter_notes = TransformationTools._apply_simple_row_filters(df, row_filters)
            if work.empty:
                return ToolResult(
                    success=False,
                    data=df,
                    message="All rows removed by row_filters; nothing to allocate."
                    + (f" Notes: {'; '.join(filter_notes)}" if filter_notes else ""),
                )

            metric_name, _ = TransformationTools._fuzzy_find_column(work, metric_col)
            if not metric_name or metric_name not in work.columns:
                return ToolResult(success=False, data=df, message=f"metric_col {metric_col!r} not found.")

            bs_resolved = []
            for k in block_start_columns or []:
                c, _ = TransformationTools._fuzzy_find_column(work, k)
                if c and c in work.columns:
                    bs_resolved.append(c)
            if not bs_resolved:
                return ToolResult(
                    success=False,
                    data=df,
                    message="block_start_columns resolved to empty; pass at least one existing column.",
                )

            mnorm = str(method or "equal").strip().lower()
            if mnorm not in ("equal", "by_weight"):
                return ToolResult(
                    success=False,
                    data=df,
                    message=f"method must be 'equal' or 'by_weight', got {method!r}",
                )

            weight_name: Optional[str] = None
            if mnorm == "by_weight":
                if weight_col is None:
                    return ToolResult(
                        success=False,
                        data=df,
                        message="weight_col is required when method='by_weight'.",
                    )
                weight_name, _ = TransformationTools._fuzzy_find_column(work, weight_col)
                if not weight_name or weight_name not in work.columns:
                    return ToolResult(success=False, data=df, message=f"weight_col {weight_col!r} not found.")

            out = work.copy()
            merged_specs = TransformationTools._resolve_merged_metric_range_specs(
                out, merged_metric_ranges
            )
            merged_segments = TransformationTools._merged_segments_for_metric(
                merged_specs, metric_name
            )
            if merged_segments:
                out, changes_made = TransformationTools._allocate_metric_on_merged_segments(
                    out,
                    metric_name,
                    merged_segments,
                    method=mnorm,
                    weight_name=weight_name,
                )
                seg_desc = f"{len(merged_segments)} merged range(s)"
            else:
                segments, bs_resolved = TransformationTools._block_segments(out, bs_resolved)
                if not segments:
                    return ToolResult(
                        success=False,
                        data=out,
                        message=(
                            "No block-start rows detected (all block_start_columns empty). "
                            "Check column choices or sheet structure."
                        ),
                    )
                out, changes_made = TransformationTools._allocate_metric_on_segments(
                    out,
                    metric_name,
                    segments,
                    method=mnorm,
                    weight_name=weight_name,
                )
                seg_desc = f"block keys: {', '.join(bs_resolved)}"

            msg = (
                f"Allocated block-level {metric_name!r} across {changes_made['blocks_processed']} segment(s) "
                f"using method={mnorm!r} ({seg_desc})."
            )
            if changes_made.get("blocks_zero_total"):
                msg += f" {changes_made['blocks_zero_total']} segment(s) had no numeric total in {metric_name!r}."
            if filter_notes:
                msg += " Filters: " + "; ".join(filter_notes)

            return ToolResult(
                success=True,
                data=out,
                message=msg,
                changes_made={
                    **changes_made,
                    "metric_col": metric_name,
                    "block_start_columns": bs_resolved,
                    "method": mnorm,
                    "weight_col": weight_name,
                    "merged_metric_ranges": merged_specs,
                },
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in allocate_block_metric: {e}")

    @staticmethod
    def expand_grouped_block(
        df: pd.DataFrame,
        dimension_columns: List[Any],
        block_start_columns: Optional[List[Any]] = None,
        allocations: Optional[List[Dict[str, Any]]] = None,
        auto_detect_block_metrics: bool = True,
        child_numeric_sparse_threshold: float = 0.25,
        parent_numeric_rate_threshold: float = 0.55,
        row_filters: Optional[List[Dict[str, Any]]] = None,
        merged_metric_ranges: Optional[List[Dict[str, Any]]] = None,
    ) -> ToolResult:
        """
        Expand grouped influencer-style blocks in one pass (correct order):

        1. Segment rows using **sparse** ``block_start_columns`` (before forward-fill).
        2. Forward-fill **dimension_columns only** within each block (never metrics).
        3. Allocate **block-level** metrics across child rows (equal or weighted).

        Layout cases:
        - **group_header**: dimensions and some metrics only on the parent row; child rows carry
          row-level metrics (e.g. impressions). Dimensions are ffilled; block metrics split.
        - **row_dense**: metrics already on every child row — dimensions may still be ffilled;
          allocation is skipped for metrics that are dense on continuation rows.
        """
        try:
            if df is None or df.empty:
                return ToolResult(success=True, data=df, message="Empty input; nothing to expand.")

            work, filter_notes = TransformationTools._apply_simple_row_filters(df, row_filters)
            if work.empty:
                return ToolResult(
                    success=False,
                    data=df,
                    message="All rows removed by row_filters; cannot expand grouped blocks."
                    + (f" Notes: {'; '.join(filter_notes)}" if filter_notes else ""),
                )

            dim_resolved, _ = TransformationTools._resolve_columns(work, dimension_columns)
            if not dim_resolved:
                return ToolResult(
                    success=False,
                    data=df,
                    message="dimension_columns resolved to empty; pass at least one dimension column.",
                )

            bs_keys = list(block_start_columns or dimension_columns)
            segments, bs_resolved = TransformationTools._block_segments(work, bs_keys)
            if not segments:
                return ToolResult(
                    success=False,
                    data=work,
                    message=(
                        "No block-start rows detected for expand_grouped_block. "
                        "Use sparse delimiter columns (filled only on section header rows)."
                    ),
                )

            multi_segments = [(s, e) for (s, e) in segments if (e - s) >= 2]
            out = TransformationTools._ffill_columns_within_segments(work, dim_resolved, segments)
            merged_specs = TransformationTools._resolve_merged_metric_range_specs(
                out, merged_metric_ranges
            )
            merged_dim_ffill: List[str] = []
            if merged_specs:
                out, merged_dim_ffill = TransformationTools._ffill_columns_within_merged_spans(
                    out, dim_resolved, merged_specs
                )

            dim_set = set(dim_resolved)
            alloc_specs: List[Dict[str, Any]] = list(allocations or [])
            metric_reports: List[Dict[str, Any]] = []

            if auto_detect_block_metrics and not alloc_specs:
                for col in work.columns:
                    cname = str(col)
                    if cname in dim_set:
                        metric_reports.append(
                            {"column": cname, "grain": "dimension", "action": "forward_fill_only"}
                        )
                        continue
                    num = pd.to_numeric(work[cname], errors="coerce")
                    if float(num.notna().mean()) < 0.04:
                        continue
                    if not multi_segments:
                        metric_reports.append(
                            {
                                "column": cname,
                                "grain": "unspecified_single_row_blocks",
                                "action": "none",
                            }
                        )
                        continue

                    parent_hits = 0
                    child_rates: List[float] = []
                    for s, e in multi_segments:
                        parent_has = bool(np.isfinite(num.iloc[s]))
                        nrow_c = max(e - s - 1, 0)
                        child_finite = int(num.iloc[s + 1 : e].notna().sum()) if nrow_c > 0 else 0
                        cr = float(child_finite / nrow_c) if nrow_c > 0 else 0.0
                        child_rates.append(cr)
                        if parent_has:
                            parent_hits += 1
                    parent_rate = float(parent_hits / max(len(multi_segments), 1))
                    mean_child = float(sum(child_rates) / max(len(child_rates), 1))

                    if (
                        parent_rate >= parent_numeric_rate_threshold
                        and mean_child <= child_numeric_sparse_threshold
                        and TransformationTools._spend_like_column_name(cname)
                    ):
                        weight = None
                        method = "equal"
                        alloc_specs.append(
                            {
                                "metric_col": cname,
                                "method": method,
                                "weight_col": weight,
                            }
                        )
                        metric_reports.append(
                            {
                                "column": cname,
                                "grain": "block_header_total",
                                "action": "allocate",
                                "method": method,
                                "weight_col": weight,
                            }
                        )
                    elif mean_child >= child_numeric_sparse_threshold:
                        metric_reports.append(
                            {"column": cname, "grain": "post_row_metric", "action": "none"}
                        )
                    else:
                        metric_reports.append(
                            {"column": cname, "grain": "ambiguous", "action": "none"}
                        )

            alloc_results: List[Dict[str, Any]] = []
            skipped_row_level: List[str] = []
            child_thresh = float(child_numeric_sparse_threshold)
            for spec in alloc_specs:
                if not isinstance(spec, dict):
                    continue
                metric_col = spec.get("metric_col")
                if not metric_col:
                    continue
                metric_name, _ = TransformationTools._fuzzy_find_column(out, metric_col)
                if not metric_name or metric_name in dim_set:
                    continue
                metric_series = pd.to_numeric(out[metric_name], errors="coerce")
                merged_segments = TransformationTools._merged_segments_for_metric(
                    merged_specs, metric_name
                )
                if (
                    not merged_segments
                    and not spec.get("force_allocate")
                    and TransformationTools._metric_has_row_level_children(
                        metric_series,
                        segments,
                        child_sparse_threshold=child_thresh,
                    )
                ):
                    skipped_row_level.append(metric_name)
                    metric_reports.append(
                        {
                            "column": metric_name,
                            "grain": "post_row_metric",
                            "action": "skipped_allocation",
                            "reason": "child rows already have numeric values",
                        }
                    )
                    continue
                method = str(spec.get("method") or "equal").strip().lower()
                if method not in ("equal", "by_weight"):
                    method = "equal"
                weight_name: Optional[str] = None
                if method == "by_weight":
                    wc = spec.get("weight_col")
                    if wc:
                        weight_name, _ = TransformationTools._fuzzy_find_column(out, wc)
                if merged_segments:
                    out, st = TransformationTools._allocate_metric_on_merged_segments(
                        out,
                        metric_name,
                        merged_segments,
                        method=method,
                        weight_name=weight_name,
                    )
                else:
                    out, st = TransformationTools._allocate_metric_on_segments(
                        out,
                        metric_name,
                        segments,
                        method=method,
                        weight_name=weight_name,
                        child_sparse_threshold=child_thresh,
                    )
                alloc_results.append(
                    {
                        "metric_col": metric_name,
                        **st,
                        "method": method,
                        "weight_col": weight_name,
                        "used_merged_ranges": bool(merged_segments),
                    }
                )

            msg = (
                f"expand_grouped_block: ffilled {len(dim_resolved)} dimension column(s) within "
                f"{len(segments)} segment(s); allocated {len(alloc_results)} block-level metric(s). "
                f"Block keys: {', '.join(bs_resolved)}."
            )
            if skipped_row_level:
                msg += (
                    " Skipped block allocation for row-level metric(s): "
                    + ", ".join(skipped_row_level)
                    + " (values already on child rows)."
                )
            if filter_notes:
                msg += " Filters: " + "; ".join(filter_notes)

            return ToolResult(
                success=True,
                data=out,
                message=msg,
                changes_made={
                    "dimension_columns": dim_resolved,
                    "block_start_columns": bs_resolved,
                    "segments_total": len(segments),
                    "segments_multi_row": len(multi_segments),
                    "allocations": alloc_results,
                    "skipped_row_level_metrics": skipped_row_level,
                    "metric_layout_report": metric_reports,
                    "merged_metric_ranges": merged_specs,
                    "merged_dimension_ffill_columns": merged_dim_ffill,
                },
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in expand_grouped_block: {e}")

    @staticmethod
    def classify_metric_level(
        df: pd.DataFrame,
        block_start_columns: Optional[List[Any]] = None,
        metric_columns: Optional[List[Any]] = None,
        child_numeric_sparse_threshold: float = 0.25,
        parent_numeric_rate_threshold: float = 0.55,
        row_numeric_dense_threshold: float = 0.42,
        min_numeric_global_rate: float = 0.04,
        row_filters: Optional[List[Dict[str, Any]]] = None,
    ) -> ToolResult:
        """
        Inspect numeric columns and estimate whether values are authored at **block header** grain
        (one total per block, continuation rows sparse) versus **post / row** grain (values commonly
        present on continuation rows). Intended to steer ``transform.allocate_block_metric``
        parameters (equal vs weighted) and validation of ``block_start_columns``.

        Pair with ``allocate_block_metric`` when expanding grouped influencer-style layouts to a
        flat, summable table.
        """
        split_reference = {
            "equal_split": "row_metric = block_total / number_of_rows_in_block",
            "weighted_split": "row_metric = block_total * (row_weight / sum_weights_in_block)",
            "executor_tool": "transform.allocate_block_metric",
            "method_equal": "equal",
            "method_weighted": "by_weight + weight_col",
        }

        def _norm_header(name: str) -> str:
            return str(name or "").strip().lower().replace(" ", "_").replace("-", "_")

        def _spend_like(name: str) -> bool:
            n = _norm_header(name)
            keys = (
                "spend",
                "spent",
                "budget",
                "cost",
                "investment",
                "invest",
                "fee",
                "payout",
                "contract",
                "billing",
                "estimate",
                "forecast",
                "charge",
                "invoice",
            )
            return any(k in n for k in keys)

        def _weight_like(name: str) -> bool:
            n = _norm_header(name)
            keys = (
                "impression",
                "impr",
                "view",
                "grp",
                "reach",
                "click",
                "engagement",
                "delivery",
                "traffic",
            )
            return any(k in n for k in keys)

        def _nonempty_cell(val: Any) -> bool:
            if val is None or (isinstance(val, float) and np.isnan(val)):
                return False
            if isinstance(val, pd.Timestamp):
                return True
            s = str(val).strip()
            if not s or s.lower() in ("nan", "none"):
                return False
            return True

        try:
            if df is None or df.empty:
                return ToolResult(
                    success=True,
                    data=df,
                    message=json.dumps({"columns": [], "split_rule_reference": split_reference}),
                    changes_made={"metric_level_report": {"columns": [], "split_rule_reference": split_reference}},
                )

            work, filter_notes = TransformationTools._apply_simple_row_filters(df, row_filters)
            if work.empty:
                return ToolResult(
                    success=False,
                    data=df,
                    message="All rows removed by row_filters; nothing to classify.",
                )

            cand_cols: List[str]
            if metric_columns:
                cand_cols = []
                for mc in metric_columns:
                    nm, _ = TransformationTools._fuzzy_find_column(work, mc)
                    if nm and nm in work.columns:
                        cand_cols.append(nm)
            else:
                cand_cols = list(work.columns)

            scored_cols: List[str] = []
            for c in cand_cols:
                ns = pd.to_numeric(work[c], errors="coerce")
                finite_rate = float(ns.notna().mean()) if len(work.index) else 0.0
                if finite_rate >= float(min_numeric_global_rate):
                    scored_cols.append(str(c))

            if not scored_cols:
                return ToolResult(
                    success=True,
                    data=work.copy(),
                    message=json.dumps(
                        {
                            "columns": [],
                            "notes": ["No numeric-ish columns exceeded min_numeric_global_rate."],
                            "split_rule_reference": split_reference,
                        },
                        indent=2,
                    ),
                    changes_made={
                        "metric_level_report": {
                            "columns": [],
                            "filter_notes": filter_notes,
                            "split_rule_reference": split_reference,
                        },
                    },
                )

            bs_resolved: List[str] = []
            boundary_error: Optional[str] = None
            if block_start_columns:
                for k in block_start_columns:
                    c, _ = TransformationTools._fuzzy_find_column(work, k)
                    if c and c in work.columns:
                        bs_resolved.append(str(c))
                if not bs_resolved:
                    boundary_error = "block_start_columns resolved to empty; pass existing columns for block-aware grain."
            else:
                boundary_error = "no_block_boundaries"

            segment_bounds: List[Tuple[int, int]] = []
            if bs_resolved:
                n = len(work)
                block_starts: List[int] = []
                for i in range(n):
                    row = work.iloc[i]
                    if any(_nonempty_cell(row[c]) for c in bs_resolved):
                        block_starts.append(i)
                if not block_starts:
                    boundary_error = (
                        "No block-start rows detected (all block_start_columns empty). "
                        "Cannot infer block header vs post-row grain."
                    )
                else:
                    if block_starts[0] > 0:
                        segment_bounds.append((0, block_starts[0]))
                    for b, next_b in zip(block_starts, block_starts[1:] + [n]):
                        segment_bounds.append((b, next_b))

            col_reports: List[Dict[str, Any]] = []
            for col in scored_cols:
                num = pd.to_numeric(work[col], errors="coerce")
                global_rate = float(num.notna().mean()) if len(work.index) else 0.0
                base: Dict[str, Any] = {
                    "column": col,
                    "finite_value_rate_global": round(global_rate, 4),
                    "spend_like_name": bool(_spend_like(col)),
                    "weight_like_name": bool(_weight_like(col)),
                }

                if boundary_error or not segment_bounds:
                    base.update(
                        {
                            "metric_grain_level": "unspecified_without_boundaries"
                            if boundary_error == "no_block_boundaries"
                            else ("ambiguous_boundary_signal" if boundary_error else "unspecified"),
                            "confidence": None,
                            "stats": {},
                            "suggested_allocate_block_metric": None,
                            "boundary_error": boundary_error,
                            "split_rule_reference": split_reference,
                        }
                    )
                    col_reports.append(base)
                    continue

                multi_segments = [(s, e) for (s, e) in segment_bounds if e > s >= 0 and (e - s) >= 2]
                single_segments = [(s, e) for (s, e) in segment_bounds if e > s and (e - s) == 1]

                if not multi_segments:
                    base.update(
                        {
                            "metric_grain_level": "insufficient_multi_row_blocks",
                            "confidence": 0.0,
                            "stats": {
                                "segments_total": len(segment_bounds),
                                "segments_multi_row": 0,
                                "segments_single_row": len(single_segments),
                            },
                            "notes": ["Need blocks with ≥2 rows to distinguish header vs continuation grain."],
                            "suggested_allocate_block_metric": None,
                            "split_rule_reference": split_reference,
                        }
                    )
                    col_reports.append(base)
                    continue

                parent_hits = 0
                child_rates: List[float] = []
                parent_child_both_rates: List[float] = []

                for s, e in multi_segments:
                    parent_has = bool(np.isfinite(num.iloc[s]))
                    nrow_c = max(e - s - 1, 0)
                    child_finite = int(num.iloc[s + 1 : e].notna().sum()) if nrow_c > 0 else 0
                    cr = float(child_finite / nrow_c) if nrow_c > 0 else 0.0
                    child_rates.append(cr)
                    if parent_has:
                        parent_hits += 1
                    if nrow_c > 0 and parent_has and child_finite > 0:
                        parent_child_both_rates.append(cr)

                n_m = len(multi_segments)
                parent_rate_f = float(parent_hits / max(n_m, 1))
                mean_child = float(sum(child_rates) / max(len(child_rates), 1))

                block_sparse = (
                    parent_rate_f >= float(parent_numeric_rate_threshold)
                    and mean_child <= float(child_numeric_sparse_threshold)
                )
                row_dense = mean_child >= float(row_numeric_dense_threshold)

                level = "ambiguous"
                confidence = 0.55
                if block_sparse and not row_dense:
                    level = "block_header_total"
                    confidence = min(0.95, 0.55 + 0.25 * (parent_rate_f - mean_child))
                elif row_dense and not block_sparse:
                    level = "post_row_metric"
                    confidence = min(0.95, 0.5 + 0.35 * mean_child)
                elif row_dense and block_sparse:
                    level = "mixed_parent_and_children"
                    confidence = 0.45

                suggested: Optional[Dict[str, Any]] = None
                if level == "block_header_total" and _spend_like(col):
                    suggested = {
                        "metric_col": col,
                        "block_start_columns": list(bs_resolved),
                        "method": "equal",
                        "weight_col": None,
                        "rationale": (
                            "Block-header total on spend-like metric; equal split across rows in block: "
                            + split_reference["equal_split"]
                        ),
                    }

                base.update(
                    {
                        "metric_grain_level": level,
                        "confidence": round(float(confidence), 3),
                        "stats": {
                            "blocks_multi_row": n_m,
                            "parent_row_numeric_rate": round(parent_rate_f, 4),
                            "mean_child_row_numeric_rate": round(mean_child, 4),
                            "both_parent_and_child_populated_mean_child_rate": (
                                round(
                                    float(sum(parent_child_both_rates) / max(len(parent_child_both_rates), 1)),
                                    4,
                                )
                                if parent_child_both_rates
                                else None
                            ),
                            "block_start_columns": list(bs_resolved),
                        },
                        "suggested_allocate_block_metric": suggested,
                        "split_rule_reference": split_reference,
                    }
                )
                col_reports.append(base)

            # Note candidate weight columns for planner review (do not auto-select by_weight).
            by_name = {r["column"]: r for r in col_reports}
            for col, rep in by_name.items():
                sug = rep.get("suggested_allocate_block_metric")
                if sug and sug.get("method") == "equal" and _spend_like(col):
                    weight_candidates = [
                        other
                        for other, orep in by_name.items()
                        if other != col
                        and orep.get("metric_grain_level") == "post_row_metric"
                        and _weight_like(other)
                    ]
                    weight_candidates.sort()
                    if weight_candidates:
                        sug["weight_col_candidates"] = weight_candidates
                        sug["rationale"] = (
                            f"Block-header total on spend-like metric; equal split is the default suggestion. "
                            f"Analyst may choose weighted split using {weight_candidates[0]!r}: "
                            f"{split_reference['weighted_split']}."
                        )

            report = {
                "columns": col_reports,
                "split_rule_reference": split_reference,
                "filter_notes": filter_notes,
            }
            if boundary_error and boundary_error != "no_block_boundaries":
                report["boundary_error"] = boundary_error

            return ToolResult(
                success=True,
                data=work.copy(),
                message=json.dumps(report, indent=2),
                changes_made={"metric_level_report": report},
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in classify_metric_level: {e}")

    @staticmethod
    def infer_block_boundary_columns(
        df: pd.DataFrame,
        exclude_columns: Optional[List[Any]] = None,
        metric_columns_hint: Optional[List[Any]] = None,
        min_multi_row_blocks: int = 2,
        min_nonempty_row_rate: float = 0.01,
        max_nonempty_row_rate: float = 0.92,
        min_mean_segment_len: float = 1.25,
        max_pair_search: int = 6,
        row_filters: Optional[List[Dict[str, Any]]] = None,
    ) -> ToolResult:
        """
        Inspect **tabular** dataframe columns using the same segmentation rule as
        ``allocate_block_metric`` / ``classify_metric_level``: a row starts a block if **any**
        listed delimiter column cell is syntactically non-empty.

        Produces ranked single-column hypotheses and a small pairwise **OR-combination**
        sweep so analysts can revise ``block_start_columns`` when raw ids (e.g. ``Ref``) were noisy
        and dropped upstream.
        """
        def _nonempty_cell(val: Any) -> bool:
            if val is None or (isinstance(val, float) and np.isnan(val)):
                return False
            if isinstance(val, pd.Timestamp):
                return True
            s = str(val).strip()
            if not s or s.lower() in ("nan", "none"):
                return False
            return True

        def _segments_for_boundary_cols(
            frame: pd.DataFrame, colnames: List[str]
        ) -> Tuple[List[Tuple[int, int]], str]:
            n = len(frame)
            bs: List[int] = []
            for i in range(n):
                row = frame.iloc[i]
                if any(_nonempty_cell(row[c]) for c in colnames):
                    bs.append(i)
            if not bs:
                return [], "No boundary rows matched (chosen columns entirely empty)."
            seg: List[Tuple[int, int]] = []
            if bs[0] > 0:
                seg.append((0, bs[0]))
            for b, nxt in zip(bs, bs[1:] + [n]):
                seg.append((b, nxt))
            return seg, ""

        def _score_hypothesis(frame: pd.DataFrame, colnames: List[str]) -> Optional[Dict[str, Any]]:
            colnames = [str(c) for c in colnames if c]
            if len(colnames) != len(set(colnames)):
                return None
            if any(c not in frame.columns for c in colnames):
                return None
            segs, err = _segments_for_boundary_cols(frame, colnames)
            if err or not segs:
                return None
            multi = [(s, e) for (s, e) in segs if (e - s) >= 2]
            singletons = [(s, e) for (s, e) in segs if (e - s) == 1]
            if len(multi) < int(min_multi_row_blocks):
                return None
            lengths = [float(e - s) for (s, e) in multi]
            mean_seg = float(sum(lengths) / max(len(lengths), 1))
            if mean_seg + 1e-6 < float(min_mean_segment_len):
                return None
            nrow = max(len(frame.index), 1)
            nonempty_rows = 0
            for i in range(len(frame)):
                row = frame.iloc[i]
                if any(_nonempty_cell(row[c]) for c in colnames):
                    nonempty_rows += 1
            frac = nonempty_rows / float(nrow)
            if frac < float(min_nonempty_row_rate) or frac > float(max_nonempty_row_rate):
                return None
            denom = max(len(multi), 1)
            orphan_ratio = len(singletons) / denom
            score = (
                float(len(multi))
                * (mean_seg**0.5)
                / (1.0 + 2.5 * orphan_ratio + 0.5 * frac)
            )
            return {
                "columns": list(colnames),
                "score": round(score, 4),
                "stats": {
                    "multi_row_blocks": len(multi),
                    "singleton_segments": len(singletons),
                    "mean_multi_segment_length": round(mean_seg, 4),
                    "nonempty_boundary_row_fraction": round(frac, 4),
                    "segments_total": len(segs),
                },
            }

        try:
            if df is None or df.empty:
                payload = {"recommended_block_start_columns": [], "candidates": [], "notes": ["Empty dataframe."]}
                return ToolResult(
                    success=True,
                    data=df,
                    message=json.dumps(payload),
                    changes_made={"block_boundary_inference": payload},
                )

            work, filter_notes = TransformationTools._apply_simple_row_filters(df, row_filters)
            if work.empty:
                return ToolResult(
                    success=False,
                    data=df,
                    message="All rows removed by row_filters; cannot infer boundaries.",
                )

            excluded = set()
            if exclude_columns:
                for k in exclude_columns:
                    try:
                        c, _ = TransformationTools._fuzzy_find_column(work, k)
                        if c and c in work.columns:
                            excluded.add(str(c))
                    except ValueError:
                        continue
            hinted_metrics = set()
            if metric_columns_hint:
                for k in metric_columns_hint:
                    try:
                        c, _ = TransformationTools._fuzzy_find_column(work, k)
                        if c and c in work.columns:
                            hinted_metrics.add(str(c))
                    except ValueError:
                        continue

            date_kw = {"date", "week", "month", "day", "year", "posting", "calendar", "time"}
            dim_candidates: List[str] = []
            for col in list(work.columns):
                scol = str(col)
                if scol in excluded or scol in hinted_metrics:
                    continue
                ser = work[col]
                num = pd.to_numeric(ser, errors="coerce")
                num_frac = float(num.notna().mean()) if len(work.index) else 0.0
                if num_frac > 0.90:
                    continue
                lc = scol.strip().lower()
                if sum(1 for kw in date_kw if kw in lc.replace("_", " ")) >= 2:
                    continue
                dim_candidates.append(scol)

            singles_ranked: List[Dict[str, Any]] = []
            for col in dim_candidates:
                ev = _score_hypothesis(work, [col])
                if ev:
                    singles_ranked.append(ev)
            singles_ranked.sort(key=lambda x: (-x["score"], x["columns"]))
            singles_ranked = singles_ranked[: max(int(max_pair_search), 12)]

            pair_candidates: List[Dict[str, Any]] = []
            cap = max(0, min(int(max_pair_search), len(singles_ranked)))
            for i in range(cap):
                for j in range(i + 1, cap):
                    a = singles_ranked[i]["columns"][0]
                    b = singles_ranked[j]["columns"][0]
                    if a == b:
                        continue
                    ev = _score_hypothesis(work, [a, b])
                    if ev:
                        ev["combine_mode"] = "or_non_empty_union"
                        pair_candidates.append(ev)
            pair_candidates.sort(key=lambda x: (-x["score"], x["columns"]))
            pair_candidates = pair_candidates[:12]

            all_ranked = list(singles_ranked) + pair_candidates
            all_ranked.sort(key=lambda x: (-x["score"], str(x["columns"])))

            recommended: List[str] = []
            if all_ranked:
                recommended = list(all_ranked[0]["columns"])

            payload: Dict[str, Any] = {
                "recommended_block_start_columns": recommended,
                "single_column_hypotheses": singles_ranked,
                "pair_hypotheses": pair_candidates[:8],
                "top_hypotheses": all_ranked[:10],
                "filter_notes": filter_notes,
                "notes": [
                    "Semantics match transform.allocate_block_metric: new block wherever ANY listed column "
                    "has a syntactically non-empty cell.",
                    "Do not rely on structure_analyzer grouped_rows labels—this scans the dataframe you pass.",
                    "If identifiers like Ref were dropped, pick the next-best delimiter listed here and update "
                    "classify_metric_level / allocate_block_metric params before running them.",
                    "Keep delimiter columns until after allocate_block_metric whenever possible.",
                ],
            }
            if not recommended:
                payload["notes"].append(
                    "No hypothesis passed thresholds—inspect sheet manually or widen max_nonempty_row_rate / "
                    "lower min_multi_row_blocks slightly."
                )

            return ToolResult(
                success=True,
                data=work.copy(),
                message=json.dumps(payload, indent=2),
                changes_made={"block_boundary_inference": payload},
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in infer_block_boundary_columns: {e}")

    @staticmethod
    def get_file_inventory(file_path: str) -> ToolResult:
        """
        Phase 1 Discovery: Lists all sheets, named ranges, hidden tabs,
        and dimensions of an Excel workbook.
        
        Use this as the FIRST tool before any analysis to understand
        the structure of the input file.
        
        Args:
            file_path: Path to the Excel file
            
        Returns:
            ToolResult with inventory dict containing:
            - sheets: list of {name, rows, cols, state, has_data}
            - named_ranges: list of named range definitions
            - file_metadata: {file_name, file_size_kb}
        """
        try:
            from openpyxl import load_workbook
            from pathlib import Path as FilePath
            
            file_info = FilePath(file_path)
            wb = load_workbook(file_path, read_only=True, data_only=True)
            
            sheets_info = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                sheets_info.append({
                    "name": sheet_name,
                    "rows": ws.max_row or 0,
                    "cols": ws.max_column or 0,
                    "state": ws.sheet_state,  # 'visible', 'hidden', 'veryHidden'
                    "has_data": (ws.max_row or 0) > 0 and (ws.max_column or 0) > 0
                })
            
            named_ranges = []
            for nr in wb.defined_names.values():
                named_ranges.append({
                    "name": nr.name,
                    "value": str(nr.value),
                    "scope": nr.localSheetId
                })
            
            wb.close()
            
            inventory = {
                "sheets": sheets_info,
                "named_ranges": named_ranges,
                "total_sheets": len(sheets_info),
                "visible_sheets": sum(1 for s in sheets_info if s["state"] == "visible"),
                "hidden_sheets": sum(1 for s in sheets_info if s["state"] != "visible"),
                "file_metadata": {
                    "file_name": file_info.name,
                    "file_size_kb": round(file_info.stat().st_size / 1024, 1)
                }
            }
            
            return ToolResult(
                success=True,
                data=inventory,
                message=f"Inventory: {len(sheets_info)} sheets ({inventory['visible_sheets']} visible, {inventory['hidden_sheets']} hidden), {len(named_ranges)} named ranges"
            )
        except Exception as e:
            return ToolResult(success=False, data=None, message=f"Error in get_file_inventory: {e}")

    @staticmethod
    def inspect_sheet_structure(file_path: str, sheet_name: str = None) -> ToolResult:
        """
        Phase 1 Discovery: Deterministic structural analysis of an Excel sheet.
        
        Produces a full "X-ray" of the sheet covering both signal (data regions,
        headers) and noise (merged cells, spacer rows, hidden elements).
        No LLM is involved — this is pure Python/openpyxl analysis.
        
        Args:
            file_path: Workbook the pipeline is running on — same path as LangGraph ``state["file_path"]``
                (often the materialized clean template, not the raw upload).
            sheet_name: Sheet to inspect (default: first visible sheet); should match ``state["sheet_name"]``.
            
        Returns:
            ToolResult with structural report:
            - merged_cells: list of {range, top_left_value, rows_spanned, cols_spanned}
            - empty_rows: list of row indices that are completely empty (spacers)
            - empty_cols: list of col indices that are completely empty (spacers)
            - header_candidates: list of row indices that look like headers
            - hidden_rows: list of hidden row indices
            - hidden_cols: list of hidden column letters
            - data_region: {first_data_row, last_data_row, first_data_col, last_data_col}
            - density: fraction of non-empty cells in the data region
            - noise_score: 0.0 (clean) to 1.0 (very messy)
        """
        try:
            from openpyxl import load_workbook
            from openpyxl.utils import get_column_letter
            
            wb = load_workbook(file_path, data_only=True)
            
            # Select sheet
            if sheet_name:
                if sheet_name not in wb.sheetnames:
                    wb.close()
                    return ToolResult(success=False, data=None,
                                     message=f"Sheet '{sheet_name}' not found. Available: {wb.sheetnames}")
                ws = wb[sheet_name]
            else:
                # First visible sheet
                ws = None
                for sn in wb.sheetnames:
                    if wb[sn].sheet_state == 'visible':
                        ws = wb[sn]
                        sheet_name = sn
                        break
                if ws is None:
                    wb.close()
                    return ToolResult(success=False, data=None, message="No visible sheets found")
            
            max_row = ws.max_row or 0
            max_col = ws.max_column or 0
            
            if max_row == 0 or max_col == 0:
                wb.close()
                return ToolResult(success=True, data={"empty": True, "sheet_name": sheet_name},
                                 message=f"Sheet '{sheet_name}' is empty")
            
            # ---- 1. Merged Cells ----
            merged_info = []
            for mr in ws.merged_cells.ranges:
                top_left = ws.cell(row=mr.min_row, column=mr.min_col).value
                merged_info.append({
                    "range": str(mr),
                    "top_left_value": str(top_left) if top_left is not None else None,
                    "rows_spanned": mr.max_row - mr.min_row + 1,
                    "cols_spanned": mr.max_col - mr.min_col + 1
                })
            
            # ---- 2. Empty Rows (spacers) ----
            empty_rows = []
            row_emptiness = {}  # row_idx -> count of non-empty cells
            for row_idx in range(1, max_row + 1):
                non_empty = 0
                for col_idx in range(1, max_col + 1):
                    val = ws.cell(row=row_idx, column=col_idx).value
                    if val is not None and str(val).strip() != "":
                        non_empty += 1
                row_emptiness[row_idx] = non_empty
                if non_empty == 0:
                    empty_rows.append(row_idx - 1)  # 0-indexed for tool consistency
            
            # ---- 3. Empty Columns (spacers) ----
            empty_cols = []
            for col_idx in range(1, max_col + 1):
                non_empty = 0
                for row_idx in range(1, max_row + 1):
                    val = ws.cell(row=row_idx, column=col_idx).value
                    if val is not None and str(val).strip() != "":
                        non_empty += 1
                if non_empty == 0:
                    empty_cols.append(col_idx - 1)  # 0-indexed
            
            # ---- 4. Header Candidates ----
            # A header row is predominantly text (non-numeric), non-empty, and in the top portion
            header_candidates = []
            scan_depth = min(max_row, 20)  # Only scan top 20 rows
            for row_idx in range(1, scan_depth + 1):
                if row_emptiness.get(row_idx, 0) == 0:
                    continue  # Skip empty rows
                text_count = 0
                numeric_count = 0
                total_filled = 0
                for col_idx in range(1, max_col + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    val = cell.value
                    if val is None or str(val).strip() == "":
                        continue
                    total_filled += 1
                    # Check if text-like (non-numeric)
                    try:
                        float(str(val).replace(",", "").replace("$", "").replace("%", ""))
                        numeric_count += 1
                    except (ValueError, TypeError):
                        text_count += 1
                
                if total_filled > 0:
                    text_ratio = text_count / total_filled
                    is_bold = False
                    try:
                        first_cell = ws.cell(row=row_idx, column=1)
                        is_bold = first_cell.font.bold if first_cell.font.bold else False
                    except Exception:
                        pass
                    
                    # Header heuristic: mostly text, covers good width
                    fill_ratio = total_filled / max_col
                    if text_ratio >= 0.6 and fill_ratio >= 0.3:
                        header_candidates.append({
                            "row": row_idx - 1,  # 0-indexed
                            "text_ratio": round(text_ratio, 2),
                            "fill_ratio": round(fill_ratio, 2),
                            "is_bold": is_bold,
                            "sample_values": [
                                str(ws.cell(row=row_idx, column=c).value)
                                for c in range(1, min(max_col + 1, 6))
                                if ws.cell(row=row_idx, column=c).value is not None
                            ][:5]
                        })
            
            # ---- 5. Hidden Rows/Cols ----
            hidden_rows = []
            for row_idx, rd in ws.row_dimensions.items():
                if rd.hidden:
                    hidden_rows.append(row_idx - 1)  # 0-indexed
            
            hidden_cols = []
            for col_letter, cd in ws.column_dimensions.items():
                if cd.hidden:
                    hidden_cols.append(col_letter)
            
            # ---- 6. Data Region Detection ----
            # Find the bounding box of non-empty cells
            first_data_row, last_data_row = max_row, 0
            first_data_col, last_data_col = max_col, 0
            total_cells = 0
            filled_cells = 0
            
            for row_idx in range(1, max_row + 1):
                for col_idx in range(1, max_col + 1):
                    total_cells += 1
                    val = ws.cell(row=row_idx, column=col_idx).value
                    if val is not None and str(val).strip() != "":
                        filled_cells += 1
                        first_data_row = min(first_data_row, row_idx)
                        last_data_row = max(last_data_row, row_idx)
                        first_data_col = min(first_data_col, col_idx)
                        last_data_col = max(last_data_col, col_idx)
            
            data_region = {
                "first_data_row": first_data_row - 1,  # 0-indexed
                "last_data_row": last_data_row - 1,
                "first_data_col": first_data_col - 1,
                "last_data_col": last_data_col - 1,
                "rows": last_data_row - first_data_row + 1,
                "cols": last_data_col - first_data_col + 1
            } if filled_cells > 0 else None
            
            # ---- 7. Density ----
            density = round(filled_cells / total_cells, 3) if total_cells > 0 else 0.0
            
            # ---- 8. Noise Score ----
            # Heuristic: count anti-patterns and normalize
            noise_signals = 0
            noise_max = 5
            if len(merged_info) > 0:
                noise_signals += min(len(merged_info) / 10, 1.0)  # Merged cells
            if len(empty_rows) > 0:
                noise_signals += min(len(empty_rows) / 5, 1.0)    # Spacer rows
            if len(empty_cols) > 0:
                noise_signals += min(len(empty_cols) / 3, 1.0)    # Spacer cols
            if len(header_candidates) > 1:
                noise_signals += min((len(header_candidates) - 1) / 3, 1.0)  # Multi-headers
            if len(hidden_rows) + len(hidden_cols) > 0:
                noise_signals += 0.5  # Hidden elements
            
            noise_score = round(min(noise_signals / noise_max, 1.0), 2)
            
            wb.close()
            
            report = {
                "sheet_name": sheet_name,
                "dimensions": {"rows": max_row, "cols": max_col},
                "merged_cells": merged_info,
                "empty_rows": empty_rows,
                "empty_cols": empty_cols,
                "header_candidates": header_candidates,
                "hidden_rows": hidden_rows,
                "hidden_cols": hidden_cols,
                "data_region": data_region,
                "density": density,
                "noise_score": noise_score,
                "summary": {
                    "merged_count": len(merged_info),
                    "spacer_rows": len(empty_rows),
                    "spacer_cols": len(empty_cols),
                    "header_row_count": len(header_candidates),
                    "hidden_elements": len(hidden_rows) + len(hidden_cols),
                    "data_density_pct": round(density * 100, 1)
                }
            }
            
            return ToolResult(
                success=True,
                data=report,
                message=(
                    f"Sheet '{sheet_name}': {max_row}x{max_col}, "
                    f"noise_score={noise_score}, "
                    f"{len(merged_info)} merged, "
                    f"{len(empty_rows)} spacer rows, "
                    f"{len(header_candidates)} header candidates, "
                    f"density={density:.0%}"
                )
            )
        except Exception as e:
            return ToolResult(success=False, data=None, message=f"Error in inspect_sheet_structure: {e}")

    @staticmethod
    def fuzzy_column_align(source_columns: List[str], target_columns: List[str]) -> ToolResult:
        """
        Phase 3 Normalization: Align column names between two schemas
        using the 3-tier fuzzy matching engine (exact, synonym, difflib).
        
        Useful when merging data from different sources where column names
        differ (e.g., 'Revenue' vs 'Sales_Amt', 'Date' vs 'Period').
        
        Args:
            source_columns: Column names from the source schema
            target_columns: Column names from the target schema
            
        Returns:
            ToolResult with mapping dict:
            - mapping: {source_col: {target: matched_col, confidence: float}}
            - unmatched: list of source columns with no match
        """
        try:
            import difflib
            
            # Load synonym dictionary for enhanced matching
            synonyms = TransformationTools._load_synonyms()
            col_synonyms = synonyms.get("column_synonyms", {})
            
            # Build reverse synonym lookup
            _synonym_reverse = {}
            for canonical, aliases in col_synonyms.items():
                _synonym_reverse[canonical.lower()] = canonical.lower()
                for alias in aliases:
                    _synonym_reverse[alias.lower().replace("_", " ").replace("-", " ")] = canonical.lower()
            
            def _normalize(name: str) -> str:
                h = name.lower().strip().replace("_", " ").replace("-", " ")
                return _synonym_reverse.get(h, h)
            
            mapping = {}
            unmatched = []
            
            target_normalized = {_normalize(t): t for t in target_columns}
            target_lower = {t.lower().strip(): t for t in target_columns}
            
            for src in source_columns:
                src_norm = _normalize(src)
                src_lower = src.lower().strip()
                
                # Tier 1: Exact match (case-insensitive)
                if src_lower in target_lower:
                    mapping[src] = {"target": target_lower[src_lower], "confidence": 1.0}
                    continue
                
                # Tier 2: Synonym normalization match
                if src_norm in target_normalized:
                    mapping[src] = {"target": target_normalized[src_norm], "confidence": 0.95}
                    continue
                
                # Tier 3: Fuzzy difflib match
                target_strs = [t.lower() for t in target_columns]
                matches = difflib.get_close_matches(src_lower, target_strs, n=1, cutoff=0.6)
                if matches:
                    match_score = difflib.SequenceMatcher(None, src_lower, matches[0]).ratio()
                    matched_original = next(t for t in target_columns if t.lower() == matches[0])
                    mapping[src] = {"target": matched_original, "confidence": round(match_score, 3)}
                else:
                    unmatched.append(src)
            
            result = {
                "mapping": mapping,
                "unmatched_source": unmatched,
                "matched_count": len(mapping),
                "unmatched_count": len(unmatched)
            }
            
            return ToolResult(
                success=True,
                data=result,
                message=f"Aligned {len(mapping)}/{len(source_columns)} columns. {len(unmatched)} unmatched."
            )
        except Exception as e:
            return ToolResult(success=False, data=None, message=f"Error in fuzzy_column_align: {e}")

    @staticmethod
    def verify_checksum(df: pd.DataFrame, checksum_column: Union[int, str],
                        expected_total: float, tolerance: float = 0.01) -> ToolResult:
        """
        Phase 5 Validation: Verify data integrity by comparing a column sum
        against an expected total.
        
        Critical for ensuring no data loss during transformations.
        For example: "Total Sales in output must equal Sum of Column G in raw file."
        
        Args:
            df: DataFrame to verify
            checksum_column: Column name or index to sum
            expected_total: Expected total value
            tolerance: Acceptable difference as a fraction (default 0.01 = 1%)
            
        Returns:
            ToolResult with verification report:
            - passed: bool
            - actual_total: float
            - expected_total: float
            - difference: float
            - difference_pct: float
        """
        try:
            # Resolve column
            col_name, col_conf = TransformationTools._fuzzy_find_column(df, checksum_column)
            
            # Calculate actual sum (coerce to numeric, skip NaN)
            numeric_series = pd.to_numeric(df[col_name], errors='coerce')
            actual_total = float(numeric_series.sum())
            
            # Calculate difference
            difference = abs(actual_total - expected_total)
            difference_pct = (difference / abs(expected_total) * 100) if expected_total != 0 else (
                0.0 if difference == 0 else 100.0
            )
            
            passed = difference_pct <= (tolerance * 100)
            
            report = {
                "passed": passed,
                "column": col_name,
                "actual_total": round(actual_total, 4),
                "expected_total": round(expected_total, 4),
                "difference": round(difference, 4),
                "difference_pct": round(difference_pct, 4),
                "tolerance_pct": round(tolerance * 100, 2),
                "column_confidence": col_conf,
                "non_numeric_count": int(numeric_series.isna().sum() - df[col_name].isna().sum())
            }
            
            status = "✓ PASSED" if passed else "✗ FAILED"
            
            return ToolResult(
                success=True,
                data=report,
                message=f"Checksum {status}: {col_name} sum={actual_total:.2f} vs expected={expected_total:.2f} (diff={difference_pct:.2f}%)",
                changes_made={"checksum_report": report}
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in verify_checksum: {e}")



    @staticmethod
    def _calculate_header_signature(row_values: pd.Series) -> List[str]:
        """Create a signature for matching header rows (list of non-empty strings)."""
        return [str(v).strip().lower() for v in row_values if pd.notna(v) and str(v).strip()]

    @staticmethod
    def merge_blocks(df: pd.DataFrame, header_row: int,
                    blocks: List[Dict] = None, direction: str = None) -> ToolResult:
        """
        Merge multiple data blocks sharing the same header structure.
        
        The plan generator must provide explicit block coordinates from structure analysis.
        
        Args:
            df: The full raw DataFrame (grid).
            header_row: Index of the row defining the master header.
            blocks: Optional list of block coordinate dicts. Each dict should have:
                    - For horizontal (side-by-side): {'col_start': int, 'col_end': int}
                    - For vertical (stacked): {'row_start': int, 'row_end': int}
                    - Combined: {'row_start': int, 'row_end': int, 'col_start': int, 'col_end': int}
                    Example: [{'col_start': 0, 'col_end': 6}, {'col_start': 6, 'col_end': 12}]
            direction: Optional legacy mode ("horizontal" or "vertical") to auto-detect blocks.
            
        Returns:
            ToolResult with the unified DataFrame.
        """
        if blocks:
            from sia.agent.layout_stack_utils import assess_stack_readiness

            ready, readiness_reason = assess_stack_readiness(df, header_row, blocks)
            if not ready:
                return ToolResult(
                    success=False,
                    data=df,
                    message=(
                        f"layout.stack refused: {readiness_reason}. "
                        "Identify each side-by-side block (col_start/col_end), align schemas, "
                        "or use layout.extract per block instead of stacking the full sheet."
                    ),
                )
            return TransformationTools._merge_blocks_impl(df, header_row, blocks)

        # Backward-compatible behavior for older callers/tests.
        if direction == "horizontal":
            split_cols = []
            for col_idx in range(len(df.columns)):
                col_vals = df.iloc[header_row + 1:, col_idx] if header_row + 1 < len(df) else pd.Series(dtype=object)
                is_blank_col = col_vals.isna().all() or (col_vals.astype(str).str.strip() == "").all()
                if is_blank_col:
                    split_cols.append(col_idx)

            boundaries = [-1] + split_cols + [len(df.columns)]
            auto_blocks = []
            for i in range(len(boundaries) - 1):
                c0 = boundaries[i] + 1
                c1 = boundaries[i + 1]
                if c0 < c1:
                    auto_blocks.append({
                        "row_start": header_row + 1,
                        "row_end": len(df),
                        "col_start": c0,
                        "col_end": c1
                    })
            return TransformationTools._merge_blocks_impl(df, header_row, auto_blocks)

        if direction == "vertical":
            # Use blank rows as separators and include all columns.
            if header_row + 1 >= len(df):
                return ToolResult(success=True, data=pd.DataFrame(), message="No data rows to merge")

            blank_rows = []
            for ridx in range(header_row + 1, len(df)):
                row = df.iloc[ridx]
                if row.isna().all() or (row.astype(str).str.strip() == "").all():
                    blank_rows.append(ridx)

            boundaries = [header_row + 1] + blank_rows + [len(df)]
            auto_blocks = []
            for i in range(len(boundaries) - 1):
                r0 = boundaries[i]
                r1 = boundaries[i + 1]
                if r0 < r1:
                    auto_blocks.append({
                        "row_start": r0,
                        "row_end": r1,
                        "col_start": 0,
                        "col_end": len(df.columns)
                    })

            merged = TransformationTools._merge_blocks_impl(df, header_row, auto_blocks)
            if merged.success and isinstance(merged.data, pd.DataFrame) and not merged.data.empty:
                # Drop rows that are exact header repeats (legacy expectation).
                hdr_vals = [str(v).strip().lower() for v in df.iloc[header_row, :].tolist()]
                mask = merged.data.apply(
                    lambda r: [str(v).strip().lower() for v in r.tolist()] != hdr_vals[:len(r.tolist())],
                    axis=1
                )
                merged.data = merged.data[mask].reset_index(drop=True)
            return merged

        return ToolResult(
            success=False,
            data=df,
            message="merge_blocks requires either 'blocks' or legacy 'direction' ('horizontal'/'vertical')"
        )

    @staticmethod
    def stack_tables(df: pd.DataFrame, header_row: int, blocks: List[Dict]) -> ToolResult:
        """Alias for merge_blocks."""
        return TransformationTools.merge_blocks(df, header_row, blocks)

    @staticmethod
    def _merge_blocks_impl(df: pd.DataFrame, header_row: int, blocks: List[Dict]) -> ToolResult:
        """Internal implementation for merging blocks with Header-Aware alignment."""
        try:
            from sia.agent.layout_stack_utils import looks_like_horizontal_block_failure

            # Validate inputs
            if header_row < 0 or header_row >= len(df):
                return ToolResult(success=False, data=df, message=f"Invalid header row {header_row}")
            
            if not blocks or not isinstance(blocks, list) or len(blocks) == 0:
                return ToolResult(success=False, data=df, 
                                 message="blocks parameter is required. Provide list of block coordinates.")
            
            # Step 1: Extract individual blocks and their specific headers
            extracted_blocks = []
            all_headers_ordered = [] # Maintain order of appearance for the final schema
            header_set = set()
            normalized_to_original = {} # Map "date" -> "Date" for the first appearance
            
            # Load synonyms for header normalization
            synonyms = TransformationTools._load_synonyms()
            col_synonyms = synonyms.get("column_synonyms", {})
            
            # Build reverse lookup: alias -> canonical_key
            _synonym_reverse = {}
            for canonical, aliases in col_synonyms.items():
                _synonym_reverse[canonical.lower()] = canonical.lower()
                for alias in aliases:
                    _synonym_reverse[alias.lower().replace("_", " ").replace("-", " ")] = canonical.lower()
            
            def _normalize_header(raw: str) -> str:
                """Normalize a header to its canonical form using synonyms."""
                h = raw.lower().strip().replace("_", " ").replace("-", " ")
                # Check synonym lookup first
                if h in _synonym_reverse:
                    return _synonym_reverse[h]
                # Try singular/plural normalization
                if h.endswith("s") and h[:-1] in _synonym_reverse:
                    return _synonym_reverse[h[:-1]]
                if h + "s" in _synonym_reverse:
                    return _synonym_reverse[h + "s"]
                # Fallback to lowercased original
                return raw.lower().strip()
            
            for i, block_def in enumerate(blocks):
                col_start = int(block_def.get('col_start', 0))
                col_end = int(block_def.get('col_end', len(df.columns)))
                row_start = int(block_def.get('row_start', header_row + 1))
                row_end = int(block_def.get('row_end', len(df)))
                
                # Extract headers for THIS specific block
                block_headers_raw = df.iloc[header_row, col_start:col_end].tolist()
                
                # Normalize and map headers WITHIN this block
                block_headers = []
                for k, h in enumerate(block_headers_raw):
                    h_raw = str(h).strip() if pd.notna(h) else ""
                    if h_raw == "": 
                        h_raw = f"Col_{k}"
                    
                    # Synonym-aware normalization for alignment
                    h_norm = _normalize_header(h_raw)
                    
                    # Track the first-seen "nice" original name for the final schema
                    if h_norm not in normalized_to_original:
                        normalized_to_original[h_norm] = h_raw
                        all_headers_ordered.append(h_norm)
                    
                    block_headers.append(h_norm)
                    
                # Extract raw data
                block_data = df.iloc[row_start:row_end, col_start:col_end].copy()
                block_data.columns = block_headers
                
                # Drop fully empty rows
                block_data = block_data.dropna(how='all')
                if not block_data.empty:
                    # Clean strings and check for non-empty content
                    mask = block_data.apply(lambda r: any(pd.notna(v) and str(v).strip() != "" for v in r), axis=1)
                    block_data = block_data[mask]
                
                if not block_data.empty:
                    extracted_blocks.append(block_data)

            if not extracted_blocks:
                return ToolResult(success=False, data=df, message="No valid data blocks extracted from provided coordinates")
            
            # Step 2: Realign all blocks to the same Master Schema (all_headers_ordered)
            aligned_blocks = []
            for block in extracted_blocks:
                # Add missing columns with empty string
                missing_cols = set(all_headers_ordered) - set(block.columns)
                if missing_cols:
                    missing_df = pd.DataFrame("", index=block.index, columns=list(missing_cols))
                    block = pd.concat([block, missing_df], axis=1)
                
                # Reorder and pick columns to match the global order (using normalized names)
                aligned_block = block[all_headers_ordered].copy()
                aligned_blocks.append(aligned_block)
            
            # Step 3: Concatenate all aligned blocks
            unified_df = pd.concat(aligned_blocks, ignore_index=True)
            
            # Step 4: Restore the "nice" original names to the final schema
            final_mapping = {norm: normalized_to_original[norm] for norm in all_headers_ordered}
            unified_df = unified_df.rename(columns=final_mapping)
            
            # Final cleanup of any rows that became empty after alignment/padding (unlikely but safe)
            if not unified_df.empty:
                mask = unified_df.apply(lambda r: any(pd.notna(v) and str(v).strip() != "" for v in r), axis=1)
                unified_df = unified_df[mask]

            if looks_like_horizontal_block_failure(unified_df):
                return ToolResult(
                    success=False,
                    data=df,
                    message=(
                        "layout.stack produced side-by-side duplicate columns (e.g. Date, Date.1) — "
                        "blocks were not vertically merged. Use explicit per-block column ranges on the "
                        "raw grid, or extract each block separately before combining."
                    ),
                )

            return ToolResult(
                success=True,
                data=unified_df, 
                message=f"Merged {len(extracted_blocks)} blocks into {len(unified_df)} rows. Aligned by header names: {', '.join(all_headers_ordered[:10])}..."
            )

        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in merge_blocks: {str(e)}")

    @staticmethod
    def skip_rows(df: pd.DataFrame, rows_to_skip: List[int], 
                  preview_only: bool = False) -> Union[ToolResult, DeletionPreview]:
        """
        Remove specified rows.
        
        Args:
            df: DataFrame to modify
            rows_to_skip: List of row indices to remove (0-based)
            preview_only: If True, return DeletionPreview instead of executing deletion
            
        Returns:
            ToolResult with filtered DataFrame, or DeletionPreview if preview_only=True
        """
        try:
            original_len = len(df)
            
            # Validate indices
            valid_indices = [i for i in rows_to_skip if 0 <= i < len(df)]
            
            # Get the data that will be deleted
            deleted_data = df.iloc[valid_indices] if valid_indices else pd.DataFrame()
            
            if preview_only:
                percent_deleted = (len(valid_indices) / original_len * 100) if original_len > 0 else 0
                return DeletionPreview(
                    tool_name="skip_rows",
                    tool_description="Remove specified rows by index",
                    rows_to_delete=valid_indices,
                    columns_to_delete=[],
                    sample_deleted_data=deleted_data.head(10).to_dict(orient='records') if not deleted_data.empty else [],
                    full_deleted_data=deleted_data,
                    reason=f"Rows explicitly marked for removal: {valid_indices[:5]}{'...' if len(valid_indices) > 5 else ''}",
                    impact_summary=f"Will remove {len(valid_indices)} rows ({percent_deleted:.1f}% of data)"
                )
            
            df = df.drop(df.index[valid_indices], errors='ignore')
            df = df.reset_index(drop=True)
            
            return ToolResult(
                success=True,
                data=df,
                message=f"Removed {original_len - len(df)} rows"
            )
        except Exception as e:
            if preview_only:
                return DeletionPreview(
                    tool_name="skip_rows",
                    tool_description="Remove specified rows by index",
                    reason=f"Error during preview: {e}",
                    impact_summary="Preview failed"
                )
            return ToolResult(success=False, data=df, message=f"Error: {e}")


    @staticmethod
    def _strip_outer_column_labels(df: pd.DataFrame) -> tuple[pd.DataFrame, Dict[str, str]]:
        """Trim leading/trailing whitespace on column names; keep internal spaces."""
        strip_map: Dict[str, str] = {}
        taken = {str(c) for c in df.columns}
        for col in list(df.columns):
            name = str(col)
            stripped = name.strip()
            if not stripped or stripped == name:
                continue
            if stripped in taken and stripped != name:
                logger.warning(
                    "rename_columns: skip strip %r -> %r (name already in use)",
                    name,
                    stripped,
                )
                continue
            strip_map[name] = stripped
            taken.discard(name)
            taken.add(stripped)
        if strip_map:
            df = df.rename(columns=strip_map)
        return df, strip_map

    @staticmethod
    def rename_columns(df: pd.DataFrame, mapping: Dict[Union[int, str], str]) -> ToolResult:
        """
        Rename columns using a name-to-name or index-to-name mapping.
        
        Args:
            df: DataFrame to modify
            mapping: Dict mapping column name OR index to new name.
                     Preferred: {"old_name": "new_name"}
                     Legacy: {0: "new_name"} (index-based)
            
        Returns:
            ToolResult with renamed DataFrame
        """
        try:
            from sia.agent.target_template_utils import is_no_match_target, sanitize_rename_mapping

            columns_before = list(df.columns)
            rename_dict: Dict[str, str] = {}
            confidences = {}
            missing = []

            raw_mapping = mapping or {}
            sanitized_targets = sanitize_rename_mapping(raw_mapping)
            for raw_key, raw_new_name in raw_mapping.items():
                if is_no_match_target(raw_new_name):
                    continue
                key = raw_key if isinstance(raw_key, int) else str(raw_key or "").strip()
                new_name = str(raw_new_name or "").strip()
                if not isinstance(raw_key, int):
                    new_name = sanitized_targets.get(str(raw_key or "").strip(), new_name)
                if key == "" or not new_name:
                    continue
                if is_no_match_target(new_name):
                    continue
                if isinstance(key, int):
                    # Legacy index-based mapping
                    if 0 <= key < len(df.columns):
                        rename_dict[df.columns[key]] = new_name
                        confidences[df.columns[key]] = 1.0
                    else:
                        missing.append(f"Index {key}")
                elif isinstance(key, str):
                    # Semantic name-based mapping via fuzzy resolver
                    try:
                        resolved, conf = TransformationTools._fuzzy_find_column(df, key)
                        rename_dict[resolved] = new_name
                        confidences[resolved] = conf
                    except ValueError:
                        missing.append(f"'{key}'")
            
            if missing:
                available = list(df.columns)[:10]
                # Log warning but proceed with partial rename
                logger.warning(f"rename_columns: Could not resolve {missing}. Proceeding with {len(rename_dict)} matches.")
                if not rename_dict:
                    return ToolResult(
                        success=False,
                        data=df,
                        message=f"Column(s) {', '.join(missing)} not found. Available: {available}"
                    )
            
            if rename_dict:
                targets = list(rename_dict.values())
                if len(targets) != len(set(targets)):
                    logger.warning(
                        "rename_columns: multiple sources map to the same target; "
                        "pandas will keep one column per target name."
                    )
                df = df.rename(columns=rename_dict)

            df, strip_map = TransformationTools._strip_outer_column_labels(df)

            preserved = [c for c in columns_before if c in df.columns]
            dropped = [c for c in columns_before if c not in df.columns]
            if dropped:
                logger.warning(
                    "rename_columns: unexpected column loss after rename: %s",
                    dropped[:8],
                )

            msg = f"Renamed {len(rename_dict)} columns; kept {len(preserved)} of {len(columns_before)}"
            if strip_map:
                msg += f"; stripped outer whitespace on {len(strip_map)} label(s)"

            return ToolResult(
                success=True,
                data=df,
                message=msg,
                changes_made={
                    "renamed_cols": rename_dict,
                    "stripped_label_whitespace": strip_map,
                    "column_confidences": confidences,
                    "preserved_unmapped": [c for c in columns_before if c not in rename_dict],
                },
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error: {e}")

    @staticmethod
    def reorder_columns_layout(
        df: pd.DataFrame,
        column_order: Optional[List[str]] = None,
    ) -> ToolResult:
        """Reorder columns; unlisted columns are appended at the end in original order."""
        try:
            from sia.tools.collation_tools import reorder_columns

            if df is None:
                return ToolResult(success=False, data=df, message="reorder_columns: missing DataFrame")
            resolved = []
            skipped = []
            for col in list(column_order or []):
                try:
                    resolved_col, _ = TransformationTools._fuzzy_find_column(df, col)
                    resolved.append(resolved_col)
                except ValueError:
                    skipped.append(str(col))
            order = [c for c in resolved if c in df.columns]
            rest = [c for c in df.columns if c not in order]
            out = reorder_columns(df, order + rest)
            skipped_msg = f"; skipped missing: {skipped[:5]}" if skipped else ""
            return ToolResult(
                success=True,
                data=out,
                message=f"Reordered {len(order)} column(s); {len(rest)} trailing column(s) unchanged{skipped_msg}",
                changes_made={"column_order": order + rest, "skipped_missing": skipped},
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in reorder_columns: {e}")

    @staticmethod
    def sort_rows(
        df: pd.DataFrame,
        sort_columns: Optional[List[str]] = None,
        ascending: bool = True,
    ) -> ToolResult:
        """Sort rows by one or more columns (stable mergesort). Date columns are coerced when possible."""
        try:
            if df is None or df.empty:
                return ToolResult(success=False, data=df, message="sort_rows: missing or empty DataFrame")
            cols, _ = TransformationTools._resolve_columns(df, list(sort_columns or []))
            if not cols:
                return ToolResult(
                    success=False,
                    data=df,
                    message="sort_rows: no sort columns resolved",
                )
            out = df.copy()
            for col in cols:
                if col not in out.columns:
                    continue
                if pd.api.types.is_datetime64_any_dtype(out[col]):
                    continue
                sample = out[col].dropna().head(5)
                if len(sample) and pd.to_datetime(sample, errors="coerce").notna().all():
                    out[col] = pd.to_datetime(out[col], errors="coerce")
            out = out.sort_values(
                by=cols,
                ascending=ascending,
                kind="mergesort",
                na_position="last",
            ).reset_index(drop=True)
            return ToolResult(
                success=True,
                data=out,
                message=f"Sorted {len(out)} rows by {', '.join(cols)}",
                changes_made={"sort_columns": cols, "ascending": ascending},
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in sort_rows: {e}")

    @staticmethod
    def filter_empty_rows(df: pd.DataFrame, check_columns: List[Union[int, str]] = None, 
                          metric_columns: List[Union[int, str]] = None,
                          all_must_be_empty: bool = True,
                          require_metrics: bool = False,
                          min_populated_columns: int = 2,
                          preview_only: bool = False, **kwargs) -> Union[ToolResult, DeletionPreview]:
        """
        Remove rows where specified columns are empty or lack metrics.
        
        Args:
            df: DataFrame to filter
            check_columns: Columns to check for general emptiness (default: all)
            metric_columns: Columns that MUST contain at least one value.
            all_must_be_empty: If True (default), remove if ALL check_columns are empty.
            require_metrics: If True, remove rows that have no metric values (heuristic).
            min_populated_columns: Minimum number of populated columns required to keep the row (default: 2).
            preview_only: If True, return DeletionPreview instead of executing deletion
        """
        try:
            # Handle variations
            if check_columns is None:
                check_columns = kwargs.get('columns', [])
                
            original_len = len(df)
            cols, _ = TransformationTools._resolve_columns(df, check_columns)
            
            if not cols:
                cols = df.columns.tolist()

            mcols = []
            if metric_columns:
                mcols, _ = TransformationTools._resolve_columns(df, metric_columns)

            def is_metric(val):
                """Check if a value looks like a metric (numeric, not a date-string or year)."""
                if pd.isna(val) or str(val).strip() == "":
                    return False
                
                # Check for numeric types
                if isinstance(val, (int, float)):
                    # Heuristic: Avoid typical year range
                    if 1900 <= val <= 2100 and (val == int(val)):
                        return False 
                    return True
                
                # If it's a string, try to parse as number (handle commas/currency)
                if isinstance(val, str):
                    clean_val = val.replace(',', '').replace('$', '').replace('%', '').strip()
                    try:
                        f_val = float(clean_val)
                        if 1900 <= f_val <= 2100 and f_val == int(f_val):
                            return False
                        return True
                    except ValueError:
                        return False
                return False

            def check_row(row):
                # 1. Metric Columns check (High Priority)
                if mcols:
                    has_metric_populated = any(pd.notna(row[c]) and str(row[c]).strip() != "" for c in mcols)
                    if not has_metric_populated:
                        return False
                
                # 2. Populated count check
                populated = [v for v in row[cols] if pd.notna(v) and str(v).strip() != ""]
                if len(populated) < min_populated_columns:
                    return False
                
                # 3. Standard emptiness check
                if all_must_be_empty and len(populated) == 0:
                    return False
                
                if not all_must_be_empty and len(populated) < len(cols):
                    return False

                # 4. Metric heuristic check
                if require_metrics:
                    # Check if at least one value looks like a metric
                    has_metric = any(is_metric(v) for v in row[cols])
                    if not has_metric:
                        return False
                
                return True
            
            mask = df.apply(check_row, axis=1)
            
            rows_to_delete = df[~mask].index.tolist()
            deleted_data = df[~mask]
            
            if preview_only:
                percent_deleted = (len(rows_to_delete) / original_len * 100) if original_len > 0 else 0
                return DeletionPreview(
                    tool_name="filter_empty_rows",
                    tool_description="Remove empty or incomplete rows",
                    rows_to_delete=rows_to_delete,
                    columns_to_delete=[],
                    sample_deleted_data=deleted_data.head(10).to_dict(orient='records'),
                    full_deleted_data=deleted_data,
                    reason=f"Rows where {'all' if all_must_be_empty else 'any'} of the columns {cols[:3]}... are empty",
                    impact_summary=f"Will remove {len(rows_to_delete)} rows ({percent_deleted:.1f}% of data)"
                )
            
            df = df[mask].reset_index(drop=True)
            
            return ToolResult(
                success=True,
                data=df,
                message=f"Removed {len(rows_to_delete)} empty rows ({original_len - len(rows_to_delete)} remaining)"
            )
        except Exception as e:
            if preview_only:
                return DeletionPreview(tool_name="filter_empty_rows", reason=f"Error: {e}", impact_summary="Preview failed")
            return ToolResult(success=False, data=df, message=f"Error: {e}")

    @staticmethod
    def type_cast_columns(df: pd.DataFrame, type_map: Dict[Union[int, str], str]) -> ToolResult:
        """
        Cast columns to specified types.
        
        Args:
            df: DataFrame to modify
            type_map: Dict mapping column name OR index to type ('date', 'numeric', 'string', 'integer').
            
        Returns:
            ToolResult with type-casted DataFrame
        """
        try:
            result_df = df.copy()
            cast_results = []
            missing = []
            confidences = {}
            
            for key, dtype in type_map.items():
                try:
                    col_name, conf = TransformationTools._fuzzy_find_column(result_df, key)
                    confidences[col_name] = conf
                    
                    dtype = str(dtype).lower().strip()
                    if dtype == 'date':
                        result_df[col_name] = pd.to_datetime(result_df[col_name], errors='coerce')
                        cast_results.append(f"'{col_name}' -> date")
                    elif dtype == 'numeric' or dtype == 'integer':
                        # Clean string values for numeric/integer (handle commas, currency, etc)
                        if result_df[col_name].dtype == object:
                            result_df[col_name] = result_df[col_name].astype(str).str.replace(',', '').str.replace('$', '').str.replace('%', '').str.strip()
                            # Replace empty/None strings with NaN
                            result_df[col_name] = result_df[col_name].replace(['nan', 'None', '', 'None'], np.nan)
                        
                        result_df[col_name] = pd.to_numeric(result_df[col_name], errors='coerce')
                        
                        if dtype == 'integer':
                            # Use Int64 to allow NaNs
                            result_df[col_name] = result_df[col_name].astype('Int64')
                            cast_results.append(f"'{col_name}' -> integer")
                        else:
                            cast_results.append(f"'{col_name}' -> numeric")
                    elif dtype == 'string':
                        result_df[col_name] = result_df[col_name].astype(str)
                        # Replace 'nan' string with actual NaN
                        result_df[col_name] = result_df[col_name].replace('nan', np.nan).replace('None', np.nan)
                        cast_results.append(f"'{col_name}' -> string")
                    else:
                        cast_results.append(f"SKIP '{col_name}' (unknown type: {dtype})")
                        
                except ValueError:
                    missing.append(str(key))
            
            message = f"Cast {len(cast_results)} columns: {'; '.join(cast_results)}"
            if missing:
                message += f" | Failed to resolve: {', '.join(missing)}"
                
            return ToolResult(
                success=True,
                data=result_df,
                message=message,
                changes_made={"cast_columns": cast_results, "column_confidences": confidences}
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in type_cast_columns: {e}")


    @staticmethod
    def extract_data_block(grid, start_row: int, end_row: int, 
                           start_col: int, end_col: int, header_row: int) -> ToolResult:
        """
        Extract a rectangular block of data from the grid as a DataFrame.
        
        Args:
            grid: VisualGrid object or pandas DataFrame
            start_row: First row of data (after headers)
            end_row: Last row of data
            start_col: First column
            end_col: Last column
            header_row: Row containing headers
            
        Returns:
            ToolResult with DataFrame
        """
        try:
            # Helper to get cell value from either VisualGrid or DataFrame
            def get_value(row, col):
                if isinstance(grid, pd.DataFrame):
                    return grid.iloc[row, col] if row < len(grid) and col < len(grid.columns) else None
                else:
                    cell = grid.get_cell(row, col)
                    return cell.value if cell else None
            
            # Extract headers
            headers = []
            for col in range(start_col, end_col + 1):
                value = get_value(header_row, col)
                headers.append(str(value) if value and pd.notna(value) else f"Column_{col}")
            
            # Extract data - ensure we start AFTER the header row
            data_start = max(start_row, header_row + 1)
            
            data = []
            for row in range(data_start, end_row + 1):
                row_data = []
                for col in range(start_col, end_col + 1):
                    value = get_value(row, col)
                    row_data.append(value)
                data.append(row_data)
            
            df = pd.DataFrame(data, columns=headers)
            
            return ToolResult(
                success=True,
                data=df,
                message=f"Extracted {len(df)} rows x {len(headers)} columns"
            )
        except Exception as e:
            return ToolResult(success=False, data=None, message=f"Error: {e}")

    @staticmethod
    def identify_metrics(df: pd.DataFrame, metric_col: Union[int, str] = None, 
                         scan_direction: str = 'left', **kwargs) -> ToolResult:
        """
        Identify metric names by scanning for labels adjacent to numeric cells.
        """
        try:
            df = df.copy()
            
            # Handle common LLM parameter hallucination
            if metric_col is None:
                metric_col = kwargs.get('metric_name_column') or kwargs.get('metric_column')
            
            # Resolve to name
            resolved_cols, _ = TransformationTools._resolve_columns(df, metric_col) if metric_col is not None else ([], [])
            metric_col_name = resolved_cols[0] if resolved_cols else None
            
            # Auto-detect metric column if not provided
            if metric_col_name is None:
                # Look for columns with low cardinality text that might be metrics
                for col in df.columns[:5]:  # Check first 5 columns
                    unique_vals = df[col].dropna().unique()
                    # Metric column typically has specific names like Reach, TARPs, etc.
                    metric_keywords = ['reach', 'tarp', 'grp', 'spend', 'budget', 'ctarp', 'cpm', 'impressions']
                    if any(any(kw in str(v).lower() for kw in metric_keywords) for v in unique_vals[:10]):
                        metric_col_name = col
                        break
            
            if metric_col_name:
                # Rename the identified metric column to 'Metric' if not already named
                if metric_col_name != 'Metric':
                    df = df.rename(columns={metric_col_name: 'Metric'})
                
                metrics_found = df['Metric'].dropna().unique().tolist()
                return ToolResult(
                    success=True,
                    data=df,
                    message=f"Identified metric column with {len(metrics_found)} unique metrics: {metrics_found[:5]}"
                )
            else:
                return ToolResult(
                    success=True,
                    data=df,
                    message="No obvious metric column found - data may already have proper labels"
                )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error: {e}")

    @staticmethod
    def densify_dataframe(df: pd.DataFrame, dimension_cols: List[Union[int, str]] = None,
                          drop_empty_value_rows: bool = True, 
                          preview_only: bool = False, **kwargs) -> Union[ToolResult, DeletionPreview]:
        """
        Ensure DataFrame is dense with no blanks.
        """
        try:
            # Handle variations
            if dimension_cols is None:
                dimension_cols = kwargs.get('columns') or kwargs.get('dim_cols')
                
            df_copy = df.copy()
            original_len = len(df_copy)
            
            # Resolve dimension columns and get confidences
            dim_col_names, dim_conf_list = TransformationTools._resolve_columns(df_copy, dimension_cols)
            confidences = {name: conf for name, conf in zip(dim_col_names, dim_conf_list)}

            # Auto-detect dimension columns if not provided
            if not dim_col_names:
                for col in df_copy.columns:
                    non_null = df_copy[col].dropna()
                    if len(non_null) > 0:
                        numeric_count = pd.to_numeric(non_null, errors='coerce').notna().sum()
                        if numeric_count < len(non_null) * 0.5:
                            dim_col_names.append(col)
                    if len(dim_col_names) >= 6:
                        break
            
            # Forward-fill dimension columns
            if dim_col_names:
                for col in dim_col_names:
                    df_copy[col] = df_copy[col].apply(lambda x: np.nan if pd.isna(x) or str(x).strip() == "" else x)
                    df_copy[col] = df_copy[col].ffill()
            
            # Identify rows to delete (those with empty values after densification)
            rows_to_delete = []
            deleted_data = pd.DataFrame()
            if drop_empty_value_rows:
                value_cols = [c for c in df_copy.columns if c not in dim_col_names]
                if value_cols:
                    mask = df_copy[value_cols].apply(
                        lambda row: row.notna().any() and any(
                            str(v).strip() != '' for v in row if pd.notna(v)
                        ), axis=1
                    )
                    rows_to_delete = df_copy[~mask].index.tolist()
                    deleted_data = df_copy[~mask]
            
            if preview_only:
                percent_deleted = (len(rows_to_delete) / original_len * 100) if original_len > 0 else 0
                return DeletionPreview(
                    tool_name="densify_dataframe",
                    tool_description="Densify data by forward-filling dimensions and removing empty values",
                    rows_to_delete=rows_to_delete,
                    columns_to_delete=[],
                    sample_deleted_data=deleted_data.head(10).to_dict(orient='records') if rows_to_delete else [],
                    full_deleted_data=deleted_data if rows_to_delete else pd.DataFrame(),
                    reason=f"Found {len(rows_to_delete)} rows with missing values in {len(value_cols)} columns",
                    impact_summary=f"Will remove {len(rows_to_delete)} rows ({percent_deleted:.1f}% of data)"
                )
            
            # Execute actual deletion
            if rows_to_delete:
                df_copy = df_copy.drop(df_copy.index[rows_to_delete]).reset_index(drop=True)
            
            return ToolResult(
                success=True,
                data=df_copy,
                message=f"Densified DataFrame: filled {len(dim_col_names)} dimensions, removed {len(rows_to_delete)} empty rows",
                changes_made={"dimension_columns": dim_col_names, "column_confidences": confidences}
            )
        except Exception as e:
            if preview_only:
                return DeletionPreview(
                    tool_name="densify_dataframe",
                    reason=f"Error during preview: {e}",
                    impact_summary="Preview failed"
                )
            return ToolResult(success=False, data=df, message=f"Error: {e}")

    @staticmethod
    def propagate_dimensions(df: pd.DataFrame, dimension_cols: List[Union[int, str]] = None) -> ToolResult:
        """
        Forward-fill dimension columns to populate parent values into child rows.
        
        Args:
            df: DataFrame to modify
            dimension_cols: List of column indices or names to ffill
            
        Returns:
            ToolResult with modified DataFrame
        """
        try:
            df_copy = df.copy()
            
            # Resolve column names and get confidences
            resolved_dim_cols, dim_conf_list = TransformationTools._resolve_columns(df_copy, dimension_cols)
            confidences = {name: conf for name, conf in zip(resolved_dim_cols, dim_conf_list)}
            
            # Auto-detect dimension columns if not provided
            if not resolved_dim_cols:
                for col in df_copy.columns:
                    non_null = df_copy[col].dropna()
                    if len(non_null) > 0:
                        numeric_count = pd.to_numeric(non_null, errors='coerce').notna().sum()
                        if numeric_count < len(non_null) * 0.5:
                            resolved_dim_cols.append(col)
                logger.info(f"Auto-detected dimension columns for propagation: {resolved_dim_cols}")

            if resolved_dim_cols:
                for col in resolved_dim_cols:
                    # Treat empty strings as NaN for ffill
                    df_copy[col] = df_copy[col].apply(lambda x: np.nan if pd.isna(x) or str(x).strip() == "" else x)
                    df_copy[col] = df_copy[col].ffill()
            
            return ToolResult(
                success=True,
                data=df_copy,
                message=f"Propagated values in {len(resolved_dim_cols)} dimension columns: {resolved_dim_cols[:5]}",
                changes_made={"propagated_cols": resolved_dim_cols, "column_confidences": confidences}
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in propagate_dimensions: {e}")

    @staticmethod
    def filter_incomplete_rows(df: pd.DataFrame, value_cols: List[Union[int, str]] = None,
                               all_must_be_empty: bool = False,
                               preview_only: bool = False) -> Union[ToolResult, DeletionPreview]:
        """
        Remove rows that are missing core values (metrics).
        
        Args:
            df: DataFrame to filter
            value_cols: Columns that must contain values
            all_must_be_empty: If True, only drops if ALL checked columns are empty. 
                               If False (default), drops if ANY checked column is empty.
            preview_only: HITL preview flag
        """
        try:
            df_copy = df.copy()
            original_len = len(df_copy)
            
            # Resolve columns and get confidences
            cols, conf_list = TransformationTools._resolve_columns(df_copy, value_cols)
            confidences = {name: conf for name, conf in zip(cols, conf_list)}
            
            if not cols:
                # Auto-detect numeric/value columns
                cols = df_copy.select_dtypes(include=[np.number]).columns.tolist()
                if not cols:
                    cols = df_copy.columns.tolist()

            if all_must_be_empty:
                # Keep rows where at least ONE column has data
                mask = df_copy[cols].apply(lambda row: row.notna().any() and 
                                        any(str(v).strip() != '' for v in row if pd.notna(v)), axis=1)
            else:
                # Keep rows where ALL specified columns have data
                mask = df_copy[cols].apply(lambda row: row.notna().all() and 
                                        all(str(v).strip() != '' for v in row if pd.notna(v)), axis=1)
            
            rows_to_delete = df_copy[~mask].index.tolist()
            deleted_data = df_copy[~mask]
            
            if preview_only:
                percent_deleted = (len(rows_to_delete) / original_len * 100) if original_len > 0 else 0
                return DeletionPreview(
                    tool_name="filter_incomplete_rows",
                    tool_description="Remove rows missing values in core columns",
                    rows_to_delete=rows_to_delete,
                    columns_to_delete=[],
                    sample_deleted_data=deleted_data.head(10).to_dict(orient='records'),
                    full_deleted_data=deleted_data,
                    reason=f"Found {len(rows_to_delete)} rows missing values in: {cols[:5]}",
                    impact_summary=f"Will remove {len(rows_to_delete)} rows ({percent_deleted:.1f}% of data)"
                )

            df_result = df_copy[mask].reset_index(drop=True)
            return ToolResult(
                success=True,
                data=df_result,
                message=f"Filtered {original_len - len(df_result)} incomplete rows based on columns: {cols[:5]}",
                changes_made={"check_columns": cols, "column_confidences": confidences}
            )
        except Exception as e:
            if preview_only:
                return DeletionPreview(tool_name="filter_incomplete_rows", impact_summary=f"Preview failed: {e}")
            return ToolResult(success=False, data=df, message=f"Error in filter_incomplete_rows: {e}")

    _SUMMARY_LABEL_COL_HINTS = (
        "name", "label", "publisher", "channel", "brand", "title", "segment",
        "category", "market", "product", "platform", "advertiser", "campaign",
        "description", "line", "media", "vendor", "supplier", "partner",
    )

    @staticmethod
    def _summary_keyword_in_text(text: Any, keywords_lower: List[str]) -> bool:
        """Match summary keywords as whole words/phrases, not arbitrary substrings."""
        val_str = str(text).strip() if pd.notna(text) else ""
        if not val_str:
            return False
        val_lower = val_str.lower()
        for kw in keywords_lower:
            pattern = r"(?<![\w])" + re.escape(kw) + r"(?![\w])"
            if re.search(pattern, val_lower, re.IGNORECASE):
                return True
        return False

    @classmethod
    def _summary_label_columns(cls, df: pd.DataFrame) -> List[str]:
        """Columns likely to hold row labels (where Total/Subtotal appears)."""
        label_cols: List[str] = []
        for i, col in enumerate(df.columns):
            if not (
                pd.api.types.is_object_dtype(df[col])
                or pd.api.types.is_string_dtype(df[col])
            ):
                continue
            col_lower = str(col).lower()
            if i == 0 or any(h in col_lower for h in cls._SUMMARY_LABEL_COL_HINTS):
                label_cols.append(col)
        if label_cols:
            return label_cols
        object_cols = df.select_dtypes(include=["object", "string"]).columns.tolist()
        return object_cols[:1] if object_cols else [df.columns[0]]

    @staticmethod
    def _numeric_prone_columns(df: pd.DataFrame) -> List[str]:
        numeric_cols: List[str] = []
        for col in df.columns[1:]:
            try:
                field_data = df[col].dropna()
                if len(field_data) == 0:
                    continue
                numeric_count = pd.to_numeric(field_data, errors="coerce").notna().sum()
                if numeric_count / len(field_data) > 0.3:
                    numeric_cols.append(col)
            except Exception:
                pass
        return numeric_cols

    @staticmethod
    def filter_summary_rows(df: pd.DataFrame, keywords: List[str] = None, 
                            check_columns: List[int] = None,
                            use_structural_detection: bool = True,
                            numeric_threshold: float = 0.7,
                            preview_only: bool = False) -> Union[ToolResult, DeletionPreview]:
        """
        Remove summary/aggregate rows using both semantic (keyword-based) and structural patterns.
        
        A row is considered a summary row if:
        (1) it contains summary keywords (e.g., 'Total', 'Grand Total', 'Subtotal') as whole
            words/phrases in likely label columns, OR
        (2) structural confirmation: label column has a summary keyword AND the row is
            numeric-dense (typical total/subtotal layout).
        
        Args:
            df: DataFrame to filter
            keywords: List of keywords to identify summary rows
            check_columns: Optional list of column indices to check for keywords
            use_structural_detection: Enable structural pattern matching (default: True)
            numeric_threshold: Fraction of numeric columns required for structural match (default: 0.7)
            preview_only: If True, return DeletionPreview instead of executing deletion
            
        Returns:
            ToolResult with filtered DataFrame, or DeletionPreview if preview_only=True
        """
        try:
            df_copy = df.copy()
            original_len = len(df_copy)
            keywords = keywords or ['Total', 'Grand Total', 'Subtotal', 'Totals']
            
            # Normalize keywords to lowercase
            keywords_lower = [k.lower() for k in keywords]

            if check_columns:
                cols, _ = TransformationTools._resolve_columns(df_copy, check_columns)
            else:
                cols = TransformationTools._summary_label_columns(df_copy)
            
            if not cols:
                cols = df_copy.columns.tolist()
                 
            # Semantic detection: keyword-based matching (word boundaries, label columns only)
            semantic_mask = df_copy[cols].apply(
                lambda row: any(
                    TransformationTools._summary_keyword_in_text(val, keywords_lower)
                    for val in row
                ),
                axis=1,
            )
            
            # Structural detection: confirm summary rows — label column has a summary keyword
            # AND the row is numeric-dense (typical total/subtotal layout).
            structural_mask = pd.Series([False] * len(df_copy), index=df_copy.index)
            
            if use_structural_detection and len(df_copy.columns) > 1:
                label_col = cols[0] if cols else df_copy.columns[0]
                numeric_cols = TransformationTools._numeric_prone_columns(df_copy)

                if numeric_cols:
                    for idx in df_copy.index:
                        label_val = df_copy.at[idx, label_col]
                        if not TransformationTools._summary_keyword_in_text(label_val, keywords_lower):
                            continue

                        numeric_count = 0
                        for col in numeric_cols:
                            val = df_copy.at[idx, col]
                            if pd.notna(val):
                                try:
                                    float(val)
                                    numeric_count += 1
                                except (ValueError, TypeError):
                                    pass

                        if len(numeric_cols) > 0 and (numeric_count / len(numeric_cols)) >= numeric_threshold:
                            structural_mask.at[idx] = True
            
            # Combine masks: row is summary if semantic OR structural detection matches
            combined_mask = semantic_mask | structural_mask
            
            # Get the rows that WILL be deleted
            rows_to_delete = df_copy[combined_mask].index.tolist()
            deleted_data = df_copy[combined_mask]
            
            semantic_count = semantic_mask.sum()
            structural_count = (structural_mask & ~semantic_mask).sum()
            
            # If preview_only, return DeletionPreview instead of executing
            if preview_only:
                percent_deleted = (len(rows_to_delete) / original_len * 100) if original_len > 0 else 0
                return DeletionPreview(
                    tool_name="filter_summary_rows",
                    tool_description="Remove summary/aggregate rows (Total, Subtotal, Grand Total, etc.)",
                    rows_to_delete=rows_to_delete,
                    columns_to_delete=[],
                    sample_deleted_data=deleted_data.head(10).to_dict(orient='records'),
                    full_deleted_data=deleted_data,
                    reason=f"Found {semantic_count} keyword matches ({keywords[:3]}...) and {structural_count} structural matches",
                    impact_summary=f"Will remove {len(rows_to_delete)} rows ({percent_deleted:.1f}% of data)"
                )
            
            # Execute deletion
            df_result = df_copy[~combined_mask].reset_index(drop=True)
            
            return ToolResult(
                success=True,
                data=df_result,
                message=f"Removed {original_len - len(df_result)} summary rows ({semantic_count} semantic, {structural_count} structural)"
            )
        except Exception as e:
            if preview_only:
                return DeletionPreview(
                    tool_name="filter_summary_rows",
                    tool_description="Remove summary/aggregate rows",
                    reason=f"Error during preview: {e}",
                    impact_summary="Preview failed"
                )
            return ToolResult(success=False, data=df, message=f"Error: {e}")



    @staticmethod
    def drop_columns(df: pd.DataFrame, columns: List[Union[int, str]]) -> ToolResult:
        """
        Drop specified columns by name or index.
        """
        try:
            df = df.copy()
            cols_to_drop, _ = TransformationTools._resolve_columns(df, columns)
            if not cols_to_drop:
                 return ToolResult(success=False, data=df, message="No valid columns found to drop")
            
            df = df.drop(columns=cols_to_drop)
            return ToolResult(
                success=True,
                data=df,
                message=f"Dropped {len(cols_to_drop)} columns: {cols_to_drop[:5]}"
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error: {e}")

    @staticmethod
    def drop_blank_columns(df: pd.DataFrame, threshold: float = 1.0, 
                           preview_only: bool = False) -> Union[ToolResult, DeletionPreview]:
        """
        Automatically remove columns that are entirely or mostly blank.
        
        Args:
            df: DataFrame to modify
            threshold: Percentage (0.0 to 1.0) of blank values required to drop a column.
                       Default 1.0 (only drop 100% blank columns).
            preview_only: If True, return DeletionPreview instead of executing deletion
        """
        try:
            df_copy = df.copy()
            total_rows = len(df_copy)
            cols_to_drop = []
            
            if total_rows == 0:
                if preview_only:
                    return DeletionPreview(tool_name="drop_blank_columns", reason="Empty DataFrame", impact_summary="Nothing to delete")
                return ToolResult(success=True, data=df, message="Empty DataFrame - no columns to drop")

            for col in df_copy.columns:
                # Count NaNs and empty strings
                blank_mask = df_copy[col].isna() | (df_copy[col].astype(str).str.strip() == "")
                blank_count = blank_mask.sum()
                blank_ratio = blank_count / total_rows
                
                if blank_ratio >= threshold:
                    cols_to_drop.append(col)
                    
            if not cols_to_drop:
                if preview_only:
                    return DeletionPreview(tool_name="drop_blank_columns", reason="No blank columns found", impact_summary="Nothing to delete")
                return ToolResult(success=True, data=df, message=f"No columns found with blank ratio >= {threshold}")
            
            if preview_only:
                sample_data = df_copy[cols_to_drop].head(10).to_dict(orient='records')
                return DeletionPreview(
                    tool_name="drop_blank_columns",
                    tool_description="Remove columns with high percentage of blank values",
                    rows_to_delete=[],
                    columns_to_delete=cols_to_drop,
                    sample_deleted_data=sample_data,
                    full_deleted_data=df_copy[cols_to_drop],
                    reason=f"Found {len(cols_to_drop)} columns with blank ratio >= {threshold}: {cols_to_drop[:5]}",
                    impact_summary=f"Will remove {len(cols_to_drop)} columns"
                )
            
            df_result = df_copy.drop(columns=cols_to_drop).reset_index(drop=True)
            return ToolResult(
                success=True,
                data=df_result,
                message=f"Dropped {len(cols_to_drop)} blank columns: {cols_to_drop[:10]}",
                changes_made={"dropped_columns": cols_to_drop}
            )
        except Exception as e:
            if preview_only:
                return DeletionPreview(tool_name="drop_blank_columns", impact_summary=f"Preview failed: {e}")
            return ToolResult(success=False, data=df, message=f"Error in drop_blank_columns: {e}")



    @staticmethod
    def filter_header_rows(df: pd.DataFrame, threshold: float = 0.5, preview_only: bool = False) -> Union[ToolResult, DeletionPreview]:
        """
        Remove rows that look like repeated headers.
        """
        try:
            original_len = len(df)
            
            # Convert headers to string set for comparison
            headers = set(str(c).lower().strip() for c in df.columns)
            
            rows_to_delete = []
            
            for idx, row in df.iterrows():
                # Count how many values in this row match a column header
                matches = sum(1 for v in row if str(v).lower().strip() in headers)
                if matches / len(df.columns) >= threshold:
                    rows_to_delete.append(idx)
                    
            if not rows_to_delete:
                 return ToolResult(success=True, data=df, message="No repeated header rows found")
            
            # Preview handling
            deleted_data = df.loc[rows_to_delete]
            
            if preview_only:
                 percent_deleted = (len(rows_to_delete) / original_len * 100)
                 return DeletionPreview(
                    tool_name="filter_header_rows",
                    tool_description="Remove rows that are repetitions of the header",
                    rows_to_delete=rows_to_delete,
                    columns_to_delete=[],
                    sample_deleted_data=deleted_data.head(10).to_dict(orient='records'),
                    full_deleted_data=deleted_data,
                    reason=f"Found {len(rows_to_delete)} rows that match the header pattern",
                    impact_summary=f"Will remove {len(rows_to_delete)} repeated header rows ({percent_deleted:.1f}% of data)"
                )
            
            df = df.drop(index=rows_to_delete).reset_index(drop=True)
            
            return ToolResult(
                success=True,
                data=df,
                message=f"Removed {len(rows_to_delete)} repeated header rows"
            )
        except Exception as e:
            if preview_only:
                 return DeletionPreview(
                    tool_name="filter_header_rows",
                    reason=f"Error: {e}",
                    impact_summary="Preview failed"
                )
            return ToolResult(success=False, data=df, message=f"Error: {e}")

    @staticmethod
    def merge_header_rows(df: pd.DataFrame, header_rows: List[int], 
                          separator: str = " - ", strip_unnamed: bool = True) -> ToolResult:
        """
        Merge multiple rows into a single header row.
        
        Args:
            df: Current DataFrame (with raw data)
            header_rows: Row indices to merge
            separator: Separator for concatenation
            strip_unnamed: If True, skips generic names like "Unnamed: 0"
            
        Returns:
            ToolResult with renamed DataFrame and merged rows removed
        """
        try:
            df = df.copy()
            if not header_rows:
                return ToolResult(success=False, data=df, message="No header rows specified")
            
            # SAFETY GUARD: Check if DataFrame already has meaningful headers
            # If columns are already named (not numeric indices or "Unnamed"), 
            # this tool would corrupt data by treating data rows as headers.
            named_cols = [c for c in df.columns if not str(c).startswith("Unnamed") 
                         and not str(c).replace(".", "").isdigit()]
            if len(named_cols) > len(df.columns) * 0.5:
                logger.warning(f"header.collapse SKIPPED: DataFrame already has {len(named_cols)} "
                             f"named columns ({named_cols[:3]}...). Headers appear to be set.")
                return ToolResult(
                    success=True,
                    data=df,
                    message=f"SKIPPED: Headers already set ({named_cols[:5]}). "
                            f"Using header.collapse would corrupt data by promoting row values to headers."
                )
            
            # Extract header data
            header_data = []
            for row_idx in header_rows:
                if 0 <= row_idx < len(df):
                    header_data.append(df.iloc[row_idx].astype(str).tolist())
            
            if not header_data:
                return ToolResult(success=False, data=df, message="Specified header rows are invalid or empty")
            
            new_headers = []
            num_cols = len(df.columns)
            
            for col_idx in range(num_cols):
                parts = []
                for row_data in header_data:
                    val = row_data[col_idx].strip()
                    # Skip empty values or 'nan'
                    if not val or val.lower() == 'nan':
                        continue
                    # Skip 'Unnamed' if requested
                    if strip_unnamed and ('unnamed' in val.lower() or 'column_' in val.lower()):
                        continue
                    parts.append(val)
                
                # Combine parts
                if parts:
                    merged_name = separator.join(parts)
                else:
                    # Fallback to existing column name if all rows were empty
                    merged_name = df.columns[col_idx]
                
                new_headers.append(merged_name)
            
            # Apply new headers
            df.columns = new_headers
            
            # Remove the rows that were merged
            # Note: We sort descending to avoid index shifting if we were using iloc drop,
            # but df.drop uses the label-based index.
            df = df.drop(df.index[header_rows]).reset_index(drop=True)
            
            return ToolResult(
                success=True,
                data=df,
                message=f"Merged {len(header_rows)} rows into header. New columns: {new_headers[:5]}..."
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error merging headers: {e}")

    @staticmethod
    def crosstab_unpivot(df: pd.DataFrame, 
                         row_header_col: int = 0,
                         col_header_row: int = 0,
                         data_start_row: int = 1,
                         data_start_col: int = 1,
                         row_dim_name: str = "Row",
                         col_dim_name: str = "Column",
                         value_name: str = "Value",
                         id_cols: List[Union[int, str]] = None) -> ToolResult:
        """
        Unpivot a crosstab/matrix where both row and column headers represent dimensions.
        Converts Brand x Region matrix into flat table with Brand, Region, Value columns.
        
        Internal Logic:
        - id_cols: Dimensions that exist as columns in the same rows but should be bypassed/carried over.
        - Automatic Naming: If generic names are used, attempts to infer from header_row.
        """
        try:
            df_copy = df.copy()
            
            # 1. Resolve id_cols and Identify Auto-IDs
            id_col_names, id_confidences = TransformationTools._resolve_columns(df_copy, id_cols) if id_cols else ([], [])
            
            # **Lossless Logic**: Automatically identify all columns that are NOT the row header
            # and NOT part of the unpivot data columns as dimension IDs to preserve.
            matrix_data_cols = set(df_copy.columns[data_start_col:])
            row_header_name = df_copy.columns[row_header_col]
            
            all_other_cols = [c for c in df_copy.columns if c != row_header_name and c not in matrix_data_cols]
            
            # Merge explicit id_cols with auto-detected ones
            final_id_cols = list(dict.fromkeys(id_col_names + all_other_cols)) # Maintain order, remove dups
            
            # Build confidence map (auto-detected are 1.0)
            id_conf_map = {name: conf for name, conf in zip(id_col_names, id_confidences)}
            for col in all_other_cols:
                if col not in id_conf_map:
                    id_conf_map[col] = 1.0
            
            if final_id_cols:
                logger.info(f"Lossless Matrix: Automatically preserved columns: {final_id_cols}")

            # 2. Automatic Naming Inference
            # If default "Row" is used, try to grab name from physical col index
            if row_dim_name == "Row" and row_header_col < len(df_copy.columns):
                h_val = df_copy.iloc[col_header_row, row_header_col]
                if pd.notna(h_val) and str(h_val).strip() != "":
                    row_dim_name = str(h_val).strip()
            
            # 3. Extract column headers (The dimension values in the columns)
            col_headers = df_copy.iloc[col_header_row, data_start_col:].tolist()
            
            # 4. Perform Unpivot
            rows = []
            for idx in range(data_start_row, len(df_copy)):
                row_data = df_copy.iloc[idx]
                row_dim_val = row_data[row_header_name]
                
                # Extract extra dimensions
                extra_dims = {col: row_data[col] for col in final_id_cols}
                
                for col_idx, col_header in enumerate(col_headers):
                    val = row_data[data_start_col + col_idx]
                    
                    new_row = {
                        row_dim_name: row_dim_val,
                        col_dim_name: col_header,
                        value_name: val
                    }
                    new_row.update(extra_dims)
                    rows.append(new_row)
            
            result_df = pd.DataFrame(rows)
            
            # Reorder to put IDs first (using final_id_cols which includes auto-detected ones)
            cols_order = final_id_cols + [row_dim_name, col_dim_name, value_name]
            result_df = result_df[cols_order]
            
            return ToolResult(
                success=True,
                data=result_df,
                message=f"Matrix unpivoted ({len(df)} original rows -> {len(result_df)} flat rows)",
                changes_made={
                    "row_dim": row_dim_name,
                    "col_dim": col_dim_name,
                    "id_cols": final_id_cols,
                    "column_confidences": id_conf_map
                }
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in crosstab_unpivot: {e}")

    @staticmethod
    def transpose_data(df: pd.DataFrame,
                       header_col: int = None,
                       new_index_name: str = "Column") -> ToolResult:
        """
        Transpose the DataFrame (swap rows and columns).
        """
        try:
            if header_col is not None:
                # Use specified column as new column headers
                new_headers = df.iloc[:, header_col].tolist()
                data_df = df.drop(df.columns[header_col], axis=1)
                transposed = data_df.T
                transposed.columns = new_headers
            else:
                transposed = df.T
            
            # Reset index and rename
            transposed = transposed.reset_index()
            transposed.columns = [new_index_name] + list(transposed.columns[1:])
            
            return ToolResult(
                success=True,
                data=transposed,
                message=f"Transposed to {transposed.shape[0]} rows x {transposed.shape[1]} columns"
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in transpose_data: {e}")

    @staticmethod
    def align_schema(
        df: pd.DataFrame,
        target_columns: List[str],
        aliases: Optional[Dict[str, Any]] = None,
        fill_value: Any = None,
    ) -> ToolResult:
        """
        Deterministic column alignment for cross-source union: exact ``target_columns`` order,
        optional per-target source-name aliases, missing columns filled with ``fill_value``.
        """
        try:
            if df is None:
                return ToolResult(success=False, data=df, message="align_schema: missing DataFrame")
            aliases = aliases or {}
            parts: List[pd.Series] = []
            used_sources: Dict[str, str] = {}
            for tgt in target_columns:
                tgt_s = str(tgt)
                src_col: Optional[str] = None
                if tgt_s in df.columns:
                    src_col = tgt_s
                else:
                    raw = aliases.get(tgt_s)
                    if raw is None:
                        for k, v in aliases.items():
                            if str(k).strip() == tgt_s:
                                raw = v
                                break
                    candidates: List[str] = []
                    if isinstance(raw, str):
                        candidates = [raw]
                    elif isinstance(raw, (list, tuple)):
                        candidates = [str(x) for x in raw]
                    for cand in candidates:
                        if cand in df.columns:
                            src_col = cand
                            break
                if src_col is not None:
                    used_sources[tgt_s] = src_col
                    ser = df[src_col].copy()
                    ser.name = tgt_s
                    parts.append(ser)
                else:
                    parts.append(
                        pd.Series([fill_value] * len(df), index=df.index, dtype=object, name=tgt_s)
                    )
            out = pd.concat(parts, axis=1)
            return ToolResult(
                success=True,
                data=out,
                message=f"align_schema: {len(target_columns)} target columns ({len(used_sources)} mapped from source)",
                changes_made={"target_columns": list(target_columns), "source_columns_used": used_sources},
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in align_schema: {e}")

    @staticmethod
    def union_resolve(
        df: pd.DataFrame,
        key_columns: List[Union[int, str]],
        policy: str = "first",
        metric_columns: Optional[List[Union[int, str]]] = None,
        source_tag_column: Optional[str] = None,
        prefer_source_id: Optional[str] = None,
    ) -> ToolResult:
        """
        Post-concat resolution: collapse rows that share the same business keys.

        Policies: ``first`` (keep first row), ``sum`` / ``max`` / ``min`` on numeric metrics,
        ``prefer_source`` (needs ``source_tag_column`` + ``prefer_source_id``).
        """
        try:
            if df is None or df.empty:
                return ToolResult(success=False, data=df, message="union_resolve: empty DataFrame")
            keys, _ = TransformationTools._resolve_columns(df, key_columns)
            if not keys:
                return ToolResult(success=False, data=df, message="union_resolve: key_columns resolved to empty")
            pol = str(policy or "first").strip().lower()
            n0 = len(df)

            if pol == "first":
                out = df.drop_duplicates(subset=keys, keep="first").reset_index(drop=True)
                return ToolResult(
                    success=True,
                    data=out,
                    message=f"union_resolve(first): {n0} → {len(out)} rows",
                    changes_made={"keys": keys, "policy": pol},
                )

            if pol == "prefer_source":
                tag = source_tag_column
                if not tag or tag not in df.columns:
                    return ToolResult(
                        success=False,
                        data=df,
                        message="union_resolve(prefer_source): source_tag_column missing or not in frame",
                    )
                if prefer_source_id is None:
                    return ToolResult(
                        success=False,
                        data=df,
                        message="union_resolve(prefer_source): prefer_source_id required",
                    )

                df2 = df.copy()
                df2["_union_pref_rank"] = np.where(df2[tag] == prefer_source_id, 0, 1)
                sort_cols = list(keys) + ["_union_pref_rank"]
                df2 = df2.sort_values(sort_cols, kind="mergesort")
                out = df2.drop_duplicates(subset=keys, keep="first").drop(
                    columns=["_union_pref_rank"], errors="ignore"
                ).reset_index(drop=True)
                return ToolResult(
                    success=True,
                    data=out,
                    message=f"union_resolve(prefer_source): {n0} → {len(out)} rows",
                    changes_made={"keys": keys, "policy": pol, "source_tag_column": tag},
                )

            if pol not in {"sum", "max", "min"}:
                return ToolResult(
                    success=False,
                    data=df,
                    message=f"union_resolve: unknown policy {policy!r} (use first|sum|max|min|prefer_source)",
                )

            metrics_resolved: List[str] = []
            if metric_columns:
                metrics_resolved, _ = TransformationTools._resolve_columns(df, metric_columns)

            agg_dict: Dict[str, str] = {}
            for c in df.columns:
                if c in keys:
                    continue
                if source_tag_column and c == source_tag_column:
                    agg_dict[c] = "first"
                    continue
                use_agg = False
                if metrics_resolved:
                    use_agg = c in metrics_resolved
                else:
                    use_agg = pd.api.types.is_numeric_dtype(df[c])
                if use_agg and pd.api.types.is_numeric_dtype(df[c]):
                    agg_dict[c] = pol
                else:
                    agg_dict[c] = "first"

            out = df.groupby(keys, dropna=False, as_index=False).agg(agg_dict)
            return ToolResult(
                success=True,
                data=out,
                message=f"union_resolve({pol}): {n0} → {len(out)} rows",
                changes_made={"keys": keys, "policy": pol, "aggregated": agg_dict},
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in union_resolve: {e}")

    @staticmethod
    def validate_cross_source_union_tool(
        df: pd.DataFrame,
        column_sets_by_source: Dict[str, List[str]],
        key_hints: Optional[List[str]] = None,
        dtype_map_by_source: Optional[Dict[str, Dict[str, str]]] = None,
        date_granularity_by_source: Optional[Dict[str, str]] = None,
    ) -> ToolResult:
        """Read-only union readiness report; input frame is returned unchanged."""
        import json

        from sia.tools.cross_source_union_validate import validate_column_sets_for_union

        try:
            report = validate_column_sets_for_union(
                column_sets_by_source,
                key_hints=key_hints,
                dtype_map_by_source=dtype_map_by_source,
                date_granularity_by_source=date_granularity_by_source,
            )
            return ToolResult(
                success=True,
                data=df,
                message=json.dumps(report, indent=2),
                changes_made={"union_readiness": report},
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in validate_cross_source_union_tool: {e}")

    @staticmethod
    def deduplicate_rows(df: pd.DataFrame,
                         subset: List[Union[int, str]] = None,
                         keep: Union[str, bool] = "first") -> ToolResult:
        """
        Remove duplicate rows based on specified columns.
        """
        try:
            # Resolve column names if indices provided
            resolved_subset, confidences = TransformationTools._resolve_columns(df, subset) if subset else ([], [])
            
            original_len = len(df)
            result_df = df.drop_duplicates(subset=resolved_subset or None, keep=keep)
            removed = original_len - len(result_df)
            
            return ToolResult(
                success=True,
                data=result_df,
                message=f"Removed {removed} duplicate rows ({len(result_df)} remaining)",
                changes_made={
                    "subset_columns": resolved_subset,
                    "column_confidences": {name: conf for name, conf in zip(resolved_subset, confidences)}
                }
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in deduplicate_rows: {e}")

    @staticmethod
    def unpivot_columns(df: pd.DataFrame,
                        id_cols: List[Union[int, str]] = None,
                        value_cols: List[Union[int, str]] = None,
                        var_name: str = "Variable",
                        value_name: str = "Value",
                        **kwargs) -> ToolResult:
        """
        Unpivot (melt) a DataFrame from wide to long format.
        
        Uses lossless logic: if value_cols are provided, all other columns
        are automatically treated as ID variables to prevent data loss.
        """
        try:
            # Handle variations in parameter names from LLM
            if id_cols is None:
                id_cols = kwargs.pop('id_columns', [])
            if value_cols is None:
                value_cols = kwargs.pop('value_columns', [])

            # Resolve id_cols and value_cols and collect confidences
            resolved_id_cols, id_conf_list = TransformationTools._resolve_columns(df, id_cols) if id_cols else ([], [])
            resolved_value_cols, val_conf_list = TransformationTools._resolve_columns(df, value_cols) if value_cols else ([], [])
            
            # **Lossless Logic**: If value_columns are provided, any column NOT in value_columns
            # should automatically be treated as an ID variable to prevent data loss.
            # This ensures dimensions like Channel, Region, Date are preserved even if AI misses them.
            if resolved_value_cols:
                all_cols = df.columns.tolist()
                val_set = set(resolved_value_cols)
                # Automatically include all non-value columns as IDs
                resolved_id_cols = [c for c in all_cols if c not in val_set]
                id_conf_list = [1.0] * len(resolved_id_cols)
                logger.info(f"Lossless Unpivot: Automatically included {len(resolved_id_cols)} columns as IDs: {resolved_id_cols[:3]}...")

            # Combine confidences
            all_resolved = resolved_id_cols + resolved_value_cols
            all_conf_list = id_conf_list + val_conf_list
            confidences = {name: conf for name, conf in zip(all_resolved, all_conf_list)}

            # Melt
            melted_df = df.melt(
                id_vars=resolved_id_cols,
                value_vars=resolved_value_cols if resolved_value_cols else None,
                var_name=var_name,
                value_name=value_name,
            )

            count = len(resolved_value_cols) if resolved_value_cols else (len(df.columns) - len(resolved_id_cols))
            return ToolResult(
                success=True,
                data=melted_df,
                message=f"Unpivoted {count} columns into '{var_name}'/'{value_name}'",
                changes_made={
                    "id_cols": resolved_id_cols, 
                    "value_cols": resolved_value_cols,
                    "column_confidences": confidences
                }
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in unpivot_columns: {e}")

    @staticmethod
    def hierarchical_unpivot(df: pd.DataFrame,
                             id_cols: List[Union[int, str]] = None,
                             value_cols: List[Union[int, str]] = None,
                             var_name: str = "Variable",
                             value_name: str = "Value") -> ToolResult:
        """Backward-compatible alias for unpivot_columns."""
        return TransformationTools.unpivot_columns(
            df, id_cols=id_cols, value_cols=value_cols, var_name=var_name, value_name=value_name
        )

    @staticmethod
    def split_column(df: pd.DataFrame, column: Union[int, str], new_names: List[str],
                     separator: str = " - ", remove_original: bool = True) -> ToolResult:
        """Split one text column into multiple columns."""
        try:
            df_copy = df.copy()
            col_name, _ = TransformationTools._fuzzy_find_column(df_copy, column)

            parts = df_copy[col_name].astype(str).str.split(separator, n=len(new_names) - 1, expand=True)
            for i, new_col in enumerate(new_names):
                df_copy[new_col] = parts[i].str.strip() if i < parts.shape[1] else None

            if remove_original and col_name not in new_names:
                df_copy = df_copy.drop(columns=[col_name])

            return ToolResult(
                success=True,
                data=df_copy,
                message=f"Split column '{col_name}' into {len(new_names)} columns"
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in split_column: {e}")

    @staticmethod
    def fuzzy_standardize_column(df: pd.DataFrame, column: Union[int, str], threshold: float = 0.85,
                                 case_sensitive: bool = False) -> ToolResult:
        """Standardize label variants by simple normalization and fuzzy grouping."""
        try:
            import difflib

            df_copy = df.copy()
            col_name, _ = TransformationTools._fuzzy_find_column(df_copy, column)

            # Normalize values to improve grouping.
            series = df_copy[col_name].fillna("").astype(str)
            normalized = series.str.strip()
            if not case_sensitive:
                normalized = normalized.str.lower()

            groups: Dict[str, str] = {}
            ordered_uniques = []
            for raw in normalized:
                if raw == "":
                    continue
                if raw not in ordered_uniques:
                    ordered_uniques.append(raw)

            for token in ordered_uniques:
                canonical = None
                for existing in groups.keys():
                    if difflib.SequenceMatcher(None, token, existing).ratio() >= threshold:
                        canonical = existing
                        break
                if canonical is None:
                    groups[token] = token
                else:
                    groups[token] = canonical

            standardized = normalized.map(lambda x: groups.get(x, x))
            if not case_sensitive:
                standardized = standardized.map(lambda x: x.title() if x else x)
            df_copy[col_name] = standardized

            return ToolResult(
                success=True,
                data=df_copy,
                message=f"Standardized column '{col_name}' to {df_copy[col_name].nunique(dropna=True)} unique values"
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in fuzzy_standardize_column: {e}")

    @staticmethod
    def filter_summary_columns(df: pd.DataFrame, keywords: List[str] = None) -> ToolResult:
        """Drop columns whose names look like total/summary columns."""
        try:
            keywords = keywords or ["total", "subtotal", "sum", "grand total"]
            lower_keywords = [k.lower() for k in keywords]

            drop_cols = [
                c for c in df.columns
                if any(k in str(c).lower() for k in lower_keywords)
            ]
            result_df = df.drop(columns=drop_cols) if drop_cols else df.copy()

            return ToolResult(
                success=True,
                data=result_df,
                message=f"Dropped {len(drop_cols)} summary columns"
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in filter_summary_columns: {e}")

    @staticmethod
    def detect_structural_summary_rows(df: pd.DataFrame, text_column: Union[int, str] = 0,
                                       keywords: List[str] = None) -> ToolResult:
        """Detect likely summary rows without deleting them."""
        try:
            keywords = keywords or ["total", "subtotal", "grand total", "sum", "totals"]
            text_col, _ = TransformationTools._fuzzy_find_column(df, text_column)
            hits = []
            for idx, val in df[text_col].items():
                sval = str(val).strip().lower() if pd.notna(val) else ""
                if sval and any(k in sval for k in keywords):
                    hits.append({"row_index": int(idx), "label": str(val), "reason": "keyword_match"})

            return ToolResult(
                success=True,
                data=hits,
                message=f"Detected {len(hits)} structural summary rows"
            )
        except Exception as e:
            return ToolResult(success=False, data=[], message=f"Error in detect_structural_summary_rows: {e}")

    @staticmethod
    def expand_hierarchical_rows(df: pd.DataFrame, parent_col: Union[int, str],
                                 child_indicator_cols: List[Union[int, str]] = None) -> ToolResult:
        """Propagate parent labels into child rows in hierarchical tables."""
        try:
            df_copy = df.copy()
            parent_name, _ = TransformationTools._fuzzy_find_column(df_copy, parent_col)

            # Normalize empty markers and forward fill parent values.
            df_copy[parent_name] = df_copy[parent_name].apply(
                lambda x: np.nan if pd.isna(x) or str(x).strip() == "" else x
            ).ffill()

            return ToolResult(
                success=True,
                data=df_copy,
                message=f"Expanded hierarchical rows using parent column '{parent_name}'"
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in expand_hierarchical_rows: {e}")


    # ===== SEMANTIC COLUMN RESOLUTION =====

    _synonyms_cache = None

    @staticmethod
    def _load_synonyms() -> dict:
        """Load the synonyms dictionary from synonyms.json."""
        if TransformationTools._synonyms_cache is not None:
            return TransformationTools._synonyms_cache
        try:
            from pathlib import Path
            syn_path = Path(__file__).parent.parent.parent / "config" / "synonyms.json"
            if syn_path.exists():
                import json
                with open(syn_path, "r", encoding="utf-8") as f:
                    TransformationTools._synonyms_cache = json.load(f)
                    return TransformationTools._synonyms_cache
        except Exception as e:
            logger.warning(f"Failed to load synonyms.json: {e}")
        TransformationTools._synonyms_cache = {}
        return {}

    @staticmethod
    def _resolve_columns(df: pd.DataFrame, columns: Any) -> tuple[List[str], List[float]]:
        """
        Resolve a mix of column indices and names to actual column names.
        Now uses _fuzzy_find_column for semantic matching via synonym dictionary.
        
        Returns:
            tuple (List[str], List[float]): Resolved names and their confidence scores.
        """
        if columns is None:
            return [], []
            
        if not isinstance(columns, list):
            columns = [columns]
            
        resolved = []
        confidences = []
        for col in columns:
            res_name, conf = TransformationTools._fuzzy_find_column(df, col)
            resolved.append(res_name)
            confidences.append(conf)
            
        return resolved, confidences

    @staticmethod
    def _fuzzy_find_column(df: pd.DataFrame, query: Union[int, str]) -> tuple[str, float]:
        """
        Resolve a column name using 3-tier matching:
        1. Exact match (1.0 confidence)
        2. Synonym dictionary lookup (0.95 confidence)
        3. Fuzzy substring/ratio matching (0.6-0.85 confidence)

        Returns:
            tuple (str, float): The best-matching column name and confidence score.
        """
        columns = list(df.columns)

        # If it's an integer index, resolve directly
        if isinstance(query, int):
            if 0 <= query < len(columns):
                return columns[query], 1.0
            raise ValueError(f"Column index {query} out of range (0-{len(columns)-1})")

        query_str = str(query).strip()
        query_lower = query_str.lower()

        # Tier 1: Exact match (case-insensitive) -> 1.0
        for col in columns:
            if str(col).lower().strip() == query_lower:
                return col, 1.0

        # Tier 2: Synonym dictionary lookup -> 0.95
        synonyms = TransformationTools._load_synonyms()
        col_synonyms = synonyms.get("column_synonyms", {})

        for canonical, aliases in col_synonyms.items():
            all_names = [canonical] + [a.lower() for a in aliases]
            query_normalized = query_lower.replace("_", " ").replace("-", " ")

            for name in all_names:
                name_normalized = name.lower().replace("_", " ").replace("-", " ")
                if query_normalized == name_normalized or query_lower == name.lower():
                    for col in columns:
                        col_normalized = str(col).lower().replace("_", " ").replace("-", " ")
                        if col_normalized == canonical or col_normalized in [a.lower().replace("_", " ").replace("-", " ") for a in [canonical] + aliases]:
                            logger.info(f"Synonym match: '{query_str}' -> '{col}' (via canonical '{canonical}')")
                            return col, 0.95

        # Tier 3: Fuzzy substring matching -> 0.85
        import difflib
        col_map = {str(c).lower().replace("_", "").replace(" ", ""): c for c in columns}
        query_compact = query_lower.replace("_", "").replace(" ", "")

        for compact, original in col_map.items():
            if query_compact in compact or compact in query_compact:
                logger.info(f"Fuzzy substring match: '{query_str}' -> '{original}'")
                return original, 0.85

        # Last resort: difflib ratio -> score
        # Normalize candidates to lowercase strings
        str_cols = [str(c).lower() for c in columns]
        matches = difflib.get_close_matches(query_lower, str_cols, n=1, cutoff=0.6)
        
        if matches:
            match_name = matches[0]
            # Calculate actual ratio
            match_score = difflib.SequenceMatcher(None, query_lower, match_name).ratio()
            
            for col in columns:
                if str(col).lower() == match_name:
                    logger.info(f"Fuzzy ratio match: '{query_str}' -> '{col}' (score: {match_score:.2f})")
                    return col, match_score

        raise ValueError(
            f"Column '{query_str}' not found. Available: {columns}. "
            f"Tried: exact match, synonym lookup, fuzzy matching."
        )

    # ===== LAST-MILE TRANSFORMATION TOOLS =====

    @staticmethod
    def map_values(df: pd.DataFrame, column: Union[int, str],
                   mapping: Dict[str, str],
                   case_insensitive: bool = True,
                   default: str = None) -> ToolResult:
        """
        Map messy source values to standardized target values.

        Use this to align free-text categories to a strict enum
        (e.g., "FB" -> "social", "Google Ads" -> "search").

        Column resolution is SEMANTIC - uses fuzzy matching and synonym dictionary.

        Args:
            df: DataFrame to modify
            column: Column name, index, or semantic alias (e.g., "channel_paid_media" -> "Channel")
            mapping: Dict of {source_value: target_value}
            case_insensitive: If True, matching ignores case
            default: If set, unmapped values are replaced with this.
                     If None, unmapped values are kept as-is.

        Returns:
            ToolResult with mapped DataFrame
        """
        try:
            # Semantic column resolution (3-tier: exact, synonym, fuzzy)
            col_name, col_conf = TransformationTools._fuzzy_find_column(df, column)

            # Enrich mapping with value synonyms from dictionary
            synonyms = TransformationTools._load_synonyms()
            value_syns = synonyms.get("value_synonyms", {})

            enriched_mapping = dict(mapping)  # Start with explicit mapping
            
            # Collect ALL known canonical values (to protect from default overwrite)
            all_canonical_values = set()
            for category, syn_groups in value_syns.items():
                for canonical, aliases in syn_groups.items():
                    all_canonical_values.add(canonical.lower())
                    
            for target_val in set(mapping.values()):
                # For each target value in the mapping, check if there are known synonyms
                for category, syn_groups in value_syns.items():
                    for canonical, aliases in syn_groups.items():
                        if target_val.lower() == canonical.lower():
                            # Add all known aliases as additional source keys
                            for alias in aliases:
                                if alias.lower() not in {k.lower() for k in enriched_mapping}:
                                    enriched_mapping[alias] = target_val

            # AUTO-PROTECT: If a value already matches a known canonical value,
            # map it to itself so it survives even when default is set.
            # e.g., "social" -> "social", "TV" -> "TV", "digital display" -> "digital display"
            for category, syn_groups in value_syns.items():
                for canonical, aliases in syn_groups.items():
                    if canonical.lower() not in {k.lower() for k in enriched_mapping}:
                        enriched_mapping[canonical] = canonical
                    # Also protect aliases that ARE canonical values in other groups
                    for alias in aliases:
                        if alias.lower() in all_canonical_values:
                            if alias.lower() not in {k.lower() for k in enriched_mapping}:
                                enriched_mapping[alias] = alias

            original_unique = df[col_name].nunique()
            result_df = df.copy()

            if case_insensitive:
                # Build a lowercase mapping
                lower_map = {k.lower().strip(): v for k, v in enriched_mapping.items()}
                protected_count = 0
                def _map_fn(val):
                    nonlocal protected_count
                    if pd.isna(val):
                        return val
                    key = str(val).lower().strip()
                    if key in lower_map:
                        return lower_map[key]
                    # SAFETY: If value matches a known canonical value, preserve it
                    if key in all_canonical_values:
                        protected_count += 1
                        return val
                    return default if default is not None else val
                result_df[col_name] = result_df[col_name].apply(_map_fn)
            else:
                if default is not None:
                    result_df[col_name] = result_df[col_name].map(enriched_mapping).fillna(default)
                else:
                    result_df[col_name] = result_df[col_name].map(enriched_mapping).fillna(result_df[col_name])

            new_unique = result_df[col_name].nunique()
            mapped_count = sum(1 for v in df[col_name] if str(v).lower().strip() in
                              {k.lower().strip() for k in enriched_mapping})
            enriched_count = len(enriched_mapping) - len(mapping)

            msg = (f"Mapped {mapped_count} values in '{col_name}' "
                   f"({original_unique} unique -> {new_unique} unique). "
                   f"Enriched with {enriched_count} synonym entries.")
            if case_insensitive and protected_count > 0:
                msg += f" Protected {protected_count} values that already matched canonical values."

            return ToolResult(
                success=True,
                data=result_df,
                message=msg,
                changes_made={
                    "mapped_col": col_name,
                    "column_confidences": {col_name: col_conf}
                }
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in map_values: {e}")

    @staticmethod
    def add_column(
        df: pd.DataFrame,
        target_column: str,
        value: Any,
        when: str = "missing",
    ) -> ToolResult:
        """
        Create or fill a column with a literal value (UID hierarchy defaults, placeholders).

        ``when``:
        - ``missing``: column absent → create filled with ``value``
        - ``null_or_blank``: column exists → fill null/blank cells only
        - ``always``: overwrite entire column
        """
        try:
            if df is None:
                return ToolResult(success=False, data=df, message="add_column: empty DataFrame")
            col = str(target_column or "").strip()
            if not col:
                return ToolResult(success=False, data=df, message="add_column: target_column is required")

            mode = str(when or "missing").strip().lower().replace("-", "_")
            if mode not in ("missing", "null_or_blank", "always"):
                return ToolResult(
                    success=False,
                    data=df,
                    message=f"add_column: unsupported when={when!r}; use missing, null_or_blank, or always",
                )

            out = df.copy()
            if mode == "missing":
                if col not in out.columns:
                    out[col] = value
                    return ToolResult(
                        success=True,
                        data=out,
                        message=f"Created column {col!r} = {value!r} ({len(out)} rows)",
                        changes_made={"target_column": col, "when": mode, "rows": len(out)},
                    )
                return ToolResult(
                    success=True,
                    data=out,
                    message=f"Column {col!r} already exists; no change (when=missing)",
                )

            if col not in out.columns:
                out[col] = value
                return ToolResult(
                    success=True,
                    data=out,
                    message=f"Created column {col!r} = {value!r} ({len(out)} rows)",
                    changes_made={"target_column": col, "when": mode},
                )

            if mode == "always":
                out[col] = value
                return ToolResult(
                    success=True,
                    data=out,
                    message=f"Set all values in {col!r} = {value!r}",
                    changes_made={"target_column": col, "when": mode},
                )

            ser = out[col]
            blank = ser.isna()
            if ser.dtype == object or str(ser.dtype) == "string":
                blank = blank | ser.astype(str).str.strip().eq("")
            if bool(blank.all()) or len(out) == 0:
                out[col] = value
                msg = f"Filled empty column {col!r} = {value!r}"
            elif bool(blank.any()):
                out.loc[blank, col] = value
                msg = f"Filled {int(blank.sum())} null/blank rows in {col!r} = {value!r}"
            else:
                msg = f"No null/blank cells in {col!r}; no change"
            return ToolResult(success=True, data=out, message=msg, changes_made={"target_column": col, "when": mode})
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in add_column: {e}")

    @staticmethod
    def apply_column_rules(
        df: pd.DataFrame,
        column_rules: Optional[List[Dict[str, Any]]] = None,
    ) -> ToolResult:
        """Apply target-template ``column_rules`` (set_value when column missing or null)."""
        try:
            from sia.agent.target_template_utils import apply_column_rules_list

            if df is None:
                return ToolResult(success=False, data=df, message="apply_column_rules: empty DataFrame")
            rules = list(column_rules or [])
            if not rules:
                return ToolResult(success=False, data=df, message="apply_column_rules: column_rules is empty")

            out, actions = apply_column_rules_list(df, rules)
            if not actions:
                return ToolResult(
                    success=True,
                    data=out,
                    message="apply_column_rules: no rules matched or no changes needed",
                )
            return ToolResult(
                success=True,
                data=out,
                message="; ".join(actions[:8]) + ("..." if len(actions) > 8 else ""),
                changes_made={"rules_applied": len(rules), "actions": actions},
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in apply_column_rules: {e}")

    @staticmethod
    def calculate_column(df: pd.DataFrame,
                         target_column: str,
                         expression: str,
                         source_columns: List[str] = None) -> ToolResult:
        """
        Create or overwrite a column using a simple arithmetic expression.

        Supports: +, -, *, /, and numeric constants.
        Column references use their exact name.

        Examples:
            expression="Gross_Spend * 0.85", target_column="Net_Spend"
            expression="Impressions * 1000", target_column="Impressions"
            expression="Clicks / Impressions", target_column="CTR"

        Args:
            df: DataFrame to modify
            target_column: Name of the column to create/overwrite
            expression: Arithmetic expression using column names and constants
            source_columns: Optional explicit list of columns used
                            (for validation; auto-detected if not provided)

        Returns:
            ToolResult with updated DataFrame
        """
        try:
            result_df = df.copy()

            # Safety: only allow arithmetic operators and column references
            import re
            # Extract identifier tokens from expression (physical column names only).
            raw_tokens = re.findall(r"[A-Za-z_][A-Za-z0-9_]*", expression)
            skip_words = {
                "and", "or", "not", "if", "else", "true", "false", "sum", "min", "max", "abs",
            }
            candidate_tokens = [
                t for t in raw_tokens if t.lower() not in skip_words
            ]
            detected_cols = [t for t in candidate_tokens if t in df.columns]

            if source_columns:
                resolved, _ = TransformationTools._resolve_columns(df, source_columns)
                detected_cols = list(set(detected_cols + resolved))

            missing = [t for t in candidate_tokens if t not in df.columns]
            if missing:
                return ToolResult(
                    success=False,
                    data=df,
                    message=(
                        f"Column(s) {missing} referenced in expression not found. "
                        f"Use physical source names (before rename), not template stubs like spends_1. "
                        f"Available: {list(df.columns[:12])}"
                    ),
                )

            # Build a safe evaluation context with only pandas Series
            eval_context = {}
            safe_expr = expression
            for col in detected_cols:
                safe_key = col.replace(' ', '_').replace('-', '_')
                eval_context[safe_key] = pd.to_numeric(df[col], errors='coerce')
                safe_expr = safe_expr.replace(col, safe_key)

            # Evaluate using pandas eval for safety (avoids Python eval injection risks)
            try:
                result_df[target_column] = pd.eval(safe_expr, local_dict=eval_context, engine='python')
            except Exception:
                # Fallback for expressions pd.eval can't handle (e.g., complex column names)
                result_df[target_column] = eval(safe_expr, {"__builtins__": {}}, eval_context)

            non_null = result_df[target_column].notna().sum()

            return ToolResult(
                success=True,
                data=result_df,
                message=f"Calculated '{target_column}' = {expression} ({non_null} non-null values)"
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in calculate_column: {e}")

    @staticmethod
    def _apply_simple_row_filters(df: pd.DataFrame, filters: Optional[List[Dict[str, Any]]]) -> Tuple[pd.DataFrame, List[str]]:
        """AND-combine simple predicates (business-rule style pre-aggregation filters)."""
        if not filters:
            return df, []
        work = df.copy()
        mask = pd.Series(True, index=work.index)
        notes: List[str] = []
        for raw in filters:
            if not isinstance(raw, dict):
                continue
            col_key = raw.get("column")
            if col_key is None:
                continue
            try:
                col_name, _conf = TransformationTools._fuzzy_find_column(work, col_key)
            except ValueError:
                notes.append(f"filter skip: unknown column {col_key!r}")
                continue
            op = str(raw.get("operator", "eq")).strip().lower()
            val = raw.get("value")
            s = work[col_name]
            if op == "eq":
                if isinstance(val, (int, float)) and not isinstance(val, bool):
                    mask &= pd.to_numeric(s, errors="coerce") == float(val)
                else:
                    mask &= s.astype(str).str.strip().str.casefold() == str(val).strip().casefold()
            elif op == "ne":
                if isinstance(val, (int, float)) and not isinstance(val, bool):
                    mask &= pd.to_numeric(s, errors="coerce") != float(val)
                else:
                    mask &= s.astype(str).str.strip().str.casefold() != str(val).strip().casefold()
            elif op in ("gt", "gte", "lt", "lte"):
                num = pd.to_numeric(s, errors="coerce")
                v = float(val) if isinstance(val, (int, float, str)) and str(val).strip() != "" else np.nan
                if op == "gt":
                    mask &= num > v
                elif op == "gte":
                    mask &= num >= v
                elif op == "lt":
                    mask &= num < v
                else:
                    mask &= num <= v
            elif op == "in":
                if not isinstance(val, (list, tuple, set)):
                    notes.append("filter skip: 'in' needs list value")
                    continue
                allowed = {str(x).strip().casefold() for x in val}
                mask &= s.astype(str).str.strip().str.casefold().isin(allowed)
            elif op == "not_in":
                if not isinstance(val, (list, tuple, set)):
                    notes.append("filter skip: 'not_in' needs list value")
                    continue
                forbidden = {str(x).strip().casefold() for x in val}
                mask &= ~s.astype(str).str.strip().str.casefold().isin(forbidden)
            else:
                notes.append(f"filter skip: unknown operator {op!r}")
        out = work.loc[mask].copy()
        return out, notes

    @staticmethod
    def _monday_week_start(ts: pd.Timestamp) -> pd.Timestamp:
        ts = pd.Timestamp(ts).normalize()
        return ts - pd.Timedelta(days=int(ts.weekday()))

    @staticmethod
    def _coerce_int_value(value: Any) -> Optional[int]:
        if pd.isna(value):
            return None
        if isinstance(value, (int, np.integer)):
            return int(value)
        if isinstance(value, (float, np.floating)):
            if np.isnan(value) or not float(value).is_integer():
                return None
            return int(value)
        text = str(value).strip()
        if not text:
            return None
        try:
            parsed = float(text)
            if not parsed.is_integer():
                return None
            return int(parsed)
        except (TypeError, ValueError):
            return None

    @staticmethod
    def _parse_month_value(value: Any) -> Optional[int]:
        as_int = TransformationTools._coerce_int_value(value)
        if as_int is not None and 1 <= as_int <= 12:
            return as_int
        if pd.isna(value):
            return None
        text = str(value).strip().lower().replace(".", "")
        if not text:
            return None
        month_map = {
            "jan": 1, "january": 1,
            "feb": 2, "february": 2,
            "mar": 3, "march": 3,
            "apr": 4, "april": 4,
            "may": 5,
            "jun": 6, "june": 6,
            "jul": 7, "july": 7,
            "aug": 8, "august": 8,
            "sep": 9, "sept": 9, "september": 9,
            "oct": 10, "october": 10,
            "nov": 11, "november": 11,
            "dec": 12, "december": 12,
        }
        return month_map.get(text)

    @staticmethod
    def _parse_quarter_value(value: Any) -> Optional[int]:
        as_int = TransformationTools._coerce_int_value(value)
        if as_int is not None and 1 <= as_int <= 4:
            return as_int
        if pd.isna(value):
            return None
        text = (
            str(value)
            .strip()
            .lower()
            .replace("quarter", "")
            .replace("q", "")
            .replace("-", "")
            .replace(" ", "")
        )
        if not text:
            return None
        try:
            quarter = int(text)
        except (TypeError, ValueError):
            return None
        return quarter if 1 <= quarter <= 4 else None

    @staticmethod
    def _period_bounds_from_date(ts: pd.Timestamp, granularity: str) -> Tuple[pd.Timestamp, pd.Timestamp]:
        g = str(granularity or "").strip().lower()
        if g in {"monthly", "month"}:
            period = ts.to_period("M")
        elif g in {"quarterly", "quarter"}:
            period = ts.to_period("Q")
        else:
            raise ValueError(f"Unsupported period granularity: {granularity!r}")
        return period.start_time.normalize(), period.end_time.normalize()

    @staticmethod
    def build_date_from_parts(
        df: pd.DataFrame,
        year_col: Union[str, int],
        month_col: Optional[Union[str, int]] = None,
        day_col: Optional[Union[str, int]] = None,
        quarter_col: Optional[Union[str, int]] = None,
        target_date_col: str = "calendar_date",
        default_day: int = 1,
        drop_source_columns: bool = False,
        row_filters: Optional[List[Dict[str, Any]]] = None,
    ) -> ToolResult:
        """Build a real date column from separate year/month/day or year/quarter parts."""
        try:
            if df is None or df.empty:
                return ToolResult(success=True, data=df, message="Empty input; nothing to build.")

            work, filter_notes = TransformationTools._apply_simple_row_filters(df, row_filters)
            if work.empty:
                return ToolResult(
                    success=False,
                    data=df,
                    message="All rows removed by row_filters; nothing to build." + (
                        f" Notes: {'; '.join(filter_notes)}" if filter_notes else ""
                    ),
                )

            year_name, _ = TransformationTools._fuzzy_find_column(work, year_col)
            month_name = None
            day_name = None
            quarter_name = None
            if month_col is not None:
                month_name, _ = TransformationTools._fuzzy_find_column(work, month_col)
            if day_col is not None:
                day_name, _ = TransformationTools._fuzzy_find_column(work, day_col)
            if quarter_col is not None:
                quarter_name, _ = TransformationTools._fuzzy_find_column(work, quarter_col)

            if not month_name and not quarter_name:
                return ToolResult(
                    success=False,
                    data=df,
                    message="build_date_from_parts requires either month_col or quarter_col.",
                )

            result_df = work.copy()
            built_dates: List[pd.Timestamp] = []
            invalid_rows = 0

            for _, row in result_df.iterrows():
                year_val = TransformationTools._coerce_int_value(row[year_name])
                if year_val is None:
                    built_dates.append(pd.NaT)
                    invalid_rows += 1
                    continue

                month_val: Optional[int] = None
                if month_name:
                    month_val = TransformationTools._parse_month_value(row[month_name])
                if month_val is None and quarter_name:
                    quarter_val = TransformationTools._parse_quarter_value(row[quarter_name])
                    if quarter_val is not None:
                        month_val = 1 + ((quarter_val - 1) * 3)
                if month_val is None:
                    built_dates.append(pd.NaT)
                    invalid_rows += 1
                    continue

                day_val = default_day
                if day_name:
                    parsed_day = TransformationTools._coerce_int_value(row[day_name])
                    if parsed_day is None:
                        built_dates.append(pd.NaT)
                        invalid_rows += 1
                        continue
                    day_val = parsed_day

                try:
                    built_dates.append(pd.Timestamp(year=year_val, month=month_val, day=day_val).normalize())
                except (TypeError, ValueError):
                    built_dates.append(pd.NaT)
                    invalid_rows += 1

            result_df[target_date_col] = built_dates
            valid_count = int(result_df[target_date_col].notna().sum())
            if valid_count == 0:
                return ToolResult(
                    success=False,
                    data=df,
                    message="No valid dates could be built from the provided parts.",
                )

            dropped_cols: List[str] = []
            if drop_source_columns:
                for col_name in [year_name, month_name, day_name, quarter_name]:
                    if col_name and col_name != target_date_col and col_name in result_df.columns:
                        dropped_cols.append(str(col_name))
                if dropped_cols:
                    result_df = result_df.drop(columns=dropped_cols)

            msg_parts = [f"Built '{target_date_col}' for {valid_count} row(s) from date parts."]
            if invalid_rows:
                msg_parts.append(f"Skipped {invalid_rows} row(s) with invalid parts.")
            if dropped_cols:
                msg_parts.append(f"Dropped source part columns: {', '.join(dropped_cols)}.")
            if filter_notes:
                msg_parts.append("Filters: " + "; ".join(filter_notes))

            return ToolResult(
                success=True,
                data=result_df,
                message=" ".join(msg_parts),
                changes_made={
                    "target_date_col": target_date_col,
                    "year_col": year_name,
                    "month_col": month_name,
                    "day_col": day_name,
                    "quarter_col": quarter_name,
                    "invalid_rows": invalid_rows,
                },
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in build_date_from_parts: {e}")

    @staticmethod
    def expand_period_to_daily(
        df: pd.DataFrame,
        date_col: Union[str, int],
        value_cols: List[Any],
        input_granularity: str,
        id_cols: Optional[List[Any]] = None,
        date_column: str = "calendar_date",
        row_filters: Optional[List[Dict[str, Any]]] = None,
    ) -> ToolResult:
        """Expand monthly or quarterly period rows into equal daily slices."""
        try:
            if df is None or df.empty:
                return ToolResult(success=True, data=df, message="Empty input; nothing to expand.")

            granularity = str(input_granularity or "").strip().lower()
            if granularity not in {"monthly", "month", "quarterly", "quarter"}:
                return ToolResult(
                    success=False,
                    data=df,
                    message=f"input_granularity must be 'monthly' or 'quarterly', got {input_granularity!r}",
                )

            work, filter_notes = TransformationTools._apply_simple_row_filters(df, row_filters)
            if work.empty:
                return ToolResult(
                    success=False,
                    data=df,
                    message="All rows removed by row_filters; nothing to expand." + (
                        f" Notes: {'; '.join(filter_notes)}" if filter_notes else ""
                    ),
                )

            date_name, _ = TransformationTools._fuzzy_find_column(work, date_col)
            resolved_values, _ = TransformationTools._resolve_columns(work, value_cols)
            if not resolved_values:
                return ToolResult(success=False, data=df, message="value_cols resolved to empty list.")

            if id_cols is not None:
                id_resolved, _ = TransformationTools._resolve_columns(work, id_cols)
            else:
                id_resolved = []
            # Daily rows are rebuilt from scratch, so anything not carried here is lost.
            # A partial id_cols list from the planner must not silently drop template
            # columns added upstream (publisher/channel defaults etc.).
            reserved = {date_name, *resolved_values}
            id_resolved = id_resolved + [
                c for c in work.columns if c not in reserved and c not in id_resolved
            ]
            id_resolved = [c for c in id_resolved if c in work.columns and c not in reserved]

            parsed_date = pd.to_datetime(work[date_name], errors="coerce").dt.normalize()
            valid_mask = parsed_date.notna()
            skipped = int((~valid_mask).sum())
            base = work.loc[valid_mask].copy()
            parsed_date = parsed_date.loc[valid_mask]

            if base.empty:
                return ToolResult(
                    success=False,
                    data=df,
                    message="No valid dates after parsing; aborting period expansion.",
                )

            daily_rows: List[Dict[str, Any]] = []
            for row_idx, (_, row) in enumerate(base.iterrows()):
                current_date = parsed_date.iloc[row_idx]
                period_start, period_end = TransformationTools._period_bounds_from_date(current_date, granularity)
                n_days = int((period_end - period_start).days) + 1
                id_vals = {c: row[c] for c in id_resolved}
                vals = {c: pd.to_numeric(row[c], errors="coerce") for c in resolved_values}
                per_day = {c: (float(vals[c]) / n_days) if pd.notna(vals[c]) else np.nan for c in resolved_values}
                for offset in range(n_days):
                    d = period_start + pd.Timedelta(days=offset)
                    daily_rows.append({**id_vals, date_column: d, **{c: per_day[c] for c in resolved_values}})

            if not daily_rows:
                return ToolResult(
                    success=False,
                    data=df,
                    message="No daily rows were generated from the input periods.",
                )

            daily_df = pd.DataFrame(daily_rows)
            period_label = "month" if granularity in {"monthly", "month"} else "quarter"
            msg_parts = [
                f"Expanded {len(base)} {period_label} period row(s) into {len(daily_df)} daily row(s).",
                f"Metrics: {', '.join(resolved_values)}.",
            ]
            if skipped:
                msg_parts.append(f"Skipped {skipped} row(s) with invalid dates.")
            if filter_notes:
                msg_parts.append("Filters: " + "; ".join(filter_notes))

            return ToolResult(
                success=True,
                data=daily_df,
                message=" ".join(msg_parts),
                changes_made={
                    "date_col": date_name,
                    "input_granularity": granularity,
                    "value_cols": resolved_values,
                    "id_cols": id_resolved,
                    "date_column": date_column,
                },
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in expand_period_to_daily: {e}")

    @staticmethod
    def _expand_week_scalar_rows_to_daily(
        work: pd.DataFrame,
        date_name: str,
        resolved_values: List[str],
        id_resolved: List[str],
        date_column: str,
    ) -> ToolResult:
        """One input row per week (any day in the week) → 7 equal daily slices (Monday–Sunday)."""
        parsed_date = pd.to_datetime(work[date_name], errors="coerce").dt.normalize()
        valid_mask = parsed_date.notna()
        base = work.loc[valid_mask].copy()
        parsed_date = parsed_date.loc[valid_mask]
        if base.empty:
            return ToolResult(success=False, data=work, message="No valid dates for weekly→daily expansion.")

        daily_rows: List[Dict[str, Any]] = []
        skipped = int((~valid_mask).sum())
        for row_idx, (_, row) in enumerate(base.iterrows()):
            anchor = parsed_date.iloc[row_idx]
            monday = TransformationTools._monday_week_start(anchor)
            id_vals = {c: row[c] for c in id_resolved}
            vals = {c: pd.to_numeric(row[c], errors="coerce") for c in resolved_values}
            per_day = {c: (float(vals[c]) / 7.0) if pd.notna(vals[c]) else np.nan for c in resolved_values}
            for offset in range(7):
                d = monday + pd.Timedelta(days=offset)
                daily_rows.append({**id_vals, date_column: d, **{c: per_day[c] for c in resolved_values}})

        daily_df = pd.DataFrame(daily_rows)
        msg = (
            f"Expanded {len(base)} weekly-scoped row(s) into {len(daily_df)} daily row(s) "
            f"(equal split across Monday–Sunday). Metrics: {', '.join(resolved_values)}."
        )
        if skipped:
            msg += f" Skipped {skipped} row(s) with invalid dates."
        return ToolResult(
            success=True,
            data=daily_df,
            message=msg,
            changes_made={
                "date_col": date_name,
                "input_granularity": "weekly_scalar",
                "value_cols": resolved_values,
                "id_cols": id_resolved,
                "date_column": date_column,
            },
        )

    @staticmethod
    def infer_granularity_expand_to_daily(
        df: pd.DataFrame,
        date_col: Union[str, int],
        value_cols: List[Any],
        id_cols: Optional[List[Any]] = None,
        date_column: str = "calendar_date",
        min_confidence: float = 0.35,
        row_filters: Optional[List[Dict[str, Any]]] = None,
    ) -> ToolResult:
        """
        Infer coarse cadence from ``date_col`` values, then emit daily rows when needed.

        Monthly/quarterly reuse ``expand_period_to_daily`` (calendar bounds).
        Weekly assumes one row per ISO week (any anchor day) and prorates /7 across Mon–Sun.
        Daily and irregular (sporadic) rows are kept as one row per date with no calendar expansion.
        """
        try:
            if df is None or df.empty:
                return ToolResult(success=True, data=df, message="Empty input; nothing to expand.")

            work, filter_notes = TransformationTools._apply_simple_row_filters(df, row_filters)
            if work.empty:
                return ToolResult(
                    success=False,
                    data=df,
                    message="All rows removed by row_filters; nothing to expand."
                    + (f" Notes: {'; '.join(filter_notes)}" if filter_notes else ""),
                )

            date_name, _ = TransformationTools._fuzzy_find_column(work, date_col)
            resolved_values, _ = TransformationTools._resolve_columns(work, value_cols)
            if not date_name or date_name not in work.columns:
                return ToolResult(success=False, data=df, message=f"date_col '{date_col}' not found.")
            if not resolved_values:
                return ToolResult(success=False, data=df, message="value_cols resolved to empty list.")

            if id_cols is not None:
                id_resolved, _ = TransformationTools._resolve_columns(work, id_cols)
            else:
                id_resolved = []
            reserved = {date_name, *resolved_values}
            id_resolved = id_resolved + [
                c for c in work.columns if c not in reserved and c not in id_resolved
            ]
            id_resolved = [c for c in id_resolved if c in work.columns and c not in reserved]

            profile = infer_date_cadence_fast(work[date_name], date_name)
            cadence = str(profile.get("inferred_cadence") or "").strip().lower()
            conf = float(profile.get("confidence") or 0.0)
            method = str(profile.get("inference_method") or "")

            meta = (
                f"inferred_cadence={cadence!r} confidence={round(conf, 3)} method={method!r} "
                f"median_gap_days={profile.get('median_gap_days')}"
            )

            if (
                str(profile.get("inference_method") or "") == "row_heuristic"
                and profile.get("likely_range_string_cells")
                and float(profile.get("range_parse_rate") or 0) >= 0.5
            ):
                return ToolResult(
                    success=False,
                    data=work,
                    message=(
                        "Date column looks like start–end range text. Use `transform.date_range_to_weekly` "
                        f"(or daily granularity) instead of infer+expand. ({meta})"
                    ),
                )

            if cadence in ("insufficient_data", "unparsed_text", ""):
                return ToolResult(
                    success=False,
                    data=work,
                    message=f"Could not infer date cadence from column {date_name!r}. ({meta})",
                )

            if cadence == "irregular" and conf < float(min_confidence or 0.35):
                return ToolResult(
                    success=False,
                    data=work,
                    message=(
                        f"Inferred cadence is irregular with low confidence ({conf:.2f} < {min_confidence}). "
                        f"Set Guided Setup date_granularity or call expand_period_to_daily with explicit "
                        f"input_granularity. ({meta})"
                    ),
                )

            if cadence in ("daily", "irregular"):
                label = "sporadic daily" if cadence == "irregular" else "daily"
                msg = f"Column {date_name!r} treated as {label}; no calendar expansion. ({meta})"
                if filter_notes:
                    msg += " Filters: " + "; ".join(filter_notes)
                return ToolResult(
                    success=True,
                    data=work.copy(),
                    message=msg,
                    changes_made={
                        "inferred_cadence": cadence,
                        "treated_as": "daily",
                        "inference_method": method,
                        "expanded": False,
                    },
                )

            if conf < float(min_confidence or 0.35) and cadence not in ("monthly", "quarterly", "weekly"):
                return ToolResult(
                    success=False,
                    data=work,
                    message=f"Inference confidence too low to expand ({conf:.2f} < {min_confidence}). ({meta})",
                )

            if cadence in ("monthly", "quarterly"):
                ig = "quarterly" if cadence == "quarterly" else "monthly"
                inner = TransformationTools.expand_period_to_daily(
                    work,
                    date_name,
                    resolved_values,
                    ig,
                    id_cols=id_resolved,
                    date_column=date_column,
                    row_filters=None,
                )
                if not inner.success:
                    return inner
                msg = inner.message + f" [{meta}]"
                cm = dict(inner.changes_made or {})
                cm["inferred_cadence"] = cadence
                cm["inference_method"] = method
                return ToolResult(success=True, data=inner.data, message=msg, changes_made=cm)

            if cadence == "weekly":
                if conf < float(min_confidence or 0.35):
                    return ToolResult(
                        success=False,
                        data=work,
                        message=(
                            f"Weekly cadence inferred with low confidence ({conf:.2f} < {min_confidence}); "
                            f"refuse automatic /7 split. ({meta})"
                        ),
                    )
                inner = TransformationTools._expand_week_scalar_rows_to_daily(
                    work, date_name, resolved_values, id_resolved, date_column
                )
                if not inner.success:
                    return inner
                msg = inner.message + f" [{meta}]"
                cm = dict(inner.changes_made or {})
                cm["inferred_cadence"] = cadence
                cm["inference_method"] = method
                return ToolResult(success=True, data=inner.data, message=msg, changes_made=cm)

            return ToolResult(
                success=False,
                data=work,
                message=(
                    f"Inferred cadence is {cadence!r}; automatic expansion is not supported. "
                    f"Use explicit `transform.expand_period_to_daily` or `transform.date_range_to_weekly`. ({meta})"
                ),
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in infer_granularity_expand_to_daily: {e}")

    @staticmethod
    def date_range_to_weekly(
        df: pd.DataFrame,
        start_date_col: Union[str, int],
        end_date_col: Union[str, int],
        value_cols: List[Any],
        id_cols: Optional[List[Any]] = None,
        granularity: str = "weekly",
        week_start_col: str = "week_start",
        week_end_col: Optional[str] = None,
        days_in_period_col: str = "days_in_week",
        date_column: str = "calendar_date",
        row_filters: Optional[List[Dict[str, Any]]] = None,
    ) -> ToolResult:
        """
        Prorate each numeric metric evenly across calendar days in [start, end], then:
        - granularity ``weekly``: sum daily slices into Monday–Sunday ISO weeks.
        - granularity ``daily``: one row per calendar day with prorated values.
        """
        try:
            if df is None or df.empty:
                return ToolResult(success=True, data=df, message="Empty input; nothing to expand.")
            g = str(granularity or "weekly").strip().lower()
            if g not in ("weekly", "daily"):
                return ToolResult(success=False, data=df, message=f"granularity must be 'weekly' or 'daily', got {granularity!r}")

            work, filter_notes = TransformationTools._apply_simple_row_filters(df, row_filters)
            if work.empty:
                return ToolResult(
                    success=False,
                    data=df,
                    message="All rows removed by row_filters; nothing to aggregate." + (
                        f" Notes: {'; '.join(filter_notes)}" if filter_notes else ""
                    ),
                )

            start_name, _ = TransformationTools._fuzzy_find_column(work, start_date_col)
            end_name, _ = TransformationTools._fuzzy_find_column(work, end_date_col)
            if start_name == end_name:
                return ToolResult(
                    success=False,
                    data=df,
                    message=(
                        f"start_date_col and end_date_col both resolved to the same column {start_name!r}. "
                        "Use two distinct columns (e.g. Start Date and End Date). "
                        "If you only aggregate on one date, the full range total is summed into one week — "
                        "use transform.date_range_to_weekly with both columns first."
                    ),
                )
            resolved_values, _ = TransformationTools._resolve_columns(work, value_cols)
            if not resolved_values:
                return ToolResult(success=False, data=df, message="value_cols resolved to empty list.")

            if id_cols is not None:
                id_resolved, _ = TransformationTools._resolve_columns(work, id_cols)
            else:
                reserved = {start_name, end_name, *resolved_values}
                id_resolved = [c for c in work.columns if c not in reserved]

            st = pd.to_datetime(work[start_name], errors="coerce").dt.normalize()
            en = pd.to_datetime(work[end_name], errors="coerce").dt.normalize()
            valid = st.notna() & en.notna() & (en >= st)
            skipped = int((~valid).sum())
            base = work.loc[valid].copy()
            base["_rs"] = st[valid]
            base["_re"] = en[valid]

            # Mis-picked start/end (e.g. numeric metrics → epoch vs real dates) can create huge day spans.
            # Cap per-row expansion so bad tool args surface as skips + log instead of runaway row counts.
            _max_inclusive_day_span = 10000  # ~27 years per source row

            daily_rows: List[Dict[str, Any]] = []
            skipped_long_span = 0
            for _, row in base.iterrows():
                s: pd.Timestamp = row["_rs"]
                e: pd.Timestamp = row["_re"]
                n_days = int((e - s).days) + 1
                if n_days <= 0:
                    continue
                if n_days > _max_inclusive_day_span:
                    skipped_long_span += 1
                    logger.warning(
                        "date_range_to_weekly: skipping row with inclusive span %s days (cap %s). "
                        "Often means start_date_col/end_date_col are wrong — single calendar columns should use "
                        "transform.aggregate_weekly, not start/end range expansion.",
                        n_days,
                        _max_inclusive_day_span,
                    )
                    continue
                id_vals = {c: row[c] for c in id_resolved}
                vals = {c: pd.to_numeric(row[c], errors="coerce") for c in resolved_values}
                per_day = {c: (float(vals[c]) / n_days) if pd.notna(vals[c]) else np.nan for c in resolved_values}
                for offset in range(n_days):
                    d = s + pd.Timedelta(days=offset)
                    rec = {**id_vals, date_column: d, **{c: per_day[c] for c in resolved_values}}
                    daily_rows.append(rec)

            if not daily_rows:
                detail = f"No daily rows built (invalid dates?). Skipped rows: {skipped}."
                if skipped_long_span:
                    detail += f" Skipped {skipped_long_span} row(s) with day-span > {_max_inclusive_day_span} (likely wrong start/end columns)."
                return ToolResult(
                    success=False,
                    data=df,
                    message=detail,
                )

            daily_df = pd.DataFrame(daily_rows)

            if g == "daily":
                msg = (
                    f"Expanded to {len(daily_df)} daily rows from {len(base)} source rows "
                    f"(equal split per day). granularity=daily."
                )
                if skipped:
                    msg += f" Skipped {skipped} rows with invalid date ranges."
                if skipped_long_span:
                    msg += (
                        f" Skipped {skipped_long_span} row(s) with span > {_max_inclusive_day_span} days "
                        f"(check start/end columns; use aggregate_weekly for a single date column)."
                    )
                if filter_notes:
                    msg += " Filters: " + "; ".join(filter_notes)
                return ToolResult(success=True, data=daily_df, message=msg)

            ws = daily_df[date_column].map(TransformationTools._monday_week_start)
            daily_df[week_start_col] = ws
            gcols = id_resolved + [week_start_col]
            week_end_name = str(week_end_col).strip() if week_end_col is not None else ""
            if week_end_name:
                daily_df[week_end_name] = ws + pd.Timedelta(days=6)
                gcols.append(week_end_name)
            grouped = daily_df.groupby(gcols, dropna=False, as_index=False).agg(
                {**{c: "sum" for c in resolved_values}, date_column: "count"}
            )
            grouped = grouped.rename(columns={date_column: days_in_period_col})

            msg = (
                f"Weekly buckets: {len(grouped)} rows from {len(base)} source rows "
                f"(Monday-aligned weekly dates, equal daily proration)."
            )
            if skipped:
                msg += f" Skipped {skipped} rows with invalid date ranges."
            if skipped_long_span:
                msg += (
                    f" Skipped {skipped_long_span} row(s) with span > {_max_inclusive_day_span} days "
                    f"(check start/end columns; use aggregate_weekly for a single date column)."
                )
            if filter_notes:
                msg += " Filters: " + "; ".join(filter_notes)
            return ToolResult(success=True, data=grouped, message=msg)
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in date_range_to_weekly: {e}")

    @staticmethod
    def expand_date_range_to_daily(
        df: pd.DataFrame,
        start_date_col: Union[str, int],
        end_date_col: Union[str, int],
        value_cols: List[Any],
        id_cols: Optional[List[Any]] = None,
        week_start_col: str = "week_start",
        week_end_col: Optional[str] = None,
        days_in_period_col: str = "days_in_week",
        date_column: str = "calendar_date",
        row_filters: Optional[List[Dict[str, Any]]] = None,
    ) -> ToolResult:
        """Inclusive start/end → one row per calendar day (metrics prorated); then use aggregate_weekly if needed."""
        return TransformationTools.date_range_to_weekly(
            df,
            start_date_col,
            end_date_col,
            value_cols,
            id_cols=id_cols,
            granularity="daily",
            week_start_col=week_start_col,
            week_end_col=week_end_col,
            days_in_period_col=days_in_period_col,
            date_column=date_column,
            row_filters=row_filters,
        )

    @staticmethod
    def expand_date_range_to_weekly(
        df: pd.DataFrame,
        start_date_col: Union[str, int],
        end_date_col: Union[str, int],
        value_cols: List[Any],
        id_cols: Optional[List[Any]] = None,
        week_start_col: str = "week_start",
        week_end_col: Optional[str] = None,
        days_in_period_col: str = "days_in_week",
        date_column: str = "calendar_date",
        row_filters: Optional[List[Dict[str, Any]]] = None,
    ) -> ToolResult:
        """Inclusive start/end → weekly buckets (internal daily split)."""
        return TransformationTools.date_range_to_weekly(
            df,
            start_date_col,
            end_date_col,
            value_cols,
            id_cols=id_cols,
            granularity="weekly",
            week_start_col=week_start_col,
            week_end_col=week_end_col,
            days_in_period_col=days_in_period_col,
            date_column=date_column,
            row_filters=row_filters,
        )

    @staticmethod
    def aggregate_weekly(
        df: pd.DataFrame,
        date_col: Union[str, int],
        group_by_cols: Optional[List[Any]] = None,
        metric_rules: Optional[Dict[str, str]] = None,
        value_cols: Optional[List[Any]] = None,
        week_start_col: Optional[str] = None,
        week_end_col: Optional[str] = None,
        drop_original_date: bool = True,
        row_filters: Optional[List[Dict[str, Any]]] = None,
    ) -> ToolResult:
        """Aggregate daily/date-grain rows into Monday–Sunday ISO weekly buckets.

        Use this tool when the dataframe already has a single date column (not a
        start/end range) and needs weekly rollup. For start/end range inputs use
        ``date_range_to_weekly`` instead.

        Args:
            date_col: Name or positional index of the date column to bucket.
            group_by_cols: Columns to group by in addition to the week bucket.
                If not provided, every non-date, non-metric column is used.
            metric_rules: Mapping of metric column to aggregation rule. Supported
                rules: sum, mean/avg, min, max, first, last, count. Defaults to
                "sum" for each ``value_cols`` entry.
            value_cols: Metric columns to aggregate when ``metric_rules`` is not
                supplied. Ignored if ``metric_rules`` is provided.
            week_start_col: Output column for the Monday week-start date. When
                omitted (or empty), uses the resolved ``date_col`` name if
                ``drop_original_date`` is True so the template date column (e.g.
                ``calendar_date``) is preserved; otherwise defaults to ``week_start``.
                When set to the literal ``"week_start"`` while ``drop_original_date``
                is True and ``date_col`` resolves to any other name, the same
                single-column behaviour applies (legacy plans passed this default
                even though the source had no ``week_start`` column).
            week_end_col: Optional output column for the Sunday week-end date.
            drop_original_date: If True, remove the source date column from the
                grouping output (the weekly date column replaces it).
            row_filters: Optional simple filter list applied before aggregation.
        """
        try:
            if df is None or df.empty:
                return ToolResult(success=True, data=df, message="Empty input; nothing to aggregate.")

            work, filter_notes = TransformationTools._apply_simple_row_filters(df, row_filters)
            if work.empty:
                return ToolResult(
                    success=False,
                    data=df,
                    message="All rows removed by row_filters; nothing to aggregate." + (
                        f" Notes: {'; '.join(filter_notes)}" if filter_notes else ""
                    ),
                )

            date_name, _ = TransformationTools._fuzzy_find_column(work, date_col)
            if not date_name or date_name not in work.columns:
                return ToolResult(
                    success=False,
                    data=df,
                    message=f"date_col '{date_col}' not found in dataframe columns.",
                )

            ws_arg = week_start_col
            if ws_arg is None or (isinstance(ws_arg, str) and not str(ws_arg).strip()):
                week_bucket_col = date_name if drop_original_date else "week_start"
            else:
                ws_str = str(ws_arg).strip()
                # Many plans pass week_start_col="week_start" even when the only date
                # column is e.g. calendar_date. Reuse the input date column name when
                # we already drop the raw daily values, so the output keeps one date key.
                if (
                    drop_original_date
                    and ws_str.lower() == "week_start"
                    and date_name
                    and str(date_name).strip().lower() != "week_start"
                ):
                    week_bucket_col = date_name
                else:
                    week_bucket_col = ws_str

            try:
                from sia.agent.date_column_inference import infer_date_column_profile

                _prof = infer_date_column_profile(work[date_name], date_name)
                logger.info(
                    "[aggregate_weekly] input_rows=%s date_col=%s inferred_cadence=%s median_gap_days=%s "
                    "median_window_days=%s confidence=%s parse_rate=%s range_parse_rate=%s",
                    len(work),
                    date_name,
                    _prof.get("inferred_cadence"),
                    _prof.get("median_gap_days"),
                    _prof.get("median_window_days"),
                    round(float(_prof.get("confidence") or 0), 3),
                    _prof.get("parse_rate"),
                    _prof.get("range_parse_rate"),
                )
            except Exception as _exc:
                logger.debug("aggregate_weekly: date cadence probe skipped: %s", _exc)

            rules: Dict[str, str] = {}
            if isinstance(metric_rules, dict) and metric_rules:
                for raw_col, raw_rule in metric_rules.items():
                    resolved_name, _ = TransformationTools._fuzzy_find_column(work, raw_col)
                    if not resolved_name or resolved_name not in work.columns:
                        continue
                    rule = str(raw_rule or "sum").strip().lower()
                    if rule in ("avg", "average"):
                        rule = "mean"
                    if rule not in {"sum", "mean", "min", "max", "first", "last", "count"}:
                        rule = "sum"
                    rules[resolved_name] = rule
            elif value_cols:
                resolved_values, _ = TransformationTools._resolve_columns(work, value_cols)
                for resolved_name in resolved_values:
                    rules[resolved_name] = "sum"

            if not rules:
                inferred_metrics: List[str] = []
                for c in work.columns:
                    if c == date_name:
                        continue
                    numeric = pd.to_numeric(work[c], errors="coerce")
                    non_null = numeric.notna().sum()
                    if non_null > 0 and non_null >= max(1, int(len(work) * 0.5)):
                        inferred_metrics.append(c)
                for c in inferred_metrics:
                    rules[c] = "sum"

            if not rules:
                return ToolResult(
                    success=False,
                    data=df,
                    message="No metrics resolved for weekly aggregation. Provide metric_rules or value_cols.",
                )

            if group_by_cols is not None:
                group_resolved, _ = TransformationTools._resolve_columns(work, group_by_cols)
            else:
                reserved = {date_name, *rules.keys()}
                group_resolved = [c for c in work.columns if c not in reserved]

            group_resolved = [c for c in group_resolved if c != date_name and c not in rules]

            parsed_date = pd.to_datetime(work[date_name], errors="coerce").dt.normalize()
            valid_mask = parsed_date.notna()
            skipped = int((~valid_mask).sum())
            work = work.loc[valid_mask].copy()
            parsed_date = parsed_date.loc[valid_mask]

            if work.empty:
                return ToolResult(
                    success=False,
                    data=df,
                    message="No valid dates after parsing; aborting weekly aggregation.",
                )

            week_starts = parsed_date.map(TransformationTools._monday_week_start)
            work[week_bucket_col] = week_starts
            group_cols = group_resolved + [week_bucket_col]
            week_end_name = str(week_end_col).strip() if week_end_col is not None else ""
            if week_end_name:
                work[week_end_name] = week_starts + pd.Timedelta(days=6)
                group_cols.append(week_end_name)

            for metric_col, rule in rules.items():
                if rule in {"sum", "mean", "min", "max"}:
                    work[metric_col] = pd.to_numeric(work[metric_col], errors="coerce")

            agg_spec: Dict[str, Any] = {col: rule for col, rule in rules.items()}

            grouped = work.groupby(group_cols, dropna=False, as_index=False).agg(agg_spec)

            if not drop_original_date and date_name in work.columns:
                first_date = work.groupby(group_cols, dropna=False, as_index=False)[date_name].min()
                grouped = grouped.merge(first_date, on=group_cols, how="left")

            msg_parts = [
                f"Weekly aggregation: {len(grouped)} rows from {len(work)} input rows.",
                f"Metrics: {', '.join(f'{m}({r})' for m, r in rules.items())}.",
            ]
            if group_resolved:
                msg_parts.append(f"Grouped by: {', '.join(group_resolved)}.")
            if skipped:
                msg_parts.append(f"Skipped {skipped} rows with unparseable dates.")
            if filter_notes:
                msg_parts.append("Filters: " + "; ".join(filter_notes))

            return ToolResult(
                success=True,
                data=grouped,
                message=" ".join(msg_parts),
                changes_made={
                    "date_col": date_name,
                    "group_by_cols": group_resolved,
                    "metric_rules": rules,
                },
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in aggregate_weekly: {e}")

    @staticmethod
    def scale_columns(
        df: pd.DataFrame,
        scales: Optional[Dict[str, float]] = None,
        columns: Optional[Dict[str, float]] = None,
    ) -> ToolResult:
        """
        Multiply numeric columns by denomination factors (e.g. thousands → ×1000).

        Args:
            df: DataFrame to modify
            scales: {column_name: factor} after rename
            columns: alias for scales (validator compatibility)
        """
        try:
            scale_map = dict(scales or columns or {})
            if not scale_map:
                return ToolResult(success=True, data=df, message="No scale factors provided")

            result_df = df.copy()
            applied: List[str] = []
            for col, raw_factor in scale_map.items():
                try:
                    factor = float(raw_factor)
                except (TypeError, ValueError):
                    continue
                if factor in (0.0, 1.0):
                    continue
                try:
                    resolved_col, _ = TransformationTools._fuzzy_find_column(result_df, str(col))
                except ValueError:
                    applied.append(f"SKIP '{col}' (not found)")
                    continue
                numeric = pd.to_numeric(result_df[resolved_col], errors="coerce")
                if numeric.notna().sum() == 0:
                    applied.append(f"SKIP '{resolved_col}' (non-numeric)")
                    continue
                result_df[resolved_col] = numeric * factor
                applied.append(f"{resolved_col}×{factor:g}")

            if not applied:
                return ToolResult(success=True, data=result_df, message="No columns scaled")
            return ToolResult(
                success=True,
                data=result_df,
                message=f"Scaled columns: {', '.join(applied)}",
                changes_made={"scaled_columns": applied, "scales": scale_map},
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in scale_columns: {e}")

    @staticmethod
    def format_columns(df: pd.DataFrame,
                       format_map: Dict[str, str]) -> ToolResult:
        """
        Apply specific formatting to columns.

        Supported format types:
            - "date:YYYY-MM-DD" -> Converts to datetime then formats as string
            - "numeric:2" -> Rounds to 2 decimal places
            - "integer" -> Converts to integer (drops decimals)
            - "lowercase" -> Converts string to lowercase
            - "uppercase" -> Converts string to uppercase

        Args:
            df: DataFrame to modify
            format_map: Dict of {column_name: format_spec}

        Returns:
            ToolResult with formatted DataFrame
        """
        try:
            result_df = df.copy()
            changes = []

            for col, fmt in format_map.items():
                # Resolve column name via fuzzy matching
                try:
                    resolved_col, _ = TransformationTools._fuzzy_find_column(result_df, col)
                except ValueError:
                    changes.append(f"SKIP '{col}' (not found)")
                    continue
                col = resolved_col

                if fmt.startswith("date:"):
                    date_fmt = fmt.split(":", 1)[1]
                    # Map common shortcuts
                    fmt_map = {"YYYY-MM-DD": "%Y-%m-%d", "DD/MM/YYYY": "%d/%m/%Y",
                               "MM/DD/YYYY": "%m/%d/%Y", "YYYY": "%Y"}
                    py_fmt = fmt_map.get(date_fmt, date_fmt)
                    result_df[col] = pd.to_datetime(result_df[col], errors='coerce').dt.strftime(py_fmt)
                    changes.append(f"'{col}' -> date({date_fmt})")

                elif fmt.startswith("numeric:"):
                    decimals = int(fmt.split(":", 1)[1])
                    result_df[col] = pd.to_numeric(result_df[col], errors='coerce').round(decimals)
                    changes.append(f"'{col}' -> numeric({decimals} dp)")

                elif fmt == "integer":
                    result_df[col] = pd.to_numeric(result_df[col], errors='coerce')
                    result_df[col] = result_df[col].where(result_df[col].isna(), result_df[col].astype('Int64'))
                    changes.append(f"'{col}' -> integer")

                elif fmt == "lowercase":
                    result_df[col] = result_df[col].astype(str).str.lower()
                    changes.append(f"'{col}' -> lowercase")

                elif fmt == "uppercase":
                    result_df[col] = result_df[col].astype(str).str.upper()
                    changes.append(f"'{col}' -> uppercase")

                else:
                    changes.append(f"SKIP '{col}' (unknown format: {fmt})")

            return ToolResult(
                success=True,
                data=result_df,
                message=f"Formatted {len(changes)} columns: {'; '.join(changes)}"
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in format_columns: {e}")

    @staticmethod
    def validate_against_schema(df: pd.DataFrame,
                                schema: Dict[str, Any]) -> ToolResult:
        """
        Validate a DataFrame against a JSON Schema (draft-07 compatible).

        Checks:
        1. Mandatory columns present
        2. Data types match (string, number, integer, date)
        3. Enum constraints satisfied
        4. Minimum/maximum value constraints

        Args:
            df: DataFrame to validate
            schema: JSON Schema dict (with 'properties' and 'x_requirement' extensions)

        Returns:
            ToolResult with validation report (data = original df, message = report)
        """
        try:
            properties = schema.get("properties", {})
            issues = []
            passed = []

            for col_key, col_spec in properties.items():
                requirement = col_spec.get("x_requirement", "mandatory")
                col_type = col_spec.get("type", "string")
                description = col_spec.get("description", "")

                # 1. Presence check via fuzzy matching
                try:
                    resolved_col, _ = TransformationTools._fuzzy_find_column(df, col_key)
                except ValueError:
                    if requirement == "mandatory":
                        issues.append({
                            "column": col_key,
                            "severity": "HIGH",
                            "issue": "MISSING_MANDATORY",
                            "detail": f"Mandatory column '{col_key}' not found. Description: {description}"
                        })
                    continue

                series = df[resolved_col]
                passed.append(resolved_col)

                # 2. Null check for mandatory
                if requirement == "mandatory":
                    null_count = series.isna().sum()
                    null_pct = null_count / len(series) * 100 if len(series) > 0 else 0
                    if null_pct > 5:
                        issues.append({
                            "column": col_key,
                            "severity": "MEDIUM",
                            "issue": "HIGH_NULL_RATE",
                            "detail": f"{null_count} nulls ({null_pct:.1f}%) in mandatory column"
                        })

                # 3. Type check
                if isinstance(col_type, str):
                    if col_type == "number" or col_type == "integer":
                        numeric_series = pd.to_numeric(series, errors='coerce')
                        non_numeric = series.notna().sum() - numeric_series.notna().sum()
                        if non_numeric > 0:
                            issues.append({
                                "column": col_key,
                                "severity": "MEDIUM",
                                "issue": "TYPE_MISMATCH",
                                "detail": f"{non_numeric} values are not numeric"
                            })

                    if col_type == "string" and col_spec.get("format") == "date":
                        date_series = pd.to_datetime(series, errors='coerce')
                        non_date = series.notna().sum() - date_series.notna().sum()
                        if non_date > 0:
                            issues.append({
                                "column": col_key,
                                "severity": "MEDIUM",
                                "issue": "DATE_FORMAT",
                                "detail": f"{non_date} values do not parse as dates"
                            })

                # 4. Enum check
                if "enum" in col_spec:
                    allowed = set(col_spec["enum"])
                    actual = set(series.dropna().unique())
                    invalid = actual - allowed
                    if invalid:
                        issues.append({
                            "column": col_key,
                            "severity": "MEDIUM",
                            "issue": "ENUM_VIOLATION",
                            "detail": f"Invalid values: {list(invalid)[:5]}. Allowed: {list(allowed)}"
                        })

                # 5. Min/Max check (MEDIUM — surfaced as Review-style HITL)
                if "minimum" in col_spec or "maximum" in col_spec:
                    numeric_series = pd.to_numeric(series, errors='coerce')
                    if "minimum" in col_spec:
                        bound = col_spec["minimum"]
                        below_mask = numeric_series.notna() & (numeric_series < bound)
                        below = int(below_mask.sum())
                        if below > 0:
                            samples = [
                                float(v) if pd.notna(v) else None
                                for v in numeric_series[below_mask].head(5).tolist()
                            ]
                            issues.append({
                                "column": col_key,
                                "severity": "MEDIUM",
                                "issue": "MIN_VIOLATION",
                                "bound": bound,
                                "minimum": bound,
                                "count": below,
                                "violation_count": below,
                                "sample_values": samples,
                                "detail": f"{below} values below minimum ({bound})",
                            })
                    if "maximum" in col_spec:
                        bound = col_spec["maximum"]
                        above_mask = numeric_series.notna() & (numeric_series > bound)
                        above = int(above_mask.sum())
                        if above > 0:
                            samples = [
                                float(v) if pd.notna(v) else None
                                for v in numeric_series[above_mask].head(5).tolist()
                            ]
                            issues.append({
                                "column": col_key,
                                "severity": "MEDIUM",
                                "issue": "MAX_VIOLATION",
                                "bound": bound,
                                "maximum": bound,
                                "count": above,
                                "violation_count": above,
                                "sample_values": samples,
                                "detail": f"{above} values above maximum ({bound})",
                            })

            # Build report
            high_count = sum(1 for i in issues if i["severity"] == "HIGH")
            med_count = sum(1 for i in issues if i["severity"] == "MEDIUM")
            low_count = sum(1 for i in issues if i["severity"] == "LOW")
            constraint_count = sum(
                1 for i in issues if i.get("issue") in ("MIN_VIOLATION", "MAX_VIOLATION")
            )
            # Constraint breaches fail validation (Review HITL pauses the job).
            is_valid = high_count == 0 and constraint_count == 0

            report = {
                "valid": is_valid,
                "columns_matched": len(passed),
                "columns_expected": len([
                    p for p, s in properties.items()
                    if (s.get("x_requirement", "mandatory") == "mandatory")
                ]),
                "issues": issues,
                "constraint_violation_count": constraint_count,
                "summary": f"{'PASS' if is_valid else 'FAIL'}: "
                           f"{high_count} HIGH, {med_count} MEDIUM"
                           + (f", {low_count} LOW" if low_count else "")
                           + " issues"
                           + (f" ({constraint_count} constraint)" if constraint_count else ""),
            }

            import json
            return ToolResult(
                success=True,
                data=df,
                message=json.dumps(report, indent=2),
                changes_made={"validation_report": report}
            )
        except Exception as e:
            return ToolResult(success=False, data=df, message=f"Error in validate_against_schema: {e}")


def execute_tool(tool_name: str, **params) -> ToolResult:
    """
    Execute a transformation tool by name.
    
    Args:
        tool_name: Name of the tool to execute
        **params: Parameters for the tool
        
    Returns:
        ToolResult from the tool execution
    """
    tools = TransformationTools()
    
    if hasattr(tools, tool_name):
        tool_func = getattr(tools, tool_name)
        return tool_func(**params)
    else:
        return ToolResult(
            success=False,
            data=None,
            message=f"Unknown tool: {tool_name}"
        )


def get_tools_description() -> str:
    """Get a formatted description of all available tools for the LLM prompt."""
    lines = ["## Available Transformation Tools\n"]
    
    for tool in AVAILABLE_TOOLS:
        lines.append(f"### {tool['name']}")
        lines.append(f"{tool['description']}")
        lines.append("\n**Parameters:**")
        for param, desc in tool['params'].items():
            lines.append(f"- `{param}`: {desc}")
        lines.append("")
    
    return "\n".join(lines)


# ===== Destructive Tool Identification =====

DESTRUCTIVE_TOOLS = {
    # Python method names
    "skip_rows",
    "filter_summary_rows",
    "filter_header_rows",
    "filter_empty_rows",
    "drop_blank_columns",
    # MCP-style names (used in plan tool_calls)
    "xls.data.skip_rows",
    "xls.data.filter_summaries",
    "xls.data.filter_empty",
    "xls.data.drop_blank_columns",
}


def is_tool_destructive(tool_name: str) -> bool:
    """Check if a tool is destructive (modifies/deletes data).
    Handles MCP names, Python method names, and planner canonical names (transform.*).
    """
    if not tool_name:
        return False
    if tool_name in DESTRUCTIVE_TOOLS:
        return True
    short_name = tool_name.rsplit(".", 1)[-1] if "." in tool_name else tool_name
    if short_name in DESTRUCTIVE_TOOLS:
        return True
    try:
        from sia.tools.tool_validator import TOOL_SCHEMAS, normalize_tool_name

        normalized, _ = normalize_tool_name(tool_name)
        schema = TOOL_SCHEMAS.get(normalized)
        if schema is not None:
            return bool(schema.destructive)
    except Exception:
        pass
    return False


def generate_deletion_preview(df: pd.DataFrame, tool_call: Dict[str, Any]) -> Optional[DeletionPreview]:
    """
    Generate a preview of what a destructive tool would delete.
    
    Args:
        df: Current DataFrame
        tool_call: Tool call dict with 'tool' and 'params'
        
    Returns:
        DeletionPreview if tool is destructive and supports preview, else None
    """
    tool_name = tool_call.get("tool", "")
    params = tool_call.get("params", {})
    
    # Normalize MCP names (e.g. 'xls.data.drop_blank_columns') to Python names ('drop_blank_columns')
    short_name = tool_name.rsplit(".", 1)[-1] if "." in tool_name else tool_name
    
    if not is_tool_destructive(tool_name):
        return None
    
    if df is None or df.empty:
        return DeletionPreview(
            tool_name=tool_name,
            tool_description=f"Cannot preview: DataFrame is empty",
            impact_summary="No data to preview"
        )
    
    # Call the tool with preview_only=True
    try:
        if short_name == "skip_rows":
            return TransformationTools.skip_rows(df, preview_only=True, **params)
        elif short_name == "filter_summary_rows" or short_name == "filter_summaries":
            return TransformationTools.filter_summary_rows(df, preview_only=True, **params)
        elif short_name == "filter_empty_rows" or short_name == "filter_empty":
            return TransformationTools.filter_empty_rows(df, preview_only=True, **params)
        elif short_name == "filter_header_rows":
            return TransformationTools.filter_header_rows(df, preview_only=True, **params)
        elif short_name == "drop_blank_columns":
            return TransformationTools.drop_blank_columns(df, preview_only=True, **params)
        else:
            return None
    except Exception as e:
        logger.error(f"Error generating preview for {tool_name}: {e}")
        return DeletionPreview(
            tool_name=tool_name,
            tool_description=f"Error generating preview: {e}",
            impact_summary="Preview failed"
        )


def generate_all_previews(df: pd.DataFrame, tool_calls: List[Dict[str, Any]]) -> List[DeletionPreview]:
    """
    Generate previews for all destructive tools in a tool call sequence.
    
    Args:
        df: Current DataFrame
        tool_calls: List of tool call dicts
        
    Returns:
        List of DeletionPreview objects for destructive tools
    """
    previews = []
    for tool_call in tool_calls:
        if is_tool_destructive(tool_call.get("tool", "")):
            preview = generate_deletion_preview(df, tool_call)
            if preview:
                previews.append(preview)
    return previews


