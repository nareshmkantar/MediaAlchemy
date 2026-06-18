"""
Visual Normalizer Module - Converts Excel files to standardized VisualGrid.
Handles merged cells, hidden rows/columns, and cell formatting.
"""
import logging
from typing import Any, Callable, Collection, Dict, List, Mapping, Optional, Tuple
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.cell.cell import Cell as OpenpyxlCell
from openpyxl.utils import get_column_letter

from ..models.cell import Cell, CellStyle, CellType, VisualGrid, DataBlock
from ..models.confidence import ConfidenceResult

logger = logging.getLogger(__name__)


class VisualNormalizer:
    """
    Converts Excel workbooks into standardized VisualGrid format.
    This is the first step in the SIA pipeline.
    """
    
    def __init__(self, include_hidden: bool = False):
        """
        Initialize the normalizer.
        
        Args:
            include_hidden: Whether to include hidden rows/columns
        """
        self.include_hidden = include_hidden
    
    def normalize_workbook(
        self,
        file_path: str,
        sheet_dimension_hints: Optional[Mapping[str, Tuple[int, int]]] = None,
        *,
        on_progress: Optional[Callable[[Dict[str, Any]], None]] = None,
        only_sheet_names: Optional[Collection[str]] = None,
    ) -> Tuple[Dict[str, VisualGrid], ConfidenceResult]:
        """
        Load and normalize sheets in a workbook.

        Args:
            file_path: Path to Excel file
            sheet_dimension_hints: Optional map ``sheet_name -> (nrows, ncols)`` from
                ``pandas.read_excel(..., header=None).shape``. When openpyxl's
                ``max_row`` / ``max_column`` are stale (common with large tables),
                hints expand the scan rectangle so Guided Setup matches the file.
            only_sheet_names: When set (e.g. job ``source_registry`` tabs the user kept
                after upload), only those sheet tabs are normalized. Other tabs are not
                read (upload "remove" does not delete tabs from the file on disk).
        """
        try:
            wb = load_workbook(file_path, data_only=True)
            grids = {}
            sheet_names = list(wb.sheetnames)
            allow: Optional[set[str]] = None
            if only_sheet_names:
                allow = {str(s).strip() for s in only_sheet_names if str(s or "").strip()}
                if not allow:
                    allow = None
            to_process = sheet_names
            if allow is not None:
                to_process = [s for s in sheet_names if s in allow]
                for name in sorted(allow):
                    if name not in sheet_names:
                        logger.warning(
                            "Requested sheet %r not found in workbook %s (tabs: %s)",
                            name,
                            file_path,
                            sheet_names[:20],
                        )
            sheets_total = len(to_process)
            if on_progress:
                try:
                    on_progress(
                        {
                            "event": "workbook_start",
                            "file_path": str(file_path),
                            "sheets_total": sheets_total,
                            "workbook_tab_count": len(sheet_names),
                            "scoped_to_registry": bool(allow),
                        }
                    )
                except Exception:
                    pass

            for idx, sheet_name in enumerate(to_process, start=1):
                ws = wb[sheet_name]

                # Skip hidden sheets - only process visible sheets
                if ws.sheet_state != 'visible':
                    logger.info(f"Skipping hidden sheet: '{sheet_name}'")
                    if on_progress:
                        try:
                            on_progress(
                                {
                                    "event": "sheet_skip_hidden",
                                    "sheet_name": sheet_name,
                                    "sheet_index": idx,
                                    "sheets_total": sheets_total,
                                }
                            )
                        except Exception:
                            pass
                    continue

                if on_progress:
                    try:
                        on_progress(
                            {
                                "event": "sheet_normalizing",
                                "sheet_name": sheet_name,
                                "sheet_index": idx,
                                "sheets_total": sheets_total,
                            }
                        )
                    except Exception:
                        pass

                hint = None
                if sheet_dimension_hints and sheet_name in sheet_dimension_hints:
                    hint = tuple(sheet_dimension_hints[sheet_name])  # type: ignore[arg-type]
                grid = self._normalize_sheet(ws, sheet_name, dimension_hint=hint)
                grids[sheet_name] = grid
                logger.info(f"Normalized sheet '{sheet_name}': {grid.total_rows}x{grid.total_cols}")
                if on_progress:
                    try:
                        on_progress(
                            {
                                "event": "sheet_normalized",
                                "sheet_name": sheet_name,
                                "sheet_index": idx,
                                "sheets_total": sheets_total,
                                "rows": grid.total_rows,
                                "cols": grid.total_cols,
                            }
                        )
                    except Exception:
                        pass
            
            wb.close()
            
            confidence = ConfidenceResult(
                score=0.95,
                rationale="Successfully loaded and normalized workbook sheets in scope",
                signals=["file_loaded", "sheets_processed"],
                module_name="visual_normalizer"
            )
            
            return grids, confidence
            
        except Exception as e:
            logger.error(f"Failed to normalize workbook: {e}")
            confidence = ConfidenceResult(
                score=0.0,
                rationale=f"Failed to load file: {str(e)}",
                signals=["file_error"],
                module_name="visual_normalizer"
            )
            return {}, confidence
    
    def _normalize_sheet(
        self,
        ws: Worksheet,
        sheet_name: str,
        dimension_hint: Optional[Tuple[int, int]] = None,
    ) -> VisualGrid:
        """
        Convert a single worksheet to VisualGrid.
        
        Args:
            ws: openpyxl Worksheet object
            sheet_name: Name of the sheet
            dimension_hint: Optional ``(nrows, ncols)`` from pandas for this sheet.
            
        Returns:
            VisualGrid representation
        """
        # Get dimensions (openpyxl max_row/max_col can lag behind real data extent)
        max_row = int(ws.max_row or 1)
        max_col = int(ws.max_column or 1)
        if dimension_hint:
            try:
                pr, pc = int(dimension_hint[0]), int(dimension_hint[1])
                if pr > max_row:
                    logger.info(
                        "Sheet %r: expanding max_row %s -> %s (pandas shape hint)",
                        sheet_name,
                        max_row,
                        pr,
                    )
                    max_row = pr
                if pc > max_col:
                    logger.info(
                        "Sheet %r: expanding max_col %s -> %s (pandas shape hint)",
                        sheet_name,
                        max_col,
                        pc,
                    )
                    max_col = pc
            except (TypeError, ValueError):
                pass
        
        # Collect merged cell ranges
        merged_ranges = [str(r) for r in ws.merged_cells.ranges]
        
        # Build merged cell value map
        merge_values = self._build_merge_map(ws)
        
        # Get hidden rows/cols
        hidden_rows = self._get_hidden_rows(ws) if not self.include_hidden else []
        hidden_cols = self._get_hidden_cols(ws) if not self.include_hidden else []
        
        # Build cell grid
        cells = []
        for row_idx in range(1, max_row + 1):
            if row_idx in hidden_rows and not self.include_hidden:
                continue
                
            row_cells = []
            for col_idx in range(1, max_col + 1):
                if col_idx in hidden_cols and not self.include_hidden:
                    continue
                
                openpyxl_cell = ws.cell(row=row_idx, column=col_idx)
                cell = self._convert_cell(openpyxl_cell, merge_values)
                row_cells.append(cell)
            
            cells.append(row_cells)
        
        return VisualGrid(
            cells=cells,
            sheet_name=sheet_name,
            total_rows=len(cells),
            total_cols=len(cells[0]) if cells else 0,
            merged_ranges=merged_ranges,
            hidden_rows=hidden_rows,
            hidden_cols=hidden_cols
        )
    
    def _build_merge_map(self, ws: Worksheet) -> Dict[str, Any]:
        """
        Build a map of merged cell coordinates to their values.
        The value of the top-left cell is propagated to all cells in the range.
        """
        merge_values = {}
        
        for merged_range in ws.merged_cells.ranges:
            # Get top-left cell value
            min_row, min_col = merged_range.min_row, merged_range.min_col
            top_left_value = ws.cell(row=min_row, column=min_col).value
            
            # Map all cells in range to this value
            for row in range(merged_range.min_row, merged_range.max_row + 1):
                for col in range(merged_range.min_col, merged_range.max_col + 1):
                    coord = f"{get_column_letter(col)}{row}"
                    merge_values[coord] = {
                        'value': top_left_value,
                        'range': str(merged_range),
                        'is_merged': True
                    }
        
        return merge_values
    
    def _convert_cell(self, openpyxl_cell: OpenpyxlCell, 
                      merge_values: Dict[str, Any]) -> Cell:
        """
        Convert an openpyxl cell to our Cell model.
        """
        coord = openpyxl_cell.coordinate
        
        # Check if this cell is part of a merged range
        merge_info = merge_values.get(coord, {})
        is_merged = merge_info.get('is_merged', False)
        
        # Get value (from merge map if merged, else from cell)
        if is_merged:
            value = merge_info.get('value')
        else:
            value = openpyxl_cell.value
        
        # Extract style
        style = self._extract_style(openpyxl_cell, merge_info)
        
        # Determine cell type
        cell_type = self._determine_type(openpyxl_cell, value)
        
        return Cell(
            value=value,
            row=openpyxl_cell.row - 1,  # Convert to 0-indexed
            col=openpyxl_cell.column - 1,
            style=style,
            original_type=cell_type,
            original_coordinate=coord
        )
    
    def _extract_style(self, cell: OpenpyxlCell, 
                       merge_info: Dict[str, Any]) -> CellStyle:
        """Extract cell formatting as CellStyle."""
        font = cell.font
        fill = cell.fill
        
        return CellStyle(
            is_bold=font.bold if font.bold else False,
            is_italic=font.italic if font.italic else False,
            bg_color=fill.start_color.rgb if fill and fill.start_color and hasattr(fill.start_color, 'rgb') else None,
            font_color=font.color.rgb if font.color and hasattr(font.color, 'rgb') else None,
            font_size=font.size if font.size else None,
            indent_level=cell.alignment.indent if cell.alignment else 0,
            is_merged=merge_info.get('is_merged', False),
            merge_range=merge_info.get('range')
        )
    
    def _determine_type(self, cell: OpenpyxlCell, value: Any) -> CellType:
        """Determine the type of cell content."""
        if value is None:
            return CellType.EMPTY
        
        if cell.data_type == 'f':
            return CellType.FORMULA
        if cell.data_type == 'e':
            return CellType.ERROR
        if cell.data_type == 'b':
            return CellType.BOOLEAN
        if cell.data_type == 'd':
            return CellType.DATE
        if cell.data_type == 'n':
            return CellType.NUMBER
        
        return CellType.TEXT
    
    def _get_hidden_rows(self, ws: Worksheet) -> List[int]:
        """Get list of hidden row indices."""
        hidden = []
        for row_idx, rd in ws.row_dimensions.items():
            if rd.hidden:
                hidden.append(row_idx)
        return hidden
    
    def _get_hidden_cols(self, ws: Worksheet) -> List[int]:
        """Get list of hidden column indices."""
        hidden = []
        for col_letter, cd in ws.column_dimensions.items():
            if cd.hidden:
                # Convert letter to index
                col_idx = 0
                for char in col_letter:
                    col_idx = col_idx * 26 + (ord(char.upper()) - ord('A') + 1)
                hidden.append(col_idx)
        return hidden
