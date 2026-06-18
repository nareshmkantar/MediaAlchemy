"""
Web Server for Structure Inference Agent
Provides REST API for file upload, processing, and results viewing.
"""
import os
import json
import copy
import uuid
import logging
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from flask import Flask, request, jsonify, send_from_directory, send_file, redirect, make_response
from flask_cors import CORS
from werkzeug.utils import secure_filename
import pandas as pd
import asyncio
import threading
from openpyxl import load_workbook

# Add project to path
import sys
sys.path.insert(0, str(Path(__file__).parent))

from sia.agent.orchestrator import StructureInferenceAgent
from datetime import datetime, date, timezone
from enum import Enum
from sia.utils.logger import setup_logging
from sia.utils.df_preview import dataframe_to_preview_records
from sia.agent.job_manager import job_manager
from sia.agent.context_packet import ContextPacket, build_context_packet
from sia.agent.hitl import CheckpointType, HITLManager
from sia.agent.parent_collation_graph import (
    build_suggested_duplicate_key_decisions,
    diagnose_stacked_union_duplicates,
    enumerate_duplicate_key_groups_for_review,
    normalize_union_duplicate_merge_mode,
    resolve_duplicate_merge_action,
    union_duplicate_merge_mode_from_relationships,
)
from sia.agent.relationships import (
    apply_relationship_decisions,
    collate_frames,
    collate_frames_detailed,
    propose_file_relationships,
    union_stack_break_row_indices,
    _collate_frames_baseline,
    _normalize_frames_keys,
    _union_stack_column_intersection,
    _union_stack_source_ids,
)
from sia.agent.artifact_exports import (
    build_pre_transform_dataframe,
    iter_artifact_paths,
    persist_processing_artifacts,
)
from sia.agent.schema_mapper import SchemaMapper
from sia.agent.demarcator import StructureDemarcator
from collections import OrderedDict

from sia.agent.scoped_source import build_scoped_source, load_raw_sheet_dataframe, load_scoped_dataframe
from sia.agent.materialized_clean_sheet import (
    refresh_materialized_clean_templates,
    resolve_processing_workbook,
)
from sia.agent.job_run_ledger import (
    clear_multi_source_ledger,
    ensure_bootstrap_and_route,
    record_source_run_ledger,
)
from sia.agent.multi_block_sheet import (
    build_scoped_source_for_block,
    enrich_context_packet_for_block_run,
    ensure_block_union_relationship,
    expand_multi_block_source_ids,
    is_block_virtual_source_id,
    resolve_block_run,
)
from sia.agent.target_template_utils import (
    mapping_target_columns_for_ui,
    normalize_target_template,
    primary_target_columns,
    validate_template_shape,
)
from sia.agent.planner import finalize_extraction_plan
from sia.agent.planner_decisions import (
    apply_resolved_approvals_to_context,
    merge_resolved_approval_items,
)
from sia.utils.export_dates import normalize_dataframe_dates_for_export
from sia.agent.date_hints import compute_date_hints_for_dataframe, normalize_date_shape_for_storage
from sia.tools.cross_source_union_validate import validate_column_sets_for_union
from sia.debug.state_snapshot import build_state_snapshot
from sia.debug.job_debug_events import (
    append_job_debug_event,
    data_files_summary,
    get_job_debug_timeline,
    job_setup_summary,
    make_job_debug_event,
)

# Setup logging
setup_logging(level=logging.INFO)
logger = logging.getLogger("sia.web")

# Workbook uploads: refuse extremely large tab counts (memory / UX).
MAX_VISIBLE_SHEETS_PER_WORKBOOK = 100
ALLOWED_DATE_GRANULARITY = frozenset({"daily", "weekly", "monthly", "quarterly", "range"})
ALLOWED_DATE_SHAPES = frozenset(
    {"unknown", "single_timestamp", "period_span", "period_text", "date_parts"}
)
# Guided Setup preview JSON: full requested range is still used for demarcation/saves; body is capped for browser performance.
DEMARCATION_PREVIEW_MAX_BODY_ROWS = 50


def _flatten_source_run_trace_steps(source_runs: Optional[List[Any]]) -> List[Dict[str, Any]]:
    """Merge per-source ``ProcessingTrace`` steps (in registry order) for combined debug timeline."""
    merged: List[Dict[str, Any]] = []
    for run in source_runs or []:
        if not isinstance(run, dict):
            continue
        sid = str(run.get("source_id") or "").strip()
        tr = run.get("trace")
        if not isinstance(tr, dict):
            continue
        for step in tr.get("steps") or []:
            if not isinstance(step, dict):
                continue
            tagged = dict(step)
            if sid:
                tagged.setdefault("source_id", sid)
            merged.append(tagged)
    return merged


def _collation_trace_to_job_steps(collation_trace: Optional[List[Dict[str, Any]]]) -> List[Dict[str, Any]]:
    """Map parent-collation graph internal steps to ``ProcessingTrace``-shaped records for the UI."""
    out: List[Dict[str, Any]] = []
    for raw in collation_trace or []:
        if not isinstance(raw, dict):
            continue
        step = str(raw.get("step") or "collation")
        bits = [step]
        for k in (
            "rows",
            "cols",
            "row_count",
            "column_count",
            "duplicate_rows_on_keys",
            "planned_tool_count",
            "tools",
            "executions",
            "sources",
        ):
            if k in raw and raw[k] not in (None, "", []):
                bits.append(f"{k}={raw[k]!r}")
        out.append(
            {
                "module": "collation",
                "input": "",
                "output": "; ".join(bits[:14]),
                "confidence": {
                    "score": 1.0,
                    "decision": "deterministic",
                    "rationale": f"collation:{step}",
                    "signals": [],
                },
            }
        )
    return out


def _collation_root_event_id(collation_events: List[Dict[str, Any]]) -> Optional[str]:
    """Top-level Post-merge collation node id from observer hierarchical events."""
    for ev in collation_events or []:
        if not isinstance(ev, dict):
            continue
        if ev.get("parent_event_id") in (None, ""):
            mod = str(ev.get("module") or "")
            if mod.startswith("collation.") or mod == "collation.merge_sources":
                return str(ev.get("event_id") or "") or None
    return None


def _link_deferred_collation_debug_events(
    collation_events: List[Dict[str, Any]],
    deferred_events: List[Dict[str, Any]],
) -> List[Dict[str, Any]]:
    """Attach deferred grain tools under the Post-merge collation tree (``source_id=__collation__``)."""
    root_id = _collation_root_event_id(collation_events)
    post_collate_id: Optional[str] = None
    linked: List[Dict[str, Any]] = []
    for raw in deferred_events or []:
        if not isinstance(raw, dict):
            continue
        ev = dict(raw)
        ev["source_id"] = "__collation__"
        mod = str(ev.get("module") or "")
        ptype = str(ev.get("process_type") or "")
        if mod == "collation.post_collate_transforms" and ptype in ("node", "chain"):
            if root_id:
                ev["parent_event_id"] = root_id
            post_collate_id = str(ev.get("event_id") or "") or post_collate_id
        elif ptype == "tool_execution":
            if post_collate_id:
                ev["parent_event_id"] = post_collate_id
            elif root_id:
                ev["parent_event_id"] = root_id
        linked.append(ev)
    return linked


def _apply_job_deferred_post_collate(
    job: Dict[str, Any],
    combined_df: Optional[pd.DataFrame],
) -> Tuple[Optional[pd.DataFrame], List[Dict[str, Any]], List[Dict[str, Any]]]:
    """After union/collation, run transforms deferred from per-source plans (e.g. ``aggregate_weekly``).

    Returns ``(dataframe, trace_events, hierarchical_debug_events)``. Grain tools are logged on an
    observer span ``collation.post_collate_transforms`` so they appear under Post-merge collation
    in the Trace tab (with CSV snapshots like other collation tools).
    """
    from sia.agent.post_collate_transforms import (
        apply_combined_frame_final_layout,
        apply_deferred_post_collate_transforms,
        merge_deferred_post_collate_tool_lists,
    )
    from sia.debug.llm_observer import LLMObserver

    if combined_df is None:
        return combined_df, [], []

    by_src = dict(job.get("_deferred_post_collate_by_source") or {})
    merged = merge_deferred_post_collate_tool_lists(by_src) if by_src else []
    target_template = load_target_template_for_job(job)
    cp = job.get("context_packet") if isinstance(job.get("context_packet"), dict) else {}
    approved_mappings = list(job.get("approved_mappings") or cp.get("approved_mappings") or [])

    if not merged and not target_template:
        return combined_df, [], []

    col_obs = LLMObserver(run_id=str(job.get("id") or "collation"), model_id="", enabled=True)
    with col_obs.trace_span(
        "collation.post_collate_transforms",
        type="node",
        metadata={
            "phase": "post_collate",
            "deferred_tool_count": len(merged),
            "final_layout": bool(target_template),
        },
    ):
        if merged:
            out, trace = apply_deferred_post_collate_transforms(
                combined_df,
                merged,
                observer=col_obs,
                target_template=target_template,
                approved_mappings=approved_mappings,
                apply_final_layout=bool(target_template),
            )
        elif target_template:
            out, trace = apply_combined_frame_final_layout(
                combined_df,
                target_template=target_template,
                approved_mappings=approved_mappings,
                observer=col_obs,
            )
        else:
            out, trace = combined_df, []
    return out, trace, col_obs.get_hierarchical_events()


def _build_multi_source_job_trace(
    *,
    source_runs: List[Dict[str, Any]],
    trace_payloads: Optional[List[Dict[str, Any]]],
    collation_trace: Optional[List[Dict[str, Any]]],
    overall_confidence: float,
    summary_message: str,
) -> Dict[str, Any]:
    per_source = _flatten_source_run_trace_steps(source_runs)
    if not per_source:
        for td in trace_payloads or []:
            if isinstance(td, dict):
                for step in td.get("steps") or []:
                    if isinstance(step, dict):
                        per_source.append(dict(step))
    col_steps = _collation_trace_to_job_steps(collation_trace)
    summary = {
        "module": "multi_source_collation",
        "input": "",
        "output": summary_message,
        "confidence": {
            "score": float(overall_confidence or 1.0),
            "decision": "unknown",
            "rationale": "",
            "signals": [],
        },
    }
    return {
        "overall_confidence": float(overall_confidence or 1.0),
        "steps": per_source + col_steps + [summary],
        "hitl_checkpoints": [],
        "source_runs": source_runs or [],
    }


def _merge_collation_into_job_debug(job: Dict[str, Any], collation_trace: Optional[List[Dict[str, Any]]]) -> None:
    """Append synthetic hierarchical debug events for post-LangGraph union/collation (Trace tab)."""
    if not collation_trace or not isinstance(collation_trace, list):
        return
    import uuid
    from datetime import datetime, timezone

    def _emit(
        events_list: List[Dict[str, Any]],
        *,
        parent: str,
        sid: str,
        ts_iso: str,
        event_id: str,
        process_type: str,
        module: str,
        operation: str,
        message: str,
        result: str = "",
        duration_ms: float = 0.0,
        params: Optional[Dict[str, Any]] = None,
        input_preview: Optional[Dict[str, Any]] = None,
        output_preview: Optional[Dict[str, Any]] = None,
    ) -> None:
        ev: Dict[str, Any] = {
            "event_id": event_id,
            "parent_event_id": parent,
            "process_type": process_type,
            "module": module,
            "operation": operation,
            "timestamp": ts_iso,
            "duration_ms": duration_ms,
            "status": "success",
            "input_summary": "",
            "message": message,
            "result": result,
            "source_id": sid,
        }
        if params:
            ev["params"] = params
        if input_preview:
            ev["input_preview"] = input_preview
        if output_preview:
            ev["output_preview"] = output_preview
        events_list.append(ev)

    ts = datetime.now(timezone.utc).isoformat()
    root_id = f"collation-root-{uuid.uuid4().hex[:10]}"
    col_sid = "__collation__"
    events: List[Dict[str, Any]] = list(job.get("debug_events") or [])
    events.append(
        {
            "event_id": root_id,
            "parent_event_id": None,
            "process_type": "node",
            "module": "collation.merge_sources",
            "operation": "NODE",
            "timestamp": ts,
            "duration_ms": 0.0,
            "status": "success",
            "input_summary": "approved_file_relationships",
            "message": "Stack union sources, align keys, dedupe / aggregate (Python collation graph)",
            "result": "",
            "source_id": col_sid,
        }
    )
    for step in collation_trace:
        if not isinstance(step, dict):
            continue
        step_name = str(step.get("step") or "collation")

        if step_name == "parent_baseline_collate":
            msg = f"rows={step.get('rows')!r}; cols={step.get('cols')!r}; sources={step.get('sources')!r}"
            if step.get("union_column_intersection"):
                msg += f"; union_column_intersection={step.get('union_column_intersection')!r}"
            _emit(
                events,
                parent=root_id,
                sid=col_sid,
                ts_iso=ts,
                event_id=f"col-{uuid.uuid4().hex[:10]}",
                process_type="tool_execution",
                module="collation.union_stack",
                operation="execute",
                message="Concat union members (baseline stack)",
                result=msg,
                params={"sources": step.get("sources")},
                input_preview={"union_column_intersection": step.get("union_column_intersection")},
                output_preview={"rows": step.get("rows"), "cols": step.get("cols")},
            )
            continue

        if step_name == "parent_analyze_merge":
            dup = step.get("duplicate_rows_on_keys")
            subset = step.get("subset_for_dedupe") or []
            conflict = step.get("numeric_conflict_on_duplicate_keys")
            _emit(
                events,
                parent=root_id,
                sid=col_sid,
                ts_iso=ts,
                event_id=f"col-{uuid.uuid4().hex[:10]}",
                process_type="tool_execution",
                module="collation.duplicate_check",
                operation="execute",
                message="Scan combined frame for duplicate keys (subset from union join_keys or column intersection)",
                result=(
                    f"duplicate_rows_on_keys={dup!r}; "
                    f"subset_for_dedupe={subset!r}; "
                    f"numeric_conflict_on_duplicate_keys={conflict!r}"
                ),
                params={
                    "union_join_keys": step.get("union_join_keys"),
                    "row_count": step.get("row_count"),
                    "column_count": step.get("column_count"),
                },
                input_preview={"subset_for_dedupe": subset},
                output_preview={
                    "duplicate_rows_on_keys": dup,
                    "numeric_conflict_on_duplicate_keys": conflict,
                },
            )
            continue

        if step_name == "parent_plan_merge_tools":
            planned = step.get("planned") or []
            tools = [t for t in (step.get("tools") or []) if t]
            pcount = int(step.get("planned_tool_count") or 0)
            if planned and isinstance(planned, list):
                for p in planned:
                    if not isinstance(p, dict):
                        continue
                    tname = str(p.get("tool") or "")
                    if not tname:
                        continue
                    tool_params = dict(p.get("params") or {})
                    _emit(
                        events,
                        parent=root_id,
                        sid=col_sid,
                        ts_iso=ts,
                        event_id=f"col-{uuid.uuid4().hex[:10]}",
                        process_type="tool_execution",
                        module="collation.plan_merge",
                        operation="execute",
                        message=f"Planned: {tname}",
                        result=f"planned_tool_count={pcount}",
                        params=tool_params,
                        output_preview={"planned_tool": tname, "planned_tool_count": pcount},
                    )
            elif tools:
                for tname in tools:
                    _emit(
                        events,
                        parent=root_id,
                        sid=col_sid,
                        ts_iso=ts,
                        event_id=f"col-{uuid.uuid4().hex[:10]}",
                        process_type="tool_execution",
                        module="collation.plan_merge",
                        operation="execute",
                        message=f"Planned: {tname}",
                        result=f"planned_tool_count={pcount}",
                        output_preview={"planned_tool": tname, "planned_tool_count": pcount},
                    )
            else:
                _emit(
                    events,
                    parent=root_id,
                    sid=col_sid,
                    ts_iso=ts,
                    event_id=f"col-{uuid.uuid4().hex[:10]}",
                    process_type="tool_execution",
                    module="collation.plan_merge",
                    operation="execute",
                    message="No merge tools planned (no duplicate key rows and column order already canonical)",
                    result=f"planned_tool_count={pcount}",
                    output_preview={"planned_tool_count": pcount, "planned_tools": []},
                )
            continue

        if step_name == "parent_execute_merge_tools":
            if step.get("skipped"):
                _emit(
                    events,
                    parent=root_id,
                    sid=col_sid,
                    ts_iso=ts,
                    event_id=f"col-{uuid.uuid4().hex[:10]}",
                    process_type="tool_execution",
                    module="collation.execute_merge_tools",
                    operation="execute",
                    message="Skipped (no dataframe)",
                    result="skipped=true",
                    output_preview={"skipped": True},
                )
                continue
            execs = step.get("executions") or []
            if not execs:
                _emit(
                    events,
                    parent=root_id,
                    sid=col_sid,
                    ts_iso=ts,
                    event_id=f"col-{uuid.uuid4().hex[:10]}",
                    process_type="tool_execution",
                    module="collation.execute_merge_tools",
                    operation="execute",
                    message="No collation tools executed (plan was empty)",
                    result="executions=[]",
                    output_preview={"executions": []},
                )
                continue
            for ev in execs:
                if not isinstance(ev, dict):
                    continue
                tool_name = str(ev.get("tool") or "collation.tool")
                rows_b = ev.get("rows_before")
                rows_a = ev.get("rows_after")
                friendly = tool_name
                if tool_name == "collation.drop_duplicate_rows":
                    friendly = "collation.remove_duplicates"
                elif tool_name == "collation.aggregate_duplicate_keys":
                    friendly = "collation.aggregate_duplicates"
                elif tool_name == "collation.resolve_duplicate_key_groups":
                    friendly = "collation.resolve_duplicate_groups"
                tool_params = dict(ev.get("params") or {})
                out_pv: Dict[str, Any] = {"rows_before": rows_b, "rows_after": rows_a}
                if rows_b is not None and rows_a is not None:
                    try:
                        out_pv["row_delta"] = int(rows_a) - int(rows_b)
                    except (TypeError, ValueError):
                        pass
                _emit(
                    events,
                    parent=root_id,
                    sid=col_sid,
                    ts_iso=ts,
                    event_id=f"col-{uuid.uuid4().hex[:10]}",
                    process_type="tool_execution",
                    module=friendly,
                    operation="execute",
                    message=f"Executed {tool_name}",
                    result=f"rows_before={rows_b!r} rows_after={rows_a!r}",
                    params=tool_params if tool_params else {"tool": tool_name},
                    input_preview={"rows_before": rows_b},
                    output_preview=out_pv,
                )
            continue

        if step_name == "parent_verify_combined":
            _emit(
                events,
                parent=root_id,
                sid=col_sid,
                ts_iso=ts,
                event_id=f"col-{uuid.uuid4().hex[:10]}",
                process_type="tool_execution",
                module="collation.verify_combined",
                operation="execute",
                message="Combined frame check after merge tools (daily grain — before deferred weekly rollup)",
                result=f"ok={step.get('ok')!r}; rows={step.get('rows')!r}",
                output_preview={"ok": step.get("ok"), "rows": step.get("rows")},
            )
            continue

        if step_name == "post_collate_transforms":
            ev_list = step.get("events") or []
            if not ev_list:
                _emit(
                    events,
                    parent=root_id,
                    sid=col_sid,
                    ts_iso=ts,
                    event_id=f"col-{uuid.uuid4().hex[:10]}",
                    process_type="node",
                    module="collation.post_collate_transforms",
                    operation="NODE",
                    message="No deferred grain tools to run",
                    result="events=[]",
                )
                continue
            post_node_id = f"col-{uuid.uuid4().hex[:10]}"
            _emit(
                events,
                parent=root_id,
                sid=col_sid,
                ts_iso=ts,
                event_id=post_node_id,
                process_type="node",
                module="collation.post_collate_transforms",
                operation="NODE",
                message=f"Deferred grain transforms after union ({len(ev_list)} tool run(s))",
                result="",
            )
            for ev in ev_list:
                if not isinstance(ev, dict):
                    continue
                tool_name = str(ev.get("tool") or "transform.deferred")
                rows_b = ev.get("rows_before")
                rows_a = ev.get("rows_after")
                ok = bool(ev.get("ok", True))
                out_pv: Dict[str, Any] = {}
                if rows_b is not None:
                    out_pv["rows_before"] = rows_b
                if rows_a is not None:
                    out_pv["rows_after"] = rows_a
                if rows_b is not None and rows_a is not None:
                    try:
                        out_pv["row_delta"] = int(rows_a) - int(rows_b)
                    except (TypeError, ValueError):
                        pass
                _emit(
                    events,
                    parent=post_node_id,
                    sid=col_sid,
                    ts_iso=ts,
                    event_id=f"col-{uuid.uuid4().hex[:10]}",
                    process_type="tool_execution",
                    module=tool_name,
                    operation="execute",
                    message=(ev.get("message") or f"Deferred post-collate: {tool_name}")[:240],
                    result=f"ok={ok!r}; rows_before={rows_b!r}; rows_after={rows_a!r}",
                    params={"deferred_post_collate": True},
                    output_preview=out_pv or None,
                )
            continue

        if step_name == "collate_baseline":
            _emit(
                events,
                parent=root_id,
                sid=col_sid,
                ts_iso=ts,
                event_id=f"col-{uuid.uuid4().hex[:10]}",
                process_type="tool_execution",
                module="collation.single_source",
                operation="execute",
                message="Single-source collate (baseline only)",
                result=f"rows={step.get('rows')!r}; sources={step.get('sources')!r}",
                input_preview={"sources": step.get("sources")},
                output_preview={"rows": step.get("rows")},
            )
            continue

        tid = f"col-{uuid.uuid4().hex[:10]}"
        msg_bits = [step_name]
        for k in ("rows", "cols", "duplicate_rows_on_keys", "planned_tool_count", "tools", "executions", "ok"):
            if k in step and step[k] not in (None, "", []):
                msg_bits.append(f"{k}={step[k]!r}")
        _emit(
            events,
            parent=root_id,
            sid=col_sid,
            ts_iso=ts,
            event_id=tid,
            process_type="node",
            module=step_name,
            operation="NODE",
            message="; ".join(msg_bits[:12]),
            result="",
        )

    job["debug_events"] = events


def _trace_deferred_post_collate_tools(trace: Any) -> List[Dict[str, Any]]:
    """Return deferred grain tools from a trace object without assuming its concrete type."""
    dft = getattr(trace, "deferred_post_collate_tools", None)
    if dft:
        return [dict(x) for x in dft if isinstance(x, dict)]
    try:
        trace_dict = trace.to_dict() if trace is not None and hasattr(trace, "to_dict") else {}
    except Exception:
        trace_dict = {}
    return [dict(x) for x in (trace_dict.get("deferred_post_collate_tools") or []) if isinstance(x, dict)]


def _record_resumed_source_deferred_tools(
    job: Dict[str, Any],
    *,
    source_id: str,
    sheet_name: str,
    processing_sheet: str,
    trace: Any,
) -> None:
    """Ledger deferred tools from a HITL-resumed source before multi-source collation."""
    sid = str(source_id or "").strip()
    if not sid:
        return
    deferred_tools = _trace_deferred_post_collate_tools(trace)
    if not deferred_tools:
        return
    acc = job.get("_deferred_post_collate_by_source") or {}
    if isinstance(acc, dict) and acc.get(sid):
        return
    record_source_run_ledger(
        job,
        sid,
        {
            "sheet_name": str(sheet_name or ""),
            "processing_sheet": str(processing_sheet or ""),
        },
        deferred_tools,
    )


def _append_job_state_snapshot(
    job: Dict[str, Any],
    *,
    label: str,
    phase: str,
    source_id: Optional[str] = None,
    sheet_name: Optional[str] = None,
    context_packet: Optional[Dict[str, Any]] = None,
    decision: Optional[Dict[str, Any]] = None,
) -> None:
    """Append a compact job/memory snapshot for the Debug State tab."""
    try:
        state = {
            "source_id": source_id,
            "sheet_name": sheet_name,
            "context_packet": context_packet or job.get("context_packet") or {},
            "multi_source_active_batch": str(job.get("processing_route") or "").startswith("multi_"),
        }
        snapshot = build_state_snapshot(
            label=label,
            phase=phase,
            state=state,
            context_packet=context_packet or job.get("context_packet") or {},
            job=job,
            decision=decision or {},
        )
        event_id = f"job-state-{uuid.uuid4().hex[:8]}"
        snapshot["event_id"] = event_id
        snapshot["snapshot_id"] = f"state-{event_id}"
        snapshot["timestamp"] = datetime.now(timezone.utc).isoformat()
        if source_id:
            snapshot["source_id"] = source_id
        if sheet_name is not None:
            snapshot["sheet_name"] = sheet_name
        existing = [s for s in (job.get("state_snapshots") or []) if isinstance(s, dict)]
        existing.append(snapshot)
        job["state_snapshots"] = existing
    except Exception as exc:
        logger.debug("Failed to append job state snapshot %s/%s: %s", label, phase, exc)


def _record_job_debug(
    job: Dict[str, Any],
    *,
    label: str,
    phase: str,
    parent_event_id: Optional[str] = None,
    parent_key: Optional[str] = None,
    set_parent_key: Optional[str] = None,
    source_id: Optional[str] = None,
    sheet_name: Optional[str] = None,
    module: str = "job",
    operation: str = "orchestration",
    status: str = "info",
    summary: str = "",
    metadata: Optional[Dict[str, Any]] = None,
) -> str:
    """Append a normalized orchestration event; optionally track parent ids on the job."""
    parents = job.setdefault("_job_debug_parents", {})
    parent = parent_event_id
    if parent is None and parent_key:
        parent = parents.get(parent_key)
    ev = make_job_debug_event(
        label=label,
        phase=phase,
        parent_event_id=parent,
        source_id=source_id,
        sheet_name=sheet_name,
        module=module,
        operation=operation,
        status=status,
        summary=summary,
        metadata=metadata,
    )
    append_job_debug_event(job, ev)
    eid = str(ev.get("event_id") or "")
    if set_parent_key and eid:
        parents[set_parent_key] = eid
    return eid


def _merge_collation_debug_into_job(
    job: Dict[str, Any],
    collation_trace: Optional[List[Dict[str, Any]]],
    collation_hierarchical_events: Optional[List[Dict[str, Any]]],
) -> None:
    """Merge collation into ``job['debug_events']``: observer events (CSV snapshots) or synthetic fallback."""
    if collation_hierarchical_events:
        fresh = [e for e in collation_hierarchical_events if isinstance(e, dict)]
        if not fresh:
            _merge_collation_into_job_debug(job, collation_trace)
            return
        _annotate_debug_rows(fresh, source_id="__collation__")
        existing = [e for e in (job.get("debug_events") or []) if isinstance(e, dict)]
        seen = {e.get("event_id") for e in existing if e.get("event_id")}
        for ev in fresh:
            eid = ev.get("event_id")
            if eid and eid in seen:
                continue
            if eid:
                seen.add(eid)
            existing.append(ev)
        job["debug_events"] = existing
        return
    _merge_collation_into_job_debug(job, collation_trace)


def _pandas_sheet_shape(file_path: str, sheet_name: str) -> Optional[Tuple[int, int]]:
    """Return (nrows, ncols) for a sheet as ``pandas.read_excel(..., header=None)`` sees it."""
    try:
        df = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
        return len(df), len(df.columns)
    except Exception:
        return None


def _dimension_hints_for_workbook_sources(
    registry: List[Dict[str, Any]],
    resolved_file_path: str,
) -> Dict[str, Tuple[int, int]]:
    """Map sheet_name -> (nrows, ncols) for sources that belong to this workbook path."""
    hints: Dict[str, Tuple[int, int]] = {}
    if str(resolved_file_path).lower().endswith(".csv"):
        return hints
    try:
        resolved = str(Path(resolved_file_path).resolve())
    except OSError:
        resolved = str(resolved_file_path)
    try:
        xf = pd.ExcelFile(resolved_file_path)
    except Exception:
        return hints
    try:
        for src in registry:
            fp = str(src.get("file_path") or "")
            if not fp:
                continue
            try:
                if str(Path(fp).resolve()) != resolved:
                    continue
            except OSError:
                if fp != resolved_file_path:
                    continue
            sn = src.get("sheet_name")
            if not sn or sn not in xf.sheet_names:
                continue
            try:
                df = pd.read_excel(xf, sn, header=None)
                hints[str(sn)] = (len(df), len(df.columns))
            except Exception:
                continue
    finally:
        try:
            xf.close()
        except Exception:
            pass
    return hints


def _registry_sheet_names_for_workbook(
    registry: List[Dict[str, Any]],
    job: Dict[str, Any],
    resolved_cache_key: str,
) -> Optional[set[str]]:
    """Tab names the user kept in ``source_registry`` for this workbook (resolved path key)."""
    names: set[str] = set()
    for src in registry:
        fp = src.get("file_path") or job.get("file_path")
        if not fp:
            continue
        try:
            ck = str(Path(str(fp)).resolve())
        except OSError:
            ck = str(fp)
        if ck != resolved_cache_key:
            continue
        sn = src.get("sheet_name")
        if sn is not None and str(sn).strip():
            names.add(str(sn))
    return names if names else None


def _debug_log(hypothesis_id: str, message: str, data: Dict[str, Any]) -> None:
    try:
        payload = {
            "sessionId": "3fc92d",
            "runId": str(data.get("job_id") or data.get("checkpoint_id") or "unknown"),
            "hypothesisId": hypothesis_id,
            "location": "web_server.py",
            "message": message,
            "data": data,
            "timestamp": int(datetime.now().timestamp() * 1000),
        }
        with open(Path(__file__).parent / "debug-3fc92d.log", "a", encoding="utf-8") as f:
            f.write(json.dumps(payload, ensure_ascii=False) + "\n")
    except Exception:
        pass

# Initialize modules
schema_mapper = SchemaMapper(prompts_dir="prompts")
demarcator = None # Will be initialized per request or once if llm_client is global


from flask.json.provider import DefaultJSONProvider

# Keys on the in-memory job dict that must never be sent to the browser (non-JSON types / large).
_STATUS_RESPONSE_EXCLUDED_KEYS = frozenset(
    {
        "_process_worker_running",
        "_multi_source_frames_by_source",
        "_pending_collation_frames",
        "_pending_collation_schema",
    }
)


class CustomJSONProvider(DefaultJSONProvider):
    """Modern Flask JSON provider that handles datetime, date, numpy, and pandas types."""

    def default(self, obj):
        if isinstance(obj, (datetime, date)):
            return obj.isoformat()
        if isinstance(obj, Enum):
            return obj.value
        if isinstance(obj, pd.DataFrame):
            return {
                "_type": "pandas.DataFrame",
                "shape": [int(len(obj)), int(obj.shape[1]) if len(obj.shape) > 1 else 0],
            }
        if isinstance(obj, pd.Series):
            return {"_type": "pandas.Series", "len": int(len(obj))}
        # numpy / pandas dtypes (e.g. preview payloads)
        try:
            import numpy as np

            if isinstance(obj, np.dtype):
                return str(obj)
        except Exception:
            pass
        # Handle numpy types
        if hasattr(obj, "item"):
            try:
                return obj.item()
            except Exception:
                pass
        if hasattr(obj, "tolist"):
            try:
                return obj.tolist()
            except Exception:
                pass
        # Handle pandas Timestamp
        if hasattr(obj, "isoformat") and not isinstance(obj, (datetime, date, pd.DataFrame, pd.Series)):
            try:
                return obj.isoformat()
            except Exception:
                pass
        return super().default(obj)


app = Flask(__name__, static_folder='web', static_url_path='')
app.json = CustomJSONProvider(app)
CORS(app)

# Disable caching for development
app.config['SEND_FILE_MAX_AGE_DEFAULT'] = 0

@app.route('/api/download/<path:filename>')
def download_output_file(filename):
    """Download a generated output file."""
    return send_from_directory(OUTPUT_FOLDER, filename, as_attachment=True)


@app.after_request
def add_cache_headers(response):
    """Add no-cache headers for development."""
    if 'text/css' in response.content_type or 'javascript' in response.content_type:
        response.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
        response.headers['Pragma'] = 'no-cache'
        response.headers['Expires'] = '0'
    return response


# Configuration
ROOT_DIR = Path(__file__).parent
from sia.utils.env_loader import (
    SECRET_CONFIG_KEYS,
    PROVIDER_ENV_MAP,
    load_project_dotenv,
    provider_env_var,
    update_dotenv,
)

load_project_dotenv()

CONFIG_DIR = ROOT_DIR / 'config'
RUNTIME_DIR = ROOT_DIR / 'runtime'
UPLOAD_FOLDER = RUNTIME_DIR / 'uploads'
OUTPUT_FOLDER = RUNTIME_DIR / 'outputs'
SNAPSHOT_FOLDER = RUNTIME_DIR / 'snapshots'
LOG_FOLDER = RUNTIME_DIR / 'logs'
# MCP/orchestrator tool snapshots may be written under repo/output/step_*.xlsx.
LEGACY_TOOL_SNAPSHOT_DIR = ROOT_DIR / 'output'
CONFIG_FILE = CONFIG_DIR / 'user_config.json'
ALLOWED_EXTENSIONS = {'xlsx', 'xls', 'csv', 'json'}

CONFIG_DIR.mkdir(exist_ok=True)
UPLOAD_FOLDER.mkdir(parents=True, exist_ok=True)
OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
SNAPSHOT_FOLDER.mkdir(parents=True, exist_ok=True)
LOG_FOLDER.mkdir(parents=True, exist_ok=True)
LEGACY_TOOL_SNAPSHOT_DIR.mkdir(parents=True, exist_ok=True)


def _resolve_debug_snapshot_file(safe_name: str) -> Optional[Path]:
    """Resolve a snapshot basename under runtime/snapshots or legacy repo/output."""
    name = Path(str(safe_name)).name
    if not name:
        return None
    for base in (SNAPSHOT_FOLDER, LEGACY_TOOL_SNAPSHOT_DIR):
        try:
            root = base.resolve()
            candidate = (base / name).resolve()
        except OSError:
            continue
        try:
            candidate.relative_to(root)
        except ValueError:
            continue
        if candidate.is_file():
            return candidate
    return None


def _persist_excel_output(job_id, job, df):
    """Save the final dataframe as the canonical Excel deliverable for a job.

    Populates ``job["full_excel_path"]``, ``job["output_file"]`` and exposes
    the same path on the attached ``trace`` so every UI surface (Processing,
    Debug, Review) points to the same downloadable Excel file.
    """
    try:
        import pandas as _pd
    except ImportError:
        _pd = None

    if df is None or (_pd is not None and getattr(df, "empty", True)):
        return None

    try:
        excel_path = OUTPUT_FOLDER / f"{job_id}_full_data.xlsx"
        df_write = normalize_dataframe_dates_for_export(df)
        if df_write is None:
            return None
        df_write.to_excel(excel_path, index=False, engine="openpyxl")
        path_str = str(excel_path.resolve())
        job["full_excel_path"] = path_str
        job["output_file"] = path_str
        output_files = job.get("output_files") or {}
        if isinstance(output_files, dict):
            output_files.setdefault("excel", path_str)
            job["output_files"] = output_files
        trace = job.get("trace")
        if isinstance(trace, dict):
            trace["output_file"] = path_str
        return path_str
    except Exception:  # noqa: BLE001
        raise


def _persist_processing_artifact_exports(
    job_id: str,
    job: Dict[str, Any],
    target_template: Dict[str, Any],
    post_df: Optional[pd.DataFrame],
    source_ids: List[str],
) -> Dict[str, Any]:
    """Persist analyst-facing pre/post transformation workbook artifacts."""
    pre_frames: Dict[str, pd.DataFrame] = {}
    scope_registry = job.get("source_scope_registry") or {}

    for source_id in source_ids or []:
        source = resolve_job_source(job, source_id=source_id)
        if not source:
            continue
        scoped_source = dict(scope_registry.get(source_id) or job.get("scoped_source") or {})
        try:
            eff_path, eff_sheet, eff_scoped, used_materialized = resolve_processing_workbook(
                job,
                str(source_id),
                source["file_path"],
                str(source.get("sheet_name") or ""),
                scoped_source or None,
            )
            if used_materialized:
                logger.info(
                    "[PROCESS] Artifacts: using materialized workbook for source %s: %s",
                    source_id,
                    eff_path,
                )
            prepared_df, _ = load_scoped_dataframe(
                eff_path,
                eff_sheet,
                eff_scoped or scoped_source or None,
            )
        except Exception as exc:
            logger.warning("[PROCESS] Failed to build pre-transform frame for %s: %s", source_id, exc)
            continue

        context_packet = build_context_packet(
            job,
            target_template=target_template,
            selected_sheet=source.get("sheet_name"),
            selected_source_id=source_id,
        )
        approved_mappings = list(context_packet.get("approved_mappings") or [])
        pre_df = build_pre_transform_dataframe(
            prepared_df,
            target_template=target_template,
            approved_mappings=approved_mappings,
        )
        if pre_df is None:
            continue
        label = "__".join(
            [
                str(source.get("file_name") or source.get("source_id") or "source"),
                str(source.get("sheet_name") or source_id or "sheet"),
            ]
        )
        pre_frames[label] = pre_df

    artifacts = persist_processing_artifacts(OUTPUT_FOLDER, job_id, pre_frames, post_df)
    job["artifact_files"] = artifacts
    if isinstance(job.get("output_files"), dict) and artifacts.get("post_transform_file"):
        job["output_files"].setdefault("post_transform", artifacts["post_transform_file"])
    return artifacts


def _materialize_clean_templates_at_process_start(
    job: Dict[str, Any], job_id: str, source_ids: List[Any]
) -> None:
    """Write cleaned workbooks once when processing starts (not on each mapping/demarcation save)."""
    ids = [str(s) for s in (source_ids or []) if s]
    try:
        mat_summary = refresh_materialized_clean_templates(
            job, job_id, source_ids=ids if ids else None
        )
        if mat_summary.get("updated"):
            logger.info(
                "[PROCESS] Materialized clean workbook(s) for pipeline: %s",
                mat_summary.get("updated"),
            )
    except Exception as mat_err:
        logger.warning(
            "[PROCESS] Materialized template refresh failed (non-fatal): %s",
            mat_err,
        )


# ===== GLOBAL STATE MOVED TO JOB_MANAGER =====
# Using the job_manager global singleton for state management


def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def safe_serialize_df(df):
    """Convert DataFrame to JSON-safe dicts with keys in frame column order."""
    _, records = dataframe_to_preview_records(df)
    return records


def assign_job_data_preview(job: Dict[str, Any], df, *, head: int = 20) -> None:
    """Set ``data_preview`` and ``data_preview_column_order`` on a job dict."""
    if df is None or getattr(df, "empty", True):
        job["data_preview"] = []
        job["data_preview_column_order"] = []
        return
    order, records = dataframe_to_preview_records(df, max_rows=head)
    job["data_preview"] = records
    job["data_preview_column_order"] = order


def serialize_job_sources(job: Dict[str, Any]) -> List[Dict[str, Any]]:
    sources = []
    for source in job.get("source_registry", []) or []:
        sources.append({
            "source_id": source.get("source_id"),
            "file_id": source.get("file_id"),
            "file_name": source.get("file_name"),
            "sheet_name": source.get("sheet_name"),
            "source_type": source.get("source_type"),
            "variable_type": source.get("variable_type"),
        })
    return sources


def resolve_job_source(job: Dict[str, Any], source_id: str | None = None, sheet_name: str | None = None) -> Dict[str, Any]:
    """Resolve a row from ``source_registry`` for the given source or sheet.

    When ``source_id`` is provided, match is **strict** (after normalizing to string).
    There is **no silent fallback** to the first registry row on mismatch — callers that
    need a default when the id is absent should pass ``source_id=None`` and handle
    ``{}`` (multi-source loops skip unknown ids and log).
    """
    registry = [s for s in (job.get("source_registry") or []) if isinstance(s, dict)]
    if not registry:
        return {}

    if source_id is not None and str(source_id).strip() != "":
        want = str(source_id).strip()
        for source in registry:
            reg_id = source.get("source_id")
            if reg_id is None:
                continue
            if str(reg_id).strip() == want:
                return source
        return {}

    if sheet_name is not None and str(sheet_name).strip() != "":
        want_sheet = str(sheet_name).strip()
        for source in registry:
            if str(source.get("sheet_name") or "").strip() == want_sheet:
                return source

    return dict(registry[0])


def _resolve_resume_source_id(job: Dict[str, Any], resume_state: Dict[str, Any]) -> Optional[str]:
    """Resolve canonical registry ``source_id`` for HITL resume.

    ``pending_state`` may omit top-level ``source_id`` (it is not part of :class:`AgentState`),
    so LangGraph can drop it between nodes. Falling back to ``sheet_name`` alone is unsafe when
    multiple workbooks share the same sheet name (e.g. ``Sheet1``): :func:`resolve_job_source`
    would return the first registry row and resume would run the wrong file.
    """
    if not isinstance(resume_state, dict):
        return None
    sid = resume_state.get("source_id")
    if sid is not None and str(sid).strip():
        return str(sid).strip()
    sm = resume_state.get("source_metadata")
    if isinstance(sm, dict):
        sid = sm.get("source_id")
        if sid is not None and str(sid).strip():
            return str(sid).strip()
    cp = resume_state.get("context_packet")
    if isinstance(cp, dict):
        csm = cp.get("source_metadata")
        if isinstance(csm, dict):
            sid = csm.get("source_id")
            if sid is not None and str(sid).strip():
                return str(sid).strip()
    fp = resume_state.get("file_path")
    if not fp:
        return None
    try:
        want = str(Path(str(fp)).resolve())
    except OSError:
        want = str(fp)
    templates = job.get("materialized_clean_templates") or {}
    for sid, meta in templates.items():
        if not isinstance(meta, dict):
            continue
        mp = meta.get("path")
        if not mp:
            continue
        try:
            if str(Path(str(mp)).resolve()) == want:
                out = str(sid).strip()
                if out:
                    return out
        except OSError:
            if str(mp) == str(fp):
                out = str(sid).strip()
                if out:
                    return out
    for row in job.get("source_registry") or []:
        if not isinstance(row, dict):
            continue
        rfp = row.get("file_path")
        if not rfp:
            continue
        try:
            got = str(Path(str(rfp)).resolve())
        except OSError:
            got = str(rfp)
        if got == want or str(rfp) == str(fp):
            out = row.get("source_id")
            if out is not None and str(out).strip():
                return str(out).strip()
    return None


def _count_non_empty_cells_in_grid_df(grid_df: pd.DataFrame) -> int:
    """Match demarcator occupancy: None, blank strings, and pandas NA are empty."""
    if grid_df is None or grid_df.empty:
        return 0
    count = 0
    for row in grid_df.itertuples(index=False, name=None):
        for value in row:
            if value is None:
                continue
            if isinstance(value, str) and not value.strip():
                continue
            if pd.isna(value):
                continue
            count += 1
    return count


def _serialize_plan_for_review(plan_obj: Any) -> Dict[str, Any]:
    """Normalize an extraction plan into a browser-safe payload."""
    if not plan_obj:
        return {}

    if isinstance(plan_obj, dict):
        source = plan_obj
    else:
        source = {
            "tool_calls": getattr(plan_obj, "tool_calls", []),
            "expected_columns": getattr(plan_obj, "expected_columns", []),
            "business_rule_actions": getattr(plan_obj, "business_rule_actions", []),
            "approval_items": getattr(plan_obj, "approval_items", []),
            "reasoning": getattr(plan_obj, "reasoning", ""),
            "confidence": getattr(plan_obj, "confidence", None),
            "review_reason": getattr(plan_obj, "review_reason", ""),
        }

    return {
        "tool_calls": list(source.get("tool_calls") or []),
        "expected_columns": list(source.get("expected_columns") or []),
        "business_rule_actions": list(source.get("business_rule_actions") or source.get("planned_rule_actions") or []),
        "approval_items": list(source.get("approval_items") or []),
        "reasoning": str(source.get("reasoning") or ""),
        "confidence": source.get("confidence"),
        "review_reason": str(source.get("review_reason") or ""),
    }


def _build_agent_from_config(config: Dict[str, Any]) -> StructureInferenceAgent:
    model_name = config.get('llm_model', 'gemini-3-flash-preview')
    api_key = '' if model_name == 'rule-based' else config.get('api_key', '')
    return StructureInferenceAgent(
        config_path=str(Path(__file__).parent / 'config' / 'semantic_config.yaml'),
        api_key=api_key,
        model_name=model_name if model_name != 'rule-based' else None,
        debug_enabled=config.get('debug_enabled', False),
        enable_llm_judge=bool(config.get('enable_llm_judge', False)),
        azure_endpoint=config.get('azure_endpoint'),
        azure_deployment=config.get('azure_deployment'),
        azure_api_version=config.get('azure_api_version'),
        azure_ssl_verify=config.get('azure_ssl_verify', True)
    )


def _apply_hitl_pause(
    job_id: str, job: Dict[str, Any], trace: Any, df: Optional[pd.DataFrame]
) -> Tuple[Dict[str, Any], int]:
    """Update job / review queue for a HITL pause. Returns JSON-serializable body and HTTP status.

    Does not call ``jsonify`` — safe from background threads (no Flask application context).
    Request handlers should wrap: ``return jsonify(payload), status``.
    """
    low_conf = getattr(trace, 'low_confidence_items', [])
    checkpoints = list(getattr(trace, "hitl_checkpoints", []) or [])
    deletion_previews = list(getattr(trace, "deletion_previews", []) or [])
    pause_reason = getattr(trace, 'review_reason', '') or (
        f"Paused for approval of {len(deletion_previews)} destructive operations"
        if deletion_previews else
        "Paused for human review"
    )

    cap_sid, cap_sheet, cap_proc = _debug_capture_source_tags(job, trace)
    capture_llm_metadata(
        job_id,
        source_id=cap_sid,
        sheet_name=cap_sheet or None,
        processing_sheet=cap_proc or None,
    )
    assign_job_data_preview(job, df, head=20)
    job["trace"] = trace.to_dict()

    if checkpoints:
        pending_state = getattr(trace, "pending_state", {}) or {}
        added = 0
        checkpoint_types = set()
        next_checkpoint_id: Optional[str] = None
        # region agent log
        _debug_log(
            "H5_H6",
            "resume returned hitl pause with checkpoints",
            {
                "job_id": job_id,
                "pause_reason": pause_reason,
                "checkpoint_types": [cp.get("checkpoint_type") for cp in checkpoints if isinstance(cp, dict)],
                "checkpoint_titles": [cp.get("title") for cp in checkpoints if isinstance(cp, dict)],
                "trace_review_reason": getattr(trace, "review_reason", ""),
            },
        )
        # endregion
        for checkpoint in checkpoints:
            if checkpoint.get("resolved", False):
                continue
            checkpoint_payload = dict(checkpoint)
            checkpoint_payload["pending_state"] = pending_state
            cp_id = job_manager.add_checkpoint_to_review(job_id, checkpoint_payload)
            checkpoint_types.add(checkpoint.get("checkpoint_type"))
            added += 1
            if next_checkpoint_id is None and checkpoint.get("checkpoint_type") == "plan_review":
                next_checkpoint_id = cp_id

        review_step = "Plan Review" if checkpoint_types == {"plan_review"} else "Human Review"
        job_manager.update_job_status(job_id, "awaiting_review", pause_reason, review_step)
        payload: Dict[str, Any] = {
            "success": True,
            "job_id": job_id,
            "status": "awaiting_review",
            "message": pause_reason,
            "checkpoint_count": added,
            "requires_review": True,
        }
        if next_checkpoint_id:
            payload["next_checkpoint_id"] = next_checkpoint_id
        try:
            from sia.debug.hitl_debug_events import record_hitl_pause_debug

            record_hitl_pause_debug(
                job,
                trace,
                source_id=cap_sid,
                sheet_name=cap_sheet or None,
                processing_sheet=cap_proc or None,
            )
        except Exception as exc:
            logger.debug("[HITL] Failed to record checkpoint pause debug for %s: %s", job_id, exc)
        return payload, 200

    if deletion_previews:
        job_manager.pause_for_destructive_approval(
            job_id,
            deletion_previews,
            getattr(trace, 'pending_state', {}),
            low_confidence_items=low_conf
        )
        try:
            from sia.debug.hitl_debug_events import record_hitl_pause_debug

            record_hitl_pause_debug(
                job,
                trace,
                source_id=cap_sid,
                sheet_name=cap_sheet or None,
                processing_sheet=cap_proc or None,
            )
        except Exception as exc:
            logger.debug("[HITL] Failed to record destructive pause debug for %s: %s", job_id, exc)

        return {
            "success": True,
            "job_id": job_id,
            "status": "awaiting_approval",
            "message": f"Processing paused - {pause_reason}",
            "deletion_previews": deletion_previews,
            "low_confidence_items": low_conf,
            "preview_count": len(deletion_previews) + len(low_conf)
        }, 200

    logger.warning("[HITL] Pause requested without checkpoints or deletion previews for job %s", job_id)
    return {
        "success": False,
        "job_id": job_id,
        "status": "error",
        "error": "HITL pause did not include any reviewable checkpoints or destructive previews."
    }, 500


def _merge_observer_hierarchical_events_into_job(
    job: Dict[str, Any],
    *,
    source_id: Optional[str] = None,
    sheet_name: Optional[str] = None,
    processing_sheet: Optional[str] = None,
) -> None:
    """Append observer debug events to the job without dropping prior sources' trees.

    Assigning only the latest observer flush to job['debug_events'] drops earlier
    multi-source passes in the Debug UI so it looks as if a later file replaced an earlier one.
    """
    try:
        from sia.debug.llm_observer import get_observer

        observer = get_observer()
        if not observer:
            return
        fresh = observer.get_hierarchical_events() or []
    except Exception:
        return
    if not isinstance(fresh, list) or not fresh:
        return
    if source_id or sheet_name or processing_sheet:
        _annotate_debug_rows(
            fresh,
            source_id=source_id,
            sheet_name=sheet_name,
            processing_sheet=processing_sheet,
        )
    existing = [e for e in (job.get("debug_events") or []) if isinstance(e, dict)]
    seen = {e.get("event_id") for e in existing if e.get("event_id")}
    for ev in fresh:
        if not isinstance(ev, dict):
            continue
        eid = ev.get("event_id")
        if eid and eid in seen:
            continue
        if eid:
            seen.add(eid)
        existing.append(ev)
    job["debug_events"] = existing


def _run_resumed_processing(job_id: str, job: Dict[str, Any], resume_state: Dict[str, Any], config: Dict[str, Any]):
    if not resume_state:
        raise ValueError("No saved state to resume from")

    resume_state = dict(resume_state)
    multi_source_remaining_ids = list(resume_state.pop("multi_source_remaining_ids", []) or [])
    agent = _build_agent_from_config(config)

    if agent.llm_client:
        resume_state["llm_client"] = agent.llm_client
        logger.info("[HITL] Refreshed LLM client in resume state")
    resume_state["enable_llm_judge"] = getattr(agent, "enable_llm_judge", True)

    logger.info(f"[HITL] Resume state keys: {list(resume_state.keys())}")
    target_template = load_target_template_for_job(job)
    selected_sheet = resume_state.get("sheet_name") or (job.get("scoped_source") or {}).get("sheet_name")
    prev_sid = str(resume_state.get("source_id") or "").strip()
    selected_source_id = _resolve_resume_source_id(job, resume_state)
    if selected_source_id:
        resume_state["source_id"] = selected_source_id
        if prev_sid != selected_source_id:
            logger.info(
                "[HITL] Normalized resume source_id to %r for job %s (was %r; avoids Sheet1→first-registry fallback)",
                selected_source_id,
                job_id,
                prev_sid or None,
            )
    source = resolve_job_source(job, source_id=selected_source_id, sheet_name=selected_sheet)

    logger.info(f"[HITL] Resuming job {job_id}")
    _materialize_clean_templates_at_process_start(
        job,
        job_id,
        [selected_source_id] if selected_source_id else [],
    )
    # region agent log
    _debug_log(
        "H3_H4_H5",
        "resume processing start",
        {
            "job_id": job_id,
            "resume_mode": resume_state.get("resume_mode"),
            "suggested_tools_count": len(resume_state.get("suggested_tools") or []),
            "tool_names": [
                t.get("tool") if isinstance(t, dict) else str(type(t))
                for t in (resume_state.get("suggested_tools") or [])
            ],
            "requires_review": bool(resume_state.get("requires_review")),
            "hitl_pending_approval": bool(resume_state.get("hitl_pending_approval")),
        },
    )
    # endregion
    eff_path, eff_sheet, eff_scoped, used_materialized = resolve_processing_workbook(
        job,
        selected_source_id,
        source.get("file_path") or job["file_path"],
        str(selected_sheet or ""),
        resume_state.get("scoped_source") or {},
    )
    resume_state = dict(resume_state)
    resume_state["sheet_name"] = eff_sheet
    resume_state["scoped_source"] = eff_scoped
    resume_state["materialized_clean_active"] = used_materialized
    if used_materialized:
        logger.info("[HITL] Resume using materialized clean workbook: %s", eff_path)

    resume_batch_ids = [str(selected_source_id).strip()] if selected_source_id else []
    if multi_source_remaining_ids:
        resume_batch_ids = list(
            dict.fromkeys(
                [s for s in resume_batch_ids + [str(x).strip() for x in multi_source_remaining_ids] if s]
            )
        )
    ensure_bootstrap_and_route(job, job_id, resume_batch_ids or [str(selected_source_id or "").strip()])
    existing_multi_frames = job.get("_multi_source_frames_by_source") or {}
    multi_source_resume = bool(
        multi_source_remaining_ids
        or len(resume_batch_ids) > 1
        or (isinstance(existing_multi_frames, dict) and len(existing_multi_frames) > 0)
        or str(job.get("processing_route") or "").startswith("multi_")
    )
    if multi_source_resume:
        # Plan-review resumes re-enter the graph with saved state; make sure grain tools
        # are still withheld from this source until the parent duplicate check runs.
        resume_state["multi_source_active_batch"] = True

    context_packet = build_context_packet(
        job,
        target_template=target_template,
        selected_sheet=selected_sheet,
        selected_source_id=selected_source_id,
    )
    job["context_packet"] = context_packet

    resume_state["context_packet"] = context_packet
    resume_state["approved_mappings"] = list(context_packet.get("approved_mappings") or [])
    from sia.agent.graph_resume import prepare_plan_review_resume_state

    resume_state = prepare_plan_review_resume_state(
        resume_state,
        use_existing_plan=str(resume_state.get("resume_mode") or "") == "use_existing_plan",
        materialized_clean_active=used_materialized,
    )

    schema, df, new_trace = agent.process_file(
        eff_path,
        resume_state=resume_state,
        context_packet=context_packet,
    )

    if resume_state.get("multi_source_active_batch"):
        _record_resumed_source_deferred_tools(
            job,
            source_id=str(selected_source_id or ""),
            sheet_name=str(source.get("sheet_name") or ""),
            processing_sheet=str(eff_sheet or ""),
            trace=new_trace,
        )

    if getattr(new_trace, 'hitl_pending', False):
        payload, code = _apply_hitl_pause(job_id, job, new_trace, df)
        return jsonify(payload), code

    # region agent log
    _debug_log(
        "H6",
        "resume processing completed without hitl pause",
        {
            "job_id": job_id,
            "rows_processed": 0 if df is None else len(df),
            "overall_confidence": getattr(new_trace, "overall_confidence", None),
            "trace_review_reason": getattr(new_trace, "review_reason", ""),
        },
    )
    # endregion

    if multi_source_remaining_ids:
        # Critical: the next process_file() will init_observer() and wipe in-memory spans.
        # Flush this source's hierarchical events now or execute_tools / verify never reach job["debug_events"].
        capture_llm_metadata(
            job_id,
            source_id=str(selected_source_id).strip() if selected_source_id else None,
            sheet_name=str(source.get("sheet_name") or ""),
            processing_sheet=str(eff_sheet or ""),
        )
        logger.info(
            "[HITL] Resuming multi-source batch: current source finished; running %d remaining: %s",
            len(multi_source_remaining_ids),
            multi_source_remaining_ids,
        )
        canonical_done = str(selected_source_id or "").strip()
        seed_frames: Dict[str, pd.DataFrame] = {}
        if df is not None and not df.empty and canonical_done:
            seed_frames[canonical_done] = df
        seed_run = {
            "source_id": canonical_done,
            "sheet_name": str(source.get("sheet_name") or ""),
            "processing_sheet": str(eff_sheet or ""),
            "trace": new_trace.to_dict(),
        }
        schema_ms, multi_res, pause_tail = _process_job_sources(
            job_id,
            job,
            config,
            target_template,
            multi_source_remaining_ids,
            _seed_frames=seed_frames or None,
            _seed_runs=[seed_run],
            _seed_schema=schema,
            _skip_materialize=True,
        )
        if pause_tail:
            j = job_manager.get_job(job_id) or job
            return jsonify(
                {
                    "success": True,
                    "job_id": job_id,
                    "status": j.get("status", "awaiting_review"),
                    "requires_review": True,
                    "message": "Multi-source run paused for human review before all workbooks finished.",
                }
            ), 200

        combined_df = (multi_res or {}).get("combined_df")
        if combined_df is None:
            combined_df = pd.DataFrame()
        job["status"] = "completed"
        out_schema = schema_ms if schema_ms is not None else schema
        job["schema"] = out_schema.to_dict() if out_schema is not None else {}
        export_df = normalize_dataframe_dates_for_export(combined_df) if combined_df is not None else None
        tabular_df = export_df if export_df is not None else combined_df
        assign_job_data_preview(job, export_df, head=20)
        job["total_rows"] = len(combined_df) if combined_df is not None else 0
        trace_payloads = (multi_res or {}).get("traces") or []
        overall_confidence = min(
            (float(p.get("overall_confidence", 1.0)) for p in trace_payloads if isinstance(p, dict)),
            default=1.0,
        )
        sr = (multi_res or {}).get("source_runs") or []
        job["source_traces"] = sr
        col_tr = (multi_res or {}).get("collation_trace")
        job["trace"] = _build_multi_source_job_trace(
            source_runs=sr,
            trace_payloads=trace_payloads,
            collation_trace=col_tr,
            overall_confidence=overall_confidence,
            summary_message=f"Collated {len(sr)} sources",
        )
        capture_llm_metadata(job_id)
        _merge_collation_debug_into_job(job, col_tr, (multi_res or {}).get("collation_debug_events"))
        job["overall_confidence"] = overall_confidence

        for step in new_trace.steps:
            job["steps"].append(
                {
                    "step": step.get("module", "Resumed"),
                    "message": step.get("output", ""),
                    "confidence": step.get("confidence", {}).get("score", 0),
                }
            )

        last_batch_sid = (
            str(multi_source_remaining_ids[-1]).strip()
            if multi_source_remaining_ids
            else None
        )
        try:
            from sia.debug.llm_observer import get_observer

            observer = get_observer()
            if observer:
                new_llm_traces = observer.export_for_debug_ui()
                if "llm_traces" in job:
                    existing_ids = {t["trace_id"] for t in job["llm_traces"] if "trace_id" in t}
                    for trace in new_llm_traces:
                        if trace.get("trace_id") not in existing_ids:
                            job["llm_traces"].append(trace)
                else:
                    job["llm_traces"] = new_llm_traces

                from sia.debug.llm_observer import refresh_job_llm_summary

                refresh_job_llm_summary(job)
                new_tools = observer.get_tool_executions()
                if "tool_executions" in job:
                    job["tool_executions"].extend(new_tools)
                else:
                    job["tool_executions"] = new_tools

                # Observer here is from the last process_file in _process_job_sources (remaining batch).
                _merge_observer_hierarchical_events_into_job(
                    job,
                    source_id=last_batch_sid,
                )
                logger.info(f"[HITL] Merged new LLM traces into job {job_id}")
        except Exception as e:
            logger.warning(f"[HITL] Failed to merge resumed LLM traces: {e}")

        if job_id in pending_deletions:
            del pending_deletions[job_id]

        output_format = config.get("output_format", "csv")
        outputs = agent.save_results(out_schema, tabular_df, str(OUTPUT_FOLDER), format=output_format)
        job["output_files"] = outputs

        try:
            _persist_excel_output(job_id, job, tabular_df)
        except Exception as excel_err:
            logger.warning(f"[HITL] Failed to persist Excel output: {excel_err}")

        artifact_ids = list(((multi_res or {}).get("frames_by_source") or {}).keys())
        try:
            _persist_processing_artifact_exports(job_id, job, target_template, tabular_df, artifact_ids)
        except Exception as art_err:
            logger.warning("[HITL] Failed to persist processing artifacts: %s", art_err)

        job_manager.update_job_status(
            job_id,
            "completed",
            f"Processed {len(combined_df)} rows after multi-source HITL resume",
            "Complete",
        )
        logger.info("[HITL] Job %s resumed and completed multi-source batch successfully", job_id)
        return jsonify(
            {
                "success": True,
                "job_id": job_id,
                "status": "completed",
                "rows_processed": len(combined_df),
                "fields_inferred": len(out_schema.fields) if out_schema is not None else 0,
                "overall_confidence": overall_confidence,
            }
        )

    # When the last source finishes inside HITL resume, ``multi_source_remaining_ids`` is empty
    # but earlier sources' frames live on the job — collate here so exports are not last-sheet-only.
    sid_done = str(selected_source_id or "").strip()
    acc = dict(job.get("_multi_source_frames_by_source") or {})
    if df is not None and not df.empty and sid_done:
        acc[sid_done] = df
    job["_multi_source_frames_by_source"] = acc
    main_ids = _expand_relationship_resume_source_ids(job, [])
    nonempty: Dict[str, pd.DataFrame] = {
        k: v for k, v in acc.items() if v is not None and not getattr(v, "empty", True)
    }

    if len(main_ids) >= 2 and len(nonempty) >= 2:
        approved_here = list(job.get("approved_file_relationships") or [])
        merged_sr = list(job.get("source_traces") or [])
        seen_sr_ids = {
            str(x.get("source_id") or "").strip()
            for x in merged_sr
            if isinstance(x, dict) and x.get("source_id") is not None
        }
        if sid_done and sid_done not in seen_sr_ids:
            merged_sr.append(
                {
                    "source_id": sid_done,
                    "sheet_name": str(source.get("sheet_name") or ""),
                    "processing_sheet": str(eff_sheet or ""),
                    "trace": new_trace.to_dict(),
                }
            )
        job["source_traces"] = merged_sr

        if not approved_here:
            _queue_post_execution_relationship_review(
                job_id, job, main_ids, nonempty, schema, merged_sr
            )
            if job.get("_post_execution_relationship_review"):
                capture_llm_metadata(
                    job_id,
                    source_id=sid_done or None,
                    sheet_name=str(source.get("sheet_name") or ""),
                    processing_sheet=str(eff_sheet or ""),
                )
                return jsonify(
                    {
                        "success": True,
                        "job_id": job_id,
                        "status": "awaiting_review",
                        "requires_review": True,
                        "message": "Per-source processing finished; review how outputs are combined before export.",
                    }
                ), 200

        combined_df, col_trace, col_dbg = collate_frames_detailed(
            nonempty, approved_here, target_template=target_template
        )
        combined_df, pc_ev, pc_dbg = _apply_job_deferred_post_collate(job, combined_df)
        if pc_ev:
            col_trace = list(col_trace or []) + [{"step": "post_collate_transforms", "events": pc_ev}]
        if pc_dbg:
            col_dbg = list(col_dbg or []) + _link_deferred_collation_debug_events(
                list(col_dbg or []), pc_dbg
            )
        export_df = normalize_dataframe_dates_for_export(combined_df) if combined_df is not None else None
        tabular_df = export_df if export_df is not None else combined_df
        job["status"] = "completed"
        job["schema"] = schema.to_dict() if schema is not None else {}
        assign_job_data_preview(job, export_df, head=20)
        job["total_rows"] = len(combined_df) if combined_df is not None else 0

        trace_payloads_ms = [x.get("trace") for x in merged_sr if isinstance(x, dict)]
        overall_confidence = min(
            (
                float(p.get("overall_confidence", 1.0))
                for p in trace_payloads_ms
                if isinstance(p, dict)
            ),
            default=float(getattr(new_trace, "overall_confidence", 1.0) or 1.0),
        )
        job["overall_confidence"] = overall_confidence

        col_ui = _collation_trace_to_job_steps(col_trace)
        summary_step = {
            "module": "multi_source_collation",
            "input": "",
            "output": f"Collated {len(nonempty)} sources after multi-source HITL resume",
            "confidence": {
                "score": float(overall_confidence or 1.0),
                "decision": "unknown",
                "rationale": "",
                "signals": [],
            },
        }

        try:
            if "trace" in job and isinstance(job["trace"], dict):
                old_trace = job["trace"]
                new_trace_dict = new_trace.to_dict()
                old_trace["steps"] = old_trace.get("steps", []) + new_trace_dict.get("steps", [])
                old_trace["errors"] = old_trace.get("errors", []) + new_trace_dict.get("errors", [])
                old_trace["warnings"] = old_trace.get("warnings", []) + new_trace_dict.get("warnings", [])
                old_trace["overall_confidence"] = overall_confidence
                old_trace["approval_items"] = new_trace_dict.get("approval_items", [])
                old_trace["planned_rule_actions"] = new_trace_dict.get("planned_rule_actions", [])
                old_trace["hitl_checkpoints"] = new_trace_dict.get("hitl_checkpoints", [])
                if "lifecycle_events" in new_trace_dict:
                    old_trace["lifecycle_events"] = (
                        old_trace.get("lifecycle_events", []) + new_trace_dict.get("lifecycle_events", [])
                    )
                if new_trace_dict.get("judge_result"):
                    old_trace["judge_result"] = new_trace_dict["judge_result"]
                if "debug_events" in new_trace_dict:
                    old_trace["debug_events"] = old_trace.get("debug_events", []) + new_trace_dict.get(
                        "debug_events", []
                    )
                old_trace["source_runs"] = merged_sr
                old_trace["steps"] = list(old_trace.get("steps") or []) + col_ui + [summary_step]
                job["trace"] = old_trace
                capture_llm_metadata(
                    job_id,
                    source_id=sid_done or None,
                    sheet_name=str(source.get("sheet_name") or ""),
                    processing_sheet=str(eff_sheet or ""),
                )
                _merge_collation_debug_into_job(job, col_trace, col_dbg)
            else:
                ntd = new_trace.to_dict()
                flat = _flatten_source_run_trace_steps(merged_sr)
                job["trace"] = {
                    "overall_confidence": overall_confidence,
                    "steps": flat + col_ui + [summary_step],
                    "errors": ntd.get("errors", []),
                    "warnings": ntd.get("warnings", []),
                    "hitl_checkpoints": ntd.get("hitl_checkpoints", []),
                    "source_runs": merged_sr,
                }
                capture_llm_metadata(
                    job_id,
                    source_id=sid_done or None,
                    sheet_name=str(source.get("sheet_name") or ""),
                    processing_sheet=str(eff_sheet or ""),
                )
                _merge_collation_debug_into_job(job, col_trace, col_dbg)
        except Exception as merge_err:
            logger.error("[HITL] Error merging traces (multi resume tail): %s", merge_err)

        for step in new_trace.steps:
            job.setdefault("steps", []).append(
                {
                    "step": step.get("module", "Resumed"),
                    "message": step.get("output", ""),
                    "confidence": step.get("confidence", {}).get("score", 0),
                }
            )

        try:
            from sia.debug.llm_observer import get_observer

            observer = get_observer()
            if observer:
                new_llm_traces = observer.export_for_debug_ui()
                if "llm_traces" in job:
                    existing_ids = {t["trace_id"] for t in job["llm_traces"] if "trace_id" in t}
                    for trace in new_llm_traces:
                        if trace.get("trace_id") not in existing_ids:
                            job["llm_traces"].append(trace)
                else:
                    job["llm_traces"] = new_llm_traces
                from sia.debug.llm_observer import refresh_job_llm_summary

                refresh_job_llm_summary(job)
                new_tools = observer.get_tool_executions()
                if "tool_executions" in job:
                    job["tool_executions"].extend(new_tools)
                else:
                    job["tool_executions"] = new_tools
                _merge_observer_hierarchical_events_into_job(
                    job,
                    source_id=sid_done or None,
                    sheet_name=str(source.get("sheet_name") or ""),
                    processing_sheet=str(eff_sheet or ""),
                )
                logger.info("[HITL] Merged new LLM traces into job %s (multi resume tail)", job_id)
        except Exception as e:
            logger.warning("[HITL] Failed to merge resumed LLM traces: %s", e)

        if job_id in pending_deletions:
            del pending_deletions[job_id]

        output_format = config.get("output_format", "csv")
        outputs = agent.save_results(schema, tabular_df, str(OUTPUT_FOLDER), format=output_format)
        job["output_files"] = outputs

        try:
            _persist_excel_output(job_id, job, tabular_df)
        except Exception as excel_err:
            logger.warning("[HITL] Failed to persist Excel output: %s", excel_err)

        artifact_ids = list(nonempty.keys())
        try:
            _persist_processing_artifact_exports(job_id, job, target_template, tabular_df, artifact_ids)
        except Exception as art_err:
            logger.warning("[HITL] Failed to persist processing artifacts: %s", art_err)

        job_manager.update_job_status(
            job_id,
            "completed",
            f"Processed {len(combined_df) if combined_df is not None else 0} combined rows after multi-source HITL resume",
            "Complete",
        )
        logger.info("[HITL] Job %s resumed and completed (multi-source collated tail)", job_id)
        return jsonify(
            {
                "success": True,
                "job_id": job_id,
                "status": "completed",
                "rows_processed": len(combined_df) if combined_df is not None else 0,
                "fields_inferred": len(schema.fields) if schema is not None else 0,
                "overall_confidence": overall_confidence,
            }
        )

    job["status"] = "completed"
    job["schema"] = schema.to_dict()
    export_df = normalize_dataframe_dates_for_export(df) if df is not None else None
    assign_job_data_preview(job, export_df, head=20)
    job["total_rows"] = len(df)

    try:
        if "trace" in job and isinstance(job["trace"], dict):
            old_trace = job["trace"]
            new_trace_dict = new_trace.to_dict()
            old_trace["steps"] = old_trace.get("steps", []) + new_trace_dict.get("steps", [])
            old_trace["errors"] = old_trace.get("errors", []) + new_trace_dict.get("errors", [])
            old_trace["warnings"] = old_trace.get("warnings", []) + new_trace_dict.get("warnings", [])
            old_trace["overall_confidence"] = new_trace_dict.get("overall_confidence", 0)
            old_trace["approval_items"] = new_trace_dict.get("approval_items", [])
            old_trace["planned_rule_actions"] = new_trace_dict.get("planned_rule_actions", [])
            old_trace["hitl_checkpoints"] = new_trace_dict.get("hitl_checkpoints", [])
            # Preserve prior lifecycle events, then append new ones from the resumed run
            if "lifecycle_events" in new_trace_dict:
                old_trace["lifecycle_events"] = (
                    old_trace.get("lifecycle_events", []) + new_trace_dict.get("lifecycle_events", [])
                )

            if new_trace_dict.get("judge_result"):
                old_trace["judge_result"] = new_trace_dict["judge_result"]
            if "debug_events" in new_trace_dict:
                old_trace["debug_events"] = old_trace.get("debug_events", []) + new_trace_dict["debug_events"]

            job["trace"] = old_trace
            capture_llm_metadata(
                job_id,
                source_id=str(selected_source_id).strip() if selected_source_id else None,
                sheet_name=str(source.get("sheet_name") or ""),
                processing_sheet=str(eff_sheet or ""),
            )
        else:
            job["trace"] = new_trace.to_dict()
            capture_llm_metadata(
                job_id,
                source_id=str(selected_source_id).strip() if selected_source_id else None,
                sheet_name=str(source.get("sheet_name") or ""),
                processing_sheet=str(eff_sheet or ""),
            )
    except Exception as merge_err:
        logger.error(f"[HITL] Error merging traces: {merge_err}")
        job["trace"] = new_trace.to_dict()

    for step in new_trace.steps:
        job["steps"].append({
            "step": step.get("module", "Resumed"),
            "message": step.get("output", ""),
            "confidence": step.get("confidence", {}).get("score", 0)
        })

    try:
        from sia.debug.llm_observer import get_observer
        observer = get_observer()
        if observer:
            new_llm_traces = observer.export_for_debug_ui()
            if "llm_traces" in job:
                existing_ids = {t["trace_id"] for t in job["llm_traces"] if "trace_id" in t}
                for trace in new_llm_traces:
                    if trace.get("trace_id") not in existing_ids:
                        job["llm_traces"].append(trace)
            else:
                job["llm_traces"] = new_llm_traces

            from sia.debug.llm_observer import refresh_job_llm_summary

            refresh_job_llm_summary(job)
            new_tools = observer.get_tool_executions()
            if "tool_executions" in job:
                job["tool_executions"].extend(new_tools)
            else:
                job["tool_executions"] = new_tools

            _merge_observer_hierarchical_events_into_job(
                job,
                source_id=str(selected_source_id).strip() if selected_source_id else None,
                sheet_name=str(source.get("sheet_name") or ""),
                processing_sheet=str(eff_sheet or ""),
            )
            logger.info(f"[HITL] Merged new LLM traces into job {job_id}")
    except Exception as e:
        logger.warning(f"[HITL] Failed to merge resumed LLM traces: {e}")

    if job_id in pending_deletions:
        del pending_deletions[job_id]

    tabular_df = export_df if export_df is not None else df
    output_format = config.get('output_format', 'csv')
    outputs = agent.save_results(schema, tabular_df, str(OUTPUT_FOLDER), format=output_format)
    job["output_files"] = outputs

    try:
        _persist_excel_output(job_id, job, tabular_df)
    except Exception as excel_err:
        logger.warning(f"[HITL] Failed to persist Excel output: {excel_err}")

    logger.info(f"[HITL] Job {job_id} resumed and completed successfully")
    return jsonify({
        "success": True,
        "job_id": job_id,
        "status": "completed",
        "rows_processed": len(df),
        "fields_inferred": len(schema.fields),
        "overall_confidence": new_trace.overall_confidence
    })


def _create_relationship_review(job_id: str, job: Dict[str, Any], source_ids: List[str]):
    target_template = load_target_template_for_job(job)
    packet = build_context_packet(
        job,
        target_template=target_template,
        selected_source_id=source_ids[0] if source_ids else None,
    )
    proposals = propose_file_relationships(packet.get("available_source_summaries") or [])
    if not proposals:
        return None

    checkpoint = HITLManager().create_checkpoint(
        CheckpointType.FILE_RELATIONSHIP_REVIEW,
        {"current_step": "relationship_review", "confidence_trajectory": [min((item.get("confidence", 0.7) for item in proposals), default=0.7)]},
        "Multiple sources were uploaded. Review the inferred file relationships before processing continues.",
        title="File Relationship Review",
        description="Review how the uploaded files should be collated before processing continues.",
        severity="medium",
        available_actions=["approve", "modify", "cancel"],
        recommended_action="approve",
        trigger_data={
            "relationship_proposals": proposals,
            "sources": packet.get("available_source_summaries") or [],
        },
    ).to_dict()
    checkpoint["pending_state"] = {
        "workflow": "multi_source",
        "source_ids": source_ids,
        "process_all_sources": True,
        "selected_source_id": source_ids[0] if source_ids else None,
    }
    job["relationship_proposals"] = proposals
    job_manager.add_checkpoint_to_review(job_id, checkpoint)
    job_manager.update_job_status(job_id, "awaiting_review", checkpoint["trigger_reason"], "Relationship Review")
    return jsonify({
        "success": True,
        "job_id": job_id,
        "status": "awaiting_review",
        "message": checkpoint["trigger_reason"],
        "checkpoint_count": 1,
        "requires_review": True,
    })


def _queue_post_execution_relationship_review(
    job_id: str,
    job: Dict[str, Any],
    source_ids: List[str],
    frames_by_source: Dict[str, pd.DataFrame],
    last_schema: Any,
    source_runs: List[Dict[str, Any]],
) -> None:
    """Per-source pipelines finished; defer union/join choice until after execution.

    Stores output frames on the job for collation-only resolve. Analyst reviews the same
    checkpoint type as the pre-process flow, but copy reflects post-execution context.
    """
    target_template = load_target_template_for_job(job)
    packet = build_context_packet(
        job,
        target_template=target_template,
        selected_source_id=source_ids[0] if source_ids else None,
    )
    summaries = list(packet.get("available_source_summaries") or [])
    proposals = propose_file_relationships(summaries)
    if not proposals:
        return

    job["_pending_collation_frames"] = {str(k): v for k, v in (frames_by_source or {}).items()}
    job["_post_execution_relationship_review"] = True
    job["_pending_collation_schema"] = last_schema
    job["source_traces"] = list(source_runs or [])

    execution_preview: Dict[str, Any] = {}
    for sid, fr in (frames_by_source or {}).items():
        if fr is None or not hasattr(fr, "shape"):
            continue
        execution_preview[str(sid)] = {
            "rows": int(len(fr)),
            "columns": [str(c) for c in list(fr.columns)[:48]],
        }

    reason = (
        "Per-source processing finished. Choose how outputs should be combined "
        "(union / join / independent) before the final Excel and CSV exports are built."
    )
    checkpoint = HITLManager().create_checkpoint(
        CheckpointType.FILE_RELATIONSHIP_REVIEW,
        {"current_step": "relationship_review_post_exec", "confidence_trajectory": [min((item.get("confidence", 0.7) for item in proposals), default=0.7)]},
        reason,
        title="Combine outputs (post-execution)",
        description=(
            "Each source has been transformed. Review how to merge the resulting tables "
            "into one deliverable. Row/column previews reflect executed outputs."
        ),
        severity="medium",
        available_actions=["approve", "modify", "cancel"],
        recommended_action="approve",
        trigger_data={
            "relationship_proposals": proposals,
            "sources": summaries,
            "post_execution_preview": execution_preview,
            "post_execution": True,
        },
    ).to_dict()
    checkpoint["pending_state"] = {
        "workflow": "multi_source_post_exec",
        "source_ids": source_ids,
        "process_all_sources": True,
        "selected_source_id": source_ids[0] if source_ids else None,
    }
    job["relationship_proposals"] = proposals
    job_manager.add_checkpoint_to_review(job_id, checkpoint)
    job_manager.update_job_status(job_id, "awaiting_review", reason, "Relationship Review")


def _execute_collation_only_after_review(job_id: str) -> None:
    """Apply approved relationships to pre-computed per-source frames; export only (no LangGraph re-run)."""
    import traceback as tb

    job = job_manager.get_job(job_id)
    if not job:
        logger.error("[REL_REVIEW] Collation-only worker: job %s not found", job_id)
        return
    try:
        frames = job.get("_pending_collation_frames") or {}
        if not isinstance(frames, dict) or not frames:
            logger.warning("[REL_REVIEW] Collation-only: no pending frames for job %s", job_id)
            job_manager.update_job_status(job_id, "error", "Missing pending output frames for collation.", "Error")
            return

        approved = list(job.get("approved_file_relationships") or [])
        target_tpl = load_target_template_for_job(job)
        df, col_trace, col_dbg = collate_frames_detailed(frames, approved, target_template=target_tpl)
        df, pc_ev, pc_dbg = _apply_job_deferred_post_collate(job, df)
        if pc_ev:
            col_trace = list(col_trace or []) + [{"step": "post_collate_transforms", "events": pc_ev}]
        if pc_dbg:
            col_dbg = list(col_dbg or []) + _link_deferred_collation_debug_events(
                list(col_dbg or []), pc_dbg
            )
        schema_obj = job.pop("_pending_collation_schema", None)
        job.pop("_pending_collation_frames", None)
        job.pop("_post_execution_relationship_review", None)

        if schema_obj is None and df is not None and not df.empty:
            from sia.models.schema import InferredSchema

            schema_obj = InferredSchema.from_dataframe(df, "combined")
        elif schema_obj is None:
            from sia.models.schema import InferredSchema

            schema_obj = InferredSchema("combined")

        export_df = normalize_dataframe_dates_for_export(df) if df is not None else None
        tabular_df = export_df if export_df is not None else df
        assign_job_data_preview(job, export_df, head=20)
        job["total_rows"] = len(df) if df is not None else 0
        job["schema"] = schema_obj.to_dict() if hasattr(schema_obj, "to_dict") else {}

        config = load_user_config()
        output_format = config.get("output_format", "csv")
        agent = _build_agent_from_config(config)
        outputs = agent.save_results(schema_obj, tabular_df, str(OUTPUT_FOLDER), format=output_format)
        job["output_files"] = outputs
        try:
            _persist_excel_output(job_id, job, tabular_df)
        except Exception as excel_err:
            logger.warning("[REL_REVIEW] Collation-only Excel export failed: %s", excel_err)

        job["trace"] = {
            "overall_confidence": float(job.get("overall_confidence") or 1.0),
            "steps": _collation_trace_to_job_steps(col_trace)
            + [
                {
                    "module": "post_execution_collation",
                    "input": "",
                    "output": f"Collated {len(frames)} sources after relationship review",
                    "confidence": {
                        "score": float(job.get("overall_confidence") or 1.0),
                        "decision": "unknown",
                        "rationale": "",
                        "signals": [],
                    },
                }
            ],
            "hitl_checkpoints": [],
            "source_runs": job.get("source_traces") or [],
        }

        _merge_collation_debug_into_job(job, col_trace, col_dbg)

        job_manager.update_job_status(
            job_id,
            "completed",
            f"Exported {len(df) if df is not None else 0} combined rows after relationship review",
            "Complete",
        )
        logger.info("[REL_REVIEW] Job %s collation-only export completed", job_id)
    except Exception as e:
        error_traceback = tb.format_exc()
        logger.error("[REL_REVIEW] Collation-only failed for job %s:\n%s", job_id, error_traceback)
        job_manager.update_job_status(job_id, "error", str(e), "Error")
        job["error_details"] = error_traceback
    finally:
        j = job_manager.get_job(job_id)
        if j is not None:
            j["_relationship_resume_running"] = False


def _process_job_sources(
    job_id: str,
    job: Dict[str, Any],
    config: Dict[str, Any],
    target_template: Dict[str, Any],
    source_ids: List[str],
    *,
    _seed_frames: Optional[Dict[str, pd.DataFrame]] = None,
    _seed_runs: Optional[List[Dict[str, Any]]] = None,
    _seed_schema: Any = None,
    _skip_materialize: bool = False,
):
    """Run ``process_file`` for each ``source_id`` and collate.

    When a source hits HITL mid-batch, ``trace.pending_state`` is tagged with
    ``multi_source_remaining_ids`` so :func:`_run_resumed_processing` can finish
    the current workbook and then continue this loop for the rest.

    Seeded parameters support that continuation after resume (skip re-materializing
    workbooks that were already refreshed for the full batch).
    """
    agent = _build_agent_from_config(config)
    template_path = job.get("template_path") or job.get("target_template_path")
    frames_by_source: Dict[str, pd.DataFrame] = dict(_seed_frames or {})
    source_runs: List[Dict[str, Any]] = list(_seed_runs or [])
    trace_dicts: List[Dict[str, Any]] = []
    for _run in source_runs:
        tr = _run.get("trace")
        if isinstance(tr, dict):
            trace_dicts.append(tr)
    last_schema = _seed_schema

    source_ids, multi_block_expanded = expand_multi_block_source_ids(job, list(source_ids or []))
    if multi_block_expanded:
        ensure_block_union_relationship(job, source_ids)
        job["_multi_block_active_batch"] = True
        logger.info(
            "[MULTI_BLOCK] job=%s processing %d block run(s) on sheet (collate after each block)",
            job_id,
            len(source_ids),
        )

    defer_graph_hitl = len(source_ids) > 1 and not (job.get("approved_file_relationships") or [])

    if not _skip_materialize:
        materialize_ids = []
        for sid in source_ids:
            if is_block_virtual_source_id(sid):
                from sia.agent.multi_block_sheet import parse_block_virtual_source_id

                parsed = parse_block_virtual_source_id(sid)
                if parsed:
                    materialize_ids.append(parsed[0])
            else:
                materialize_ids.append(sid)
        _materialize_clean_templates_at_process_start(
            job, job_id, list(dict.fromkeys(materialize_ids))
        )

    ensure_bootstrap_and_route(job, job_id, source_ids)

    if len(source_ids) > 1:
        job["_multi_source_frames_by_source"] = dict(frames_by_source)

    for idx, source_id in enumerate(source_ids):
        if idx == 0 and len(source_ids) > 1 and _seed_frames is None and not _seed_runs:
            clear_multi_source_ledger(job)
        block_run = resolve_block_run(job, str(source_id or "").strip())
        if block_run:
            parent_id = str(block_run.get("parent_source_id") or "").strip()
            source = resolve_job_source(job, source_id=parent_id)
            if not source:
                logger.warning(
                    "[PROCESS] job=%s skipping block run %r — parent source %r not in registry",
                    job_id,
                    source_id,
                    parent_id,
                )
                continue
            canonical_id = str(source_id).strip()
            block = dict(block_run.get("block") or {})
            scoped_source = build_scoped_source_for_block(
                job,
                parent_id,
                block,
                sheet_name=str(source.get("sheet_name") or ""),
            )
            job.setdefault("source_scope_registry", {})[canonical_id] = scoped_source
        else:
            source = resolve_job_source(job, source_id=source_id)
            if not source:
                logger.warning(
                    "[PROCESS] Skipping unknown source_id %r for job %s (not in source_registry).",
                    source_id,
                    job_id,
                )
                continue

            canonical_id = str(source.get("source_id") or source_id).strip()

            scope_registry = job.get("source_scope_registry") or {}
            scoped_raw = scope_registry.get(source_id) or scope_registry.get(canonical_id)
            if scoped_raw is None and source_id is not None:
                sid = str(source_id).strip()
                scoped_raw = scope_registry.get(sid)
                if scoped_raw is None:
                    for k, v in scope_registry.items():
                        if str(k).strip() == sid or str(k).strip() == canonical_id:
                            scoped_raw = v
                            break
            scoped_source = dict(scoped_raw or {})

        context_packet = build_context_packet(
            job,
            target_template=target_template,
            selected_sheet=source.get("sheet_name"),
            selected_source_id=canonical_id,
            defer_file_relationship_graph_hitl=defer_graph_hitl,
        )
        if block_run:
            context_packet = enrich_context_packet_for_block_run(
                context_packet,
                virtual_source_id=canonical_id,
                parent_source_id=str(block_run.get("parent_source_id") or ""),
                block=dict(block_run.get("block") or {}),
            )
        job["context_packet"] = context_packet
        _record_job_debug(
            job,
            label="Context packet built",
            phase="context",
            parent_key="process_root",
            source_id=canonical_id,
            sheet_name=str(source.get("sheet_name") or ""),
            module="context_packet",
            operation="build",
            status="success",
            summary=f"Context ready for source {canonical_id}",
            metadata={
                "defer_file_relationship_graph_hitl": defer_graph_hitl,
                "multi_source_active_batch": len(source_ids) > 1,
            },
        )
        run_parent = _record_job_debug(
            job,
            label=f"Agent run: {canonical_id}",
            phase="agent_run",
            parent_key="process_root",
            set_parent_key=f"agent_run:{canonical_id}",
            source_id=canonical_id,
            sheet_name=str(source.get("sheet_name") or ""),
            module="agent",
            operation="process_file",
            status="info",
            summary=f"Starting LangGraph for {str(source.get('sheet_name') or canonical_id)}",
        )
        _append_job_state_snapshot(
            job,
            label="job.multi_source.source_start",
            phase="before_process_file",
            source_id=canonical_id,
            sheet_name=str(source.get("sheet_name") or ""),
            context_packet=context_packet,
            decision={
                "source_ids": [str(x) for x in source_ids],
                "multi_source_active_batch": len(source_ids) > 1,
                "defer_file_relationship_graph_hitl": defer_graph_hitl,
                "processing_route": job.get("processing_route"),
            },
        )
        eff_path, eff_sheet, eff_scoped, used_materialized = resolve_processing_workbook(
            job,
            canonical_id,
            source["file_path"],
            str(source.get("sheet_name") or ""),
            scoped_source,
        )
        if used_materialized:
            logger.info(
                "[PROCESS] Multi-source: using materialized clean workbook for %s: %s",
                canonical_id,
                eff_path,
            )

        schema, df, trace = agent.process_file(
            eff_path,
            target_template_path=template_path,
            resume_state={
                "sheet_name": eff_sheet,
                "source_id": canonical_id,
                "scoped_source": eff_scoped or scoped_source or None,
                "approved_relationships": list(job.get("approved_file_relationships") or []),
                "materialized_clean_active": used_materialized,
                "multi_source_active_batch": len(source_ids) > 1,
                "multi_block_active_batch": bool(job.get("_multi_block_active_batch")),
                "process_blocks_separately": bool(block_run),
            },
            context_packet=context_packet,
        )
        _record_job_debug(
            job,
            label=f"Agent run finished: {canonical_id}",
            phase="agent_run",
            parent_event_id=run_parent,
            source_id=canonical_id,
            sheet_name=str(source.get("sheet_name") or ""),
            module="agent",
            operation="process_file",
            status="success" if not getattr(trace, "hitl_pending", False) else "pending",
            summary=(
                "HITL pause"
                if getattr(trace, "hitl_pending", False)
                else f"Rows: {len(df) if df is not None else 0}"
            ),
            metadata={
                "hitl_pending": bool(getattr(trace, "hitl_pending", False)),
                "overall_confidence": getattr(trace, "overall_confidence", None),
                "rows": len(df) if df is not None else 0,
            },
        )
        if getattr(trace, "hitl_pending", False):
            remainder = [str(s) for s in source_ids[idx + 1 :]]
            if remainder:
                ps = dict(getattr(trace, "pending_state", None) or {})
                ps["multi_source_remaining_ids"] = remainder
                trace.pending_state = ps
                logger.info(
                    "[PROCESS] Multi-source HITL pause on %s; after resume, %d more source(s) will run: %s",
                    canonical_id,
                    len(remainder),
                    remainder,
                )
            job["_multi_source_frames_by_source"] = dict(frames_by_source)
            pause_payload, pause_code = _apply_hitl_pause(job_id, job, trace, df)
            if pause_code >= 400:
                err_msg = str(pause_payload.get("error") or pause_payload)
                job["error_details"] = err_msg
                job_manager.update_job_status(job_id, "error", err_msg, "Error")
            return None, None, True

        trace_dict = trace.to_dict()
        if len(source_ids) > 1:
            dft = list(trace_dict.get("deferred_post_collate_tools") or [])
            record_source_run_ledger(
                job,
                canonical_id,
                {
                    "sheet_name": str(source.get("sheet_name") or ""),
                    "processing_sheet": str(eff_sheet or ""),
                },
                dft,
            )
            _append_job_state_snapshot(
                job,
                label="job.multi_source.source_done",
                phase="after_process_file",
                source_id=canonical_id,
                sheet_name=str(source.get("sheet_name") or ""),
                context_packet=context_packet,
                decision={
                    "deferred_post_collate_tools": [
                        t.get("tool") for t in dft if isinstance(t, dict)
                    ],
                    "source_execution_registry_count": len(job.get("source_execution_registry") or []),
                    "frame_recorded": df is not None and not getattr(df, "empty", True),
                },
            )
        trace_dicts.append(trace_dict)
        source_runs.append(
            {
                "source_id": canonical_id,
                "sheet_name": str(source.get("sheet_name") or ""),
                "processing_sheet": str(eff_sheet or ""),
                "trace": trace_dict,
            }
        )
        last_schema = schema
        if df is not None and not df.empty:
            frames_by_source[canonical_id] = df
        if len(source_ids) > 1:
            job["_multi_source_frames_by_source"] = dict(frames_by_source)

        capture_llm_metadata(
            job_id,
            source_id=canonical_id,
            sheet_name=str(source.get("sheet_name") or ""),
            processing_sheet=str(eff_sheet or ""),
        )

    if not source_runs:
        reg_ids = [
            str(s.get("source_id"))
            for s in (job.get("source_registry") or [])
            if isinstance(s, dict) and s.get("source_id") is not None
        ]
        msg = (
            f"No registry rows matched requested source_ids {source_ids!r}. "
            f"Registry source_ids: {reg_ids}."
        )
        logger.error("[PROCESS] job=%s %s", job_id, msg)
        job["error_details"] = msg
        job_manager.update_job_status(job_id, "error", msg, "Error")
        return None, None, True

    if len(source_ids) > 1 and not (job.get("approved_file_relationships") or []):
        if len(frames_by_source) >= 2:
            _queue_post_execution_relationship_review(
                job_id, job, source_ids, frames_by_source, last_schema, source_runs
            )
            if job.get("_post_execution_relationship_review"):
                logger.info(
                    "[PROCESS] job=%s queued post-execution relationship review (%d sources)",
                    job_id,
                    len(frames_by_source),
                )
                return (
                    last_schema,
                    {
                        "frames_by_source": frames_by_source,
                        "traces": trace_dicts,
                        "source_runs": source_runs,
                        "combined_df": None,
                    },
                    True,
                )

    _record_job_debug(
        job,
        label="Collation started",
        phase="collation",
        parent_key="process_root",
        set_parent_key="collation_root",
        source_id="__collation__",
        module="collation",
        operation="start",
        status="info",
        summary=f"Merging {len(frames_by_source)} source frame(s)",
        metadata={"frame_sources": list(frames_by_source.keys())},
    )
    combined_df, collation_trace, collation_debug_events = collate_frames_detailed(
        frames_by_source,
        job.get("approved_file_relationships") or [],
        target_template=target_template,
    )
    _record_job_debug(
        job,
        label="Collation completed (pre-deferred transforms)",
        phase="collation",
        parent_key="collation_root",
        source_id="__collation__",
        module="collation",
        operation="union_and_dedupe",
        status="success",
        summary="Union stack and duplicate handling finished",
        metadata={
            "collation_steps": [
                s.get("step") for s in (collation_trace or []) if isinstance(s, dict)
            ],
        },
    )
    _append_job_state_snapshot(
        job,
        label="job.multi_source.collation",
        phase="after_duplicate_check_before_deferred",
        source_id="__collation__",
        sheet_name="",
        context_packet=job.get("context_packet") or {},
        decision={
            "frame_sources": list(frames_by_source.keys()),
            "collation_steps": [
                s.get("step") for s in (collation_trace or []) if isinstance(s, dict)
            ],
            "deferred_by_source": {
                str(k): [t.get("tool") for t in v if isinstance(t, dict)]
                for k, v in (job.get("_deferred_post_collate_by_source") or {}).items()
                if isinstance(v, list)
            },
        },
    )
    combined_df, pc_ev, pc_dbg = _apply_job_deferred_post_collate(job, combined_df)
    if pc_ev:
        collation_trace = list(collation_trace or []) + [{"step": "post_collate_transforms", "events": pc_ev}]
    if pc_dbg:
        collation_debug_events = list(collation_debug_events or []) + _link_deferred_collation_debug_events(
            list(collation_debug_events or []), pc_dbg
        )
        _record_job_debug(
            job,
            label="Deferred post-collate transforms",
            phase="collation",
            parent_key="collation_root",
            source_id="__collation__",
            module="collation.post_collate_transforms",
            operation="execute",
            status="success",
            summary=f"Applied {len(pc_ev)} deferred grain tool run(s)",
            metadata={
                "post_collate_events": pc_ev,
                "tools": [e.get("tool") for e in pc_ev if isinstance(e, dict)],
            },
        )
    _append_job_state_snapshot(
        job,
        label="job.multi_source.collation",
        phase="after_deferred_transforms",
        source_id="__collation__",
        sheet_name="",
        context_packet=job.get("context_packet") or {},
        decision={
            "post_collate_events": pc_ev,
            "combined_rows": len(combined_df) if combined_df is not None else 0,
        },
    )
    return (
        last_schema,
        {
            "frames_by_source": frames_by_source,
            "traces": trace_dicts,
            "source_runs": source_runs,
            "combined_df": combined_df,
            "collation_trace": collation_trace,
            "collation_debug_events": collation_debug_events,
        },
        None,
    )


def _expand_relationship_resume_source_ids(job: Dict[str, Any], seed_ids: List[Any]) -> List[str]:
    """Union every source that must run after relationship review.

    Prefer IDs from the checkpoint, then every ``source_id`` mentioned in approved
    relationships (so a union of two files cannot collapse to one stale ID), then
    main rows from ``source_registry`` if still short.
    """
    out: List[str] = []
    seen: set[str] = set()

    def add(raw: Any) -> None:
        if raw is None:
            return
        s = str(raw).strip()
        if not s or s in seen:
            return
        seen.add(s)
        out.append(s)

    for x in seed_ids or []:
        add(x)
    for rel in job.get("approved_file_relationships") or []:
        if not isinstance(rel, dict):
            continue
        for x in rel.get("source_ids") or []:
            add(x)
    if len(out) < 2:
        for reg in job.get("source_registry") or []:
            if not isinstance(reg, dict):
                continue
            if reg.get("contains_main_data", True) or not reg.get("contains_reference_data"):
                add(reg.get("source_id"))
    return out


def _execute_relationship_review_resume(job_id: str, source_ids: List[str]) -> None:
    """Apply approved relationships and run multi-source collation + export. Background thread (HTTP returns immediately)."""
    import traceback as tb

    job = job_manager.get_job(job_id)
    if not job:
        logger.error("[REL_REVIEW] Worker: job %s not found", job_id)
        return
    try:
        config = load_user_config()
        target_template = load_target_template_for_job(job)
        sid_list = [s for s in (source_ids or []) if s]
        if not sid_list:
            registry = list(job.get("source_registry") or [])
            sid_list = [
                s.get("source_id")
                for s in registry
                if s.get("contains_main_data", True) or not s.get("contains_reference_data")
            ] or [s.get("source_id") for s in registry if s.get("source_id")]
        sid_list = _expand_relationship_resume_source_ids(job, sid_list)
        logger.info(
            "[REL_REVIEW] Worker job=%s expanded source_ids for resume: %s",
            job_id,
            sid_list,
        )

        if len(sid_list) > 1:
            try:
                from sia.debug.llm_observer import clear_observer

                clear_observer()
            except Exception:
                pass
            job["debug_events"] = []
            job["llm_traces"] = []
            job["tool_executions"] = []

        schema, multi_result, pause_response = _process_job_sources(job_id, job, config, target_template, sid_list)
        if pause_response is not None:
            logger.info("[REL_REVIEW] Worker: HITL pause or error during resume for job %s", job_id)
            return

        df = (multi_result or {}).get("combined_df")
        if df is None:
            df = pd.DataFrame()
        job["status"] = "completed"
        job["schema"] = schema.to_dict() if schema is not None else {}
        export_df = normalize_dataframe_dates_for_export(df) if df is not None else None
        tabular_df = export_df if export_df is not None else df
        assign_job_data_preview(job, export_df, head=20)
        job["total_rows"] = len(df) if df is not None else 0
        overall_confidence = min(
            (item.get("overall_confidence", 1.0) for item in (multi_result or {}).get("traces", [])),
            default=1.0,
        )
        sr = (multi_result or {}).get("source_runs") or []
        job["source_traces"] = sr
        col_tr = (multi_result or {}).get("collation_trace")
        job["trace"] = _build_multi_source_job_trace(
            source_runs=sr,
            trace_payloads=(multi_result or {}).get("traces") or [],
            collation_trace=col_tr,
            overall_confidence=overall_confidence,
            summary_message=f"Collated {len(sid_list)} sources after relationship review",
        )
        capture_llm_metadata(job_id)
        _merge_collation_debug_into_job(job, col_tr, (multi_res or {}).get("collation_debug_events"))
        job["overall_confidence"] = overall_confidence

        output_format = config.get("output_format", "csv")
        agent = _build_agent_from_config(config)
        outputs = agent.save_results(schema, tabular_df, str(OUTPUT_FOLDER), format=output_format)
        job["output_files"] = outputs

        try:
            _persist_excel_output(job_id, job, tabular_df)
        except Exception as excel_err:
            logger.warning("[REL_REVIEW] Failed to save Excel export: %s", excel_err)

        job_manager.update_job_status(
            job_id,
            "completed",
            f"Processed {len(df)} rows after relationship review",
            "Complete",
        )
        logger.info("[REL_REVIEW] Job %s completed after relationship review (background)", job_id)
    except Exception as e:
        error_traceback = tb.format_exc()
        logger.error("[REL_REVIEW] Resume failed for job %s:\n%s", job_id, error_traceback)
        job_manager.update_job_status(job_id, "error", str(e), "Error")
        job["error_details"] = error_traceback
    finally:
        j = job_manager.get_job(job_id)
        if j is not None:
            j["_relationship_resume_running"] = False


def index_to_excel_column(n):
    """Convert 0-indexed column number to Excel column name (A, B, C...)."""
    result = ""
    while n >= 0:
        result = chr(n % 26 + ord('A')) + result
        n = n // 26 - 1
    return result

def index_to_excel_range(start_row, end_row, start_col, end_col):
    """Convert 0-indexed coordinates to Excel range (e.g., A1:C10)."""
    # Excel is 1-indexed for rows
    s_col = index_to_excel_column(start_col)
    e_col = index_to_excel_column(end_col)
    return f"{s_col}{start_row+1}:{e_col}{end_row+1}"

def make_unique_columns(columns):
    """Ensure a list of column names is unique by adding suffixes."""
    new_cols = []
    seen = {}
    for col in columns:
        # Normalize to string and handle nulls
        col_str = str(col) if pd.notna(col) else "Unnamed"
        if col_str not in seen:
            seen[col_str] = 0
            new_cols.append(col_str)
        else:
            seen[col_str] += 1
            new_cols.append(f"{col_str}.{seen[col_str]}")
    return new_cols

def load_target_columns_for_job(job: Dict[str, Any]) -> List[str]:
    """Load semantic-mapping target columns: date, uid hierarchy (rest), metrics."""
    candidate_paths = []
    if job.get("template_path"):
        candidate_paths.append(Path(job["template_path"]))
    if job.get("target_template_path"):
        candidate_paths.append(Path(job["target_template_path"]))
    default_template = CONFIG_DIR / "target_template.json"
    candidate_paths.append(default_template)

    for path in candidate_paths:
        try:
            if path.exists():
                with open(path, "r", encoding="utf-8") as f:
                    tpl = normalize_target_template(json.load(f))
                ok, err = validate_template_shape(tpl)
                if ok:
                    return mapping_target_columns_for_ui(tpl)
                logger.warning("Invalid target template at %s: %s", path, err)
        except Exception as e:
            logger.warning(f"Failed to load target columns from {path}: {e}")
    return []


def load_target_template_for_job(job: Dict[str, Any]) -> Dict[str, Any]:
    """Load the full target template for a job (first file that passes shape validation)."""
    candidate_paths = []
    if job.get("template_path"):
        candidate_paths.append(Path(job["template_path"]))
    if job.get("target_template_path"):
        candidate_paths.append(Path(job["target_template_path"]))
    candidate_paths.append(CONFIG_DIR / "target_template.json")

    for path in candidate_paths:
        try:
            if path.exists():
                with open(path, "r", encoding="utf-8") as f:
                    tpl = normalize_target_template(json.load(f))
                ok, err = validate_template_shape(tpl)
                if ok:
                    return tpl
                logger.warning("Skipping invalid target template at %s: %s", path, err)
        except Exception as e:
            logger.warning(f"Failed to load target template from {path}: {e}")
    return {}


def layout_registry_row_is_main_data(row: Dict[str, Any]) -> bool:
    """Include only Main Data blocks the user approved as data (Treat as Data), not Ignore/Metadata."""
    cat_compact = str(row.get("block_category") or "").strip().lower().replace(" ", "").replace("_", "")
    if cat_compact != "maindata":
        return False
    dec = str(row.get("decision") or "").strip().lower()
    if dec in ("discard", "context"):
        return False
    return dec in ("keep", "approved") or dec == ""


def layout_registry_row_to_block(row: Dict[str, Any]) -> Dict[str, Any]:
    hr = row.get("header_row")
    if hr is None:
        hr = row.get("start_row", 0)
    nested = row.get("coordinates") or {}
    hre = row.get("header_row_end", nested.get("header_row_end"))
    hm = row.get("header_mode", nested.get("header_mode"))
    coords: Dict[str, Any] = {
        "start_row": int(row.get("start_row", 0)),
        "end_row": int(row.get("end_row", 0)),
        "start_col": int(row.get("start_col", 0)),
        "end_col": int(row.get("end_col", 0)),
        "header_row": int(hr),
    }
    if hre is not None:
        coords["header_row_end"] = int(hre)
    if hm is not None:
        coords["header_mode"] = str(hm).strip().lower()
    return {
        "id": str(row.get("block_id") or row.get("layout_id") or "block"),
        "category": "Main Data",
        "decision": "Keep",
        "coordinates": coords,
    }


def resolve_registry_source(job: Dict[str, Any], source_id: Optional[str]) -> Dict[str, Any]:
    if not source_id:
        return {}
    for entry in job.get("source_registry") or []:
        if entry.get("source_id") == source_id:
            return entry
    return {}


def resolve_demarcation_blocks_for_source(
    job: Dict[str, Any],
    job_id: str,
    source_id: Optional[str],
) -> List[Dict[str, Any]]:
    """Prefer per-source batch proposals, then saved scope, then legacy single pending demarcation."""
    batch = job.get("demarcation_batch_proposals") or {}
    if source_id and source_id in batch:
        return list(batch[source_id].get("blocks") or [])
    existing = ((job.get("source_scope_registry") or {}).get(source_id or "") or {})
    main_blocks = existing.get("main_blocks") or []
    if main_blocks:
        return list(main_blocks)
    dem = job_manager.pending_demarcations.get(job_id, {})
    pending_sid = dem.get("source_id")
    if source_id and pending_sid and pending_sid != source_id:
        return []
    return list(dem.get("blocks") or [])


def _decorate_demarcation_proposal_excel_ranges(proposal: Dict[str, Any]) -> None:
    for block in proposal.get("blocks") or []:
        coords = block.get("coordinates", {})
        if coords:
            block["excel_range"] = index_to_excel_range(
                coords.get("start_row", 0),
                coords.get("end_row", 0),
                coords.get("start_col", 0),
                coords.get("end_col", 0),
            )


def invert_column_mappings_to_targets(mappings: List[Dict[str, Any]]) -> Dict[str, str]:
    """Map target_column -> source column_name for the mapping matrix."""
    out: Dict[str, str] = {}
    for m in mappings or []:
        if not isinstance(m, dict):
            continue
        tcol = m.get("target_column")
        src = m.get("column_name")
        if tcol and tcol != "No match" and src:
            out[str(tcol)] = str(src)
    return out


def blocks_to_layout_registry(job: Dict[str, Any], blocks: List[Dict[str, Any]], sheet_name: str = None, source_id: str = None) -> List[Dict[str, Any]]:
    """Normalize UI demarcation blocks into persistent layout registry rows."""
    source_id = job_manager.get_source_id(job["id"], sheet_name or (job.get("scoped_source") or {}).get("sheet_name"), source_id=source_id)
    registry_rows = []
    for idx, block in enumerate(blocks or []):
        coords = block.get("coordinates", block) if isinstance(block, dict) else {}
        hr = int(coords.get("header_row", coords.get("start_row", 0)))
        hre = coords.get("header_row_end")
        hm = coords.get("header_mode")
        row_out: Dict[str, Any] = {
            "layout_id": str(block.get("layout_id") or f"{source_id}:layout:{idx}"),
            "source_id": source_id,
            "sheet_name": sheet_name,
            "block_id": str(block.get("id") or f"block_{idx}"),
            "block_label": block.get("label", f"Block {idx + 1}"),
            "block_category": block.get("category", "main_data"),
            "start_row": int(coords.get("start_row", 0)),
            "end_row": int(coords.get("end_row", 0)),
            "start_col": int(coords.get("start_col", 0)),
            "end_col": int(coords.get("end_col", 0)),
            "header_row": hr,
            "decision": block.get("decision", block.get("ai_suggestion", "approved")),
            "confidence": float(block.get("confidence", 1.0) or 1.0),
            "approved_by": "user",
            "approved_at": datetime.now().isoformat(),
            "layout_version": 1,
            "coordinates": {
                "start_row": int(coords.get("start_row", 0)),
                "end_row": int(coords.get("end_row", 0)),
                "start_col": int(coords.get("start_col", 0)),
                "end_col": int(coords.get("end_col", 0)),
                "header_row": hr,
            },
        }
        if hre is not None:
            try:
                ire = int(hre)
                row_out["header_row_end"] = ire
                row_out["coordinates"]["header_row_end"] = ire
            except (TypeError, ValueError):
                pass
        if hm is not None:
            sm = str(hm).strip().lower()
            if sm in ("single", "multi"):
                row_out["header_mode"] = sm
                row_out["coordinates"]["header_mode"] = sm
        registry_rows.append(row_out)
    return registry_rows

def load_user_config():
    """Load user configuration. Secrets come from env vars / .env only."""
    config = {
        "api_key": "",
        "llm_model": "gemini-3-flash-preview",
        "debug_enabled": True,
        # Paused by default: LLM judge adds latency and can force review from FAIL/REVIEW verdicts.
        "enable_llm_judge": False,
    }
    
    if CONFIG_FILE.exists():
        try:
            with open(CONFIG_FILE, 'r') as f:
                persisted = json.load(f)
            for key in SECRET_CONFIG_KEYS:
                persisted.pop(key, None)
            config.update(persisted)
        except Exception as e:
            logger.error(f"Failed to load user config: {e}")
            
    config['keys'] = {}
        
    for provider, env_var in PROVIDER_ENV_MAP.items():
        val = os.environ.get(env_var)
        if val:
            config['keys'][provider] = val
                
    current_model = config.get('llm_model', '')
    provider = 'groq' if current_model.startswith('groq-') else \
               'azure' if current_model.startswith('azure-') else \
               'openai' if current_model.startswith('gpt-') else 'gemini'
        
    config['api_key'] = config['keys'].get(provider) or \
                        os.environ.get('OPENAI_API_KEY') or \
                        os.environ.get('AZURE_OPENAI_API_KEY') or \
                        os.environ.get('GEMINI_API_KEY') or \
                        os.environ.get('GROQ_API_KEY')

    if not config.get('azure_endpoint'):
        config['azure_endpoint'] = os.environ.get('AZURE_OPENAI_ENDPOINT')
    if not config.get('azure_api_version'):
        config['azure_api_version'] = os.environ.get('AZURE_OPENAI_API_VERSION')
    if not config.get('azure_deployment'):
        config['azure_deployment'] = os.environ.get('AZURE_OPENAI_DEPLOYMENT')
    if 'azure_ssl_verify' not in config:
        config['azure_ssl_verify'] = True # Default to True

    _ej = os.environ.get("SIA_ENABLE_LLM_JUDGE")
    if _ej is not None and str(_ej).strip().lower() in ("1", "true", "yes", "on"):
        config["enable_llm_judge"] = True
    elif _ej is not None and str(_ej).strip().lower() in ("0", "false", "no", "off"):
        config["enable_llm_judge"] = False

    return config


def save_user_config(config):
    """Save non-secret user preferences (never persists API keys)."""
    safe = {k: v for k, v in config.items() if k not in SECRET_CONFIG_KEYS}
    CONFIG_DIR.mkdir(exist_ok=True)
    with open(CONFIG_FILE, 'w') as f:
        json.dump(safe, f, indent=2)


@app.route('/')
def index():
    """Redirect to upload page."""
    return redirect('/pages/upload.html')


@app.route('/pages/<path:filename>')
def serve_page(filename):
    """Serve pages from web/pages directory."""
    from flask import make_response
    response = make_response(send_from_directory('web/pages', filename))
    if filename.endswith('.html'):
        response.headers['Content-Type'] = 'text/html; charset=utf-8'
        response.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
        response.headers['Pragma'] = 'no-cache'
        response.headers['Expires'] = '0'
    return response


@app.route('/js/<path:filename>')
def serve_js(filename):
    """Serve JavaScript from web/js directory."""
    from flask import make_response
    response = make_response(send_from_directory('web/js', filename))
    response.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
    response.headers['Pragma'] = 'no-cache'
    response.headers['Expires'] = '0'
    return response


@app.route('/api/config', methods=['GET'])
def get_config():
    """Get current configuration."""
    config = load_user_config()
    if config.get('api_key'):
        config['api_key_set'] = True
        config['api_key_preview'] = config['api_key'][:10] + '...' if len(config['api_key']) > 10 else '***'
    else:
        config['api_key_set'] = False
        config['api_key_preview'] = ''
    for key in SECRET_CONFIG_KEYS:
        config.pop(key, None)
    return jsonify(config)


@app.route('/api/config', methods=['POST'])
def update_config():
    """Update configuration including API key and model."""
    data = request.json
    config = load_user_config()
    
    if 'api_key' in data and data['api_key']:
        config['api_key'] = data['api_key']
    # Multi-provider key management
    if 'api_key' in data and data['api_key']:
        if 'keys' not in config:
            config['keys'] = {}
            
        provider = 'azure' if data.get('llm_model', '').startswith('azure-') else \
                   'openai' if data.get('llm_model', '').startswith('gpt-') else 'gemini'
        config['keys'][provider] = data['api_key']
        config['api_key'] = data['api_key']

    if 'azure_endpoint' in data:
        config['azure_endpoint'] = data['azure_endpoint']
    if 'azure_api_version' in data:
        config['azure_api_version'] = data['azure_api_version']
    if 'azure_deployment' in data:
        config['azure_deployment'] = data['azure_deployment']
    if 'azure_ssl_verify' in data:
        config['azure_ssl_verify'] = data['azure_ssl_verify']

    if 'enable_llm_judge' in data:
        config['enable_llm_judge'] = bool(data['enable_llm_judge'])

    if 'llm_model' in data:
        config['llm_model'] = data['llm_model']
        
        target_provider = 'azure' if config['llm_model'].startswith('azure-') else \
                          'openai' if config['llm_model'].startswith('gpt-') else \
                          'groq' if config['llm_model'].startswith('groq-') else 'gemini'
        if target_provider in config.get('keys', {}):
            config['api_key'] = config['keys'][target_provider]

    config['debug_enabled'] = True

    env_updates = {}
    if 'api_key' in data and data['api_key']:
        provider = 'azure' if data.get('llm_model', config.get('llm_model', '')).startswith('azure-') else \
                   'openai' if data.get('llm_model', config.get('llm_model', '')).startswith('gpt-') else \
                   'groq' if data.get('llm_model', config.get('llm_model', '')).startswith('groq-') else 'gemini'
        env_var = provider_env_var(provider)
        if env_var:
            env_updates[env_var] = data['api_key']
    if env_updates:
        update_dotenv(env_updates)
        config = load_user_config()
    
    save_user_config(config)
    
    # Verify API key with selected model
    if config.get('api_key'):
        try:
            model_name = config.get('llm_model', '')
            is_azure = model_name.lower().startswith('azure-')
            is_openai = model_name.lower().startswith('gpt-') or config.get('api_key', '').startswith('sk-')
            
            if is_azure:
                # Verify with Azure OpenAI
                from openai import AzureOpenAI
                import httpx
                endpoint = config.get('azure_endpoint', '').strip().rstrip('/')
                if '/openai/' in endpoint:
                    endpoint = endpoint.split('/openai/')[0]
                
                try:
                    # Try with SSL verification first
                    client = AzureOpenAI(
                        api_key=config['api_key'],
                        api_version=config.get('azure_api_version', '2024-02-15-preview'),
                        azure_endpoint=endpoint
                    )
                    client.models.list()
                    return jsonify({"success": True, "message": f"Settings saved! Verified with Azure OpenAI"})
                except Exception as first_err:
                    full_err_msg = f"{str(first_err)} {str(getattr(first_err, '__cause__', ''))}".lower()
                    if any(term in full_err_msg for term in ["certificate verify failed", "ssl", "cert", "handshake"]):
                        logger.warning(f"SSL/Connection issue detected for {endpoint}. Retrying without verification...")
                        try:
                            # Fallback: Disable SSL verification for corporate proxies
                            client = AzureOpenAI(
                                api_key=config['api_key'],
                                api_version=config.get('azure_api_version', '2024-02-15-preview'),
                                azure_endpoint=endpoint,
                                http_client=httpx.Client(verify=False)
                            )
                            client.models.list()
                            return jsonify({
                                "success": True, 
                                "message": f"Settings saved! Verified with Azure OpenAI"
                            })
                        except Exception as second_err:
                             return jsonify({"success": False, "message": f"Azure Connection failed even without SSL verification: {str(second_err)}"})
                    
                    import traceback
                    logger.error(f"Azure verification failed for {endpoint}: {str(first_err)}")
                    logger.error(traceback.format_exc())
                    return jsonify({"success": False, "message": f"Azure Connection failed (Endpoint: {endpoint}). Details: {str(first_err)}"})
            elif is_openai:
                # Verify with OpenAI
                from openai import OpenAI
                client = OpenAI(api_key=config['api_key'])
                client.models.list() # Lightweight check
                return jsonify({"success": True, "message": f"Settings saved! Verified with OpenAI"})
            else:
                # Verify with Google Gemini
                import google.generativeai as genai
                genai.configure(api_key=config['api_key'])
                model = genai.GenerativeModel(model_name)
                # Quick test
                response = model.generate_content("Say 'API key valid' in 3 words")
                return jsonify({"success": True, "message": f"Settings saved! Verified with {model_name}"})
        except Exception as e:
            return jsonify({"success": False, "message": f"Settings saved but verification failed: {str(e)}"})
    
    return jsonify({"success": True, "message": "Configuration saved"})


# --- Helper for Trace Metadata ---
def _debug_capture_source_tags(
    job: Dict[str, Any], trace: Any = None
) -> Tuple[Optional[str], str, str]:
    """Resolve source_id / sheet for tagging observer rows (Debug multi-source tree filter).

    HITL pause and resume paths used to call ``capture_llm_metadata`` without tags; the Debug UI
    then drops those rows when rendering per-source sections (``source_id === sid``).
    """
    sid: Optional[str] = None
    sheet = ""
    proc = ""
    ps: Any = None
    if trace is not None:
        ps = getattr(trace, "pending_state", None) or {}
    if isinstance(ps, dict):
        sm = ps.get("source_metadata") or {}
        if isinstance(sm, dict):
            cand = str(sm.get("source_id") or "").strip()
            if cand:
                sid = cand
        if not sid:
            cand = str(ps.get("source_id") or "").strip()
            if cand:
                sid = cand
        sm2 = ps.get("source_metadata") if isinstance(ps.get("source_metadata"), dict) else {}
        sheet = str(
            ps.get("sheet_name")
            or (sm2.get("sheet_name") if isinstance(sm2, dict) else "")
            or "",
        ).strip()
    cp = job.get("context_packet")
    if isinstance(cp, dict):
        if not sid:
            sm3 = cp.get("source_metadata") or {}
            if isinstance(sm3, dict):
                cand = str(sm3.get("source_id") or "").strip()
                if cand:
                    sid = cand
        if not sheet:
            sheet = str(
                cp.get("selected_source_sheet")
                or cp.get("sheet_name")
                or "",
            ).strip()
    if not sheet:
        for key in ("active_sheet", "sheet_name"):
            v = job.get(key)
            if v:
                sheet = str(v).strip()
                break
    proc = sheet or proc
    return sid, sheet, proc


def _annotate_debug_rows(
    rows: List[Dict[str, Any]],
    *,
    source_id: Optional[str] = None,
    sheet_name: Optional[str] = None,
    processing_sheet: Optional[str] = None,
) -> None:
    """In-place tags for multi-source debug filtering (UI + exports)."""
    for row in rows:
        if not isinstance(row, dict):
            continue
        if source_id:
            row["source_id"] = source_id
        if sheet_name is not None:
            row["sheet_name"] = sheet_name
        if processing_sheet is not None:
            row["processing_sheet"] = processing_sheet
        meta = row.get("metadata")
        if isinstance(meta, dict) and source_id:
            meta.setdefault("source_id", source_id)


def _job_llm_summary_for_api(job: Dict[str, Any]) -> Dict[str, Any]:
    """Return Run Summary metrics aggregated across all persisted ``llm_traces``."""
    from sia.debug.llm_observer import summarize_llm_traces

    traces = job.get("llm_traces") or []
    computed = summarize_llm_traces(traces)
    if computed.get("total_calls"):
        return computed
    stored = job.get("llm_summary")
    return stored if isinstance(stored, dict) else computed


def capture_llm_metadata(
    job_id: str,
    *,
    source_id: Optional[str] = None,
    sheet_name: Optional[str] = None,
    processing_sheet: Optional[str] = None,
) -> None:
    """Extract LLM traces, summary, and events from observer and merge into the job.

    When ``source_id`` is set (multi-source runs), each appended row is tagged so the
    Debug UI can split trees per sheet/source.
    """
    job = job_manager.get_job(job_id)
    if not job:
        return

    if not source_id:
        inferred_sid, inferred_sh, inferred_ps = _debug_capture_source_tags(job, trace=None)
        source_id = inferred_sid
        if not sheet_name and inferred_sh:
            sheet_name = inferred_sh
        if not processing_sheet and inferred_ps:
            processing_sheet = inferred_ps
        
    try:
        from sia.debug.llm_observer import get_observer
        observer = get_observer()
        if observer:
            new_traces = observer.export_for_debug_ui()
            new_tools = observer.get_tool_executions()
            new_events = observer.get_hierarchical_events()
            new_state_snapshots = (
                observer.get_state_snapshots()
                if hasattr(observer, "get_state_snapshots")
                else []
            )

            _annotate_debug_rows(new_traces, source_id=source_id, sheet_name=sheet_name, processing_sheet=processing_sheet)
            _annotate_debug_rows(new_tools, source_id=source_id, sheet_name=sheet_name, processing_sheet=processing_sheet)
            _annotate_debug_rows(new_events, source_id=source_id, sheet_name=sheet_name, processing_sheet=processing_sheet)
            _annotate_debug_rows(new_state_snapshots, source_id=source_id, sheet_name=sheet_name, processing_sheet=processing_sheet)

            existing_traces = job.get("llm_traces", []) or []
            existing_trace_ids = {t.get("trace_id") for t in existing_traces if isinstance(t, dict)}
            for trace in new_traces:
                if trace.get("trace_id") not in existing_trace_ids:
                    existing_traces.append(trace)
            job["llm_traces"] = existing_traces

            existing_tools = job.get("tool_executions", []) or []
            existing_tool_ids = {t.get("event_id") for t in existing_tools if isinstance(t, dict)}
            for tool in new_tools:
                if tool.get("event_id") not in existing_tool_ids:
                    existing_tools.append(tool)
            job["tool_executions"] = existing_tools

            existing_events = job.get("debug_events", []) or []
            existing_event_ids = {e.get("event_id") for e in existing_events if isinstance(e, dict)}
            for event in new_events:
                if event.get("event_id") not in existing_event_ids:
                    existing_events.append(event)
            job["debug_events"] = existing_events

            existing_snapshots = job.get("state_snapshots", []) or []
            existing_snapshot_ids = {
                s.get("snapshot_id") for s in existing_snapshots if isinstance(s, dict)
            }
            for snapshot in new_state_snapshots:
                if snapshot.get("snapshot_id") not in existing_snapshot_ids:
                    existing_snapshots.append(snapshot)
            job["state_snapshots"] = existing_snapshots

            from sia.debug.llm_observer import refresh_job_llm_summary

            refresh_job_llm_summary(job)
            logger.info(
                "[PROCESS] Metadata captured for job %s (%s events, source=%s)",
                job_id,
                len(job.get("debug_events", [])),
                source_id or "single",
            )
    except Exception as e:
        logger.warning(f"[PROCESS] Failed to capture LLM metadata for {job_id}: {e}")

@app.route('/api/upload', methods=['POST'])
def upload_file():
    """Upload file (Data or Schema)."""
    if 'file' not in request.files:
        return jsonify({"error": "No file provided"}), 400
    
    file = request.files['file']
    upload_type = request.form.get('type', 'data')
    job_id = request.form.get('job_id')

    if file.filename == '':
        return jsonify({"error": "No file selected"}), 400
    
    if not allowed_file(file.filename):
        return jsonify({"error": f"Invalid file type: {file.filename}"}), 400

    filename = secure_filename(file.filename)

    if upload_type == 'data':
        if not job_id:
            job_id = str(uuid.uuid4())[:8]
        file_path = UPLOAD_FOLDER / f"{job_id}_{filename}"
        file.save(str(file_path))
        
        # Detect sheets if Excel
        sheets = []
        if filename.endswith(('.xlsx', '.xls')):
            try:
                wb = load_workbook(str(file_path), read_only=True)
                sheets = [s for s in wb.sheetnames if wb[s].sheet_state == 'visible']
                wb.close()
                if len(sheets) > MAX_VISIBLE_SHEETS_PER_WORKBOOK:
                    try:
                        file_path.unlink(missing_ok=True)
                    except OSError:
                        pass
                    return jsonify({
                        "error": (
                            f"This workbook has {len(sheets)} visible sheets; "
                            f"the maximum supported per file is {MAX_VISIBLE_SHEETS_PER_WORKBOOK}."
                        ),
                    }), 400
            except Exception as e:
                logger.warning(f"Failed to read sheets from {filename}: {e}")
        
        job = job_manager.get_job(job_id)
        if job is None:
            job = job_manager.create_job(job_id, file.filename, str(file_path), sheets=sheets)
        else:
            job_manager.add_data_file(job_id, file.filename, str(file_path), sheets=sheets)
            job = job_manager.get_job(job_id)

        if job is not None:
            dfs = list(job.get("data_files") or [])
            total_sheets = sum(len(df.get("sheets") or []) for df in dfs)
            _record_job_debug(
                job,
                label="Data file uploaded",
                phase="upload",
                module="upload",
                operation="data_file",
                status="success",
                summary=f"Added {file.filename} ({len(sheets)} sheet(s))",
                metadata={
                    "filename": file.filename,
                    "file_count": len(dfs),
                    "total_sheets": total_sheets,
                    "source_count": len(job.get("source_registry") or []),
                },
            )
        
        return jsonify({
            "success": True, 
            "message": "Data file uploaded",
            "job_id": job_id,
            "filename": file.filename,
            "sheets": sheets,
            "sources": serialize_job_sources(job or {}),
            "data_files": list((job or {}).get("data_files") or []),
        })

    elif upload_type == 'schema':
        if not job_id:
            return jsonify({"error": "Job ID required for schema upload"}), 400
        
        # Save schema associated with job
        schema_path = UPLOAD_FOLDER / f"{job_id}_schema_{filename}"
        file.save(str(schema_path))
        
        # Update job with schema path
        job = job_manager.get_job(job_id)
        if job:
            job['template_path'] = str(schema_path)
            job['target_template_path'] = str(schema_path)
            _record_job_debug(
                job,
                label="Target template uploaded",
                phase="upload",
                module="upload",
                operation="target_template",
                status="success",
                summary=f"Schema template: {filename}",
                metadata={"template_path": str(schema_path)},
            )
            
        return jsonify({
            "success": True,
            "message": "Schema uploaded",
            "job_id": job_id
        })
    
    return jsonify({"error": "Invalid upload type"}), 400


def _execute_process_file_impl(job_id: str, data: Dict[str, Any]) -> None:
    """Run the main processing pipeline (LangGraph + exports). Invoked from a background thread."""
    import traceback as tb
    job = job_manager.get_job(job_id)
    if not job:
        logger.error("[PROCESS] Worker: job %s not found", job_id)
        return
    try:
        _record_job_debug(
            job,
            label="Processing started",
            phase="process",
            module="process",
            operation="start",
            status="info",
            set_parent_key="process_root",
            summary="Background worker started",
            metadata={
                "source_id": data.get("source_id"),
                "sheet_name": data.get("sheet_name"),
                "process_all_sources": bool(data.get("process_all_sources")),
            },
        )
        config = load_user_config()

        if config.get("langsmith_api_key"):
            os.environ["LANGCHAIN_TRACING_V2"] = "true"
            os.environ["LANGCHAIN_API_KEY"] = config["langsmith_api_key"]
            os.environ["LANGCHAIN_PROJECT"] = "SchemaAgent-LG"

        selected_source_id = data.get("source_id")
        selected_sheet = data.get("sheet_name")
        process_all_sources = bool(data.get("process_all_sources"))
        source_registry = list(job.get("source_registry") or [])
        if not source_registry:
            raise RuntimeError("No uploaded sources found for this job")

        selected_source = resolve_job_source(job, source_id=selected_source_id, sheet_name=selected_sheet)
        if not selected_source:
            selected_source = source_registry[0]

        selected_source_id = selected_source.get("source_id")
        selected_sheet = selected_source.get("sheet_name")
        existing_scope = dict(((job.get("source_scope_registry") or {}).get(selected_source_id)) or job.get("scoped_source") or {})
        scope_blocks = resolve_demarcation_blocks_for_source(job, job_id, selected_source_id)
        if not scope_blocks:
            scope_blocks = existing_scope.get("main_blocks") or []
        scoped_source = build_scoped_source(
            sheet_name=selected_sheet,
            blocks=scope_blocks,
            user_selected_header_row=job.get("user_selected_header_row"),
        )
        job["scoped_source"] = scoped_source
        job.setdefault("source_scope_registry", {})[selected_source_id] = scoped_source

        target_template = load_target_template_for_job(job)
        if job_id in job_manager.pending_schema_mappings:
            logger.info(f"[MAPPING] Using graph-native mapping stage for job {job_id}")

        source_ids = [selected_source_id]
        if process_all_sources:
            source_ids = [
                source.get("source_id")
                for source in source_registry
                if source.get("contains_main_data", True) or not source.get("contains_reference_data")
            ] or [selected_source_id]

        _record_job_debug(
            job,
            label="Sources selected for processing",
            phase="process",
            parent_key="process_root",
            module="process",
            operation="select_sources",
            status="info",
            summary=f"{len(source_ids)} source(s): {', '.join(str(x) for x in source_ids[:5])}",
            metadata={"source_ids": [str(x) for x in source_ids], "process_all_sources": process_all_sources},
        )

        source_ids, multi_block_expanded = expand_multi_block_source_ids(job, list(source_ids or []))
        if multi_block_expanded:
            ensure_block_union_relationship(job, source_ids)
            job["_multi_block_active_batch"] = True
            logger.info(
                "[MULTI_BLOCK] job=%s expanded sheet into %d block pipeline run(s)",
                job_id,
                len(source_ids),
            )

        if (process_all_sources and len(source_ids) > 1) or multi_block_expanded:
            job.pop("_multi_source_frames_by_source", None)
            schema, multi_result, pause_response = _process_job_sources(job_id, job, config, target_template, source_ids)
            if pause_response is not None:
                return
            df = multi_result.get("combined_df") if multi_result else pd.DataFrame()
            trace_payloads = multi_result.get("traces") if multi_result else []
            overall_confidence = min((payload.get("overall_confidence", 1.0) for payload in trace_payloads), default=1.0)
            trace = None
        else:
            agent = _build_agent_from_config(config)
            template_path = job.get("template_path") or job.get("target_template_path")
            _materialize_clean_templates_at_process_start(job, job_id, [selected_source_id])
            ensure_bootstrap_and_route(job, job_id, [selected_source_id])
            context_packet = build_context_packet(
                job,
                target_template=target_template,
                selected_sheet=selected_sheet,
                selected_source_id=selected_source_id,
            )
            job["context_packet"] = context_packet
            _record_job_debug(
                job,
                label="Context packet built",
                phase="context",
                parent_key="process_root",
                source_id=str(selected_source_id or ""),
                sheet_name=str(selected_sheet or ""),
                module="context_packet",
                operation="build",
                status="success",
                summary=f"Context ready for source {selected_source_id}",
            )
            run_parent = _record_job_debug(
                job,
                label=f"Agent run: {selected_source_id}",
                phase="agent_run",
                parent_key="process_root",
                set_parent_key=f"agent_run:{selected_source_id}",
                source_id=str(selected_source_id or ""),
                sheet_name=str(selected_sheet or ""),
                module="agent",
                operation="process_file",
                status="info",
                summary=f"Starting LangGraph for {selected_sheet or selected_source_id}",
            )
            eff_path, eff_sheet, eff_scoped, used_materialized = resolve_processing_workbook(
                job,
                selected_source_id,
                selected_source["file_path"],
                selected_sheet,
                scoped_source,
            )
            if used_materialized:
                logger.info(
                    "[PROCESS] Using materialized clean workbook for source %s: %s",
                    selected_source_id,
                    eff_path,
                )

            schema, df, trace = agent.process_file(
                eff_path,
                target_template_path=template_path,
                resume_state={
                    "sheet_name": eff_sheet,
                    "source_id": selected_source_id,
                    "scoped_source": eff_scoped,
                    "approved_relationships": list(job.get("approved_file_relationships") or []),
                    "materialized_clean_active": used_materialized,
                },
                context_packet=context_packet,
            )
            _record_job_debug(
                job,
                label=f"Agent run finished: {selected_source_id}",
                phase="agent_run",
                parent_event_id=run_parent,
                source_id=str(selected_source_id or ""),
                sheet_name=str(selected_sheet or ""),
                module="agent",
                operation="process_file",
                status="success" if not getattr(trace, "hitl_pending", False) else "pending",
                summary=(
                    "HITL pause"
                    if getattr(trace, "hitl_pending", False)
                    else f"Rows: {len(df) if df is not None else 0}"
                ),
                metadata={"hitl_pending": bool(getattr(trace, "hitl_pending", False))},
            )

            if getattr(trace, "hitl_pending", False):
                pause_payload, pause_code = _apply_hitl_pause(job_id, job, trace, df)
                if pause_code >= 400:
                    err_msg = str(pause_payload.get("error") or pause_payload)
                    job["error_details"] = err_msg
                    job_manager.update_job_status(job_id, "error", err_msg, "Error")
                return

            # Live steps already appended via _emit_job_progress during the graph; do not replay
            # trace.steps here (generic "Executed N tools" lines would duplicate and bury detail).
            overall_confidence = trace.overall_confidence

        export_df = normalize_dataframe_dates_for_export(df) if df is not None else None
        tabular_df = export_df if export_df is not None else df
        job["status"] = "completed"
        job["schema"] = schema.to_dict() if schema is not None else {}
        assign_job_data_preview(job, export_df, head=20)
        job["total_rows"] = len(df) if df is not None else 0
        if trace is not None:
            job["trace"] = trace.to_dict()
            job.pop("source_traces", None)
        else:
            sr = (multi_result or {}).get("source_runs") or []
            job["source_traces"] = sr
            col_tr = (multi_result or {}).get("collation_trace")
            job["trace"] = _build_multi_source_job_trace(
                source_runs=sr,
                trace_payloads=trace_payloads,
                collation_trace=col_tr,
                overall_confidence=overall_confidence,
                summary_message=f"Collated {len(source_ids)} sources",
            )
        job["overall_confidence"] = overall_confidence

        capture_llm_metadata(job_id)
        if trace is None and multi_result:
            _merge_collation_debug_into_job(
                job,
                (multi_result or {}).get("collation_trace"),
                (multi_result or {}).get("collation_debug_events"),
            )

        output_format = config.get("output_format", "csv")
        agent_for_save = _build_agent_from_config(config)
        outputs = agent_for_save.save_results(schema, tabular_df, str(OUTPUT_FOLDER), format=output_format)
        job["output_files"] = outputs

        try:
            _persist_excel_output(job_id, job, tabular_df)
        except Exception as e:
            logger.warning(f"[PROCESS] Failed to save Excel export: {e}")

        try:
            _persist_processing_artifact_exports(job_id, job, target_template, tabular_df, source_ids)
        except Exception as e:
            logger.warning(f"[PROCESS] Failed to save processing artifacts: {e}")

        _record_job_debug(
            job,
            label="Export completed",
            phase="export",
            parent_key="process_root",
            module="export",
            operation="complete",
            status="success",
            summary=f"Processed {len(df) if df is not None else 0} rows",
            metadata={
                "output_files": list(job.get("output_files") or []),
                "total_rows": job.get("total_rows"),
                "overall_confidence": overall_confidence,
            },
        )
        job_manager.update_job_status(
            job_id,
            "completed",
            f"Processed {len(df)} rows with {overall_confidence:.1%} confidence",
            "Complete",
        )
        logger.info("[PROCESS] Job %s completed successfully (background)", job_id)

    except Exception as e:
        error_traceback = tb.format_exc()
        logger.error(f"[PROCESS] Processing failed for job {job_id}:\n{error_traceback}")
        job_manager.update_job_status(job_id, "error", str(e), "Error")
        job["error_details"] = error_traceback
    finally:
        j = job_manager.get_job(job_id)
        if j is not None:
            j["_process_worker_running"] = False


@app.route('/api/process/<job_id>', methods=['POST'])
def process_file(job_id):
    """Process an uploaded file. Runs the heavy pipeline in a background thread and returns immediately."""
    logger.info(f"[PROCESS] Starting processing for job: {job_id}")

    job = job_manager.get_job(job_id)
    if not job:
        logger.error(f"[PROCESS] Job not found: {job_id}")
        return jsonify({"error": "Job not found"}), 404

    data = request.get_json(silent=True) or {}
    process_all_sources = bool(data.get("process_all_sources"))
    if process_all_sources:
        ux_summary = job_manager.build_ux_stepper_summary(job)
        if not ux_summary.get("stage1_complete"):
            logger.info(
                "[PROCESS] Job %s: process_all_sources with incomplete guided setup "
                "(stage1_complete=false); continuing with partial-job run.",
                job_id,
            )

    source_registry = list(job.get("source_registry") or [])
    if not source_registry:
        return jsonify({"error": "No uploaded sources found for this job"}), 400

    selected_source_id = data.get("source_id")
    selected_sheet = data.get("sheet_name")

    selected_source = resolve_job_source(job, source_id=selected_source_id, sheet_name=selected_sheet)
    if not selected_source:
        selected_source = source_registry[0]

    selected_source_id = selected_source.get("source_id")
    selected_sheet = selected_source.get("sheet_name")
    existing_scope = dict(((job.get("source_scope_registry") or {}).get(selected_source_id)) or job.get("scoped_source") or {})
    scope_blocks = resolve_demarcation_blocks_for_source(job, job_id, selected_source_id)
    if not scope_blocks:
        scope_blocks = existing_scope.get("main_blocks") or []
    scoped_source = build_scoped_source(
        sheet_name=selected_sheet,
        blocks=scope_blocks,
        user_selected_header_row=job.get("user_selected_header_row"),
    )
    job["scoped_source"] = scoped_source
    job.setdefault("source_scope_registry", {})[selected_source_id] = scoped_source

    source_ids = [selected_source_id]
    if process_all_sources:
        source_ids = [
            source.get("source_id")
            for source in source_registry
            if source.get("contains_main_data", True) or not source.get("contains_reference_data")
        ] or [selected_source_id]

    if job.get("_process_worker_running"):
        return jsonify({"error": "Processing is already running for this job."}), 409

    job["_process_worker_running"] = True
    job["current_ux_stage"] = max(int(job.get("current_ux_stage") or 0), 2)
    job_manager.update_job_status(job_id, "processing", "Initializing agent...", "Starting")

    payload = dict(data)
    threading.Thread(target=_execute_process_file_impl, args=(job_id, payload), daemon=True).start()

    return jsonify(
        {
            "success": True,
            "job_id": job_id,
            "status": "processing",
            "async": True,
        }
    )


def _download_availability(job: Dict[str, Any]) -> Dict[str, bool]:
    """Return which export artifacts exist on disk (used to enable Processing page downloads)."""

    def _is_file(path: Any) -> bool:
        if not path:
            return False
        try:
            return Path(str(path)).expanduser().is_file()
        except OSError:
            return False

    out = job.get("output_files") or {}
    art = job.get("artifact_files") or {}
    excel_p = job.get("full_excel_path") or job.get("output_file") or out.get("excel") or out.get("xlsx")
    pre_list = list(art.get("pre_transform_files") or [])
    pre_p = pre_list[0] if len(pre_list) == 1 else art.get("pre_transform_bundle")
    post_p = art.get("post_transform_file")
    bundle_p = art.get("artifacts_bundle")

    excel_ok = _is_file(excel_p)
    post_ok = _is_file(post_p) and excel_ok and str(post_p) != str(excel_p)

    return {
        "excel": excel_ok,
        "schema": _is_file(out.get("schema")),
        "data": _is_file(out.get("data")),
        "pre_transform": _is_file(pre_p),
        "post_transform": post_ok,
        "artifacts_bundle": _is_file(bundle_p),
    }


@app.route('/api/status/<job_id>', methods=['GET'])
def get_status(job_id):
    """Get processing status and intermediate results."""
    job = job_manager.get_job(job_id)
    if not job:
        return jsonify({"error": "Job not found"}), 404

    payload = {
        k: v
        for k, v in job.items()
        if k not in _STATUS_RESPONSE_EXCLUDED_KEYS
    }
    try:
        payload["_download_availability"] = _download_availability(job)
    except Exception:
        payload["_download_availability"] = {}
    try:
        payload["ux_stepper_summary"] = job_manager.build_ux_stepper_summary(job)
    except Exception:
        payload["ux_stepper_summary"] = {}
    try:
        from sia.agent.transformation_phase_banner import build_transformation_phase_banner

        payload["transformation_phase_banner"] = build_transformation_phase_banner(
            hints=job.get("structure_transform_hints"),
            plan_tool_names=job.get("transformation_plan_tools"),
            tool_executions=job.get("tool_executions"),
            job_status=job.get("status"),
        )
    except Exception:
        payload["transformation_phase_banner"] = {}
    return jsonify(payload)


@app.route("/api/jobs/<job_id>/ux-state", methods=["PATCH"])
def patch_job_ux_state(job_id):
    """Update UX stage fields (current_ux_stage 0-4, relationships gate, per-source progress)."""
    job = job_manager.get_job(job_id)
    if not job:
        return jsonify({"error": "Job not found"}), 404
    data = request.get_json(silent=True) or {}
    summary = job_manager.patch_ux_state(job_id, data)
    return jsonify({"success": True, "ux_stepper_summary": summary})


@app.route('/api/download/<job_id>/<file_type>', methods=['GET'])
def download_file(job_id, file_type):
    """Download output files."""
    job = job_manager.get_job(job_id)
    if not job:
        return jsonify({"error": "Job not found"}), 404
    
    if file_type == 'schema':
        file_path = job.get("output_files", {}).get("schema")
    elif file_type == 'data':
        file_path = job.get("output_files", {}).get("data")
    elif file_type == 'excel':
        file_path = job.get("full_excel_path") or job.get("output_file")
        if not file_path:
            outputs = job.get("output_files") or {}
            if isinstance(outputs, dict):
                file_path = outputs.get("excel") or outputs.get("xlsx")
    elif file_type == 'pre-transform':
        artifacts = job.get("artifact_files") or {}
        pre_files = list(artifacts.get("pre_transform_files") or [])
        file_path = pre_files[0] if len(pre_files) == 1 else artifacts.get("pre_transform_bundle")
    elif file_type == 'post-transform':
        artifacts = job.get("artifact_files") or {}
        file_path = artifacts.get("post_transform_file") or job.get("full_excel_path") or job.get("output_file")
    elif file_type == 'artifacts':
        artifacts = job.get("artifact_files") or {}
        file_path = artifacts.get("artifacts_bundle")
    else:
        return jsonify({"error": "Invalid file type"}), 400
    
    if not file_path or not Path(file_path).exists():
        return jsonify({"error": "File not found"}), 404
    
    return send_file(file_path, as_attachment=True)


@app.route('/api/download/snapshot/<path:filename>', methods=['GET'])
def download_snapshot(filename):
    """Download a tool execution snapshot file."""
    safe = Path(str(filename)).name
    file_path = _resolve_debug_snapshot_file(safe)
    if not file_path:
        return jsonify({"error": "Snapshot not found"}), 404
    return send_file(str(file_path), as_attachment=True, download_name=safe)


@app.route('/api/jobs/<job_id>/snapshot-preview', methods=['GET'])
def snapshot_preview(job_id: str):
    """Return a small JSON preview of a snapshot workbook (first sheet, capped rows) for Results UI."""
    if not job_id or not str(job_id).strip():
        return jsonify({"error": "Invalid job id"}), 400
    raw = request.args.get("file") or request.args.get("snapshot") or ""
    safe_name = Path(str(raw)).name
    if not safe_name:
        return jsonify({"error": "Invalid snapshot name"}), 400
    file_path = _resolve_debug_snapshot_file(safe_name)
    if not file_path:
        return jsonify({"error": "Snapshot not found"}), 404
    job = job_manager.get_job(job_id)
    if not job:
        logger.info("[SNAPSHOT_PREVIEW] job %s not in memory; preview allowed for resolved file %s", job_id, safe_name)
    try:
        suffix = file_path.suffix.lower()
        if suffix == ".csv":
            df = pd.read_csv(file_path, nrows=100)
        else:
            df = pd.read_excel(file_path, sheet_name=0, nrows=100, engine="openpyxl")
    except Exception as ex:
        logger.warning("[SNAPSHOT_PREVIEW] read failed %s: %s", safe_name, ex)
        return jsonify({"error": str(ex)}), 400
    if df is None or df.empty:
        return jsonify({"columns": [], "rows": [], "row_count": 0, "column_count": 0, "sheet": 0})
    cols = [str(c) for c in df.columns]
    records = safe_serialize_df(df)
    return jsonify(
        {
            "columns": cols,
            "rows": records,
            "row_count": int(len(df)),
            "column_count": int(len(cols)),
            "sheet": 0,
            "filename": safe_name,
        }
    )


@app.route('/api/jobs', methods=['GET'])
def list_jobs():
    """List all processing jobs."""
    jobs = []
    for job_id, job in job_manager.jobs.items():
        jobs.append({
            "job_id": job_id,
            "filename": job.get("filename"),
            "status": job.get("status"),
            "created_at": job.get("created_at"),
            "confidence": job.get("overall_confidence", 0.0)
        })
    
    # Sort by created_at descending
    jobs.sort(key=lambda x: x.get("created_at", ""), reverse=True)
    
    return jsonify({"jobs": jobs})


@app.route('/api/jobs/<job_id>', methods=['DELETE'])
def delete_job(job_id):
    """Delete a job and related uploaded/output files."""
    job = job_manager.get_job(job_id)
    if not job:
        return jsonify({"error": "Job not found"}), 404

    if job.get("status") == "processing":
        return jsonify({"error": "Cannot delete a job while processing is active"}), 409

    candidate_paths = []
    for path_str in [
        job.get("file_path"),
        job.get("template_path"),
        job.get("target_template_path"),
        job.get("full_excel_path"),
    ]:
        if path_str:
            candidate_paths.append(Path(path_str))

    for path_str in (job.get("output_files") or {}).values():
        if path_str:
            candidate_paths.append(Path(path_str))

    for data_file in (job.get("data_files") or []):
        if data_file.get("file_path"):
            candidate_paths.append(Path(data_file["file_path"]))

    for artifact_path in iter_artifact_paths(job.get("artifact_files")):
        candidate_paths.append(artifact_path)

    seen_paths = set()
    deleted_files = []
    for path in candidate_paths:
        try:
            resolved = path.resolve()
        except Exception:
            continue
        normalized = str(resolved)
        if normalized in seen_paths:
            continue
        seen_paths.add(normalized)

        try:
            if resolved.exists() and resolved.is_file():
                resolved.unlink()
                deleted_files.append(str(resolved))
        except Exception as e:
            logger.warning(f"[JOBS] Failed to delete file for job {job_id}: {resolved} ({e})")

    removed_job = job_manager.remove_job(job_id)
    if not removed_job:
        return jsonify({"error": "Job not found"}), 404

    return jsonify({
        "success": True,
        "job_id": job_id,
        "deleted_files_count": len(deleted_files),
        "message": "Job deleted"
    })


def _job_locked_for_edit(job: Dict[str, Any]) -> bool:
    return str(job.get("status") or "").lower() == "processing"


@app.route("/api/jobs/<job_id>/sources/remove", methods=["POST"])
def remove_job_source(job_id):
    """Remove one sheet (source) from an in-memory job."""
    job = job_manager.get_job(job_id)
    if not job:
        return jsonify({"error": "Job not found"}), 404
    if _job_locked_for_edit(job):
        return jsonify({"error": "Cannot modify sources while processing is active"}), 409
    data = request.get_json(silent=True) or {}
    source_id = data.get("source_id")
    if not source_id:
        return jsonify({"error": "source_id required"}), 400
    ok = job_manager.remove_source(job_id, str(source_id))
    if not ok:
        return jsonify({"error": "Source not found"}), 404
    job = job_manager.get_job(job_id) or {}
    return jsonify({
        "success": True,
        "sources": serialize_job_sources(job),
        "data_files": list(job.get("data_files") or []),
    })


@app.route("/api/jobs/<job_id>/data-files/remove", methods=["POST"])
def remove_job_data_file(job_id):
    """Remove one uploaded workbook (all its sheets) from the job."""
    job = job_manager.get_job(job_id)
    if not job:
        return jsonify({"error": "Job not found"}), 404
    if _job_locked_for_edit(job):
        return jsonify({"error": "Cannot modify sources while processing is active"}), 409
    data = request.get_json(silent=True) or {}
    file_id = data.get("file_id")
    if not file_id:
        return jsonify({"error": "file_id required"}), 400
    ok = job_manager.remove_data_file(job_id, str(file_id))
    if not ok:
        return jsonify({"error": "Data file not found"}), 404
    job = job_manager.get_job(job_id) or {}
    return jsonify({
        "success": True,
        "sources": serialize_job_sources(job),
        "data_files": list(job.get("data_files") or []),
    })


@app.route('/api/evals', methods=['GET'])
def list_evals():
    """List all job evaluations."""
    evals = []
    for job_id, job in job_manager.jobs.items():
        # Check if job has trace with judge result
        trace = job.get("trace", {})
        judge_result = trace.get("judge_result") if trace else None
        
        if judge_result or job.get("status") == "completed":
            evals.append({
                "job_id": job_id,
                "filename": job.get("filename"),
                "status": job.get("status"),
                "created_at": job.get("created_at"),
                "judge_result": judge_result,
                "overall_confidence": job.get("overall_confidence", 0.0)
            })
    
    # Sort by created_at descending
    evals.sort(key=lambda x: x.get("created_at", ""), reverse=True)
    
    return jsonify({"evals": evals})


@app.route('/api/health', methods=['GET'])
def health_check():
    """Health check endpoint."""
    return jsonify({
        "status": "ok",
        "server": "SIA Web Server",
        "jobs_count": job_manager.get_job_count()
    })


@app.route('/api/pipeline-catalog', methods=['GET'])
def get_pipeline_catalog():
    """Pipeline stages, contracts, and tool membership (same source as planner/replanner injection)."""
    from sia.tools.pipeline_catalog import serialize_catalog_json

    return jsonify({"stages": serialize_catalog_json()})


@app.route('/api/debug/<job_id>', methods=['GET'])
def get_debug_info(job_id):
    """Get full debug information for a job."""
    # Use JobManager as primary source
    job = job_manager.get_job(job_id)
    if not job:
        # Fallback to legacy dictionary
        if job_id not in processing_results:
            return jsonify({"error": "Job not found"}), 404
        job = processing_results[job_id]
    trace = job.get("trace") or {}
    lifecycle_events = trace.get("lifecycle_events") if isinstance(trace, dict) else None
    source_traces = job.get("source_traces")
    if source_traces is None and isinstance(trace, dict):
        source_traces = trace.get("source_runs")
    if not isinstance(source_traces, list):
        source_traces = []
    reg = job.get("source_registry") or []
    source_registry_summary = []
    if isinstance(reg, list):
        for s in reg:
            if not isinstance(s, dict):
                continue
            source_registry_summary.append(
                {
                    "source_id": s.get("source_id"),
                    "sheet_name": s.get("sheet_name"),
                    "file_name": s.get("file_name") or job.get("filename"),
                }
            )

    return jsonify({
        "job_id": job_id,
        "status": job.get("status"),
        "filename": job.get("filename"),
        "file_path": job.get("file_path"),
        "output_file": job.get("output_file") or (trace.get("output_file") if isinstance(trace, dict) else None),
        "steps": job.get("steps", []),
        "error_details": job.get("error_details"),
        "trace": trace,
        "source_traces": source_traces,
        "source_registry": source_registry_summary,
        "schema": job.get("schema"),
        "llm_traces": job.get("llm_traces"),
        "llm_summary": _job_llm_summary_for_api(job),
        "tool_executions": job.get("tool_executions"),
        "debug_events": job.get("debug_events"),
        "job_debug_events": get_job_debug_timeline(job),
        "data_files_summary": data_files_summary(job),
        "setup_summary": job_setup_summary(job),
        "state_snapshots": job.get("state_snapshots", []),
        "data_preview": job.get("data_preview", []),
        "data_preview_column_order": job.get("data_preview_column_order", []),
        "lifecycle_events": lifecycle_events or [],
        "hitl_pause_type": trace.get("hitl_pause_type") if isinstance(trace, dict) else None,
        "hitl_pending": trace.get("hitl_pending") if isinstance(trace, dict) else False,
    })


# ===== HITL Review API Endpoints =====

@app.route('/api/review/pending', methods=['GET'])
def get_pending_reviews():
    """List all items pending human review (both low confidence and destructive approvals)."""
    pending = job_manager.get_all_pending_reviews()
    
    return jsonify({
        "pending_count": len(pending),
        "items": pending
    })


@app.route('/api/review/<job_id>', methods=['GET'])
def get_review_details(job_id):
    """Get detailed review information for a specific job."""
    # Check regular review queue
    if job_id in job_manager.review_queue:
        item = job_manager.review_queue[job_id]
        job = job_manager.get_job(job_id) or {}
        
        return jsonify({
            **item,
            "type": "confidence_review",
            "full_schema": job.get("schema"),
            "full_data_preview": job.get("data_preview"),
            "tool_history": job.get("trace", {}).get("steps", [])  # NEW: Pass execution history
        })
    
    # Check pending deletions
    if job_id in job_manager.pending_deletions:
        item = job_manager.pending_deletions[job_id]
        job = job_manager.get_job(job_id) or {}
        return jsonify({
            "job_id": job_id,
            "filename": job.get("filename"),
            "status": "awaiting_approval",
            "type": "destructive_approval",
            "reason": f"Approval required for {len(item.get('previews', []))} destructive operations",
            "previews": item.get("previews", []),
            "low_confidence_items": item.get("low_confidence_items", []),
            "confidence": job.get("trace", {}).get("overall_confidence", 0.5),
            "created_at": item.get("created_at")
        })
        
    return jsonify({"error": "Review item not found"}), 404


def _planning_summary_for_review_job(job: Dict[str, Any]) -> Dict[str, Any]:
    """When context_packet has no planning_summary (e.g. HTTP workflow), derive it from current mapping/layout."""
    try:
        target_template = load_target_template_for_job(job)
        selected_sheet = (job.get("scoped_source") or {}).get("sheet_name")
        selected_source_id = job.get("mapping_source_id")
        pkt = build_context_packet(
            job,
            target_template=target_template,
            selected_sheet=selected_sheet,
            selected_source_id=selected_source_id,
        )
        return ContextPacket(**pkt).planning_summary()
    except Exception as ex:
        logger.warning("Could not build planning_summary for review checkpoint: %s", ex)
        return {}


def _preview_scalar_for_json(val: Any) -> Any:
    if val is None or isinstance(val, (bool, int, float, str)):
        return val
    return str(val)


def _find_preview_column_key(rows: List[Dict[str, Any]], name: str) -> Optional[str]:
    """Match column name case-insensitively against first preview row keys."""
    if not rows or not name:
        return None
    first = next((r for r in rows if isinstance(r, dict)), None)
    if not first:
        return None
    if name in first:
        return name
    nl = str(name).strip().lower()
    for k in first.keys():
        if str(k).strip().lower() == nl:
            return k
    return None


def _sample_column_from_preview_rows(rows: List[Dict[str, Any]], col_key: str, limit: int = 8) -> List[Any]:
    if not rows or not col_key:
        return []
    out: List[Any] = []
    seen = set()
    for r in rows:
        if not isinstance(r, dict) or col_key not in r:
            continue
        val = r[col_key]
        key = repr(val)
        if key in seen:
            continue
        seen.add(key)
        out.append(val)
        if len(out) >= limit:
            break
    return out


def _source_column_for_target(job: Dict[str, Any], target_column: str) -> Optional[str]:
    sid = job.get("mapping_source_id")
    tc = str(target_column).strip()
    if not tc:
        return None
    for row in job.get("mapping_registry") or []:
        if not isinstance(row, dict):
            continue
        if sid and row.get("source_id") and row.get("source_id") != sid:
            continue
        if str(row.get("target_column") or "").strip() == tc:
            sc = str(row.get("source_column") or "").strip()
            if sc:
                return sc
    return None


def _enrich_approval_items_with_previews(job: Dict[str, Any], approval_items: List[Any]) -> List[Dict[str, Any]]:
    """
    Attach sample values for each approval item by reading job.data_preview rows.
    Tries target column name first (post-mapping frame), then mapped source column name.
    """
    rows = job.get("data_preview") or []
    if not isinstance(rows, list):
        rows = []
    enriched: List[Dict[str, Any]] = []
    for raw in approval_items:
        item = dict(raw) if isinstance(raw, dict) else {"summary": str(raw)}
        tgt = (item.get("target_column") or item.get("target") or "").strip()
        preview_values: List[Any] = []
        label = ""
        if tgt:
            key = _find_preview_column_key(rows, tgt)
            if key:
                preview_values = _sample_column_from_preview_rows(rows, key, limit=8)
                label = key
            if not preview_values:
                src = _source_column_for_target(job, tgt)
                if src:
                    sk = _find_preview_column_key(rows, src)
                    if sk:
                        preview_values = _sample_column_from_preview_rows(rows, sk, limit=8)
                        label = f"{tgt} ← {sk}"
        item["preview_column_label"] = label
        item["preview_values"] = [_preview_scalar_for_json(v) for v in preview_values]
        if tgt and not preview_values:
            item["preview_note"] = (
                "No sample values found in the job data preview for this field. "
                "The column may be missing, renamed after mapping, or the preview may need a fresh run from Guided Setup."
            )
        enriched.append(item)
    return enriched


@app.route('/api/review/checkpoint/<checkpoint_id>', methods=['GET'])
def get_checkpoint_details(checkpoint_id):
    """Get detailed review information for a specific checkpoint."""
    checkpoint = job_manager.pending_checkpoints.get(checkpoint_id)
    if not checkpoint:
        return jsonify({"error": "Checkpoint not found"}), 404

    job = job_manager.get_job(checkpoint.get("job_id")) or {}
    pending_state = checkpoint.get("pending_state") or {}
    trigger_data = checkpoint.get("trigger_data") or {}
    context_packet = pending_state.get("context_packet") or job.get("context_packet") or {}
    planning_summary = context_packet.get("planning_summary") or {}
    if not planning_summary:
        planning_summary = _planning_summary_for_review_job(job)
    plan_payload = _serialize_plan_for_review(pending_state.get("extraction_plan"))
    if not plan_payload:
        plan_payload = _serialize_plan_for_review(trigger_data)
    else:
        if not plan_payload.get("tool_calls"):
            plan_payload["tool_calls"] = list(trigger_data.get("tool_calls") or [])
        if not plan_payload.get("expected_columns"):
            plan_payload["expected_columns"] = list(trigger_data.get("expected_columns") or [])
        if not plan_payload.get("reasoning"):
            plan_payload["reasoning"] = str(trigger_data.get("reasoning") or "")
        if not plan_payload.get("business_rule_actions"):
            plan_payload["business_rule_actions"] = list(trigger_data.get("planned_rule_actions") or [])
        if not plan_payload.get("approval_items"):
            plan_payload["approval_items"] = list(trigger_data.get("approval_items") or [])

    if plan_payload.get("approval_items"):
        plan_payload["approval_items"] = _enrich_approval_items_with_previews(job, plan_payload["approval_items"])

    checkpoint_payload = {
        key: value
        for key, value in checkpoint.items()
        if key != "pending_state"
    }
    return jsonify({
        **checkpoint_payload,
        "job_status": job.get("status"),
        "steps": job.get("steps", []),
        "trace": job.get("trace", {}),
        "data_preview": job.get("data_preview", []),
        "schema_preview": job.get("schema"),
        "plan": plan_payload,
        "planning_summary": planning_summary,
        "lineage": context_packet.get("lineage") or {},
    })


def _relationship_duplicate_visual_sample(
    stacked: Optional[pd.DataFrame],
    subset: List[str],
    post_merge_action: str,
    *,
    max_rows: int = 28,
    max_measure_cols: int = 5,
) -> Optional[Dict[str, Any]]:
    """Build a small table of rows involved in duplicate-key groups with human-readable outcomes."""
    if stacked is None or stacked.empty or not subset:
        return None
    use = [c for c in subset if c in stacked.columns]
    if not use:
        return None
    in_group = stacked.duplicated(subset=use, keep=False) | stacked.duplicated(subset=use, keep="last")
    if not bool(in_group.any()):
        return None
    view = stacked.loc[in_group].copy()
    if post_merge_action == "drop_duplicate_rows":
        remove = stacked.duplicated(subset=use, keep="first")
        view["Outcome"] = [
            (
                "Removed — duplicate key; first stacked row kept"
                if bool(remove.loc[idx])
                else "Kept — first in stack for this key"
            )
            for idx in view.index
        ]
    elif post_merge_action == "aggregate_duplicate_keys":
        view["Outcome"] = "Combined — numerics summed per key"
    else:
        view["Outcome"] = "No duplicate-key removal in this sample"

    measures = [c for c in stacked.columns if c not in use][:max_measure_cols]
    disp_cols = use + measures
    disp_cols = [c for c in disp_cols if c in view.columns] + ["Outcome"]
    view = view.loc[:, [c for c in disp_cols if c in view.columns]]
    view = view.sort_values(by=use).head(int(max_rows))
    return {
        "title": "Duplicate key rows (sample)",
        "columns": [str(c) for c in view.columns],
        "rows": safe_serialize_df(view),
    }


def _patch_union_duplicate_merge_mode_on_unions(applied: List[Dict[str, Any]], mode: str) -> None:
    m = normalize_union_duplicate_merge_mode(mode)
    for r in applied:
        if str(r.get("relationship_kind") or "").lower() == "union":
            r["duplicate_merge_mode"] = m


def _set_union_duplicate_key_decisions(applied: List[Dict[str, Any]], decisions: Dict[str, Any]) -> None:
    for r in applied:
        if str(r.get("relationship_kind") or "").lower() == "union":
            r["duplicate_key_group_decisions"] = dict(decisions or {})
            return


def _normalize_duplicate_key_decisions_payload(raw: Any) -> Dict[str, Dict[str, str]]:
    if not isinstance(raw, dict) or not raw:
        return {}
    out: Dict[str, Dict[str, str]] = {}
    for k, v in raw.items():
        if not isinstance(v, dict):
            continue
        gid = str(k).strip()
        if not gid:
            continue
        cat = str(v.get("category") or "").strip()
        act = str(v.get("action") or "").strip()
        if cat == "exact" and act in ("treat_duplicate", "treat_separate"):
            out[gid] = {"category": cat, "action": act}
        elif cat == "partial" and act in ("keep_source_1", "keep_source_2", "keep_both"):
            out[gid] = {"category": cat, "action": act}
    return out


@app.route("/api/review/checkpoint/<checkpoint_id>/relationship-samples", methods=["GET", "POST"])
def get_relationship_preview_samples(checkpoint_id: str):
    """First rows per source + illustrative combined head for file-relationship review (analyst preview)."""
    body = request.get_json(silent=True) if request.method == "POST" and request.is_json else {}
    body = body or {}
    checkpoint = job_manager.pending_checkpoints.get(checkpoint_id)
    if not checkpoint or checkpoint.get("type") != "file_relationship_review":
        return jsonify({"error": "Not found"}), 404
    job_id = checkpoint.get("job_id")
    job = job_manager.get_job(job_id) if job_id else None
    if not job:
        return jsonify({"error": "Job not found"}), 404
    proposals = list(
        (checkpoint.get("trigger_data") or {}).get("relationship_proposals")
        or job.get("relationship_proposals")
        or []
    )
    summaries = list((checkpoint.get("trigger_data") or {}).get("sources") or [])
    sid_order: List[str] = []
    for prop in proposals:
        for sid in prop.get("source_ids") or []:
            s = str(sid)
            if s and s not in sid_order:
                sid_order.append(s)
    if not sid_order:
        for row in job.get("source_registry") or []:
            s = str(row.get("source_id") or "")
            if s and s not in sid_order:
                sid_order.append(s)
    summary_by_id = {str(s.get("source_id")): s for s in summaries if s.get("source_id")}
    inputs: List[Dict[str, Any]] = []
    frames_by_source: Dict[str, pd.DataFrame] = {}
    sample_limit = 3
    frame_cap = 120
    for sid in sid_order:
        source = resolve_job_source(job, source_id=sid)
        if not source:
            continue
        sheet_name = str(source.get("sheet_name") or "")
        scoped = dict(((job.get("source_scope_registry") or {}).get(str(sid))) or {})
        try:
            eff_path, eff_sheet, eff_scoped, _used = resolve_processing_workbook(
                job,
                str(sid),
                str(source.get("file_path") or job.get("file_path") or ""),
                sheet_name,
                scoped,
            )
            raw_df = load_raw_sheet_dataframe(str(eff_path), str(eff_sheet or sheet_name))
            df, _ = load_scoped_dataframe(str(eff_path), str(eff_sheet or sheet_name), eff_scoped or scoped, raw_df=raw_df)
        except Exception as ex:
            logger.warning("[REL_PREVIEW] skip source %s: %s", sid, ex)
            continue
        if df is None or df.empty:
            continue
        head = df.head(sample_limit)
        sm = summary_by_id.get(str(sid)) or {}
        inputs.append(
            {
                "source_id": str(sid),
                "file_name": str(sm.get("file_name") or source.get("file_name") or ""),
                "sheet_name": str(sm.get("sheet_name") or source.get("sheet_name") or eff_sheet or sheet_name),
                "columns": [str(c) for c in head.columns],
                "rows": safe_serialize_df(head),
            }
        )
        frames_by_source[str(sid)] = df.head(frame_cap)
    applied = apply_relationship_decisions(proposals, [])
    query_mode = str(body.get("duplicate_merge_mode") or request.args.get("duplicate_merge_mode") or "").strip()
    if query_mode:
        base_mode = normalize_union_duplicate_merge_mode(query_mode)
    else:
        base_mode = union_duplicate_merge_mode_from_relationships(applied)
    _patch_union_duplicate_merge_mode_on_unions(applied, base_mode)
    norm_frames = _normalize_frames_keys(frames_by_source)

    stacked: Optional[pd.DataFrame] = None
    diag: Dict[str, Any] = {}
    dup = 0
    conflict = False
    subset: List[str] = []
    exact_groups: List[Dict[str, Any]] = []
    partial_groups: List[Dict[str, Any]] = []
    union_source_labels: List[str] = []
    duplicate_decisions: Dict[str, Dict[str, str]] = {}
    applied_for_collate = applied

    if len(norm_frames) >= 2:
        try:
            stacked = _collate_frames_baseline(norm_frames, applied)
            uci = _union_stack_column_intersection(norm_frames, applied)
            diag = diagnose_stacked_union_duplicates(
                stacked if stacked is not None and not stacked.empty else None,
                applied,
                uci,
            )
            dup = int(diag.get("duplicate_rows_on_keys") or 0)
            conflict = bool(diag.get("numeric_conflict_on_duplicate_keys"))
            subset = list(diag.get("subset_for_dedupe") or [])
            if stacked is not None and not stacked.empty and dup > 0 and subset:
                uid_rows = _union_stack_source_ids(norm_frames, applied)
                union_source_labels = []
                for sid in uid_rows:
                    inp = next((x for x in inputs if str(x.get("source_id")) == str(sid)), None)
                    if inp:
                        lab = f"{inp.get('file_name') or ''} · {inp.get('sheet_name') or ''}".strip()
                        union_source_labels.append(lab or str(sid))
                    else:
                        union_source_labels.append(str(sid))
                breaks = union_stack_break_row_indices(frames_by_source, applied)
                exact_groups, partial_groups = enumerate_duplicate_key_groups_for_review(
                    stacked,
                    subset,
                    union_break_before_row=breaks,
                    source_labels=union_source_labels,
                )
                for grp in exact_groups + partial_groups:
                    if grp.get("rows"):
                        grp["rows"] = safe_serialize_df(pd.DataFrame(grp["rows"]))
                suggested = build_suggested_duplicate_key_decisions(
                    exact_groups, partial_groups, base_mode, conflict
                )
                user_norm = _normalize_duplicate_key_decisions_payload(body.get("duplicate_key_group_decisions"))
                duplicate_decisions = {**suggested, **user_norm}
                if exact_groups or partial_groups:
                    applied_for_collate = copy.deepcopy(applied)
                    _set_union_duplicate_key_decisions(applied_for_collate, duplicate_decisions)
        except Exception as ex_enum:
            logger.warning("[REL_PREVIEW] duplicate enumeration failed: %s", ex_enum)

    output_sample: Optional[Dict[str, Any]] = None
    combined: Optional[pd.DataFrame] = None
    try:
        target_tpl = load_target_template_for_job(job)
        combined = collate_frames(frames_by_source, applied_for_collate, target_template=target_tpl)
        if combined is not None and not combined.empty:
            preview_cap = min(120, int(len(combined)))
            out = combined.iloc[:preview_cap]
            union_break_before_row = union_stack_break_row_indices(frames_by_source, applied_for_collate)
            output_sample = {
                "columns": [str(c) for c in out.columns],
                "rows": safe_serialize_df(out),
                "union_break_before_row": [i for i in union_break_before_row if 0 < i < preview_cap],
                "row_count_total": int(len(combined)),
                "row_count_shown": int(len(out)),
            }
    except Exception as ex:
        logger.warning("[REL_PREVIEW] collate failed: %s", ex)

    union_duplicate_preview: Optional[Dict[str, Any]] = None
    if len(norm_frames) >= 2:
        try:
            action = resolve_duplicate_merge_action(dup, conflict, base_mode)
            stacked_n = int(len(stacked)) if stacked is not None and not stacked.empty else 0
            after_n = int(len(combined)) if combined is not None and not combined.empty else stacked_n
            delta = max(0, stacked_n - after_n)
            use_per_group = bool(duplicate_decisions)
            rows_removed_duplicate = 0 if use_per_group else (int(delta) if action == "drop_duplicate_rows" else 0)
            rows_collapsed_by_aggregate = 0 if use_per_group else (int(delta) if action == "aggregate_duplicate_keys" else 0)
            uid_rows = _union_stack_source_ids(norm_frames, applied)
            row_count_before_stack = int(sum(len(norm_frames[s]) for s in uid_rows if s in norm_frames))
            if row_count_before_stack <= 0 and norm_frames:
                row_count_before_stack = int(sum(len(f) for f in norm_frames.values()))
            union_duplicate_preview = {
                **diag,
                "post_merge_action": "resolve_duplicate_key_groups" if use_per_group else action,
                "duplicate_merge_mode": base_mode,
                "duplicate_key_group_decisions": duplicate_decisions,
                "exact_duplicate_groups": exact_groups,
                "partial_duplicate_groups": partial_groups,
                "source_1_label": union_source_labels[0] if len(union_source_labels) > 0 else "Source 1",
                "source_2_label": union_source_labels[1] if len(union_source_labels) > 1 else "Source 2",
                "row_count_before_stack": row_count_before_stack,
                "row_count_stacked": stacked_n,
                "rows_removed_duplicate": rows_removed_duplicate,
                "rows_collapsed_by_aggregate": rows_collapsed_by_aggregate,
                "preview_max_rows_per_source": int(frame_cap),
            }
            if use_per_group and delta:
                union_duplicate_preview["rows_net_reduced"] = int(delta)
            if combined is not None and not combined.empty:
                union_duplicate_preview["row_count_after_merge_tools"] = after_n
            if stacked is not None and not stacked.empty and subset and not (exact_groups or partial_groups):
                vis = _relationship_duplicate_visual_sample(stacked, subset, action)
                if vis:
                    union_duplicate_preview["visual_sample"] = vis
        except Exception as ex2:
            logger.warning("[REL_PREVIEW] duplicate diagnostics failed: %s", ex2)

    payload: Dict[str, Any] = {"inputs": inputs, "output_sample": output_sample}
    if union_duplicate_preview is not None:
        payload["union_duplicate_preview"] = union_duplicate_preview
    return jsonify(payload)


@app.route('/api/review/<job_id>/approve', methods=['POST'])
def approve_review(job_id):
    """Approve the proposed output as-is."""
    if job_id not in job_manager.review_queue:
        return jsonify({"error": "Review item not found"}), 404
    
    job_manager.approve_review(job_id)
    return jsonify({"success": True, "job_id": job_id, "message": "Review approved successfully"})


@app.route('/api/review/<job_id>/reject', methods=['POST'])
def reject_review(job_id):
    """Reject and discard the output."""
    if job_id not in job_manager.review_queue:
        return jsonify({"error": "Review item not found"}), 404
    
    data = request.json or {}
    reason = data.get("reason", "Rejected by user")
    
    job_manager.reject_review(job_id, reason)
    return jsonify({"success": True, "job_id": job_id, "message": "Review rejected"})


@app.route('/api/review/<job_id>/correct', methods=['POST'])
def apply_correction(job_id):
    """Apply human corrections to the output."""
    if job_id not in job_manager.review_queue:
        return jsonify({"error": "Review item not found"}), 404
    
    data = request.json or {}
    corrections = data.get("corrections", {})
    
    job_manager.apply_review_correction(job_id, corrections)
    
    # Apply schema corrections to job metadata if provided
    job = job_manager.get_job(job_id)
    if job and "schema" in corrections and job.get("schema"):
        for field_name, field_updates in corrections["schema"].items():
            schema = job["schema"]
            for field in schema.get("fields", []):
                if field.get("name") == field_name:
                    field.update(field_updates)
    
    return jsonify({"success": True, "job_id": job_id, "message": "Corrections applied successfully"})


@app.route('/api/review/count', methods=['GET'])
def get_review_count():
    """Get count of pending reviews (for badge)."""
    return jsonify({"pending_count": job_manager.get_pending_review_count()})


@app.route('/api/review/checkpoint/<checkpoint_id>/resolve', methods=['POST'])
def resolve_checkpoint_review(checkpoint_id):
    """Resolve a checkpoint-based review item (structural, checksum, etc.)."""
    # region agent log
    _debug_log(
        "H5_H6",
        "resolve_checkpoint_review entry",
        {
            "checkpoint_id": checkpoint_id,
            "checkpoint_exists": checkpoint_id in job_manager.pending_checkpoints,
        },
    )
    # endregion
    if checkpoint_id not in job_manager.pending_checkpoints:
        return jsonify({"error": "Checkpoint not found"}), 404
    
    data = request.json or {}
    action = data.get("action", "approve")
    resolution_data = data.get("resolution_data", {})
    
    cp = job_manager.pending_checkpoints[checkpoint_id]
    cp_type = cp.get("type")
    job_id = cp.get("job_id")
    job = job_manager.get_job(job_id) if job_id else None
    # region agent log
    _debug_log(
        "H5_H6",
        "resolve_checkpoint_review loaded checkpoint",
        {
            "checkpoint_id": checkpoint_id,
            "job_id": job_id,
            "checkpoint_type": cp_type,
            "job_exists": job is not None,
        },
    )
    # endregion
    if job is None:
        # Job was removed or never registered; keep the Review queue from listing a dead checkpoint.
        logger.warning(
            "[HITL] Orphan checkpoint %s (job_id=%s): clearing stale review item",
            checkpoint_id,
            job_id,
        )
        job_manager.pending_checkpoints.pop(checkpoint_id, None)
        return jsonify(
            {
                "success": True,
                "checkpoint_id": checkpoint_id,
                "action": action,
                "status": "resolved",
                "message": "Stale review item removed (job no longer exists).",
            }
        )

    if cp_type == "structural_review" and action == "select_header_row" and "header_row" not in resolution_data:
        return jsonify({"error": "header_row is required for select_header_row"}), 400

    job_manager.resolve_checkpoint(checkpoint_id, action, resolution_data)
    
    # If user selected a header row, store it for potential re-processing
    if action == "select_header_row" and "header_row" in resolution_data:
        if job_id in job_manager.jobs:
            job_manager.jobs[job_id]["user_selected_header_row"] = resolution_data["header_row"]

    if cp_type == "structural_review":
        if action == "select_header_row" and "header_row" in resolution_data:
            header_row = resolution_data["header_row"]
            source_id = (
                resolution_data.get("source_id")
                or (cp.get("trigger_data") or {}).get("source_id")
                or ((job.get("context_packet") or {}).get("source_metadata") or {}).get("source_id")
            )
            if isinstance(job.get("scoped_source"), dict):
                job["scoped_source"]["header_row"] = header_row
            if source_id:
                scope_registry = job.setdefault("source_scope_registry", {})
                existing_scope = dict(scope_registry.get(source_id) or job.get("scoped_source") or {})
                existing_scope["header_row"] = header_row
                scope_registry[source_id] = existing_scope
            return jsonify({
                "success": True,
                "checkpoint_id": checkpoint_id,
                "action": action,
                "status": "resolved",
                "job_id": job_id,
                "message": f"Header row set to {int(header_row) + 1} for this sheet.",
            })
        if action == "approve":
            return jsonify({
                "success": True,
                "checkpoint_id": checkpoint_id,
                "action": action,
                "status": "resolved",
                "job_id": job_id,
                "message": "Structural review approved.",
            })
        if action in {"skip_sheet", "retry_with_hints", "reject"}:
            return jsonify({
                "success": True,
                "checkpoint_id": checkpoint_id,
                "action": action,
                "status": "resolved",
                "job_id": job_id,
                "message": f"Structural review resolved: {action}",
            })

    if cp_type == "file_relationship_review":
        if action in {"approve", "modify"}:
            proposals = list((cp.get("trigger_data") or {}).get("relationship_proposals") or job.get("relationship_proposals") or [])
            overrides = list(resolution_data.get("relationships") or [])
            approved_relationships = apply_relationship_decisions(proposals, overrides=overrides)
            job_manager.save_file_relationships(job_id, approved_relationships)

            source_ids = list((cp.get("pending_state") or {}).get("source_ids") or [item.get("source_id") for item in job.get("source_registry", []) or []])
            if job.get("_relationship_resume_running"):
                return jsonify({"error": "Resume already in progress for this job."}), 409
            if job.get("_process_worker_running"):
                return jsonify({"error": "Processing is already running for this job."}), 409

            if job.get("_pending_collation_frames"):
                job["_relationship_resume_running"] = True
                job["requires_review"] = False
                job_manager.update_job_status(
                    job_id,
                    "processing",
                    "Building combined export from per-source outputs…",
                    "Processing",
                )
                threading.Thread(
                    target=_execute_collation_only_after_review,
                    args=(job_id,),
                    daemon=True,
                ).start()
                return jsonify({
                    "success": True,
                    "checkpoint_id": checkpoint_id,
                    "action": action,
                    "status": "processing",
                    "job_id": job_id,
                    "message": "Applying relationship choice and writing combined Excel/CSV.",
                })

            job["_relationship_resume_running"] = True
            job["requires_review"] = False
            job_manager.update_job_status(
                job_id,
                "processing",
                "Combining sources after relationship review…",
                "Processing",
            )
            threading.Thread(
                target=_execute_relationship_review_resume,
                args=(job_id, source_ids),
                daemon=True,
            ).start()

            return jsonify({
                "success": True,
                "checkpoint_id": checkpoint_id,
                "action": action,
                "status": "processing",
                "job_id": job_id,
                "message": "Combining sources and finishing the job. Open Processing to watch progress.",
            })
        if action in {"cancel", "reject"}:
            job.pop("_pending_collation_frames", None)
            job.pop("_post_execution_relationship_review", None)
            job.pop("_pending_collation_schema", None)
            cancelled_job = job_manager.cancel_job(job_id, resolution_data.get("reason") or "Relationship review cancelled by user")
            return jsonify({
                "success": True,
                "checkpoint_id": checkpoint_id,
                "action": action,
                "status": "cancelled",
                "job_id": job_id,
                "message": "Job cancelled from relationship review",
                "job_status": (cancelled_job or {}).get("status", "cancelled"),
            })

    if cp_type == "plan_review":
        feedback_summary = job_manager.apply_plan_feedback(job_id, checkpoint_id, action, resolution_data)
        if action in {"approve", "modify", "regenerate"}:
            pending_state = dict(cp.get("pending_state") or {})
            if not pending_state:
                return jsonify({"error": "No saved planner state to resume from"}), 500

            inferred_sid = _resolve_resume_source_id(job, pending_state)
            if inferred_sid:
                pending_state["source_id"] = inferred_sid

            original_approval = list(
                (cp.get("trigger_data") or {}).get("approval_items")
                or pending_state.get("approval_items")
                or []
            )
            resolved_approval = merge_resolved_approval_items(
                original_approval,
                list(resolution_data.get("approval_items") or []),
            )
            if resolved_approval:
                context_packet = apply_resolved_approvals_to_context(
                    dict(pending_state.get("context_packet") or {}),
                    resolved_approval,
                )
                pending_state["context_packet"] = context_packet
                pending_state["approved_mappings"] = list(
                    context_packet.get("approved_mappings")
                    or pending_state.get("approved_mappings")
                    or []
                )
                job["resolved_planner_decisions"] = resolved_approval

            if action == "approve":
                current_plan = _serialize_plan_for_review(pending_state.get("extraction_plan"))
                if not current_plan:
                    current_plan = _serialize_plan_for_review(cp.get("trigger_data") or {})

                updated_tool_calls = resolution_data.get("updated_tool_calls")
                updated_expected_columns = resolution_data.get("updated_expected_columns")
                updated_reasoning = resolution_data.get("updated_reasoning")
                updated_rule_actions = resolution_data.get("updated_rule_actions")

                if isinstance(updated_tool_calls, list) and updated_tool_calls:
                    current_plan["tool_calls"] = updated_tool_calls
                if isinstance(updated_expected_columns, list) and updated_expected_columns:
                    current_plan["expected_columns"] = updated_expected_columns
                if isinstance(updated_reasoning, str) and updated_reasoning.strip():
                    current_plan["reasoning"] = updated_reasoning.strip()
                if isinstance(updated_rule_actions, list) and updated_rule_actions:
                    current_plan["business_rule_actions"] = updated_rule_actions

                current_plan = finalize_extraction_plan(
                    current_plan,
                    pending_state.get("context_packet"),
                    pending_state.get("target_template"),
                    pending_state.get("structure_analysis"),
                )

                pending_state["extraction_plan"] = current_plan
                pending_state["suggested_tools"] = list(current_plan.get("tool_calls") or [])
                pending_state["expected_columns"] = list(current_plan.get("expected_columns") or [])
                pending_state["planned_rule_actions"] = list(current_plan.get("business_rule_actions") or [])
                pending_state["approval_items"] = list(current_plan.get("approval_items") or [])
                # region agent log
                _debug_log(
                    "H3_H4",
                    "plan review approve prepared pending_state",
                    {
                        "job_id": job_id,
                        "checkpoint_id": checkpoint_id,
                        "action": action,
                        "updated_tool_calls_count": len(updated_tool_calls or []),
                        "current_plan_tool_count": len(current_plan.get("tool_calls") or []),
                        "suggested_tools_count": len(pending_state.get("suggested_tools") or []),
                        "suggested_tool_names": [
                            t.get("tool") if isinstance(t, dict) else str(type(t))
                            for t in (pending_state.get("suggested_tools") or [])
                        ],
                        "tool_valid_flags": [
                            None if not isinstance(t, dict) else t.get("_valid")
                            for t in (pending_state.get("suggested_tools") or [])
                        ],
                        "requires_extraction": bool((pending_state.get("scoped_source") or {}).get("requires_extraction")),
                    },
                )
                # endregion

            filtered_checkpoints = []
            for item in pending_state.get("hitl_checkpoints", []) or []:
                if item.get("checkpoint_id") == checkpoint_id:
                    resolved_item = dict(item)
                    resolved_item["resolved"] = True
                    resolved_item["resolution_action"] = action
                    resolved_item["resolution_data"] = resolution_data
                    filtered_checkpoints.append(resolved_item)
                else:
                    filtered_checkpoints.append(item)

            pending_state["hitl_checkpoints"] = filtered_checkpoints
            pending_state["hitl_resume_from"] = cp_type or "plan_review"
            pending_state["hitl_pending_approval"] = False
            pending_state["hitl_pause_type"] = None
            pending_state["requires_review"] = False
            pending_state["review_reason"] = ""
            pending_state["escalation_reason"] = ""

            if action == "approve":
                pending_state["resume_mode"] = "use_existing_plan"
            else:
                pending_state["resume_mode"] = "replan"
                pending_state["extraction_plan"] = None
                pending_state["suggested_tools"] = []
                pending_state["approval_items"] = []
                pending_state["planned_rule_actions"] = []

            try:
                from sia.debug.hitl_debug_events import record_hitl_resume_debug

                cap_sid, cap_sheet, cap_proc = _debug_capture_source_tags(job, trace=None)
                record_hitl_resume_debug(
                    job,
                    checkpoint_id=checkpoint_id,
                    pause_type=cp_type or "plan_review",
                    action=action,
                    resolution_data=resolution_data,
                    pending_state=pending_state,
                    source_id=cap_sid,
                    sheet_name=cap_sheet or None,
                    processing_sheet=cap_proc or None,
                )
            except Exception as exc:
                logger.debug("[HITL] Failed to record resume debug for %s: %s", job_id, exc)

            config = load_user_config()
            resume_response = _run_resumed_processing(job_id, job, pending_state, config)
            if isinstance(resume_response, tuple) and len(resume_response) == 2:
                resp_obj, status_code = resume_response
                payload = (resp_obj.get_json(silent=True) if resp_obj is not None else None) or {}
            else:
                resp_obj = resume_response
                payload = (resp_obj.get_json(silent=True) if resp_obj is not None else None) or {}
                status_code = getattr(resp_obj, "status_code", 200) or 200
            payload["feedback"] = feedback_summary
            return jsonify(payload), status_code
        if action in {"cancel", "reject"}:
            cancelled_job = job_manager.cancel_job(job_id, resolution_data.get("reason") or "Planner review cancelled by user")
            return jsonify({
                "success": True,
                "checkpoint_id": checkpoint_id,
                "action": action,
                "status": "cancelled",
                "job_id": job_id,
                "message": "Job cancelled from planner review",
                "job_status": (cancelled_job or {}).get("status", "cancelled"),
            })
    
    logger.info(f"[HITL] Checkpoint {checkpoint_id} resolved: {action}")
    return jsonify({
        "success": True,
        "checkpoint_id": checkpoint_id,
        "action": action,
        "message": f"Checkpoint resolved: {action}"
    })


@app.route('/api/demarcation/scan-progress/<job_id>', methods=['GET'])
def get_demarcation_scan_progress(job_id):
    """Lightweight poll while ``propose-all`` runs (long normalize + batch demarcation)."""
    job = job_manager.get_job(job_id)
    if not job:
        return jsonify({"error": "Job not found"}), 404
    return jsonify(job.get("demarcation_scan_progress") or {})


@app.route('/api/demarcation/propose-all/<job_id>', methods=['GET'])
def propose_demarcation_all(job_id):
    """
    Run Python demarcation for every source in the registry, then one batched LLM call
    for all ambiguous blocks across sheets. Caches per-source proposals on the job.
    """
    use_ai_param = (request.args.get("use_ai") or "true").strip().lower()
    use_ai_classification = use_ai_param not in ("0", "false", "no", "off")
    active_source_id = request.args.get("source_id")
    active_sheet = request.args.get("sheet_name")

    job = job_manager.get_job(job_id)
    if not job:
        return jsonify({"error": "Job not found"}), 404

    registry = list(job.get("source_registry") or [])
    if not registry:
        return jsonify({"error": "No sources in registry"}), 400

    uniq_workbooks: List[Tuple[str, str]] = []
    seen_wb: set = set()
    for src in registry:
        fp0 = src.get("file_path") or job.get("file_path")
        if not fp0:
            continue
        try:
            ck0 = str(Path(str(fp0)).resolve())
        except OSError:
            ck0 = str(fp0)
        if ck0 not in seen_wb:
            seen_wb.add(ck0)
            uniq_workbooks.append((ck0, fp0))
    files_total_workbooks = max(len(uniq_workbooks), 1)
    file_index_by_key = {ck: i + 1 for i, (ck, _) in enumerate(uniq_workbooks)}

    def _scan_msg_for_event(payload: Dict[str, Any], fi: int, fname: str) -> str:
        ev = payload.get("event")
        sn = payload.get("sheet_name") or ""
        si = payload.get("sheet_index")
        st = payload.get("sheets_total")
        tab = f"{si}/{st}" if si is not None and st is not None else ""
        if ev == "workbook_start":
            wt = payload.get("workbook_tab_count")
            sc = payload.get("scoped_to_registry")
            base = (
                f"File {fi}/{files_total_workbooks} · {fname} — opened workbook"
                f"{f', {wt} tab(s) on disk' if wt is not None else ''}"
                f"{', normalizing only sheet(s) from your upload selection' if sc else ', normalizing visible tab(s) in scope'}…"
            )
            return base
        if ev == "sheet_skip_hidden":
            return f"File {fi}/{files_total_workbooks} · {fname} — skipping hidden sheet «{sn}» ({tab} tabs)."
        if ev == "sheet_normalizing":
            return f"File {fi}/{files_total_workbooks} · {fname} — normalizing «{sn}» ({tab})… large grids can take a minute."
        if ev == "sheet_normalized":
            return (
                f"File {fi}/{files_total_workbooks} · {fname} — finished «{sn}»: "
                f"{payload.get('rows')}×{payload.get('cols')} ({tab})."
            )
        return f"File {fi}/{files_total_workbooks} · {fname} — {ev or 'working'}…"

    try:
        config = load_user_config()
        agent = _build_agent_from_config(config)
        from sia.agent.demarcator import StructureDemarcator
        from sia.modules.visual_normalizer import VisualNormalizer

        demarcator_tool = StructureDemarcator(agent.llm_client)
        visual_normalizer = VisualNormalizer(include_hidden=True)
        file_grids_cache: Dict[str, Any] = {}

        job["demarcation_scan_progress"] = {
            "active": True,
            "phase": "normalize_workbook",
            "message": (
                f"Scanning {len(registry)} source(s) across {files_total_workbooks} workbook file(s). "
                "Each file: only sheet tabs you kept after upload are normalized (other tabs stay on disk but are skipped). "
                "Hidden sheet tabs are skipped; large selected grids take the most time."
            ),
            "sources_in_registry": len(registry),
            "files_total": files_total_workbooks,
            "updated_at": datetime.now(timezone.utc).isoformat(),
        }

        sheet_payloads: List[Dict[str, Any]] = []
        for registry_index, src in enumerate(registry):
            fp = src.get("file_path") or job.get("file_path")
            sheet_name = src.get("sheet_name")
            if not fp or not sheet_name:
                continue
            try:
                cache_key = str(Path(str(fp)).resolve())
            except OSError:
                cache_key = str(fp)
            if cache_key not in file_grids_cache:
                fi = file_index_by_key.get(cache_key, 1)
                fname = Path(fp).name

                def _make_sheet_progress_handler(file_index: int, file_label: str):
                    def _on_sheet_progress(payload: Dict[str, Any]) -> None:
                        prog = job.setdefault("demarcation_scan_progress", {})
                        prog["active"] = True
                        prog["phase"] = "normalize_workbook"
                        prog["file_index"] = file_index
                        prog["files_total"] = files_total_workbooks
                        prog["file_name"] = file_label
                        prog["sheet_name"] = payload.get("sheet_name")
                        prog["sheet_index"] = payload.get("sheet_index")
                        prog["sheets_in_workbook"] = payload.get("sheets_total")
                        prog["rows"] = payload.get("rows")
                        prog["cols"] = payload.get("cols")
                        prog["message"] = _scan_msg_for_event(payload, file_index, file_label)
                        prog["updated_at"] = datetime.now(timezone.utc).isoformat()

                    return _on_sheet_progress

                try:
                    dim_hints = _dimension_hints_for_workbook_sources(registry, cache_key)
                    allowed_sheets = _registry_sheet_names_for_workbook(registry, job, cache_key)
                    grids, _ = visual_normalizer.normalize_workbook(
                        fp,
                        sheet_dimension_hints=dim_hints or None,
                        on_progress=_make_sheet_progress_handler(fi, fname),
                        only_sheet_names=allowed_sheets,
                    )
                    file_grids_cache[cache_key] = grids
                except Exception as exc:
                    logger.warning("normalize_workbook failed for %s: %s", fp, exc)
                    file_grids_cache[cache_key] = {}
            grids = file_grids_cache[cache_key]
            if sheet_name not in grids:
                logger.warning("Sheet %r missing in workbook %s", sheet_name, fp)
                continue
            grid = grids[sheet_name]
            grid_df = grid.to_dataframe()
            visual_patterns = {
                "sheet_name": sheet_name,
                "total_rows": grid.total_rows,
                "total_cols": grid.total_cols,
                "merged_ranges": list(grid.merged_ranges or []),
                "hidden_rows": len(grid.hidden_rows or []),
                "hidden_cols": len(grid.hidden_cols or []),
            }
            sheet_payloads.append({
                "registry_index": registry_index,
                "source_id": src.get("source_id"),
                "sheet_name": sheet_name,
                "filename": src.get("file_name") or job.get("filename"),
                "grid_df": grid_df,
                "visual_patterns": visual_patterns,
            })

        job["demarcation_scan_progress"] = {
            "active": True,
            "phase": "demarcation_batch",
            "message": (
                f"Workbooks normalized. Running block detection on {len(sheet_payloads)} source sheet(s)"
                + (" (batched AI only for uncertain regions)…" if use_ai_classification else "…")
            ),
            "sources_total": len(sheet_payloads),
            "files_total": files_total_workbooks,
            "updated_at": datetime.now(timezone.utc).isoformat(),
        }

        import asyncio
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        try:
            proposals_by_id, shared_llm_error, batch_meta = loop.run_until_complete(
                demarcator_tool.batch_propose_demarcation(
                    sheet_payloads,
                    use_ai_classification=use_ai_classification,
                )
            )
        finally:
            loop.close()

        for src in registry:
            sid = src.get("source_id")
            if not sid or sid in proposals_by_id:
                continue
            proposals_by_id[sid] = demarcator_tool._assemble_sheet_proposal(
                [],
                visual_patterns={"sheet_name": src.get("sheet_name"), "total_rows": 0, "total_cols": 0},
                sheet_name=src.get("sheet_name") or "",
                filename=src.get("file_name") or job.get("filename") or "",
                source_id=sid,
                use_ai_classification=use_ai_classification,
                llm_error=None,
                gating_meta={"llm_skipped_no_ambiguous": True, "llm_candidate_count": 0},
                batch_note="[Sheet could not be loaded for batch demarcation.]",
            )

        for prop in proposals_by_id.values():
            _decorate_demarcation_proposal_excel_ranges(prop)

        job["demarcation_batch_proposals"] = proposals_by_id
        job["demarcation_batch_meta"] = {
            "shared_llm_error": shared_llm_error,
            **batch_meta,
        }

        for sid, prop in proposals_by_id.items():
            src = resolve_registry_source(job, sid)
            fp = str(src.get("file_path") or job.get("file_path") or "")
            sn = str(prop.get("sheet_name") or src.get("sheet_name") or "")
            grid = None
            hints: Dict[str, Tuple[int, int]] = {}
            if fp:
                try:
                    cache_key = str(Path(fp).resolve())
                except OSError:
                    cache_key = fp
                grids = file_grids_cache.get(cache_key) or {}
                grid = grids.get(sn) if sn else None
                hints = _dimension_hints_for_workbook_sources(registry, cache_key)
            ps = hints.get(sn) if sn else None
            prop["sheet_shape_audit"] = {
                "pandas_rows": int(ps[0]) if ps else None,
                "pandas_cols": int(ps[1]) if ps else None,
                "visual_grid_rows": int(grid.total_rows) if grid else None,
                "visual_grid_cols": int(grid.total_cols) if grid else None,
                "dimension_hint_used": bool(hints),
                "include_hidden_rows_cols": True,
                "note": (
                    "pandas_* is the raw sheet extent (header=None). visual_grid_* is the grid used for "
                    "block detection. Saved layout uses absolute coordinates; mapping and processing load "
                    "the full sheet via pandas — not limited to preview row caps."
                ),
            }
            sheet_nm = prop.get("sheet_name") or src.get("sheet_name")
            job.setdefault("source_scope_registry", {})[sid] = build_scoped_source(
                sheet_name=sheet_nm,
                blocks=prop.get("blocks", []),
                user_selected_header_row=job.get("user_selected_header_row"),
            )
            job_manager.update_source_metadata(job_id, sheet_nm, {
                "sheet_name": sheet_nm,
                "contains_main_data": True,
                "status": "demarcation_ready",
            }, source_id=sid)

        pick_sid = None
        if active_source_id and active_source_id in proposals_by_id:
            pick_sid = active_source_id
        elif active_sheet:
            for src in registry:
                if src.get("sheet_name") == active_sheet:
                    cand = src.get("source_id")
                    if cand and cand in proposals_by_id:
                        pick_sid = cand
                        break
        if not pick_sid and registry:
            pick_sid = registry[0].get("source_id")
        if pick_sid and pick_sid in proposals_by_id:
            job_manager.pending_demarcations[job_id] = proposals_by_id[pick_sid]
            job["scoped_source"] = dict(job["source_scope_registry"].get(pick_sid) or {})

        capture_llm_metadata(job_id)
        job_manager.update_job_status(job_id, "demarcation_ready", "Batch demarcation generated", "Demarcation")
        _record_job_debug(
            job,
            label="Batch demarcation proposed",
            phase="demarcation",
            module="demarcation",
            operation="propose_all",
            status="success",
            summary=f"Proposals for {len(proposals_by_id)} source(s)",
            metadata={
                "source_count": len(proposals_by_id),
                "shared_llm_error": bool(shared_llm_error),
            },
        )

        sources_list = [{"source_id": k, "proposal": v} for k, v in proposals_by_id.items()]
        job["demarcation_scan_progress"] = {
            "active": False,
            "phase": "done",
            "message": "Scan complete. Review blocks below or switch source tabs.",
            "sources_in_registry": len(registry),
            "files_total": files_total_workbooks,
            "updated_at": datetime.now(timezone.utc).isoformat(),
        }
        resp = make_response(jsonify({
            "job_id": job_id,
            "batch": {"shared_llm_error": shared_llm_error, **batch_meta},
            "proposals_by_source_id": proposals_by_id,
            "sources": sources_list,
            "proposal": proposals_by_id.get(pick_sid) if pick_sid else {},
        }))
        resp.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        resp.headers["Pragma"] = "no-cache"
        return resp
    except Exception as e:
        capture_llm_metadata(job_id)
        logger.exception(f"Batch demarcation failed: {e}")
        err_job = job_manager.get_job(job_id)
        if err_job is not None:
            err_job["demarcation_scan_progress"] = {
                "active": False,
                "phase": "error",
                "message": str(e),
                "updated_at": datetime.now(timezone.utc).isoformat(),
            }
        return jsonify({"error": str(e)}), 500


@app.route('/api/demarcation/propose/<job_id>', methods=['GET'])
def propose_demarcation(job_id):
    """Generate an AI-powered structure demarcation proposal.

    Roadmap (not implemented here): parallel small LLM calls per visible sheet with shared global
    context (all file/sheet names, shapes, noise/header hints, target schema constraints), then a
    deterministic merge pass (no LLM) for conflicting mappings, duplicate schemas, and grain
    reconciliation (stack / union / isolate). This route stays one active sheet per request.
    """
    selected_sheet = request.args.get("sheet_name")
    selected_source_id = request.args.get("source_id")
    use_ai_param = (request.args.get("use_ai") or "true").strip().lower()
    use_ai_classification = use_ai_param not in ("0", "false", "no", "off")
    
    job = job_manager.get_job(job_id)
    if not job:
        return jsonify({"error": "Job not found"}), 404
        
    try:
        # Load Config for LLM credentials
        config = load_user_config()
        
        # We need an agent instance to get the LLM client
        agent = _build_agent_from_config(config)
        
        from sia.agent.demarcator import StructureDemarcator
        from sia.modules.visual_normalizer import VisualNormalizer
        
        demarcator_tool = StructureDemarcator(agent.llm_client)
        visual_normalizer = VisualNormalizer(include_hidden=True)

        source = resolve_job_source(job, source_id=selected_source_id, sheet_name=selected_sheet)
        file_path = source.get("file_path") or job["file_path"]

        dim_hints: Dict[str, Tuple[int, int]] = {}
        if not str(file_path).lower().endswith(".csv"):
            try:
                with pd.ExcelFile(file_path) as xf:
                    for sn in xf.sheet_names:
                        try:
                            df0 = pd.read_excel(xf, sn, header=None)
                            dim_hints[str(sn)] = (len(df0), len(df0.columns))
                        except Exception:
                            continue
            except Exception as _hint_exc:
                logger.warning("pandas dimension hints for demarcation failed: %s", _hint_exc)

        allowed_one = {str(selected_sheet)} if selected_sheet else None
        grids, _ = visual_normalizer.normalize_workbook(
            file_path,
            sheet_dimension_hints=dim_hints or None,
            only_sheet_names=allowed_one,
        )

        target_sheet = selected_sheet
        if not target_sheet or target_sheet not in grids:
            target_sheet = list(grids.keys())[0] if grids else None

        if not target_sheet:
            return jsonify({"error": "No sheets found in file"}), 400

        grid = grids[target_sheet]
        grid_df = grid.to_dataframe()
        pandas_shape = dim_hints.get(target_sheet) if dim_hints else None
        if pandas_shape is None and not str(file_path).lower().endswith(".csv"):
            pandas_shape = _pandas_sheet_shape(str(file_path), target_sheet)
        visual_patterns = {
            "sheet_name": target_sheet,
            "total_rows": grid.total_rows,
            "total_cols": grid.total_cols,
            "merged_ranges": list(grid.merged_ranges or []),
            "hidden_rows": len(grid.hidden_rows or []),
            "hidden_cols": len(grid.hidden_cols or []),
        }
        
        # Run async proposal in sync context
        import asyncio
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        proposal = loop.run_until_complete(
            demarcator_tool.propose_demarcation(
                grid_df,
                visual_patterns=visual_patterns,
                use_ai_classification=use_ai_classification,
            )
        )
        
        # Add metadata to proposal for UI display
        proposal["filename"] = source.get("file_name") or job["filename"]
        proposal["sheet_name"] = target_sheet
        proposal["source_id"] = source.get("source_id")
        _decorate_demarcation_proposal_excel_ranges(proposal)

        proposal["sheet_shape_audit"] = {
            "pandas_rows": int(pandas_shape[0]) if pandas_shape else None,
            "pandas_cols": int(pandas_shape[1]) if pandas_shape else None,
            "visual_grid_rows": int(grid.total_rows),
            "visual_grid_cols": int(grid.total_cols),
            "dimension_hint_used": bool(dim_hints),
            "include_hidden_rows_cols": True,
            "note": (
                "pandas_* is the raw sheet extent (header=None). The visual grid drives block detection; "
                "it is expanded to at least the pandas extent and includes hidden rows/columns for Guided Setup."
            ),
        }

        if source.get("source_id"):
            job.setdefault("demarcation_batch_proposals", {})[source["source_id"]] = dict(proposal)

        demarcation_debug = None
        if len(proposal.get("blocks") or []) == 0:
            fp = Path(str(file_path))
            file_size = fp.stat().st_size if fp.is_file() else 0
            demarcation_debug = {
                "resolved_path": str(file_path),
                "file_size_bytes": file_size,
                "file_exists": fp.is_file(),
                "data_file_name": source.get("file_name") or job.get("filename"),
                "target_sheet": target_sheet,
                "grid_shape_rows": int(len(grid_df.index)),
                "grid_shape_cols": int(len(grid_df.columns)),
                "pandas_shape_rows": int(pandas_shape[0]) if pandas_shape else None,
                "pandas_shape_cols": int(pandas_shape[1]) if pandas_shape else None,
                "non_empty_cells": _count_non_empty_cells_in_grid_df(grid_df),
                "workbook_sheets": list(grids.keys()),
                "source_id_used": source.get("source_id"),
                "proposal_error": proposal.get("error"),
                "llm_status": proposal.get("llm_status"),
                "llm_error": proposal.get("llm_error"),
            }
            proposal["demarcation_debug"] = demarcation_debug
        
        job_manager.pending_demarcations[job_id] = proposal
        job["scoped_source"] = build_scoped_source(
            sheet_name=target_sheet,
            blocks=proposal.get("blocks", []),
            user_selected_header_row=job.get("user_selected_header_row"),
        )
        if source.get("source_id"):
            job.setdefault("source_scope_registry", {})[source["source_id"]] = dict(job["scoped_source"])
        job_manager.update_source_metadata(job_id, target_sheet, {
            "sheet_name": target_sheet,
            "contains_main_data": True,
            "status": "demarcation_ready",
        }, source_id=source.get("source_id"))
        capture_llm_metadata(job_id)
        job_manager.update_job_status(job_id, "demarcation_ready", "Demarcation proposal generated", "Demarcation")

        resp = make_response(jsonify({
            "job_id": job_id,
            "proposal": proposal,
            # Duplicate at root so clients never miss it if nested serialization/caching is odd
            "demarcation_debug": demarcation_debug,
        }))
        resp.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        resp.headers["Pragma"] = "no-cache"
        return resp
    except Exception as e:
        capture_llm_metadata(job_id)
        logger.exception(f"Demarcation proposal failed: {e}")
        return jsonify({"error": str(e)}), 500


@app.route('/api/demarcation/submit/<job_id>', methods=['POST'])
def submit_demarcation(job_id):
    """Submit user-validated block demarcations."""
    data = request.json or {}
    blocks = data.get("blocks", [])
    source_id = data.get("source_id")
    job = job_manager.get_job(job_id)

    if not job_id or not blocks:
        return jsonify({"error": "Missing data"}), 400

    # Prefer client-provided sheet for this source; avoids wrong sheet when switching sources.
    sheet_name = data.get("sheet_name")
    if not sheet_name and source_id and job is not None:
        sheet_name = ((job.get("demarcation_batch_proposals") or {}).get(source_id) or {}).get("sheet_name")
    if not sheet_name and job is not None:
        for reg in job.get("source_registry") or []:
            if reg.get("source_id") == source_id:
                sheet_name = reg.get("sheet_name")
                break
    if not sheet_name:
        sheet_name = (
            (job_manager.pending_demarcations.get(job_id) or {}).get("sheet_name")
            if job_id in job_manager.pending_demarcations
            else None
        ) or ((job or {}).get("scoped_source", {}) or {}).get("sheet_name")

    pending = job_manager.pending_demarcations.setdefault(job_id, {})
    pending["blocks"] = blocks
    if sheet_name:
        pending["sheet_name"] = sheet_name
    if job is not None and source_id:
        prev = (job.get("demarcation_batch_proposals") or {}).get(source_id, {})
        merged_prop = dict(prev)
        merged_prop.update(dict(job_manager.pending_demarcations.get(job_id, {})))
        merged_prop["blocks"] = blocks
        merged_prop["source_id"] = source_id
        if sheet_name:
            merged_prop["sheet_name"] = sheet_name
        job.setdefault("demarcation_batch_proposals", {})[source_id] = merged_prop
    if job is not None:
        job["scoped_source"] = build_scoped_source(
            sheet_name=sheet_name,
            blocks=blocks,
            user_selected_header_row=job.get("user_selected_header_row"),
        )
        if source_id:
            job.setdefault("source_scope_registry", {})[source_id] = dict(job["scoped_source"])
        job_manager.save_layout_registry(job_id, blocks_to_layout_registry(job, blocks, sheet_name=sheet_name, source_id=source_id), source_id=source_id)
        if source_id:
            job_manager.mark_ux_source_layout(job_id, str(source_id), True)
            main_kept = sum(
                1
                for row in (job.get("layout_registry") or [])
                if str(row.get("source_id")) == str(source_id) and layout_registry_row_is_main_data(row)
            )
            # No treat-as-data Main Data blocks → semantic column mapping does not apply; treat as done.
            if main_kept == 0:
                job_manager.mark_ux_source_mapping(job_id, str(source_id), True)
            else:
                job_manager.mark_ux_source_mapping(job_id, str(source_id), False)
    logger.info(f"[DEMARCATION] User submitted {len(blocks)} blocks for job {job_id}")
    if job is not None:
        _record_job_debug(
            job,
            label="Demarcation saved",
            phase="demarcation",
            source_id=str(source_id) if source_id else None,
            sheet_name=str(sheet_name) if sheet_name else None,
            module="demarcation",
            operation="submit",
            status="success",
            summary=f"Saved {len(blocks)} block(s) for {sheet_name or source_id or 'source'}",
            metadata={"block_count": len(blocks), "source_id": source_id},
        )

    skip_semantic_mapping = False
    if job is not None and source_id:
        skip_semantic_mapping = not any(
            layout_registry_row_is_main_data(row)
            for row in (job.get("layout_registry") or [])
            if str(row.get("source_id")) == str(source_id)
        )

    return jsonify({"success": True, "skip_semantic_mapping": skip_semantic_mapping})


@app.route('/api/demarcation/preview/<job_id>', methods=['GET'])
def get_demarcation_preview(job_id):
    """Get a data preview for a specific grid range."""
    sheet_name = request.args.get("sheet_name")
    source_id = request.args.get("source_id")
    start_row = int(request.args.get("start_row", 0))
    end_row = int(request.args.get("end_row", 0))
    start_col = int(request.args.get("start_col", 0))
    end_col = int(request.args.get("end_col", 0))
    
    job = job_manager.get_job(job_id)
    if not job:
        return jsonify({"error": "Job not found"}), 404
        
    try:
        source = resolve_job_source(job, source_id=source_id, sheet_name=sheet_name)
        file_path = source.get("file_path") or job["file_path"]
        # Load data
        if str(file_path).endswith('.csv'):
            df = pd.read_csv(file_path, header=None)
        else:
            df = pd.read_excel(file_path, sheet_name=sheet_name, header=None) if sheet_name else pd.read_excel(file_path, header=None)
        
        # Slice data (clipping to bounds) — full rectangle is authoritative; JSON body may be capped.
        preview_df = df.iloc[start_row:end_row+1, start_col:end_col+1]
        range_row_count = int(len(preview_df.index))
        range_col_count = int(len(preview_df.columns))
        truncated = range_row_count > DEMARCATION_PREVIEW_MAX_BODY_ROWS
        body_df = (
            preview_df.iloc[:DEMARCATION_PREVIEW_MAX_BODY_ROWS]
            if truncated
            else preview_df
        )
        preview_data = safe_serialize_df(body_df)

        # Column keys must match preview record keys (stringified labels; often 0..N for header=None).
        preview_columns = [str(c) for c in preview_df.columns]
        return jsonify({
            "job_id": job_id,
            "preview": preview_data,
            "columns": preview_columns,
            "excel_range": index_to_excel_range(start_row, end_row, start_col, end_col),
            "preview_truncated": truncated,
            "preview_row_count_returned": int(len(body_df.index)),
            "range_row_count": range_row_count,
            "range_col_count": range_col_count,
            "preview_row_cap": DEMARCATION_PREVIEW_MAX_BODY_ROWS,
        })
    except Exception as e:
        logger.error(f"Preview failed: {e}")
        return jsonify({"error": str(e)}), 500


@app.route('/api/jobs/<job_id>/source-inventory', methods=['GET'])
def get_source_inventory(job_id):
    """Compact file/sheet list with layout and mapping counts for navigation."""
    job = job_manager.get_job(job_id)
    if not job:
        return jsonify({"error": "Job not found"}), 404

    layout_registry = list(job.get("layout_registry") or [])
    mapping_registry = list(job.get("mapping_registry") or [])
    active_source_id = request.args.get("active_source_id") or request.args.get("source_id")
    active_sheet = request.args.get("active_sheet_name") or request.args.get("sheet_name") or ""

    sources_out: List[Dict[str, Any]] = []
    batch_props = job.get("demarcation_batch_proposals") or {}
    for src in job.get("source_registry") or []:
        sid = src.get("source_id")
        sheet_name = src.get("sheet_name") or ""
        layouts = [x for x in layout_registry if x.get("source_id") == sid]
        maps = [x for x in mapping_registry if x.get("source_id") == sid]
        pending_prop = batch_props.get(sid) if sid else None
        bl_scan = (pending_prop or {}).get("blocks") or []
        status_bits = []
        if layouts:
            status_bits.append(f"{len(layouts)} layout block(s)")
            block_count = len(layouts)
            main_n = sum(1 for x in layouts if layout_registry_row_is_main_data(x))
        elif bl_scan:
            status_bits.append(f"{len(bl_scan)} proposed block(s) (batch scan)")
            block_count = len(bl_scan)
            main_n = sum(1 for b in bl_scan if b.get("category") == "Main Data")
        else:
            status_bits.append("no saved layout")
            block_count = 0
            main_n = 0
        if maps:
            status_bits.append(f"{len(maps)} mapping col(s)")
        sources_out.append({
            "source_id": sid,
            "file_name": src.get("file_name") or job.get("filename"),
            "sheet_name": sheet_name,
            "block_count": block_count,
            "main_block_count": main_n,
            "mapping_column_count": len(maps),
            "status": "; ".join(status_bits),
            "is_active": sid == active_source_id and sheet_name == active_sheet,
            "has_saved_layout": len(layouts) > 0,
        })

    return jsonify({"job_id": job_id, "sources": sources_out})


def column_sets_by_source_from_job_mapping(job: Dict[str, Any]) -> Dict[str, List[str]]:
    """Distinct raw source column names per ``source_id`` from ``mapping_registry``."""
    out: Dict[str, List[str]] = {}
    for row in job.get("mapping_registry") or []:
        if not isinstance(row, dict):
            continue
        sid = str(row.get("source_id") or "").strip()
        if not sid:
            continue
        col = str(
            row.get("source_column")
            or row.get("column_name")
            or row.get("raw_column")
            or ""
        ).strip()
        if not col:
            continue
        bucket = out.setdefault(sid, [])
        if col not in bucket:
            bucket.append(col)
    return out


@app.route('/api/jobs/<job_id>/validate-cross-source-union', methods=['GET'])
def get_validate_cross_source_union(job_id):
    """Union readiness from saved column mappings (read-only; safe to call before process-all)."""
    job = job_manager.get_job(job_id)
    if not job:
        return jsonify({"error": "Job not found"}), 404

    column_sets = column_sets_by_source_from_job_mapping(job)
    raw_hints = request.args.get("key_hints") or ""
    key_hints = [h.strip() for h in str(raw_hints).split(",") if h.strip()]

    date_granularity_by_source: Dict[str, str] = {}
    for s in job.get("source_registry") or []:
        if not isinstance(s, dict):
            continue
        sid = str(s.get("source_id") or "").strip()
        dg = str(s.get("date_granularity") or "").strip().lower()
        if sid and dg:
            date_granularity_by_source[sid] = dg
    dg_arg: Optional[Dict[str, str]] = date_granularity_by_source or None

    report = validate_column_sets_for_union(
        column_sets,
        key_hints=key_hints or None,
        dtype_map_by_source=None,
        date_granularity_by_source=dg_arg,
    )
    n_sources = len(column_sets)
    return jsonify(
        {
            "job_id": job_id,
            "source_count": n_sources,
            "columns_mapped_per_source": {k: len(v) for k, v in column_sets.items()},
            **report,
        }
    )


@app.route('/api/mapping/matrix/<job_id>', methods=['GET'])
def get_mapping_matrix(job_id):
    """Target columns × main data blocks matrix (raw column picks per target)."""
    job = job_manager.get_job(job_id)
    if not job:
        return jsonify({"error": "Job not found"}), 404

    target_columns = load_target_columns_for_job(job)
    draft = job.get("mapping_matrix_draft") or {}
    matrix_rows: List[Dict[str, Any]] = []

    # One pd.read_excel per (file, sheet); previously every block re-read the full sheet (very slow on large sheets).
    buckets: "OrderedDict[tuple, List[Dict[str, Any]]]" = OrderedDict()
    for row in job.get("layout_registry") or []:
        if not layout_registry_row_is_main_data(row):
            continue
        src = resolve_registry_source(job, row.get("source_id"))
        fp = src.get("file_path") or job.get("file_path")
        sheet = row.get("sheet_name") or src.get("sheet_name")
        key = (str(fp or ""), sheet)
        buckets.setdefault(key, []).append(row)

    raw_cache: Dict[tuple, Optional[pd.DataFrame]] = {}
    for key, reg_rows in buckets.items():
        fp, sheet = key
        try:
            raw_cache[key] = load_raw_sheet_dataframe(fp, sheet)
        except Exception as ex:
            logger.warning("Mapping matrix could not load sheet %s %s: %s", fp, sheet, ex)
            raw_cache[key] = None

    for row in job.get("layout_registry") or []:
        if not layout_registry_row_is_main_data(row):
            continue
        src = resolve_registry_source(job, row.get("source_id"))
        fp = src.get("file_path") or job.get("file_path")
        sheet = row.get("sheet_name") or src.get("sheet_name")
        block = layout_registry_row_to_block(row)
        scoped = {"main_blocks": [block], "sheet_name": sheet}
        raw_cols: List[str] = []
        cache_key = (str(fp or ""), sheet)
        raw_df = raw_cache.get(cache_key)
        hdr_meta: Dict[str, Any] = {"derived": False}
        try:
            if raw_df is not None:
                df, _scope = load_scoped_dataframe(str(fp), sheet, scoped, raw_df=raw_df)
                raw_cols = [str(c) for c in df.columns]
                hdr_meta = (_scope or {}).get("header_derivation") or {"derived": False}
        except Exception as ex:
            logger.warning("Mapping matrix row skipped: %s", ex)

        lid = row.get("layout_id")
        row_draft = draft.get(lid) or draft.get(str(lid)) or {}
        matrix_rows.append({
            "layout_id": lid,
            "source_id": row.get("source_id"),
            "file_name": src.get("file_name") or job.get("filename"),
            "sheet_name": sheet,
            "block_id": row.get("block_id"),
            "excel_range": index_to_excel_range(
                int(row.get("start_row", 0)),
                int(row.get("end_row", 0)),
                int(row.get("start_col", 0)),
                int(row.get("end_col", 0)),
            ),
            "raw_columns": raw_cols,
            "cell_values": row_draft,
            "header_derivation": hdr_meta,
        })

    return jsonify({
        "job_id": job_id,
        "target_columns": target_columns,
        "rows": matrix_rows,
        "draft": draft,
    })


@app.route('/api/mapping/matrix/<job_id>', methods=['POST'])
def save_mapping_matrix(job_id):
    """Persist mapping matrix draft on the job (in-memory)."""
    job = job_manager.get_job(job_id)
    if not job:
        return jsonify({"error": "Job not found"}), 404
    data = request.json or {}
    matrix = data.get("matrix")
    if matrix is None:
        return jsonify({"error": "Missing matrix"}), 400
    job["mapping_matrix_draft"] = matrix
    return jsonify({"success": True})


@app.route('/api/mapping/matrix/<job_id>/suggest', methods=['POST'])
def suggest_mapping_matrix(job_id):
    """Fill target→source suggestions per main block (LLM or heuristic-only)."""
    job = job_manager.get_job(job_id)
    if not job:
        return jsonify({"error": "Job not found"}), 404

    data = request.json or {}
    use_llm = bool(data.get("use_llm", True))
    target_columns = load_target_columns_for_job(job)

    try:
        config = load_user_config()
        agent = _build_agent_from_config(config)
        llm_client = agent.llm_client if use_llm else None
        ai_mapper = SchemaMapper(llm_client=llm_client, prompts_dir="prompts")

        suggestions: Dict[str, Dict[str, str]] = {}

        buckets2: "OrderedDict[tuple, List[Dict[str, Any]]]" = OrderedDict()
        for row in job.get("layout_registry") or []:
            if not layout_registry_row_is_main_data(row):
                continue
            src = resolve_registry_source(job, row.get("source_id"))
            fp = src.get("file_path") or job.get("file_path")
            sheet = row.get("sheet_name") or src.get("sheet_name")
            key = (str(fp or ""), sheet)
            buckets2.setdefault(key, []).append(row)

        raw_cache2: Dict[tuple, Optional[pd.DataFrame]] = {}
        for key, _rows in buckets2.items():
            fp, sheet = key
            try:
                raw_cache2[key] = load_raw_sheet_dataframe(fp, sheet)
            except Exception as ex:
                logger.warning("Matrix suggest could not load sheet %s %s: %s", fp, sheet, ex)
                raw_cache2[key] = None

        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        try:
            for row in job.get("layout_registry") or []:
                if not layout_registry_row_is_main_data(row):
                    continue
                src = resolve_registry_source(job, row.get("source_id"))
                fp = src.get("file_path") or job.get("file_path")
                sheet = row.get("sheet_name") or src.get("sheet_name")
                block = layout_registry_row_to_block(row)
                scoped = {"main_blocks": [block], "sheet_name": sheet}
                cache_key = (str(fp or ""), sheet)
                raw_df = raw_cache2.get(cache_key)
                try:
                    if raw_df is None:
                        continue
                    df, _ = load_scoped_dataframe(str(fp), sheet, scoped, raw_df=raw_df)
                    if df is None or len(df.columns) == 0:
                        continue
                    mapping = loop.run_until_complete(
                        ai_mapper.propose_mapping(
                            df,
                            target_columns=target_columns,
                            allow_heuristic_fallback=True,
                        )
                    )
                    lid = row.get("layout_id")
                    if lid:
                        suggestions[str(lid)] = invert_column_mappings_to_targets(mapping)
                except Exception as ex:
                    logger.warning("Matrix suggest failed for layout %s: %s", row.get("layout_id"), ex)
        finally:
            loop.close()

        capture_llm_metadata(job_id)
        return jsonify({"suggestions": suggestions})
    except Exception as e:
        logger.error(f"Mapping matrix suggest failed: {e}")
        capture_llm_metadata(job_id)
        return jsonify({"error": str(e)}), 500


def _date_hints_for_job_source(
    job: Dict[str, Any],
    job_id: str,
    source: Dict[str, Any],
    selected_sheet: Optional[str],
    pending_mapping: List[Dict[str, Any]],
) -> Dict[str, Any]:
    """Sample scoped dataframe and return ``compute_date_hints_for_dataframe`` for Guided Setup defaults."""
    try:
        file_path = Path(source.get("file_path") or job.get("file_path") or "")
        if not file_path.exists():
            return compute_date_hints_for_dataframe(None)
        sid = source.get("source_id")
        blocks = resolve_demarcation_blocks_for_source(job, job_id, sid)
        if not blocks:
            blocks = (job.get("scoped_source") or {}).get("main_blocks") or []
        sheet = selected_sheet or source.get("sheet_name") or (job.get("scoped_source") or {}).get("sheet_name")
        scoped = build_scoped_source(
            sheet_name=str(sheet or ""),
            blocks=blocks,
            user_selected_header_row=job.get("user_selected_header_row"),
        )
        df_map, _ = load_scoped_dataframe(str(file_path), str(sheet or ""), scoped)
        tt = load_target_template_for_job(job)
        return compute_date_hints_for_dataframe(
            df_map,
            approved_mappings=list(pending_mapping or []),
            target_template=tt,
        )
    except Exception as exc:
        logger.warning("date_hints_for_job_source failed: %s", exc)
        return {
            "suggested_date_shape": "unknown",
            "suggested_date_granularity": "",
            "suggested_range_start_col": "",
            "suggested_range_end_col": "",
            "confidence": 0.0,
            "rationale": f"error:{exc}",
        }


def _normalize_mapping_sheet_key(name: Optional[str]) -> str:
    return str(name or "").strip().casefold()


def apply_same_sheet_name_mapping_reuse(
    job: Dict[str, Any],
    *,
    current_source_id: str,
    sheet_name: Optional[str],
    mapping_cols: List[Dict[str, Any]],
) -> Optional[Dict[str, Any]]:
    """
    When another source has saved mapping for the same sheet name (case-insensitive),
    copy target_column, role, and date_semantic onto current columns by matching source column name.
    """
    if not mapping_cols or not current_source_id:
        return None
    key = _normalize_mapping_sheet_key(sheet_name)
    if not key:
        return None
    reg = list(job.get("mapping_registry") or [])
    by_sid: Dict[str, List[Dict[str, Any]]] = {}
    for r in reg:
        if not isinstance(r, dict):
            continue
        sid = str(r.get("source_id") or "")
        if not sid or sid == str(current_source_id):
            continue
        if _normalize_mapping_sheet_key(r.get("sheet_name")) != key:
            continue
        by_sid.setdefault(sid, []).append(r)
    if not by_sid:
        return None
    donor_sid = max(by_sid.keys(), key=lambda s: len(by_sid[s]))
    donor_by_col: Dict[str, Dict[str, Any]] = {}
    for r in by_sid[donor_sid]:
        col = str(r.get("source_column") or r.get("column_name") or "").strip()
        if col:
            donor_by_col[col] = r
    n = 0
    for m in mapping_cols:
        if not isinstance(m, dict):
            continue
        cn = str(m.get("column_name") or "").strip()
        if cn not in donor_by_col:
            continue
        d = donor_by_col[cn]
        tc = str(d.get("target_column") or "").strip()
        if tc:
            m["target_column"] = tc
        role = str(d.get("role") or "").strip().lower()
        if role in ("primary", "supporting", "exclude"):
            m["role"] = role
        ds = str(d.get("date_semantic") or "").strip()
        if ds:
            m["date_semantic"] = ds
        m["mapping_reused_from_source_id"] = donor_sid
        m["mapping_reused_from_sheet"] = str(sheet_name or "")
        n += 1
    if n == 0:
        return None
    return {
        "donor_source_id": donor_sid,
        "sheet_key": key,
        "columns_reused": n,
    }


def _enrich_mapping_ui_rows_for_source(
    job: Dict[str, Any],
    job_id: str,
    source: Dict[str, Any],
    sheet_name: Optional[str],
    source_id: str,
    rows: List[Dict[str, Any]],
) -> List[Dict[str, Any]]:
    """Re-attach column samples when mapping is loaded from registry / draft without unique_values."""
    from sia.agent.multi_block_sheet import (
        build_scoped_source_for_block,
        get_main_blocks_for_source,
        should_process_main_blocks_separately,
    )
    from sia.agent.schema_mapper import (
        enrich_mapping_rows_from_dataframe,
        mapping_rows_need_sample_enrichment,
    )
    from sia.agent.scoped_source import build_scoped_source, load_scoped_dataframe

    if not rows or not mapping_rows_need_sample_enrichment(rows):
        return list(rows or [])

    file_path = str(source.get("file_path") or job.get("file_path") or "")
    if not file_path or not sheet_name:
        return list(rows or [])

    sid = str(source_id or "").strip()
    multi_block = should_process_main_blocks_separately(job, sid) if sid else False
    main_blocks = get_main_blocks_for_source(job, sid) if sid and multi_block else []
    block_by_id: Dict[str, Dict[str, Any]] = {}
    for block in main_blocks:
        if not isinstance(block, dict):
            continue
        bid = str(block.get("id") or block.get("block_id") or "").strip()
        if bid:
            block_by_id[bid] = block

    df_cache: Dict[str, Optional[pd.DataFrame]] = {}
    enriched_rows: List[Dict[str, Any]] = []

    def _df_for_block(block_id: str) -> Optional[pd.DataFrame]:
        if block_id in df_cache:
            return df_cache[block_id]
        block = block_by_id.get(block_id)
        if not block:
            df_cache[block_id] = None
            return None
        scoped_block = build_scoped_source_for_block(
            job,
            sid,
            block,
            sheet_name=str(sheet_name or ""),
        )
        try:
            df_block, _ = load_scoped_dataframe(file_path, sheet_name, scoped_block)
        except Exception as exc:
            logger.warning("Sample enrichment failed for block %s: %s", block_id, exc)
            df_block = None
        df_cache[block_id] = df_block
        return df_block

    if multi_block and block_by_id and any(str(r.get("block_id") or "").strip() for r in rows if isinstance(r, dict)):
        for row in rows:
            if not isinstance(row, dict):
                continue
            bid = str(row.get("block_id") or "").strip()
            df_use = _df_for_block(bid) if bid else None
            if df_use is None and "_sheet" not in df_cache:
                blocks = resolve_demarcation_blocks_for_source(job, job_id, sid or None)
                scoped = build_scoped_source(
                    sheet_name=sheet_name,
                    blocks=blocks,
                    user_selected_header_row=job.get("user_selected_header_row"),
                )
                try:
                    df_cache["_sheet"], _ = load_scoped_dataframe(file_path, sheet_name, scoped)
                except Exception as exc:
                    logger.warning("Sample enrichment failed for sheet %s: %s", sheet_name, exc)
                    df_cache["_sheet"] = None
                df_use = df_cache["_sheet"]
            enriched_rows.extend(enrich_mapping_rows_from_dataframe([row], df_use))
        return enriched_rows

    if "_sheet" not in df_cache:
        blocks = resolve_demarcation_blocks_for_source(job, job_id, sid or None)
        scoped = build_scoped_source(
            sheet_name=sheet_name,
            blocks=blocks,
            user_selected_header_row=job.get("user_selected_header_row"),
        )
        try:
            df_cache["_sheet"], _ = load_scoped_dataframe(file_path, sheet_name, scoped)
        except Exception as exc:
            logger.warning("Sample enrichment failed for sheet %s: %s", sheet_name, exc)
            df_cache["_sheet"] = None
    return enrich_mapping_rows_from_dataframe(rows, df_cache.get("_sheet"))


def mapping_registry_rows_for_ui(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Turn ``mapping_registry`` rows into the per-column dict shape used by Guided Setup (matches LLM propose)."""
    out: List[Dict[str, Any]] = []
    for row in rows or []:
        if not isinstance(row, dict):
            continue
        col_name = str(row.get("source_column") or row.get("column_name") or "").strip()
        if not col_name:
            continue
        item = dict(row)
        item["column_name"] = col_name
        item.setdefault("semantic_status", "saved")
        item.setdefault("classification", str(row.get("classification") or "") or "Unclassified")
        item.setdefault("column_type", str(row.get("source_column_type") or row.get("column_type") or "Other"))
        item.setdefault("target_match_confidence", float(row.get("target_match_confidence") or 0.0))
        item.setdefault("target_match_method", str(row.get("target_match_method") or "saved"))
        item.setdefault("confidence", float(row.get("confidence") or 0.0))
        item.setdefault("target_column", str(row.get("target_column") or "No match"))
        item.setdefault("decision", str(row.get("decision") or "Keep"))
        if row.get("block_id"):
            item["block_id"] = str(row.get("block_id"))
        if row.get("block_label"):
            item["block_label"] = str(row.get("block_label"))
        if row.get("unique_values"):
            item["unique_values"] = list(row.get("unique_values") or [])
        else:
            item.setdefault("unique_values", [])
        if row.get("stats"):
            item["stats"] = dict(row.get("stats") or {})
        else:
            item.setdefault("stats", {})
        role_raw = str(item.get("role") or "").strip().lower()
        if role_raw not in ("primary", "supporting", "exclude"):
            dec = str(item.get("decision") or "").strip().lower()
            if dec == "discard":
                item["role"] = "exclude"
        out.append(item)
    return out


def _set_mapping_propose_progress(
    job_id: str,
    message: str,
    *,
    phase: str = "working",
    done: bool = False,
    error: Optional[str] = None,
    sheet_name: Optional[str] = None,
    source_id: Optional[str] = None,
) -> None:
    """In-memory status for Guided Setup mapping UI (polled while ``/api/mapping/propose`` runs)."""
    j = job_manager.get_job(job_id)
    if not j:
        return
    j["mapping_propose_progress"] = {
        "phase": phase,
        "message": message,
        "done": bool(done),
        "error": error,
        "sheet_name": sheet_name or "",
        "source_id": source_id or "",
        "updated_at": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
    }


@app.route('/api/mapping/propose-progress/<job_id>', methods=['GET'])
def get_mapping_propose_progress(job_id):
    """Poll mapping proposal progress (same pattern as demarcation scan-progress)."""
    job = job_manager.get_job(job_id)
    if not job:
        return jsonify({"error": "Job not found"}), 404
    return jsonify(job.get("mapping_propose_progress") or {})


@app.route('/api/mapping/propose/<job_id>', methods=['GET'])
def propose_mapping(job_id):
    """Generate an AI-powered column mapping proposal."""
    selected_sheet = request.args.get("sheet_name")
    selected_source_id = request.args.get("source_id")
    refresh_requested = request.args.get("refresh", "").lower() in {"1", "true", "yes"}
    
    job = job_manager.get_job(job_id)
    if not job:
        return jsonify({"error": "Job not found"}), 404
        
    try:
        source = resolve_job_source(job, source_id=selected_source_id, sheet_name=selected_sheet)
        file_path = Path(source.get("file_path") or job["file_path"])
        existing_scope = job.get("scoped_source", {}) or {}
        selected_sheet = selected_sheet or source.get("sheet_name") or existing_scope.get("sheet_name")
        progress_src = str(source.get("source_id") or selected_source_id or "")
        _set_mapping_propose_progress(
            job_id,
            "Preparing semantic column mapping (resolving sheet scope and template fields)…",
            phase="starting",
            sheet_name=str(selected_sheet or ""),
            source_id=progress_src,
        )
        target_columns = load_target_columns_for_job(job)
        target_template = load_target_template_for_job(job)
        primary_targets = primary_target_columns(target_template)

        resolved_sid = str(source.get("source_id") or selected_source_id or "")

        # Prefer persisted mapping_registry over in-memory pending so returning to a sheet
        # always reflects the last saved submit (roles, excludes, targets).
        from sia.agent.multi_block_sheet import (
            attach_block_metadata_to_mapping_rows,
            block_ui_label,
            build_scoped_source_for_block,
            flatten_mapping_blocks,
            get_main_blocks_for_source,
            group_mapping_rows_by_block,
            should_process_main_blocks_separately,
        )

        multi_block_sheet = should_process_main_blocks_separately(job, resolved_sid)
        main_blocks_for_mapping = get_main_blocks_for_source(job, resolved_sid) if multi_block_sheet else []

        if not refresh_requested and resolved_sid:
            registry_rows = [
                r
                for r in (job.get("mapping_registry") or [])
                if str(r.get("source_id") or "") == resolved_sid
            ]
            if registry_rows and multi_block_sheet and any(
                str(r.get("block_id") or "").strip() for r in registry_rows
            ):
                sections = group_mapping_rows_by_block(
                    mapping_registry_rows_for_ui(registry_rows),
                    main_blocks_for_mapping,
                )
                for sec in sections:
                    sec["mapping"] = _enrich_mapping_ui_rows_for_source(
                        job,
                        job_id,
                        source,
                        selected_sheet,
                        resolved_sid,
                        list(sec.get("mapping") or []),
                    )
                ui_mapping = flatten_mapping_blocks(sections)
                if ui_mapping and all(len(s.get("mapping") or []) > 0 for s in sections):
                    job_manager.pending_schema_mappings[job_id] = ui_mapping
                    job["mapping_sheet"] = selected_sheet
                    job["mapping_source_id"] = source.get("source_id")
                    matched_targets = {
                        m.get("target_column")
                        for m in ui_mapping
                        if isinstance(m, dict) and m.get("target_column") and m.get("target_column") != "No match"
                    }
                    unmatched_targets = [t for t in target_columns if t not in matched_targets]
                    _set_mapping_propose_progress(
                        job_id,
                        "Loaded saved per-block column mapping — skipping a new AI call.",
                        phase="registry_cached",
                        done=True,
                        sheet_name=str(selected_sheet or ""),
                        source_id=progress_src,
                    )
                    dh = _date_hints_for_job_source(
                        job, job_id, source, selected_sheet, ui_mapping,
                    )
                    return jsonify({
                        "job_id": job_id,
                        "multi_block_mapping": True,
                        "mapping_blocks": sections,
                        "mapping": ui_mapping,
                        "target_columns": target_columns,
                        "primary_target_columns": sorted(primary_targets),
                        "unmatched_target_columns": unmatched_targets,
                        "filename": source.get("file_name") or job["filename"],
                        "sheets": job.get("sheets", []),
                        "current_sheet": selected_sheet,
                        "current_source_id": source.get("source_id"),
                        "header_derivation": job.get("mapping_header_derivation") or {"derived": False},
                        "date_hints": dh,
                    })
            if registry_rows and not multi_block_sheet:
                ui_mapping = _enrich_mapping_ui_rows_for_source(
                    job,
                    job_id,
                    source,
                    selected_sheet,
                    resolved_sid,
                    mapping_registry_rows_for_ui(registry_rows),
                )
                if ui_mapping:
                    job_manager.pending_schema_mappings[job_id] = ui_mapping
                    job["mapping_sheet"] = selected_sheet
                    job["mapping_source_id"] = source.get("source_id")
                    matched_targets = {
                        m.get("target_column")
                        for m in ui_mapping
                        if isinstance(m, dict) and m.get("target_column") and m.get("target_column") != "No match"
                    }
                    unmatched_targets = [t for t in target_columns if t not in matched_targets]
                    _set_mapping_propose_progress(
                        job_id,
                        "Loaded saved column mapping for this sheet — skipping a new AI call.",
                        phase="registry_cached",
                        done=True,
                        sheet_name=str(selected_sheet or ""),
                        source_id=progress_src,
                    )
                    dh = _date_hints_for_job_source(
                        job, job_id, source, selected_sheet, ui_mapping,
                    )
                    return jsonify({
                        "job_id": job_id,
                        "mapping": ui_mapping,
                        "target_columns": target_columns,
                        "primary_target_columns": sorted(primary_targets),
                        "unmatched_target_columns": unmatched_targets,
                        "filename": source.get("file_name") or job["filename"],
                        "sheets": job.get("sheets", []),
                        "current_sheet": selected_sheet,
                        "current_source_id": source.get("source_id"),
                        "header_derivation": job.get("mapping_header_derivation") or {"derived": False},
                        "date_hints": dh,
                    })

        saved_mapping = job_manager.pending_schema_mappings.get(job_id)
        saved_mapping_sheet = job.get("mapping_sheet")
        saved_mapping_source_id = job.get("mapping_source_id")
        saved_ok_for_multi_block = saved_mapping and any(
            isinstance(m, dict) and str(m.get("block_id") or "").strip()
            for m in (saved_mapping or [])
        )
        if (
            saved_mapping
            and not refresh_requested
            and saved_mapping_sheet == selected_sheet
            and (not selected_source_id or str(saved_mapping_source_id or "") == str(selected_source_id or ""))
            and (not multi_block_sheet or saved_ok_for_multi_block)
        ):
            matched_targets = {
                m.get("target_column")
                for m in saved_mapping
                if isinstance(m, dict) and m.get("target_column") and m.get("target_column") != "No match"
            }
            unmatched_targets = [t for t in target_columns if t not in matched_targets]
            _set_mapping_propose_progress(
                job_id,
                "Using saved mapping draft on the server — skipping a new AI call.",
                phase="cached",
                done=True,
                sheet_name=str(selected_sheet or ""),
                source_id=progress_src,
            )
            saved_mapping = _enrich_mapping_ui_rows_for_source(
                job,
                job_id,
                source,
                selected_sheet,
                str(resolved_sid or ""),
                list(saved_mapping or []),
            )
            dh = _date_hints_for_job_source(
                job, job_id, source, selected_sheet, saved_mapping or [],
            )
            payload_saved: Dict[str, Any] = {
                "job_id": job_id,
                "mapping": saved_mapping,
                "target_columns": target_columns,
                "primary_target_columns": sorted(primary_targets),
                "unmatched_target_columns": unmatched_targets,
                "filename": source.get("file_name") or job["filename"],
                "sheets": job.get("sheets", []),
                "current_sheet": selected_sheet,
                "current_source_id": selected_source_id,
                "header_derivation": job.get("mapping_header_derivation") or {"derived": False},
                "date_hints": dh,
            }
            if multi_block_sheet and saved_ok_for_multi_block:
                payload_saved["multi_block_mapping"] = True
                payload_saved["mapping_blocks"] = group_mapping_rows_by_block(
                    mapping_registry_rows_for_ui(saved_mapping),
                    main_blocks_for_mapping,
                )
            return jsonify(payload_saved)

        blocks = resolve_demarcation_blocks_for_source(job, job_id, selected_source_id)
        if not blocks:
            blocks = existing_scope.get("main_blocks") or []

        config = load_user_config()
        agent = StructureInferenceAgent(
            config_path=str(Path(__file__).parent / 'config' / 'semantic_config.yaml'),
            api_key=config.get('api_key') or os.environ.get('GEMINI_API_KEY'),
            model_name=config.get('llm_model') or config.get('model_name'),
            azure_endpoint=config.get('azure_endpoint'),
            azure_deployment=config.get('azure_deployment'),
            azure_api_version=config.get('azure_api_version'),
            azure_ssl_verify=config.get('azure_ssl_verify', True),
        )
        from sia.debug.llm_observer import init_observer
        init_observer(run_id=f"mapping_{job_id}", model_id=getattr(agent, "model_name", "") or "")
        ai_mapper = SchemaMapper(llm_client=agent.llm_client, prompts_dir="prompts")

        import asyncio
        try:
            loop = asyncio.get_event_loop()
        except RuntimeError:
            loop = asyncio.new_event_loop()
            asyncio.set_event_loop(loop)

        def _run_propose_on_df(df_block):
            return loop.run_until_complete(
                ai_mapper.propose_mapping(
                    df_block,
                    target_columns=target_columns,
                    allow_heuristic_fallback=False,
                    primary_targets=primary_targets,
                )
            )

        if multi_block_sheet and len(main_blocks_for_mapping) >= 2:
            mapping_blocks_sections: List[Dict[str, Any]] = []
            all_flat: List[Dict[str, Any]] = []
            last_header_derivation: Dict[str, Any] = {"derived": False}
            for i, block in enumerate(main_blocks_for_mapping):
                if not isinstance(block, dict):
                    continue
                bid = str(block.get("id") or block.get("block_id") or f"block_{i}").strip()
                _set_mapping_propose_progress(
                    job_id,
                    f"Mapping block {i + 1} of {len(main_blocks_for_mapping)} ({bid})…",
                    phase="prepare",
                    sheet_name=str(selected_sheet or ""),
                    source_id=progress_src,
                )
                scoped_block = build_scoped_source_for_block(
                    job,
                    resolved_sid,
                    block,
                    sheet_name=str(selected_sheet or ""),
                )
                df_block, resolved_scope = load_scoped_dataframe(
                    str(file_path),
                    selected_sheet,
                    scoped_block,
                )
                n_rows, n_cols = (
                    (len(df_block), len(df_block.columns)) if df_block is not None else (0, 0)
                )
                _set_mapping_propose_progress(
                    job_id,
                    f"Block {i + 1}/{len(main_blocks_for_mapping)}: {n_rows:,} rows × {n_cols} cols — AI mapping…",
                    phase="llm",
                    sheet_name=str(selected_sheet or ""),
                    source_id=progress_src,
                )
                block_mapping = _run_propose_on_df(df_block)
                block_mapping = attach_block_metadata_to_mapping_rows(
                    block_mapping, block, index=i
                )
                excel_rng = str(block.get("excel_range") or "").strip()
                if not excel_rng:
                    c = block.get("coordinates", block)
                    if isinstance(c, dict):
                        excel_rng = index_to_excel_range(
                            c.get("start_row", 0),
                            c.get("end_row", 0),
                            c.get("start_col", 0),
                            c.get("end_col", 0),
                        )
                mapping_blocks_sections.append(
                    {
                        "block_id": bid,
                        "block_label": block_ui_label(block, i),
                        "excel_range": excel_rng,
                        "mapping": block_mapping,
                    }
                )
                all_flat.extend(block_mapping)
                last_header_derivation = (
                    resolved_scope.get("header_derivation") or last_header_derivation
                )
            capture_llm_metadata(job_id)
            job["mapping_header_derivation"] = last_header_derivation
            matched_targets = {
                m.get("target_column")
                for m in all_flat
                if isinstance(m, dict) and m.get("target_column") and m.get("target_column") != "No match"
            }
            unmatched_targets = [t for t in target_columns if t not in matched_targets]
            job_manager.update_job_status(
                job_id, "mapping_ready", "Per-block semantic mapping generated", "Schema Mapping"
            )
            _reuse_meta = apply_same_sheet_name_mapping_reuse(
                job,
                current_source_id=str(resolved_sid),
                sheet_name=selected_sheet,
                mapping_cols=all_flat,
            )
            job_manager.pending_schema_mappings[job_id] = all_flat
            job["mapping_sheet"] = selected_sheet
            job["mapping_source_id"] = source.get("source_id")
            _tt_hints = load_target_template_for_job(job)
            date_hints = compute_date_hints_for_dataframe(
                None,
                approved_mappings=all_flat,
                target_template=_tt_hints,
            )
            _set_mapping_propose_progress(
                job_id,
                f"Semantic mapping ready ({len(main_blocks_for_mapping)} blocks, {len(all_flat)} columns).",
                phase="complete",
                done=True,
                sheet_name=str(selected_sheet or ""),
                source_id=progress_src,
            )
            return jsonify({
                "job_id": job_id,
                "multi_block_mapping": True,
                "mapping_blocks": mapping_blocks_sections,
                "mapping": all_flat,
                "mapping_reuse": _reuse_meta or {},
                "target_columns": target_columns,
                "primary_target_columns": sorted(primary_targets),
                "unmatched_target_columns": unmatched_targets,
                "filename": source.get("file_name") or job["filename"],
                "sheets": job.get("sheets", []),
                "current_sheet": selected_sheet,
                "current_source_id": source.get("source_id"),
                "header_derivation": last_header_derivation,
                "date_hints": date_hints,
            })

        scoped_source = build_scoped_source(
            sheet_name=selected_sheet,
            blocks=blocks,
            user_selected_header_row=job.get("user_selected_header_row"),
        )
        context_blocks = scoped_source.get("context_blocks", [])
        job["context_metadata_blocks"] = context_blocks
        df_to_map, resolved_scope = load_scoped_dataframe(
            str(file_path),
            selected_sheet,
            scoped_source,
        )
        job["scoped_source"] = resolved_scope
        job["mapping_header_derivation"] = resolved_scope.get("header_derivation") or {"derived": False}
        if source.get("source_id"):
            job.setdefault("source_scope_registry", {})[source["source_id"]] = resolved_scope

        n_rows, n_cols = (len(df_to_map), len(df_to_map.columns)) if df_to_map is not None else (0, 0)
        _set_mapping_propose_progress(
            job_id,
            f"Loaded scoped data ({n_rows:,} rows × {n_cols} columns). Building prompts and initializing the mapping model…",
            phase="prepare",
            sheet_name=str(selected_sheet or ""),
            source_id=progress_src,
        )

        _set_mapping_propose_progress(
            job_id,
            "Waiting for AI column-mapping response (model latency can be several minutes on large sheets)…",
            phase="llm",
            sheet_name=str(selected_sheet or ""),
            source_id=progress_src,
        )
        mapping = _run_propose_on_df(df_to_map)
        _set_mapping_propose_progress(
            job_id,
            "AI mapping response received; validating and merging with column samples…",
            phase="merge",
            sheet_name=str(selected_sheet or ""),
            source_id=progress_src,
        )
        capture_llm_metadata(job_id)
        matched_targets = {
            m.get("target_column")
            for m in mapping
            if m.get("target_column") and m.get("target_column") != "No match"
        }
        unmatched_targets = [t for t in target_columns if t not in matched_targets]
        job_manager.update_job_status(job_id, "mapping_ready", "Semantic column mapping generated", "Schema Mapping")

        _reuse_meta = apply_same_sheet_name_mapping_reuse(
            job,
            current_source_id=str(resolved_sid),
            sheet_name=selected_sheet,
            mapping_cols=mapping,
        )

        job_manager.pending_schema_mappings[job_id] = mapping
        job["mapping_sheet"] = selected_sheet
        job["mapping_source_id"] = source.get("source_id")

        _pending_map = job_manager.pending_schema_mappings.get(job_id) or []
        _tt_hints = load_target_template_for_job(job)
        date_hints = compute_date_hints_for_dataframe(
            df_to_map,
            approved_mappings=_pending_map,
            target_template=_tt_hints,
        )

        _set_mapping_propose_progress(
            job_id,
            f"Semantic mapping ready ({len(mapping)} source columns).",
            phase="complete",
            done=True,
            sheet_name=str(selected_sheet or ""),
            source_id=progress_src,
        )
        return jsonify({
            "job_id": job_id,
            "mapping": mapping,
            "mapping_reuse": _reuse_meta or {},
            "target_columns": target_columns,
            "primary_target_columns": sorted(primary_targets),
            "unmatched_target_columns": unmatched_targets,
            "filename": source.get("file_name") or job["filename"],
            "sheets": job.get("sheets", []),
            "current_sheet": selected_sheet,
            "current_source_id": source.get("source_id"),
            "header_derivation": job.get("mapping_header_derivation") or {"derived": False},
            "date_hints": date_hints,
        })
        
    except Exception as e:
        capture_llm_metadata(job_id)
        logger.error(f"Mapping proposal failed: {e}")
        _set_mapping_propose_progress(
            job_id,
            f"Mapping proposal failed: {e}",
            phase="error",
            done=True,
            error=str(e),
            sheet_name=str(request.args.get("sheet_name") or ""),
            source_id=str(request.args.get("source_id") or ""),
        )
        return jsonify({"error": str(e)}), 500
    finally:
        try:
            from sia.debug.llm_observer import clear_observer
            clear_observer()
        except Exception:
            pass


@app.route('/api/mapping/submit/<job_id>', methods=['POST'])
def submit_mapping(job_id):
    """Submit user-validated column classifications."""
    data = request.json or {}
    mapping = data.get("mapping", [])
    current_sheet = data.get("sheet_name")
    current_source_id = data.get("source_id")

    _ephemeral = ("mapping_reused_from_source_id", "mapping_reused_from_sheet")
    for row in mapping or []:
        if isinstance(row, dict):
            for k in _ephemeral:
                row.pop(k, None)
    
    if not job_id or not mapping:
        return jsonify({"error": "Missing data"}), 400
        
    job_manager.pending_schema_mappings[job_id] = mapping
    job_manager.save_mapping_registry(job_id, mapping, sheet_name=current_sheet, source_id=current_source_id)
    job = job_manager.get_job(job_id)
    if job is not None:
        job["mapping_sheet"] = current_sheet or job.get("mapping_sheet") or job.get("scoped_source", {}).get("sheet_name")
        job["mapping_source_id"] = current_source_id or job.get("mapping_source_id")
        resolved_sid = current_source_id or job_manager.get_source_id(
            job_id, current_sheet or (job.get("scoped_source") or {}).get("sheet_name")
        )
        if resolved_sid:
            job_manager.mark_ux_source_mapping(job_id, str(resolved_sid), True)
        unresolved = sum(
            1
            for row in mapping or []
            if isinstance(row, dict)
            and str(row.get("target_column") or "").strip().lower() in ("", "no match", "nomatch", "no-match")
        )
        _record_job_debug(
            job,
            label="Schema mapping saved",
            phase="mapping",
            source_id=str(resolved_sid or current_source_id or ""),
            sheet_name=str(current_sheet or ""),
            module="mapping",
            operation="submit",
            status="success",
            summary=f"Saved {len(mapping)} column mapping(s); unresolved: {unresolved}",
            metadata={
                "mapping_rows": len(mapping),
                "unresolved_targets": unresolved,
                "source_id": resolved_sid or current_source_id,
            },
        )
    
    return jsonify({
        "success": True,
        "message": "Mapping saved. The shared pipeline will use these rules during the native mapping stage."
    })


@app.route('/api/context/source/<job_id>', methods=['POST'])
def save_source_context(job_id):
    """Save user-authored source metadata for a sheet."""
    job = job_manager.get_job(job_id)
    if not job:
        return jsonify({"error": "Job not found"}), 404

    data = request.json or {}
    sheet_name = data.get("sheet_name")
    source_id = data.get("source_id")
    metadata = dict(data.get("metadata") or {})
    dg = str(metadata.get("date_granularity") or "").strip().lower()
    if not dg:
        return jsonify({"error": "date_granularity is required (daily, weekly, monthly, quarterly, or range)."}), 400
    if dg not in ALLOWED_DATE_GRANULARITY:
        return jsonify({
            "error": f"Invalid date_granularity {metadata.get('date_granularity')!r}. "
            f"Allowed: {', '.join(sorted(ALLOWED_DATE_GRANULARITY))}.",
        }), 400
    metadata["date_granularity"] = dg
    shape_raw = metadata.get("date_shape", "unknown")
    shape = normalize_date_shape_for_storage(shape_raw)
    if str(shape_raw or "").strip().lower() and shape == "unknown" and str(shape_raw or "").strip().lower() not in ("", "unknown"):
        return jsonify({"error": f"Invalid date_shape {shape_raw!r}. Allowed: {', '.join(sorted(ALLOWED_DATE_SHAPES))}."}), 400
    metadata["date_shape"] = shape
    rs = str(metadata.get("range_start_column") or "").strip()
    re = str(metadata.get("range_end_column") or "").strip()
    if shape == "period_span":
        if not rs or not re:
            return jsonify(
                {"error": "For Date span (two columns), range_start_column and range_end_column are required."}
            ), 400
        if rs == re:
            return jsonify({"error": "range_start_column and range_end_column must be different columns."}), 400
        if dg != "range":
            metadata["date_granularity"] = "range"
            dg = "range"
    metadata["range_start_column"] = rs
    metadata["range_end_column"] = re
    # UID ordering was removed from Guided Setup; join keys come from mappings and the template.
    metadata["uid"] = []

    updated = job_manager.update_source_metadata(job_id, sheet_name, metadata, source_id=source_id)
    if updated is None:
        return jsonify({"error": "Failed to update source metadata"}), 400

    user_notes = metadata.get("user_notes")
    if isinstance(user_notes, list):
        job["user_notes"] = user_notes

    return jsonify({
        "success": True,
        "source_metadata": updated
    })


@app.route('/api/context/rules/<job_id>', methods=['POST'])
def save_business_rules(job_id):
    """Save user-authored business rules for the active job context."""
    job = job_manager.get_job(job_id)
    if not job:
        return jsonify({"error": "Job not found"}), 404

    data = request.json or {}
    rules = data.get("rules", [])
    sheet_name = data.get("sheet_name")
    source_id = data.get("source_id")
    normalized = job_manager.save_business_rules(job_id, rules, sheet_name=sheet_name, source_id=source_id)

    return jsonify({
        "success": True,
        "rules": normalized,
        "count": len(normalized),
    })


@app.route('/api/context/<job_id>', methods=['GET'])
def get_context_packet(job_id):
    """Assemble and return the current runtime context packet for inspection."""
    job = job_manager.get_job(job_id)
    if not job:
        return jsonify({"error": "Job not found"}), 404

    selected_sheet = request.args.get("sheet_name") or (job.get("scoped_source") or {}).get("sheet_name")
    selected_source_id = request.args.get("source_id") or (job.get("mapping_source_id"))
    target_template = load_target_template_for_job(job)
    packet = build_context_packet(job, target_template=target_template, selected_sheet=selected_sheet, selected_source_id=selected_source_id)
    job["context_packet"] = packet

    return jsonify({
        "job_id": job_id,
        "context_packet": packet,
    })


@app.route('/api/review/resolve-columns', methods=['POST'])
def resolve_columns_decision():
    """Resolve a column-level Keep/Discard decision."""
    data = request.json or {}
    job_id = data.get("job_id")
    checkpoint_id = data.get("checkpoint_id")
    decisions = data.get("decisions", {})  # {col_name: "keep" | "discard"}
    
    if not job_id or not checkpoint_id:
        return jsonify({"error": "Missing job_id or checkpoint_id"}), 400
        
    # Mark the checkpoint as resolved in the job_manager
    job_manager.resolve_checkpoint(checkpoint_id, "resolve", decisions)
    
    # Store decisions in the job state so the agent can use them when resuming
    job = job_manager.get_job(job_id)
    if job:
        if "column_decisions" not in job:
            job["column_decisions"] = {}
        job["column_decisions"].update(decisions)
        
    return jsonify({
        "success": True,
        "message": "Column decisions saved. Resuming processing..."
    })


# ===== NEW: Destructive Operation Pre-Approval Endpoints =====

# Storage for pending deletion approvals
pending_deletions = {}  # job_id -> deletion preview data


@app.route('/api/jobs/<job_id>/deletion-preview', methods=['GET'])
def get_deletion_preview(job_id):
    """Get pending deletion preview for a job awaiting HITL approval."""
    if job_id not in job_manager.pending_deletions:
        return jsonify({"error": "No pending deletions for this job", "has_pending": False}), 404
    
    preview_data = job_manager.pending_deletions[job_id]
    return jsonify({
        "has_pending": True,
        "job_id": job_id,
        "previews": preview_data.get("previews", []),
        "pending_tools": preview_data.get("pending_tools", []),
        "created_at": preview_data.get("created_at")
    })


@app.route('/api/jobs/<job_id>/approve-deletions', methods=['POST'])
def approve_deletions(job_id):
    """Approve or selectively approve pending deletions."""
    if job_id not in job_manager.pending_deletions:
        return jsonify({"error": "No pending deletions for this job"}), 404
    
    data = request.json or {}
    approve_all = data.get("approve_all", False)
    approved_indices = data.get("approved_indices", None)
    
    # Delegate to job_manager
    job_manager.approve_deletions(job_id, approved_indices if not approve_all else None)
    
    return jsonify({
        "success": True,
        "job_id": job_id,
        "message": "Deletions approved. Resume processing to continue."
    })


@app.route('/api/jobs/<job_id>/reject-deletions', methods=['POST'])
def reject_deletions(job_id):
    """Reject all pending deletions and continue processing without them."""
    if job_id not in job_manager.pending_deletions:
        return jsonify({"error": "No pending deletions for this job"}), 404
    
    job_manager.reject_deletions(job_id)
    
    return jsonify({
        "success": True,
        "job_id": job_id,
        "message": "Deletions rejected. Processing will continue without removing data."
    })


@app.route('/api/jobs/<job_id>/deletion-preview/<int:preview_index>/export', methods=['GET'])
def export_deletion_preview(job_id, preview_index):
    """Export full deletion preview as Excel for detailed review."""
    if job_id not in job_manager.pending_deletions:
        return jsonify({"error": "No pending deletions for this job"}), 404
    
    previews = job_manager.pending_deletions[job_id].get("previews", [])
    if preview_index >= len(previews):
        return jsonify({"error": "Preview index out of range"}), 404
    
    preview = previews[preview_index]
    
    # Create Excel file with the data to be deleted
    import pandas as pd
    try:
        # Get full data from the preview
        full_data = preview.get("full_deleted_data")
        if full_data is None:
            # Reconstruct from sample if full data not available
            sample_data = preview.get("sample_deleted_data", [])
            full_data = pd.DataFrame(sample_data)
        
        # Save to temp file
        export_path = OUTPUT_FOLDER / f"{job_id}_deletion_preview_{preview_index}.xlsx"
        full_data.to_excel(export_path, index=True, index_label="Original Row")
        
        return send_file(
            export_path,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f"deletion_preview_{preview.get('tool_name', 'unknown')}.xlsx"
        )
    except Exception as e:
        logger.error(f"[HITL] Failed to export deletion preview: {e}")
        return jsonify({"error": f"Export failed: {e}"}), 500


@app.route('/api/jobs/<job_id>/resume', methods=['POST'])
def resume_processing(job_id):
    """
    Resume processing after HITL approval/rejection of deletions.
    
    Call this after /approve-deletions or /reject-deletions to continue.
    """
    if job_id not in job_manager.jobs:
        return jsonify({"error": "Job not found"}), 404
    
    if job_id not in job_manager.pending_deletions:
        return jsonify({"error": "No pending HITL decisions for this job"}), 400
    
    job = job_manager.get_job(job_id)
    config = load_user_config()
    pending = job_manager.pending_deletions[job_id]
    
    if not pending.get("approved") and pending.get("rejected_at") is None:
        return jsonify({"error": "Must approve or reject deletions before resuming"}), 400
    
    try:
        resume_state = dict(pending.get("pending_state", {}))
        if not resume_state:
            return jsonify({"error": "No saved state to resume from"}), 500

        resume_state["destructive_approved"] = True
        resume_state["hitl_resume_from"] = "execute_pause"
        resume_state["hitl_pending_approval"] = False
        resume_state["hitl_pause_type"] = None
        resume_state["approved_tool_indices"] = pending.get("approved_tool_indices")
        return _run_resumed_processing(job_id, job, resume_state, config)
        
    except Exception as e:
        import traceback
        error_tb = traceback.format_exc()
        logger.error(f"[HITL] Resume failed for job {job_id}: {error_tb}")
        
        # Write to file for CLI debugging
        with open("resume_outer_error.log", "w") as f:
            f.write(f"Resume Failed: {e}\n\n{error_tb}")
            
        job["status"] = "error"
        job["error_details"] = error_tb
        return jsonify({"error": str(e), "traceback": error_tb}), 500




if __name__ == '__main__':
    print("\n" + "="*50)
    print("Structure Inference Agent - Web UI")
    print("="*50)
    print("\nOpen your browser to: http://localhost:5000")
    print("="*50 + "\n")
    
    app.run(debug=True, port=5000)
